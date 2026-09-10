from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools" / "rd_intake_windows"))
sys.path.insert(0, str(ROOT / "api"))

from rd_intake_core import (
    LocalStore,
    ReferenceEntry,
    attach_material_rows,
    build_review_rows,
    export_xlsx,
    format_ru_number,
    final_row_matches_filter,
    final_rows_filter_summary,
    final_rows_quality,
    is_material_like,
    is_temporary_culvert_work,
    load_reference_workbook,
    merge_rows,
    linked_materials_summary,
    parse_localized_number,
    parse_source_index_spec,
    parse_split_volumes,
    rd_review_progress_summary,
    rd_review_next_action_state,
    selected_target_applies_to_source_row,
    source_row_db_target,
    source_row_has_match,
    source_row_has_volume,
    source_row_is_material,
    source_row_needs_review,
    source_row_ready_for_bulk_add,
    source_hint_button_label,
    source_action_status_summary,
    source_position_block_indices,
    source_position_block_summary,
    source_scope_summary,
    source_selection_action_summary,
    source_selection_preview_hint,
    source_selection_preview_rows,
    source_raw_cell_detail_rows,
    source_raw_cells_summary,
    source_selection_cell_detail_rows,
    source_table_action_summary,
    source_table_progress_state,
    source_table_overview,
    source_table_position_summary,
    source_selection_summary,
    source_usage_from_final_rows,
    source_usage_with_ignored,
    split_row,
    suggest_entity_code,
)
from rd_parser import ParsedRdRow, _rows_from_table


def _sample_reference_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "objects"
    ws.append(["code", "name", "type", "section_code", "section_name", "pk_start", "pk_end"])
    ws.append(["OBJ-1", "Объект 1", "ROAD", "UCH_4", "Участок 4", 1000, 1100])
    ws = wb.create_sheet("work_types")
    ws.append(["code", "name", "default_unit", "analytics_tag"])
    ws.append(["WORK-1", "Устройство насыпи", "м3", "earthwork"])
    ws = wb.create_sheet("materials")
    ws.append(["code", "name", "default_unit"])
    ws.append(["MAT-1", "Песок средней крупности", "м3"])
    ws = wb.create_sheet("aliases")
    ws.append(["kind", "alias_text", "canonical_code"])
    ws.append(["work_type", "Насыпь", "WORK-1"])
    wb.save(path)


def test_pdf_table_rows_merge_continuation_bboxes_for_wrapped_positions() -> None:
    rows = _rows_from_table(
        [
            ["№", "Наименование", "Ед. изм.", "Объем"],
            ["1", "Устройство насыпи", "", ""],
            ["", "внутри региона", "м3", "12"],
        ],
        table_bbox={"x0": 10.0, "top": 10.0, "x1": 300.0, "bottom": 80.0},
        table_row_bboxes=[
            {"x0": 10.0, "top": 10.0, "x1": 300.0, "bottom": 25.0},
            {"x0": 10.0, "top": 25.0, "x1": 300.0, "bottom": 44.0},
            {"x0": 10.0, "top": 44.0, "x1": 300.0, "bottom": 63.0},
        ],
    )

    assert len(rows) == 1
    assert rows[0].raw_name == "Устройство насыпи внутри региона"
    assert rows[0].row_bbox == {"x0": 10.0, "top": 25.0, "x1": 300.0, "bottom": 63.0}


def test_parse_source_index_spec_supports_ranges_spaces_and_reverse_order() -> None:
    assert parse_source_index_spec("5-7, 10; 12 14") == {5, 6, 7, 10, 12, 14}
    assert parse_source_index_spec("9 — 7") == {7, 8, 9}
    assert parse_source_index_spec("") == set()


def test_source_table_progress_state_labels_current_completion() -> None:
    row = {"rows": "5", "unused": "3", "ready": "2", "review": "1", "unmatched": "1", "no_volume": "0", "materials": "1"}
    state = source_table_progress_state(row)

    assert state["closed"] == 2
    assert state["closed_pct"] == 40
    assert state["tone"] == "review"
    assert state["label"] == "закрыто 40% · проверить 2"

    assert source_table_progress_state({**row, "unused": "0", "ready": "0", "review": "0", "unmatched": "0"})["tone"] == "done"
    assert source_table_progress_state({**row, "review": "0", "unmatched": "0"})["tone"] == "ready"
    assert source_table_progress_state(None)["label"] == "нет таблицы"


def test_source_table_action_summary_recommends_next_table_action() -> None:
    ready = {"page_label": "Лист 6", "table_label": "Табл 2", "rows": "5", "unused": "3", "ready": "2", "review": "1", "unmatched": "1", "no_volume": "0", "materials": "1", "skipped": "1", "volume": "12 м3"}
    summary = source_table_action_summary(ready)

    assert "Текущая таблица: Лист 6 / Табл 2" in summary
    assert "закрыто 2/5 (40%)" in summary
    assert "готово к переносу 2" in summary
    assert "без БД 1" in summary
    assert "материалов 1" in summary
    assert "действие: Готовые + проверка (2)" in summary

    ready_only = {**ready, "review": "0", "unmatched": "0", "no_volume": "0"}
    assert "действие: Выбрать готовые табл. (2) или В итог готовые табл." in source_table_action_summary(ready_only)

    review = {**ready, "ready": "0", "review": "2"}
    assert "действие: К проверке таблицу" in source_table_action_summary(review)

    materials_only = {**ready, "unused": "2", "ready": "0", "review": "0", "unmatched": "0", "no_volume": "0", "materials": "2"}
    assert "действие: Материалы таблицы (2)" in source_table_action_summary(materials_only)

    closed = {**ready, "unused": "0", "ready": "0", "review": "0", "unmatched": "0", "no_volume": "0"}
    closed_summary = source_table_action_summary(closed)
    assert "закрыто 5/5 (100%)" in closed_summary
    assert closed_summary.endswith("таблица закрыта; действие: След. табл. готовые")
    assert "выберите строку или таблицу" in source_table_action_summary(None)


def test_source_table_position_summary_shows_current_table_ordinal_and_counts() -> None:
    overview = [
        {"page_label": "Лист 1", "table_label": "Табл 1", "rows": "4", "unused": "0", "ready": "0", "review": "0"},
        {"page_label": "Лист 6", "table_label": "Табл 2", "rows": "5", "unused": "3", "ready": "2", "review": "1", "skipped": "1"},
    ]

    summary = source_table_position_summary(overview, overview[1])

    assert "Таблица 2 из 2" in summary
    assert "Лист 6 / Табл 2" in summary
    assert "не в итог 3" in summary
    assert "готово 2" in summary
    assert "проверить 1" in summary
    assert "пропущено 1" in summary
    assert "закрыта" in source_table_position_summary(overview, overview[0])
    assert "всего таблиц 2" in source_table_position_summary(overview, None)
    assert "откройте PDF/DOCX/XLSX РД" in source_table_position_summary([], None)


def test_rd_review_progress_summary_combines_source_tables_and_final_quality() -> None:
    source_rows = [
        {"source_index": 1, "page_number": 1, "table_index": 1, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"},
        {"source_index": 2, "page_number": 1, "table_index": 1, "item_kind": "work", "work_code": "", "unit": "м3", "volume": 2, "status": "проверить"},
    ]
    final_rows = [
        {"source_index": 1, "object_code": "OBJ", "work_code": "WORK-1", "work_name": "Работа", "unit": "м3", "volume": 10, "status": "готово", "item_kind": "work"},
        {"source_index": 2, "object_code": "OBJ", "work_code": "", "work_name": "Не сопоставлено", "unit": "м3", "volume": 2, "status": "проверить", "item_kind": "work"},
    ]

    summary = rd_review_progress_summary(source_rows, final_rows)

    assert "таблиц 1" in summary
    assert "табл. готовые 0" in summary
    assert "готовых строк 0" in summary
    assert "табл. проверить 1" in summary
    assert "исходник 2/2 в итоге" in summary
    assert "итог 2" in summary
    assert "готово 1" in summary
    assert "проверить 1" in summary


def test_rd_review_next_action_state_prioritizes_operator_flow() -> None:
    mixed_current = {"ready": "2", "review": "1", "unmatched": "1", "no_volume": "0", "unused": "3", "materials": "1"}
    state = rd_review_next_action_state([mixed_current], mixed_current, [], has_document=True)
    assert state["action"] == "add_ready_then_review"
    assert state["label"] == "Далее: готовые + проверка (2)"
    assert "проблемный хвост" in state["hint"]
    assert state["enabled"] is True

    ready_current = {"ready": "3", "review": "0", "unmatched": "0", "no_volume": "0", "unused": "3"}
    state = rd_review_next_action_state([ready_current], ready_current, [], has_document=True)
    assert state["action"] == "add_ready_current_and_next"
    assert state["label"] == "Далее: в итог готовые табл. (3)"

    materials_current = {"ready": "0", "review": "2", "unmatched": "0", "no_volume": "0", "unused": "2", "materials": "2"}
    state = rd_review_next_action_state([materials_current], materials_current, [], has_document=True)
    assert state["action"] == "show_current_materials"
    assert state["label"] == "Далее: материалы табл. (2)"

    review_elsewhere = {"ready": "0", "review": "1", "unmatched": "0", "no_volume": "0", "unused": "1"}
    state = rd_review_next_action_state([review_elsewhere], None, [], has_document=True)
    assert state["action"] == "focus_next_review_table"
    assert state["label"] == "Далее: след. проверка (1)"

    issue_final = {"item_kind": "work", "object_code": "OBJ", "work_code": "", "work_name": "", "unit": "м3", "volume": 5}
    state = rd_review_next_action_state([], None, [issue_final])
    assert state["action"] == "show_final_issues"
    assert state["label"] == "Далее: итог проверить (1)"

    ready_final = {"item_kind": "work", "object_code": "OBJ", "work_code": "WORK", "work_name": "Работа", "unit": "м3", "volume": 5}
    state = rd_review_next_action_state([], None, [ready_final])
    assert state["action"] == "export"
    assert state["label"] == "Далее: экспорт XLSX"

    state = rd_review_next_action_state([], None, [])
    assert state["action"] == "open_document"
    assert state["enabled"] is False


def test_rd_review_progress_summary_treats_skipped_source_rows_as_processed() -> None:
    source_rows = [
        {"source_index": 1, "page_number": 1, "table_index": 1, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"},
        {"source_index": 2, "page_number": 1, "table_index": 1, "item_kind": "work", "work_code": "", "unit": "м3", "volume": "", "status": "проверить"},
    ]
    final_rows = [
        {"source_index": 1, "object_code": "OBJ", "work_code": "WORK-1", "work_name": "Работа", "unit": "м3", "volume": 10, "status": "готово", "item_kind": "work"},
    ]

    summary = rd_review_progress_summary(source_rows, final_rows, ignored_source_indices={2})

    assert "закрыто 1" in summary
    assert "табл. не в итог 0" in summary
    assert "исходник 2/2 обработано, пропущено 1, проверить 0" in summary
    assert "итог 1" in summary


def test_source_table_overview_groups_tables_with_review_counts() -> None:
    rows = [
        {"source_index": 1, "page_number": 1, "table_index": 1, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"},
        {"source_index": 2, "page_number": 1, "table_index": 1, "item_kind": "work", "work_code": "", "unit": "м3", "volume": "2,5", "status": "проверить"},
        {"source_index": 3, "page_number": 2, "table_index": 1, "item_kind": "material", "material_code": "MAT-1", "unit": "т", "volume": "", "status": "материал / проверить"},
    ]

    overview = source_table_overview(rows, source_usage={1: "final"})

    assert len(overview) == 2
    assert overview[0]["page_label"] == "Лист 1"
    assert overview[0]["table_label"] == "Табл 1"
    assert overview[0]["rows"] == "2"
    assert overview[0]["unused"] == "1"
    assert overview[0]["used"] == "1"
    assert overview[0]["ready"] == "0"
    assert overview[0]["review"] == "1"
    assert overview[0]["unmatched"] == "1"
    assert overview[0]["no_volume"] == "0"
    assert overview[0]["volume"] == "12,5 м3"
    assert overview[0]["status"] == "проверить"
    assert overview[1]["materials"] == "1"
    assert overview[1]["no_volume"] == "1"


def test_source_table_overview_uses_excel_sheet_labels() -> None:
    rows = [
        {"source_index": 1, "page_number": 1, "page_label": "Лист 1 · ВОР", "table_index": 1, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"},
        {"source_index": 2, "page_number": 2, "page_label": "Лист 2 · Ведомость", "table_index": 1, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 20, "status": "готово"},
    ]

    overview = source_table_overview(rows)

    assert [row["page_label"] for row in overview] == ["Лист 1 · ВОР", "Лист 2 · Ведомость"]
    assert source_table_position_summary(overview, overview[1]).startswith("Таблица 2 из 2 | Лист 2 · Ведомость / Табл 1")


def test_source_usage_with_ignored_and_overview_close_skipped_rows() -> None:
    rows = [
        {"source_index": 1, "page_number": 1, "table_index": 1, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"},
        {"source_index": 2, "page_number": 1, "table_index": 1, "item_kind": "work", "work_code": "", "unit": "м3", "volume": "", "status": "проверить"},
    ]
    final_rows = [
        {"source_index": 1, "object_code": "OBJ", "work_code": "WORK-1", "work_name": "Работа", "unit": "м3", "volume": 10, "status": "готово", "item_kind": "work"},
    ]

    usage = source_usage_with_ignored(final_rows, {2, "x"})
    overview = source_table_overview(rows, source_usage=usage)

    assert usage == {1: "final", 2: "skipped"}
    assert overview[0]["rows"] == "2"
    assert overview[0]["used"] == "2"
    assert overview[0]["unused"] == "0"
    assert overview[0]["skipped"] == "1"
    assert overview[0]["review"] == "0"
    assert overview[0]["unmatched"] == "0"
    assert overview[0]["no_volume"] == "0"
    assert overview[0]["status"] == "готово"


def test_source_table_overview_counts_ready_unused_work_rows() -> None:
    rows = [
        {"source_index": 1, "page_number": 1, "table_index": 1, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"},
        {"source_index": 2, "page_number": 1, "table_index": 1, "item_kind": "work", "work_code": "", "unit": "м3", "volume": 5, "status": "проверить"},
        {"source_index": 3, "page_number": 1, "table_index": 1, "item_kind": "material", "material_code": "MAT-1", "unit": "м3", "volume": 2, "status": "материал / проверить"},
    ]

    overview = source_table_overview(rows, source_usage={})

    assert overview[0]["rows"] == "3"
    assert overview[0]["unused"] == "3"
    assert overview[0]["ready"] == "1"
    assert overview[0]["review"] == "2"
    assert overview[0]["status"] == "проверить"


def test_source_selection_summary_shows_volume_types_and_locations() -> None:
    rows = [
        {"source_index": 5, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "page_number": 2, "table_index": 1, "status": "готово"},
        {"source_index": 6, "item_kind": "work", "work_code": "", "unit": "м3", "volume": "2,5", "page_number": 2, "table_index": 1, "status": "проверить"},
        {"source_index": 7, "item_kind": "material", "material_code": "MAT-1", "unit": "т", "volume": 1.25, "page_number": 3, "table_index": 2, "status": "материал / проверить"},
    ]

    summary = source_selection_summary(rows)

    assert "Выбрано строк: 3" in summary
    assert "исх.: 5, 6, 7" in summary
    assert "12,5 м3" in summary
    assert "1,25 т" in summary
    assert "работ 2" in summary
    assert "материалов 1" in summary
    assert "готово 1" in summary
    assert "проверить 2" in summary
    assert "без БД 1" in summary
    assert "л.2/т.1" in summary
    assert "л.3/т.2" in summary

    xlsx_summary = source_selection_summary([
        {"source_index": 8, "page_number": 2, "page_label": "Лист 2 · Ведомость", "table_index": 1, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"},
    ])
    assert "Лист 2 · Ведомость/т.1" in xlsx_summary


def test_source_selection_action_summary_recommends_safe_selection_action() -> None:
    rows = [
        {"source_index": 5, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"},
        {"source_index": 6, "item_kind": "work", "work_code": "", "unit": "м3", "volume": "", "status": "проверить"},
        {"source_index": 7, "item_kind": "material", "material_code": "MAT-1", "unit": "т", "volume": 1.25, "status": "материал / проверить"},
    ]

    summary = source_selection_action_summary(rows)

    assert "Выборка: строк 3" in summary
    assert "готово к итогу 1" in summary
    assert "к проверке 2" in summary
    assert "без БД 1" in summary
    assert "без объема 1" in summary
    assert "материалов 1" in summary
    assert "действие: Выбрать готовые выборки или Оставить к проверке" in summary
    assert "действие: В итог готовые выборки" in source_selection_action_summary(rows[:1])
    assert "действие: Оставить к проверке или исправить БД/объем" in source_selection_action_summary(rows[1:])


def test_source_action_status_summary_guides_current_selection_and_table() -> None:
    ready = {"source_index": 5, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"}
    issue = {"source_index": 6, "item_kind": "work", "work_code": "", "unit": "м3", "volume": "", "status": "проверить"}
    material = {"source_index": 7, "item_kind": "material", "material_code": "MAT-1", "unit": "т", "volume": 1.25, "status": "материал / проверить"}

    selected = source_action_status_summary([ready, issue, material])
    assert "Действия: выбрано 3" in selected
    assert "готово 1" in selected
    assert "проверить 2" in selected
    assert "материалы 1" in selected
    assert "без БД 1" in selected
    assert "без объема 1" in selected
    assert "возьмите готовые или откройте выборку" in selected

    table = source_action_status_summary(None, {"ready": "4", "review": "1", "unmatched": "1", "no_volume": "0", "materials": "1", "unused": "6"})
    assert table == "Действия: таблица готово 4, проверить 2 | безопасно: готовые табл. или готовые+проверка"
    assert "фильтр не показывает строк" in source_action_status_summary(None, None, table_count=3)
    assert "готовых видно 2" in source_action_status_summary(None, None, visible_count=5, ready_visible=2)
    assert source_action_status_summary() == "Действия: откройте PDF/DOCX/XLSX РД"


def test_source_selection_preview_hint_is_short_and_status_focused() -> None:
    ready = {"source_index": 5, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"}
    issue = {"source_index": 6, "item_kind": "work", "work_code": "", "unit": "м3", "volume": "", "status": "проверить"}
    material = {"source_index": 7, "item_kind": "material", "material_code": "MAT-1", "unit": "т", "volume": 1.25, "status": "материал / проверить"}

    assert source_selection_preview_hint([]) == "выберите строки таблицы"
    assert source_selection_preview_hint([ready]) == "выбрано 1 | готово 1"
    assert source_selection_preview_hint([ready, issue, material]) == "выбрано 3 | готово 1 | проверить 2 | материалы 1 | без БД 1 | без объема 1"


def test_source_position_block_selects_continuation_until_quantity_inside_table() -> None:
    rows = [
        {"source_index": 1, "page_number": 1, "table_index": 1, "row_number": 1, "raw_name": "Предыдущая позиция", "raw_unit": "м3", "raw_volume": 5},
        {"source_index": 2, "page_number": 1, "table_index": 1, "row_number": 2, "raw_name": "Устройство насыпи внутри региона"},
        {"source_index": 3, "page_number": 1, "table_index": 1, "row_number": 3, "raw_name": "на расстояние до 1 км"},
        {"source_index": 4, "page_number": 1, "table_index": 1, "row_number": 4, "raw_name": "Итого по позиции", "raw_unit": "м3", "raw_volume": 12},
        {"source_index": 5, "page_number": 1, "table_index": 2, "row_number": 1, "raw_name": "Другая таблица", "raw_unit": "м3", "raw_volume": 99},
    ]

    assert source_position_block_indices(rows, 2) == [2, 3, 4]
    assert source_position_block_indices(rows, 4) == [2, 3, 4]
    assert source_position_block_indices(rows, 5) == [5]

    summary = source_position_block_summary(rows, 2)
    assert "Позиция: исх. 2-4" in summary
    assert "строк 3" in summary
    assert "л.1 / т.1" in summary
    assert "объем 12 м3" in summary
    assert "действие: Выбрать позицию" in summary


def test_source_position_block_selects_page_break_continuations() -> None:
    rows = [
        {"source_index": 10, "page_number": 5, "table_index": 1, "row_number": 42, "raw_name": "Устройство насыпи внутри региона", "raw_unit": "м3", "raw_volume": 12},
        {"source_index": 11, "page_number": 6, "table_index": 1, "row_number": 1, "raw_name": "на расстояние до 1 км"},
        {"source_index": 12, "page_number": 6, "table_index": 1, "row_number": 2, "raw_name": "Следующая позиция", "raw_unit": "т", "raw_volume": 20},
    ]

    assert source_position_block_indices(rows, 10) == [10, 11]
    assert source_position_block_indices(rows, 11) == [10, 11]
    assert source_position_block_indices(rows, 12) == [12]

    summary = source_position_block_summary(rows, 11)
    assert "Позиция: исх. 10-11" in summary
    assert "строк 2" in summary
    assert "л.5 / т.1" in summary
    assert "объем 12 м3" in summary


def test_source_position_block_finds_quantity_after_page_break() -> None:
    rows = [
        {"source_index": 20, "page_number": 6, "table_index": 1, "row_number": 39, "raw_name": "Устройство насыпи внутри региона"},
        {"source_index": 21, "page_number": 7, "table_index": 1, "row_number": 1, "raw_name": "на расстояние свыше 1 км", "raw_unit": "м3", "raw_volume": 70292.74},
        {"source_index": 22, "page_number": 7, "table_index": 1, "row_number": 2, "raw_name": "Материал насыпи", "raw_unit": "м3", "raw_volume": 100},
    ]

    assert source_position_block_indices(rows, 20) == [20, 21]
    assert source_position_block_indices(rows, 21) == [20, 21]
    assert source_position_block_indices(rows, 22) == [22]


def test_source_hint_button_label_matches_safe_next_action() -> None:
    ready = {"source_index": 5, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"}
    issue = {"source_index": 6, "item_kind": "work", "work_code": "", "unit": "м3", "volume": "", "status": "проверить"}

    assert source_hint_button_label([ready, issue]) == "Действие: готовые (1)"
    assert source_hint_button_label([ready]) == "Действие: в итог (1)"
    assert source_hint_button_label([issue]) == "Действие: к проверке"
    assert source_hint_button_label([], {"ready": "2", "review": "1", "unused": "3"}) == "Действие: готовые+проверка (2)"
    assert source_hint_button_label([], {"ready": "2", "review": "0", "unmatched": "0", "no_volume": "0", "unused": "2"}) == "Действие: готовые табл. (2)"
    assert source_hint_button_label([], {"ready": "0", "review": "0", "unmatched": "0", "no_volume": "0", "materials": "2", "unused": "2"}) == "Действие: материалы табл. (2)"
    assert source_hint_button_label([], {"ready": "0", "review": "1", "unmatched": "0", "no_volume": "0"}) == "Действие: проверка табл."
    assert source_hint_button_label([], {"ready": "0", "review": "0", "unmatched": "0", "no_volume": "0", "unused": "4"}) == "Действие: остаток табл. (4)"
    assert source_hint_button_label([], {"ready": "0", "review": "0", "unmatched": "0", "no_volume": "0", "unused": "0"}) == "Действие: след. готовая"
    assert source_hint_button_label() == "Действие: след. таблица"


def test_source_row_review_classification_catches_missing_match_and_volume() -> None:
    ready_work = {"item_kind": "work", "work_code": "WORK-1", "volume": 12, "status": "готово"}
    unmatched_work = {"item_kind": "work", "work_code": "", "volume": 12, "status": "готово"}
    no_volume_material = {"item_kind": "material", "material_code": "MAT-1", "volume": "", "raw_volume": "", "status": "материал"}
    explicit_review = {"item_kind": "work", "work_code": "WORK-1", "volume": 1, "status": "проверить"}

    assert source_row_is_material(no_volume_material) is True
    assert source_row_has_match(ready_work) is True
    assert source_row_has_volume(ready_work) is True
    assert source_row_needs_review(ready_work) is False
    assert source_row_has_match(unmatched_work) is False
    assert source_row_needs_review(unmatched_work) is True
    assert source_row_has_volume(no_volume_material) is False
    assert source_row_needs_review(no_volume_material) is True
    assert source_row_needs_review(explicit_review) is True


def test_source_raw_cells_summary_labels_table_cells() -> None:
    row = {
        "raw_cells": ["1", "Устройство насыпи", "м3", "1 234,5"],
        "raw_name": "Устройство насыпи",
        "raw_unit": "м3",
        "raw_volume": 1234.5,
    }

    assert source_raw_cells_summary(row) == "№: 1 | Наим.: Устройство насыпи | Ед.: м3 | Объем: 1 234,5"
    assert source_raw_cells_summary({**row, "raw_cells": ["1", "Устройство насыпи", "м3", "1"], "raw_volume": 1}) == "№: 1 | Наим.: Устройство насыпи | Ед.: м3 | Объем: 1"
    assert source_raw_cells_summary({"raw_name": "Строка без ячеек"}) == "Строка без ячеек"


def test_source_raw_cell_detail_rows_return_field_value_table() -> None:
    row = {
        "raw_cells": ["1", "Устройство насыпи", "м3", "1 234,5"],
        "raw_name": "Устройство насыпи",
        "raw_unit": "м3",
        "raw_volume": 1234.5,
    }

    assert source_raw_cell_detail_rows(row) == [
        {"field": "№", "value": "1"},
        {"field": "Наим.", "value": "Устройство насыпи"},
        {"field": "Ед.", "value": "м3"},
        {"field": "Объем", "value": "1 234,5"},
    ]
    assert source_raw_cell_detail_rows({"raw_cells": ["Доп. признак"], "raw_name": "Позиция"}) == [
        {"field": "Ячейка 1", "value": "Доп. признак"},
    ]
    assert source_raw_cell_detail_rows({"raw_name": "Строка без ячеек"}) == [
        {"field": "Позиция", "value": "Строка без ячеек"},
    ]


def test_source_selection_cell_detail_rows_summarize_selected_group() -> None:
    rows = [
        {
            "source_index": 5,
            "raw_cells": ["1", "Устройство насыпи", "м3", "10"],
            "raw_name": "Устройство насыпи",
            "raw_unit": "м3",
            "raw_volume": 10,
            "status": "готово",
        },
        {
            "source_index": 6,
            "raw_cells": ["2", "Песок", "т", "1,25"],
            "raw_name": "Песок",
            "raw_unit": "т",
            "raw_volume": 1.25,
            "status": "материал / проверить",
        },
    ]

    assert source_selection_cell_detail_rows(rows) == [
        {"field": "исх. 5", "value": "№: 1 | Наим.: Устройство насыпи | Ед.: м3 | Объем: 10 | готово"},
        {"field": "исх. 6", "value": "№: 2 | Наим.: Песок | Ед.: т | Объем: 1,25 | материал / проверить"},
    ]
    capped = source_selection_cell_detail_rows(rows + [{"source_index": 7, "raw_name": "Хвост"}], max_rows=2)
    assert capped[-1] == {"field": "...", "value": "еще 1 строка"}
    capped_many = source_selection_cell_detail_rows(rows + [{"source_index": idx, "raw_name": f"Хвост {idx}"} for idx in range(7, 11)], max_rows=2)
    assert capped_many[-1] == {"field": "...", "value": "еще 4 строки"}


def test_source_selection_preview_rows_show_exact_rows_for_selection_basket() -> None:
    rows = [
        {
            "source_index": 5,
            "item_kind": "work",
            "raw_name": "Устройство насыпи",
            "raw_cells": ["1", "Устройство насыпи", "м3", "10"],
            "work_code": "WORK-1",
            "work_name": "Устройство насыпи по БД",
            "unit": "м3",
            "volume": 10,
            "page_number": 2,
            "table_index": 1,
            "row_number": 11,
            "status": "готово",
        },
        {
            "source_index": 6,
            "item_kind": "material",
            "raw_name": "Песок",
            "raw_cells": ["2", "Песок", "т", "1,25"],
            "material_code": "MAT-1",
            "material_name": "Песок средней крупности",
            "unit": "т",
            "volume": 1.25,
            "page_number": 2,
            "table_index": 1,
            "row_number": 12,
            "status": "материал / проверить",
        },
    ]

    preview = source_selection_preview_rows(rows)

    assert preview[0]["source_index"] == "5"
    assert preview[0]["location"] == "л.2/т.1/стр.11"
    xlsx_preview = source_selection_preview_rows([{**rows[0], "page_label": "Лист 2 · Ведомость"}])
    assert xlsx_preview[0]["location"] == "Лист 2 · Ведомость/т.1/стр.11"
    assert preview[0]["kind"] == "Работа"
    assert preview[0]["raw_cells"] == "№: 1 | Наим.: Устройство насыпи | Ед.: м3 | Объем: 10"
    assert preview[0]["volume"] == "10 м3"
    assert preview[0]["target"] == "WORK-1 — Устройство насыпи по БД"
    assert preview[1]["kind"] == "Материал"
    assert preview[1]["raw_cells"] == "№: 2 | Наим.: Песок | Ед.: т | Объем: 1,25"
    assert preview[1]["volume"] == "1,25 т"
    assert preview[1]["target"] == "MAT-1 — Песок средней крупности"


def test_source_selection_preview_rows_caps_long_selection() -> None:
    rows = [{"source_index": idx, "raw_name": f"Строка {idx}", "unit": "м3", "volume": idx} for idx in range(1, 11)]

    preview = source_selection_preview_rows(rows, max_rows=3)

    assert [row["source_index"] for row in preview] == ["1", "2", "3", "..."]
    assert preview[-1]["raw_name"] == "еще 7 строк"
    assert preview[-1]["raw_cells"] == ""


def test_source_row_ready_for_bulk_add_accepts_only_ready_work_rows() -> None:
    ready_work = {"item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 12, "status": "готово"}
    material = {"item_kind": "material", "material_code": "MAT-1", "unit": "м3", "volume": 12, "status": "материал / проверить"}
    no_match = {"item_kind": "work", "work_code": "", "unit": "м3", "volume": 12, "status": "готово"}
    no_volume = {"item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": "", "status": "готово"}
    explicit_review = {"item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 12, "status": "проверить"}

    assert source_row_ready_for_bulk_add(ready_work) is True
    assert source_row_ready_for_bulk_add(material) is False
    assert source_row_ready_for_bulk_add(no_match) is False
    assert source_row_ready_for_bulk_add(no_volume) is False
    assert source_row_ready_for_bulk_add(explicit_review) is False


def test_source_scope_summary_shows_current_table_state_usage_and_totals() -> None:
    rows = [
        {"source_index": 5, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"},
        {"source_index": 6, "item_kind": "work", "work_code": "", "unit": "м3", "volume": "2,5", "status": "проверить"},
        {"source_index": 7, "item_kind": "material", "material_code": "MAT-1", "unit": "т", "volume": 1.25, "status": "материал / проверить"},
    ]

    summary = source_scope_summary(
        rows,
        source_usage={5: "final"},
        page_label="Лист 6",
        table_label="Табл 2",
        filter_label="Проверить",
        search_text="песок",
    )

    assert "Текущая выборка: Лист 6 / Табл 2" in summary
    assert "видно 3" in summary
    assert "не в итоге 2" in summary
    assert "уже взято 1" in summary
    assert "работ 2" in summary
    assert "материалов 1" in summary
    assert "сопоставлено 2" in summary
    assert "проверить 2" in summary
    assert "12,5 м3" in summary
    assert "1,25 т" in summary
    assert "фильтр: Проверить" in summary
    assert "поиск: песок" in summary


def test_source_scope_summary_reports_skipped_rows_separately_from_exported_rows() -> None:
    rows = [
        {"source_index": 5, "item_kind": "work", "work_code": "WORK-1", "unit": "м3", "volume": 10, "status": "готово"},
        {"source_index": 6, "item_kind": "work", "work_code": "", "unit": "м3", "volume": "2,5", "status": "проверить"},
    ]

    summary = source_scope_summary(rows, source_usage={5: "final", 6: "skipped"})

    assert "видно 2" in summary
    assert "не в итоге 0" in summary
    assert "уже взято 1" in summary
    assert "пропущено 1" in summary
    assert "проверить" not in summary


def test_selected_target_policy_preserves_mixed_table_row_matches() -> None:
    work_row = {"item_kind": "work", "work_code": "WORK-1", "material_code": "MAT-1"}
    material_row = {"item_kind": "material", "work_code": "", "material_code": "MAT-1"}
    unmatched_row = {"item_kind": "work", "work_code": "", "material_code": ""}

    assert source_row_db_target(work_row) == ("work_type", "WORK-1")
    assert source_row_db_target(material_row) == ("material", "MAT-1")
    assert source_row_db_target(unmatched_row) == ("", "")

    assert selected_target_applies_to_source_row(work_row, target_kind="work_type", target_code="WORK-2", selected_count=1) is True
    assert selected_target_applies_to_source_row(work_row, target_kind="work_type", target_code="WORK-1", selected_count=3) is True
    assert selected_target_applies_to_source_row(work_row, target_kind="work_type", target_code="WORK-2", selected_count=3) is False
    assert selected_target_applies_to_source_row(material_row, target_kind="work_type", target_code="WORK-1", selected_count=3) is False
    assert selected_target_applies_to_source_row(unmatched_row, target_kind="work_type", target_code="WORK-1", selected_count=3) is True


def test_reference_import_alias_matching_and_export(tmp_path: Path) -> None:
    ref_path = tmp_path / "refs.xlsx"
    _sample_reference_xlsx(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")

    counts = load_reference_workbook(ref_path, store)
    assert counts == {"object": 1, "work_type": 1, "material": 1, "alias": 1}
    object_ref = store.find_reference("object", "OBJ-1")
    assert object_ref is not None
    assert object_ref.meta["section_code"] == "UCH_4"
    assert object_ref.meta["section_name"] == "Участок 4"

    rows = build_review_rows(
        [ParsedRdRow(raw_name="Насыпь", unit="м3", volume=12.5)],
        store=store,
        object_code="OBJ-1",
        rd_date="2026-07-23",
        rd_cipher="960-03-5-1213-ПЖ",
        file_name="rd.docx",
    )
    assert rows[0]["work_code"] == "WORK-1"
    assert rows[0]["status"] == "готово"

    out = export_xlsx(rows, tmp_path / "export.xlsx")
    ws = load_workbook(out, data_only=True).active
    headers = [cell.value for cell in ws[1]]
    assert headers[:5] == ["Код объекта", "Наименование работы", "Ед. изм.", "Объем", "Код работы по БД"]
    assert ws[2][0].value == "OBJ-1"
    assert ws[2][headers.index("Объем")].value == "12,5"
    assert ws[2][4].value == "WORK-1"
    assert "Шифр РД" in headers
    assert ws[2][headers.index("Шифр РД")].value == "960-03-5-1213-ПЖ"


def test_final_row_filters_match_export_quality_issues() -> None:
    ready = {"item_kind": "work", "object_code": "OBJ", "work_code": "WORK", "work_name": "Работа", "unit": "м3", "volume": 10, "status": "готово"}
    no_db = {"item_kind": "work", "object_code": "OBJ", "work_code": "", "work_name": "Работа", "unit": "м3", "volume": 10, "status": "проверить"}
    no_volume = {"item_kind": "work", "object_code": "OBJ", "work_code": "WORK", "work_name": "Работа", "unit": "м3", "volume": 0, "status": "готово"}
    material = {"item_kind": "material", "material_code": "MAT", "material_name": "Песок", "unit": "м3", "volume": 2, "status": "материал / привязать"}
    with_materials = {**ready, "linked_materials": [material]}

    assert final_row_matches_filter(ready, "Готово") is True
    assert final_row_matches_filter(no_db, "Проверить") is True
    assert final_row_matches_filter(no_db, "Без БД") is True
    assert final_row_matches_filter(no_volume, "Без объема") is True
    assert final_row_matches_filter(material, "Материалы") is True
    assert final_row_matches_filter(with_materials, "С материалами") is True
    summary = final_rows_filter_summary([ready, no_db, no_volume, material, with_materials], "Проверить")
    assert "итог 5" in summary
    assert "видно 3" in summary
    assert "без БД 1" in summary
    assert "без объема 1" in summary
    assert "материалов 1" in summary
    assert "с материалами 1" in summary


def test_final_rows_quality_flags_unfinished_rows(tmp_path: Path) -> None:
    rows = [
        {
            "source_index": 1,
            "object_code": "OBJ-1",
            "work_code": "WORK-1",
            "work_name": "Устройство насыпи",
            "unit": "м3",
            "volume": 10,
            "status": "готово",
            "item_kind": "work",
        },
        {
            "source_index": 2,
            "object_code": "OBJ-1",
            "work_code": "",
            "work_name": "Не сопоставлено",
            "unit": "м3",
            "volume": 0,
            "status": "проверить",
            "item_kind": "work",
        },
        {
            "source_index": 3,
            "material_code": "MAT-1",
            "material_name": "Песок",
            "unit": "м3",
            "volume": 5,
            "status": "материал / привязать",
            "item_kind": "material",
        },
    ]

    quality = final_rows_quality(rows)

    assert quality["total"] == 3
    assert quality["ready"] == 1
    assert quality["issue_count"] == 2
    assert "не выбрана работа БД" in quality["issues"][0]["issues"]
    assert "объем должен быть больше нуля" in quality["issues"][0]["issues"]
    assert quality["issues"][1]["issues"] == ["материал не привязан к работе"]


def test_review_rows_keep_source_table_cells_and_usage_state(tmp_path: Path) -> None:
    ref_path = tmp_path / "refs.xlsx"
    _sample_reference_xlsx(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [
            ParsedRdRow(
                raw_name="Устройство насыпи",
                unit="м3",
                volume=10,
                page_number=6,
                page_label="Лист 6 · ВОР",
                table_index=2,
                row_number=11,
                bbox={"x0": 10.0, "top": 20.0, "x1": 300.0, "bottom": 90.0},
                row_bbox={"x0": 10.0, "top": 55.0, "x1": 300.0, "bottom": 90.0},
                raw_cells=["1", "Устройство насыпи", "м3", "10"],
            ),
            ParsedRdRow(
                raw_name="Песок средней крупности",
                unit="м3",
                volume=12.2,
                page_number=6,
                table_index=2,
                row_number=12,
                raw_cells=["2", "Песок средней крупности", "м3", "12,2"],
            ),
        ],
        store=store,
        object_code="OBJ-1",
    )

    assert rows[0]["raw_cells"] == ["1", "Устройство насыпи", "м3", "10"]
    assert rows[0]["page_number"] == 6
    assert rows[0]["page_label"] == "Лист 6 · ВОР"
    assert rows[0]["bbox"] == {"x0": 10.0, "top": 20.0, "x1": 300.0, "bottom": 90.0}
    assert rows[0]["row_bbox"] == {"x0": 10.0, "top": 55.0, "x1": 300.0, "bottom": 90.0}
    attach_material_rows(rows[0], [rows[1]])
    assert source_usage_from_final_rows([rows[0]]) == {1: "final", 2: "material"}
    assert rows[0]["linked_materials"][0]["material_ratio"] == 1.22
    assert "Кмат=1,22" in linked_materials_summary(rows[0])


def test_localized_number_formatting_and_merge_accepts_comma() -> None:
    assert parse_localized_number("1 234,5") == 1234.5
    assert format_ru_number(12.5) == "12,5"
    rows = [
        {"source_index": 1, "source_item_no": "1", "raw_name": "Работа 1", "unit": "м3", "raw_unit": "м3", "volume": "10,5", "raw_volume": "10,5", "item_kind": "work"},
        {"source_index": 2, "source_item_no": "2", "raw_name": "Работа 2", "unit": "м3", "raw_unit": "м3", "volume": "1,25", "raw_volume": "1,25", "item_kind": "work"},
    ]
    merged = merge_rows(rows, [0, 1])

    assert format_ru_number(merged["volume"]) == "11,75"


def test_merge_material_rows_keeps_material_as_final_item() -> None:
    rows = [
        {"source_index": 1, "source_item_no": "1", "raw_name": "Песок", "unit": "м3", "raw_unit": "м3", "volume": "10,5", "raw_volume": "10,5", "item_kind": "material", "material_code": "MAT-1", "material_name": "Песок"},
        {"source_index": 2, "source_item_no": "2", "raw_name": "Песок", "unit": "м3", "raw_unit": "м3", "volume": "1,25", "raw_volume": "1,25", "item_kind": "material", "material_code": "MAT-1", "material_name": "Песок"},
    ]

    merged = merge_rows(rows, [0, 1])

    assert merged["item_kind"] == "material"
    assert merged["material_code"] == "MAT-1"
    assert format_ru_number(merged["volume"]) == "11,75"
    assert merged["source_index"] == "1,2"
    assert merged["source_item_no"] == "1,2"
    assert "linked_materials" not in merged


def test_merge_tracks_source_positions_and_exports_work_alias(tmp_path: Path) -> None:
    ref_path = tmp_path / "refs.xlsx"
    _sample_reference_xlsx(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)
    rows = build_review_rows(
        [
            ParsedRdRow(raw_name="Устройство", unit="м3", volume=10),
            ParsedRdRow(raw_name="насыпи", unit="м3", volume=15),
        ],
        store=store,
        object_code="OBJ-1",
        rd_cipher="RD-001",
    )
    merged = merge_rows(rows, [0, 1])
    merged["work_code"] = "WORK-1"
    merged["work_name"] = "Устройство насыпи"
    out = export_xlsx([merged], tmp_path / "merged.xlsx")
    wb = load_workbook(out, data_only=True)
    ws = wb["ВОР"]
    headers = [cell.value for cell in ws[1]]
    assert ws[2][headers.index("Исходные позиции")].value == "1,2"
    assert ws[2][headers.index("Номер позиции ВОР")].value == "1,2"
    alias_ws = wb["_vsm_aliases"]
    alias_rows = list(alias_ws.iter_rows(values_only=True))
    assert ("work_type", "Устройство | насыпи", "WORK-1", "rd_intake_result") in alias_rows


def test_work_actions_with_material_words_stay_work_review(tmp_path: Path) -> None:
    store = LocalStore(tmp_path / "local.sqlite3")
    rows = build_review_rows(
        [
            ParsedRdRow(raw_name="Устройство насыпи из песка средней крупности", unit="м3", volume=10),
            ParsedRdRow(raw_name="Транспортировка цемента ПЦ 500 на площадку", unit="т", volume=2),
            ParsedRdRow(raw_name="Песок средней крупности по ГОСТ 25100", unit="м3", volume=3),
        ],
        store=store,
    )
    assert is_material_like("Песок средней крупности по ГОСТ 25100")
    assert rows[0]["item_kind"] == "work"
    assert rows[0]["status"] == "проверить"
    assert rows[1]["item_kind"] == "work"
    assert rows[2]["item_kind"] == "material"
    assert rows[2]["status"] == "материал / проверить"


def test_merge_split_and_material_review(tmp_path: Path) -> None:
    ref_path = tmp_path / "refs.xlsx"
    _sample_reference_xlsx(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [
            ParsedRdRow(raw_name="Устройство насыпи", unit="м3", volume=10),
            ParsedRdRow(raw_name="Устройство насыпи", unit="м3", volume=15),
            ParsedRdRow(raw_name="Песок средней крупности", unit="м3", volume=3),
        ],
        store=store,
        object_code="OBJ-1",
    )
    merged = merge_rows(rows, [0, 1])
    assert merged["volume"] == 25
    split = split_row(merged, [10, 15])
    assert [row["volume"] for row in split] == [10, 15]
    assert rows[2]["status"] == "материал / проверить"


def test_parse_split_volumes_keeps_decimal_comma() -> None:
    assert parse_split_volumes("10,5") == [10.5]
    assert parse_split_volumes("10,5; 15,25") == [10.5, 15.25]
    assert parse_split_volumes("10, 15") == [10.0, 15.0]




def test_reference_import_filters_temporary_culvert_work_types(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "work_types"
    ws.append(["code", "name", "default_unit"])
    ws.append(["PIPE_ARRANGEMENT_264477", "Труба Вр. пр. а/д ПК2644+77", "%"])
    ws.append(["WORK-1", "Устройство насыпи", "м3"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")

    counts = load_reference_workbook(ref_path, store)
    assert is_temporary_culvert_work("PIPE_ARRANGEMENT_264477", "Труба Вр. пр. а/д ПК2644+77")
    assert counts["work_type"] == 1
    assert store.find_reference("work_type", "PIPE_ARRANGEMENT_264477") is None
    assert store.find_reference("work_type", "WORK-1") is not None


def test_short_inner_work_name_does_not_auto_match_long_rd_position(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "work_types"
    ws.append(["code", "name", "default_unit"])
    ws.append(["CONSOLIDATION", "Уплотнение", "м2"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [
            ParsedRdRow(
                raw_name="Устройство второго защитного слоя из ПГС. Уплотнение ПГС катком массой 25 т",
                unit="м3",
                volume=57617,
            )
        ],
        store=store,
    )
    assert rows[0]["work_code"] == ""
    assert rows[0]["status"] == "проверить"


def test_linked_materials_export_in_audit_column(tmp_path: Path) -> None:
    ref_path = tmp_path / "refs.xlsx"
    _sample_reference_xlsx(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)
    rows = build_review_rows(
        [
            ParsedRdRow(raw_name="Устройство насыпи", unit="м3", volume=10),
            ParsedRdRow(raw_name="Песок средней крупности", unit="м3", volume=12.2),
        ],
        store=store,
        object_code="OBJ-1",
    )
    attach_material_rows(rows[0], [rows[1]])
    out = export_xlsx([rows[0]], tmp_path / "materials.xlsx")
    ws = load_workbook(out, data_only=True).active
    headers = [cell.value for cell in ws[1]]
    materials = ws[2][headers.index("Материалы")].value
    assert "MAT-1" in materials
    assert "Песок средней крупности" in materials
    assert "12,2" in materials


def test_rule_based_work_matching_for_real_vor_phrases(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "objects"
    ws.append(["code", "name"])
    ws.append(["OBJ-1", "Объект 1"])
    ws = wb.create_sheet("work_types")
    ws.append(["code", "name", "default_unit"])
    ws.append(["ZS2", "ЗС2", "м3"])
    ws.append(["ASPHALT_PAVEMENT", "Устройство дорожного покрытия из а/б", "м2"])
    ws.append(["WT_AF0D31D9", "Замена слабого грунта грунтом", "м3"])
    ws.append(["CONSOLIDATION", "Уплотнение", "м2"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [
            ParsedRdRow(
                raw_name="Устройство второго защитного слоя из природной песчано-гравийной смеси. Уплотнение ПГС катком",
                unit="м3",
                volume=57617,
            ),
            ParsedRdRow(raw_name="Устройство асфальтобетонного покрытия толщиной 12 см", unit="м2", volume=21067),
            ParsedRdRow(
                raw_name="Выборка недостаточно прочного грунта в основании насыпи, погрузка на а/самосвалы",
                unit="м3",
                volume=28447,
            ),
        ],
        store=store,
        object_code="OBJ-1",
    )

    assert rows[0]["work_code"] == "ZS2"
    assert rows[0]["work_code"] != "CONSOLIDATION"
    assert rows[1]["work_code"] == "ASPHALT_PAVEMENT"
    assert rows[2]["work_code"] == "WT_AF0D31D9"


def test_transport_of_crushed_stone_keeps_transport_code_before_slope_rules(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "work_types"
    ws.append(["code", "name", "default_unit"])
    ws.append(["WT_D746ECC7", "Транспортировка готовой смеси", "м3"])
    ws.append(["SLOPE_FORMATION_CC", "Укрепление откосов щебнем", "м3"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [
            ParsedRdRow(
                raw_name="Дальность возки щебня фр. 10-20 мм для устройства укрепления дна канавы",
                unit="т",
                volume=102.816,
            ),
            ParsedRdRow(
                raw_name="Укрепление дна канавы с щебневанием (фр. 10-20мм)",
                unit="м3",
                volume=60,
            ),
        ],
        store=store,
    )

    assert rows[0]["work_code"] == "WT_D746ECC7"
    assert rows[0]["work_code"] != "SLOPE_FORMATION_CC"
    assert rows[1]["work_code"] == "SLOPE_FORMATION_CC"


def test_material_phrase_with_protective_layer_is_not_rule_matched_as_work(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "work_types"
    ws.append(["code", "name", "default_unit"])
    ws.append(["ZS2", "ЗС2", "м3"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [ParsedRdRow(raw_name="Песчано-гравийная смесь по ГОСТ 25100 для второго защитного слоя", unit="м3", volume=70292.74)],
        store=store,
    )

    assert rows[0]["item_kind"] == "material"
    assert rows[0]["work_code"] == ""
    assert rows[0]["status"] == "материал / проверить"


def test_unit_mismatch_keeps_rule_match_on_manual_review(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "objects"
    ws.append(["code", "name"])
    ws.append(["OBJ-1", "Объект 1"])
    ws = wb.create_sheet("work_types")
    ws.append(["code", "name", "default_unit"])
    ws.append(["WT_D746ECC7", "Перемещение инертного материала", "м3"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [ParsedRdRow(raw_name="Дальность возки песка внутри региона", unit="т", volume=100)],
        store=store,
        object_code="OBJ-1",
    )

    assert rows[0]["work_code"] == "WT_D746ECC7"
    assert rows[0]["status"] == "проверить"
    assert "ед.изм исходная т, БД м3" in rows[0]["comment"]


def test_merge_links_material_rows_without_summing_material_volume(tmp_path: Path) -> None:
    ref_path = tmp_path / "refs.xlsx"
    _sample_reference_xlsx(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)
    rows = build_review_rows(
        [
            ParsedRdRow(raw_name="Насыпь", unit="м3", volume=10),
            ParsedRdRow(raw_name="Песок средней крупности", unit="м3", volume=12.2),
        ],
        store=store,
        object_code="OBJ-1",
    )

    merged = merge_rows(rows, [0, 1])

    assert merged["volume"] == 10
    assert merged["source_index"] == "1,2"
    assert len(merged.get("linked_materials") or []) == 1
    assert merged["linked_materials"][0]["material_code"] == "MAT-1"


def test_single_word_alias_is_exact_only_and_does_not_steal_ground_loading(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "work_types"
    ws.append(["code", "name", "default_unit"])
    ws.append(["WT_73AD9101", "Погрузка ЩПГС в самосвалы", "м3"])
    ws.append(["WT_70E50F83", "Погрузка местного грунта в самосвалы", "м3"])
    ws = wb.create_sheet("aliases")
    ws.append(["kind", "alias_text", "canonical_code"])
    ws.append(["work_type", "Погрузка", "WT_73AD9101"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [ParsedRdRow(raw_name="Погрузка растительного грунта на временной площадке", unit="м3", volume=10)],
        store=store,
    )

    assert rows[0]["work_code"] == "WT_70E50F83"


def test_geotextile_alias_matches_and_anchor_trench_exact_alias_overrides_manual_guard(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "work_types"
    ws.append(["code", "name", "default_unit"])
    ws.append(["ZS1", "ЗС1", "м3"])
    ws.append(["EARTH_EXCAVATION", "Разработка грунта", "м3"])
    ws.append(["ANCHOR_TRENCH", "Анкерная траншея для геоматов", "м3"])
    ws.append(["GEOTEXTILE_LAYER", "Устройство прослойки из геотекстиля 300 кН/м", "м2"])
    ws = wb.create_sheet("aliases")
    ws.append(["kind", "alias_text", "canonical_code"])
    ws.append(["work_type", "Укладка геотекстиля", "GEOTEXTILE_LAYER"])
    ws.append(["work_type", "Разработка грунта", "EARTH_EXCAVATION"])
    ws.append([
        "work_type",
        "Разработка анкерной траншеи для геоматов ручным способом в грунте первого защитного слоя",
        "ANCHOR_TRENCH",
    ])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [
            ParsedRdRow(raw_name="Укладка нетканного геотекстиля по основанию 2 защитного слоя", unit="м2", volume=8596),
            ParsedRdRow(
                raw_name="Разработка анкерной траншеи для геоматов ручным способом в грунте первого защитного слоя",
                unit="м3",
                volume=1015,
            ),
            ParsedRdRow(
                raw_name="Разработка анкерной траншеи для геоматов в I группе грунта, с обратной засыпкой",
                unit="м3",
                volume=273,
            ),
            ParsedRdRow(
                raw_name="Разработка анкерной траншеи для геоматов ручным способом в грунте первого защитного слоя",
                unit="м3",
                volume=1015,
            ),
        ],
        store=store,
    )

    assert rows[0]["work_code"] == "GEOTEXTILE_LAYER"
    assert rows[1]["work_code"] == "ANCHOR_TRENCH"
    assert rows[2]["work_code"] == ""
    assert rows[3]["work_code"] == "ANCHOR_TRENCH"


def test_system_material_alias_matches_long_shpgs_phrase(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "materials"
    ws.append(["code", "name", "default_unit"])
    ws.append(["SHPGS", "ЩПГС", "м3"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    counts = load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [ParsedRdRow(raw_name="Смесь щебеночно-песчаная типа 0/31,5, категории К90", unit="м3", volume=100)],
        store=store,
    )

    assert counts["alias"] >= 1
    assert rows[0]["item_kind"] == "material"
    assert rows[0]["material_code"] == "SHPGS"


def test_work_rows_keep_work_kind_but_show_material_hint_for_geomat_and_concrete(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "work_types"
    ws.append(["code", "name", "default_unit"])
    ws.append(["WT_D746ECC7", "Транспортировка готовой смеси", "м3"])
    ws = wb.create_sheet("materials")
    ws.append(["code", "name", "default_unit"])
    ws.append(["GEOMAT", "Геомат противоэрозионный", "м2"])
    ws.append(["CONCRETE_M150", "Бетон М150", "м3"])
    ws = wb.create_sheet("aliases")
    ws.append(["kind", "alias_text", "canonical_code"])
    ws.append(["material", "Противоэрозионного геомата", "GEOMAT"])
    ws.append(["material", "Для геоматов", "GEOMAT"])
    ws.append(["material", "Бетона М150", "CONCRETE_M150"])
    ws.append(["material", "Бетонирование М150", "CONCRETE_M150"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [
            ParsedRdRow(
                raw_name="Укладка противоэрозионного геомата по откосу, плотностью не менее 550г/м2",
                unit="м2",
                volume=27333.72,
            ),
            ParsedRdRow(
                raw_name="Разработка анкерной траншеи для геоматов ручным способом в насыпном грунте",
                unit="м3",
                volume=1691,
            ),
            ParsedRdRow(raw_name="Бетонирование откосов и дна канав, М150, h-0,08м", unit="м3", volume=336),
            ParsedRdRow(
                raw_name="Дальность возки бетона М150 для устройства укрепления откосов и дна канавы",
                unit="т",
                volume=806.4,
            ),
        ],
        store=store,
    )

    assert [row["item_kind"] for row in rows] == ["work", "work", "work", "work"]
    assert rows[0]["work_code"] == ""
    assert rows[0]["material_code"] == "GEOMAT"
    assert rows[1]["work_code"] == ""
    assert rows[1]["material_code"] == "GEOMAT"
    assert rows[2]["work_code"] == ""
    assert rows[2]["material_code"] == "CONCRETE_M150"
    assert rows[3]["work_code"] == "WT_D746ECC7"
    assert rows[3]["material_code"] == "CONCRETE_M150"
    assert all("material:" in row["comment"] for row in rows)


def test_seeded_rd_work_aliases_match_remaining_real_vor_phrases(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "work_types"
    ws.append(["code", "name", "default_unit"])
    ws.append(["GEOMAT_PLACEMENT", "Укладка противоэрозионного геомата", "м2"])
    ws.append(["GEOMAT_ANCHOR_TRENCH", "Разработка анкерной траншеи для геоматов", "м3"])
    ws.append(["DITCH_CONCRETE_M150", "Бетонирование откосов и дна канав М150", "м3"])
    ws = wb.create_sheet("materials")
    ws.append(["code", "name", "default_unit"])
    ws.append(["GEOMAT", "Геомат противоэрозионный", "м2"])
    ws.append(["CONCRETE_M150", "Бетон М150", "м3"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [
            ParsedRdRow(
                raw_name="Укладка противоэрозионного геомата по откосу, плотностью не менее 550г/ м2",
                unit="м2",
                volume=27333.72,
            ),
            ParsedRdRow(
                raw_name="Разработка анкерной траншеи для геоматов ручным способом в насыпном грунте, с обратной засыпкой",
                unit="м3",
                volume=1691,
            ),
            ParsedRdRow(
                raw_name="Разработка анкерной траншеи для геоматов в I группе грунта, с обратной засыпкой",
                unit="м3",
                volume=273,
            ),
            ParsedRdRow(raw_name="Бетонирование откосов и дна канав, М150, h-0,08м", unit="м3", volume=336),
        ],
        store=store,
    )

    assert [row["work_code"] for row in rows] == [
        "GEOMAT_PLACEMENT",
        "GEOMAT_ANCHOR_TRENCH",
        "GEOMAT_ANCHOR_TRENCH",
        "DITCH_CONCRETE_M150",
    ]
    assert rows[0]["material_code"] == "GEOMAT"
    assert rows[1]["material_code"] == "GEOMAT"
    assert rows[2]["material_code"] == "GEOMAT"
    assert rows[3]["material_code"] == "CONCRETE_M150"


def test_seeded_rd_aliases_cover_pdf_inspector_vor_phrases(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "work_types"
    ws.append(["code", "name", "default_unit"])
    ws.append(["EMBANKMENT_CONSTRUCTION", "Устройство насыпи", "м3"])
    ws.append(["EARTH_EXCAVATION", "Разработка выемки", "м3"])
    ws.append(["AREA_GRADING", "Планировка", "м2"])
    ws.append(["ASPHALT_PAVEMENT", "Устройство дорожного покрытия из а/б", "м2"])
    ws.append(["CRUSHED_STONE_PLACEMENT", "Устройство ЩПС", "м3"])
    ws.append(["TOPSOIL_STRIPPING", "Снятие ПРС", "м3"])
    ws.append(["DITCH_CONSTRUCTION", "Устройство кюветов", "м3"])
    ws.append(["WT_B505C78", "Транспортировка грунта", "т"])
    ws.append(["WT_FA506691", "Погрузка песка", "т"])
    ws = wb.create_sheet("materials")
    ws.append(["code", "name", "default_unit"])
    ws.append(["SAND", "Песок", "м3"])
    ws.append(["SOIL", "Грунт", "м3"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [
            ParsedRdRow(raw_name="Устройство насыпи из песка средней крупности по ГОСТ 25100 из карьера", unit="м3", volume=10),
            ParsedRdRow(raw_name="Дальность возки грунта внутри региона на расстояние 24 км", unit="т", volume=11),
            ParsedRdRow(raw_name="Погрузка песка экскаватором с объемом ковша 1м3", unit="т", volume=12),
            ParsedRdRow(raw_name="Устройство участков переменной жесткости материала ЩПГС для мостов", unit="м3", volume=13),
            ParsedRdRow(raw_name="Устройство асфальтобетонного покрытия толщиной 12 см асфальтоукладчиком", unit="м2", volume=14),
            ParsedRdRow(raw_name="Планировка откосов насыпи грунта I группы", unit="м2", volume=15),
            ParsedRdRow(raw_name="Снятие почвенно-растительного слоя толщиной 30 см", unit="м3", volume=16),
            ParsedRdRow(raw_name="Разработка водоотводных канав", unit="м3", volume=17),
            ParsedRdRow(raw_name="Песок средней крупности по ГОСТ 25100 для насыпи", unit="м3", volume=18),
            ParsedRdRow(raw_name="Грунт ИГЭ 14д I группы грунта", unit="м3", volume=19),
        ],
        store=store,
    )

    assert [row["work_code"] for row in rows[:8]] == [
        "EMBANKMENT_CONSTRUCTION",
        "WT_B505C78",
        "WT_FA506691",
        "CRUSHED_STONE_PLACEMENT",
        "ASPHALT_PAVEMENT",
        "AREA_GRADING",
        "TOPSOIL_STRIPPING",
        "DITCH_CONSTRUCTION",
    ]
    assert rows[8]["item_kind"] == "material"
    assert rows[8]["material_code"] == "SAND"
    assert rows[9]["item_kind"] == "material"
    assert rows[9]["material_code"] == "SOIL"


def test_seeded_rd_material_aliases_match_real_vor_material_phrases(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "materials"
    ws.append(["code", "name", "default_unit"])
    ws.append(["PGS", "Песчано-гравийная смесь", "м3"])
    ws.append(["ASPHALT_MIX", "Асфальтобетонная смесь", "м3"])
    ws.append(["CEMENT", "Цемент", "т"])
    ws.append(["CRUSHED_STONE", "Щебень", "м3"])
    ref_path = tmp_path / "refs.xlsx"
    wb.save(ref_path)
    store = LocalStore(tmp_path / "local.sqlite3")
    load_reference_workbook(ref_path, store)

    rows = build_review_rows(
        [
            ParsedRdRow(raw_name="Песчано-гравийная смесь по ГОСТ 25100 для второго защитного слоя", unit="м3", volume=1),
            ParsedRdRow(raw_name="В том числе: Асфальтобетонная смесь А32Нт по ГОСТ Р 58406.2", unit="м3", volume=2),
            ParsedRdRow(raw_name="Цемент ПЦ 500", unit="т", volume=3),
            ParsedRdRow(raw_name="Щебень фр. 10-20 мм М600 F200", unit="м3", volume=4),
        ],
        store=store,
    )

    assert [row["material_code"] for row in rows] == ["PGS", "ASPHALT_MIX", "CEMENT", "CRUSHED_STONE"]
    assert all(row["item_kind"] == "material" for row in rows)
    assert all(row["status"] == "материал / проверить" for row in rows)


def test_local_entity_drafts_match_export_and_reimport(tmp_path: Path) -> None:
    store = LocalStore(tmp_path / "local.sqlite3")
    store.upsert_entity_draft(
        ReferenceEntry("object", "RDO_001", "Новый объект РД", meta={"type": "ROAD", "pk_raw_text": "ПК 1 - ПК 2"}),
        source_text="960-TEST",
        source_file="rd.pdf",
    )
    store.upsert_entity_draft(
        ReferenceEntry(
            "work_type",
            "RDW_001",
            "Устройство опытной насыпи",
            unit="м3",
            meta={"analytics_tag": "earthwork", "analytics_group": "Земляное полотно", "productivity_enabled": True},
        ),
        source_text="Опытная насыпь из местного грунта",
        source_file="rd.pdf",
    )
    store.upsert_entity_draft(
        ReferenceEntry("material", "RDM_001", "Опытный грунт", unit="м3"),
        source_text="Опытный грунт по ведомости",
        source_file="rd.pdf",
    )

    rows = build_review_rows(
        [ParsedRdRow(raw_name="Опытная насыпь из местного грунта", unit="м3", volume=42)],
        store=store,
        object_code="RDO_001",
    )

    assert rows[0]["work_code"] == "RDW_001"
    assert rows[0]["object_code"] == "RDO_001"
    assert rows[0]["status"] == "готово"

    out = export_xlsx(rows, tmp_path / "rd_with_drafts.xlsx", store=store)
    wb = load_workbook(out, data_only=True)
    assert wb["_vsm_new_work_types"].sheet_state == "visible"
    assert wb["_vsm_new_materials"].sheet_state == "visible"
    assert wb["_vsm_new_objects"].sheet_state == "visible"
    work_rows = list(wb["_vsm_new_work_types"].iter_rows(values_only=True))
    material_rows = list(wb["_vsm_new_materials"].iter_rows(values_only=True))
    object_rows = list(wb["_vsm_new_objects"].iter_rows(values_only=True))
    alias_rows = list(wb["_vsm_aliases"].iter_rows(values_only=True))
    assert any(row[0] == "RDW_001" and row[1] == "Устройство опытной насыпи" for row in work_rows[1:])
    assert any(row[0] == "RDM_001" and row[1] == "Опытный грунт" for row in material_rows[1:])
    assert any(row[0] == "RDO_001" and row[1] == "Новый объект РД" for row in object_rows[1:])
    assert ("work_type", "Опытная насыпь из местного грунта", "RDW_001", "rd_intake_result") in alias_rows

    imported = LocalStore(tmp_path / "imported.sqlite3")
    counts = load_reference_workbook(out, imported)
    assert counts["object"] == 1
    assert counts["work_type"] == 1
    assert counts["material"] == 1
    assert imported.find_reference("work_type", "RDW_001").name == "Устройство опытной насыпи"
    assert imported.find_reference("material", "RDM_001").name == "Опытный грунт"
    assert imported.find_reference("object", "RDO_001").name == "Новый объект РД"


def test_suggest_entity_code_avoids_existing_local_reference(tmp_path: Path) -> None:
    store = LocalStore(tmp_path / "local.sqlite3")
    code = suggest_entity_code("work_type", "Устройство новой насыпи", store)
    store.upsert_entity_draft(ReferenceEntry("work_type", code, "Устройство новой насыпи", unit="м3"))

    assert suggest_entity_code("work_type", "Устройство новой насыпи", store) == f"{code}_2"


def test_gui_reference_search_controls_are_present() -> None:
    source = (ROOT / "tools" / "rd_intake_windows" / "app.py").read_text(encoding="utf-8")

    ui_block = source.split("    def _build_ui(self) -> None:", 1)[1].split("    def _responsive_label", 1)[0]
    menu_block = source.split("    def build_app_menu(self) -> None:", 1)[1].split("    def focus_source_search", 1)[0]

    for token in ["PDF_ENGINE_PDF_INSPECTOR", "PDF_ENGINE_STABLE", "ParsedRdRow", "RdParserError", "parse_vor_document"]:
        assert token in source
    assert "self.main_pane = root" in ui_block
    assert "self.tables_pane = tables" in ui_block
    assert "root = tk.PanedWindow" in ui_block
    assert "orient=tk.HORIZONTAL" in ui_block
    assert "tables = tk.PanedWindow" in ui_block
    assert "orient=tk.VERTICAL" in ui_block
    assert "yscrollcommand=y_scroll.set" in ui_block
    assert "xscrollcommand=x_scroll.set" in ui_block
    assert "self.source_tree.container.pack(fill=tk.BOTH, expand=True" in ui_block
    assert "self.final_tree.container.pack(fill=tk.BOTH, expand=True" in ui_block
    assert "Поиск объекта" in ui_block
    assert "object_search_var" in source
    assert "references_matching_query" in source
    assert "edit_ref_search_var" in source
    assert 'edit_ref_search.bind("<KeyRelease>"' in source
    assert "layout_settings.json" in source
    assert "restore_layout_panes" in source
    assert "Сбросить раскладку окна" in menu_block
    assert "Справочник БД" in source
    assert "show_reference_browser" in source
    assert "estimated_preview_row_bbox" in source
    assert "preview_use_bbox_fallback_var" not in source
    assert "pdf_engine_var" in source
    assert "toggle_preview_bbox_fallback" not in source
    assert "Fallback рамки" not in source
    assert "обведено" in source
    assert "section_code" in source
    assert "Пикетаж" in source
    assert "Файл" in menu_block
    assert "Действия" in menu_block
    assert "Справка" in menu_block
    assert "Обзор таблиц" not in menu_block
    assert "План проверки" not in ui_block
    assert "Легенда" not in ui_block
    assert "Текущая выборка перед добавлением" not in ui_block


def test_app_embeds_source_table_navigation_checklist() -> None:
    source = (ROOT / "tools" / "rd_intake_windows" / "app.py").read_text(encoding="utf-8")
    readme = (ROOT / "tools" / "rd_intake_windows" / "README.md").read_text(encoding="utf-8")
    portable_builder = (ROOT / "tools" / "rd_intake_windows" / "build_portable_package.py").read_text(encoding="utf-8")

    ui_block = source.split("    def _build_ui(self) -> None:", 1)[1].split("    def _responsive_label", 1)[0]

    for label in ["Открыть файл", "Экспорт итогового XLSX", "Пред. лист", "След. лист", "Лист", "Масштаб"]:
        assert label in ui_block
    for label in ["Объект по БД", "Добавить объект", "Шифр РД", "Дата РД"]:
        assert label in ui_block
    assert "go_to_pdf_page_from_entry" in source
    assert "preview_page_entry_var" in source
    assert "parse_vor_document(path, pdf_engine=self.current_pdf_engine())" in source
    assert "def current_pdf_engine" in source
    assert "source_page_combo" in ui_block
    assert "source_table_combo" in ui_block
    assert "Распознанные позиции текущего листа" in ui_block
    assert "Итог на экспорт" in ui_block
    assert "source_table_nav_frame" not in ui_block
    assert "source_status_legend_frame" not in ui_block
    assert "selection_preview_frame" not in ui_block
    assert "Ведомости сейчас не обрабатываются" in readme
    assert "Ведомости сейчас не обрабатываются" in portable_builder
    assert "План проверки" not in readme
    assert "Обзор таблиц" not in readme
    assert "Легенда" not in readme


def test_app_exposes_selection_basket_panel_for_table_rows() -> None:
    source = (ROOT / "tools" / "rd_intake_windows" / "app.py").read_text(encoding="utf-8")
    core = (ROOT / "tools" / "rd_intake_windows" / "rd_intake_core.py").read_text(encoding="utf-8")
    readme = (ROOT / "tools" / "rd_intake_windows" / "README.md").read_text(encoding="utf-8")
    portable_builder = (ROOT / "tools" / "rd_intake_windows" / "build_portable_package.py").read_text(encoding="utf-8")

    ui_block = source.split("    def _build_ui(self) -> None:", 1)[1].split("    def _responsive_label", 1)[0]
    source_buttons_block = source.split("source_action_names = [", 1)[1].split("self.source_tree = self._tree", 1)[0]
    create_work_block = source.split('if kind == "work_type":', 1)[1].split("else:", 1)[0]

    assert "source_item_no" in source
    assert "Номер позиции ВОР" in core
    for name in ["add_selected", "merge", "split_source", "attach_material", "manual_source", "create_work", "create_material"]:
        assert f'"{name}"' in source_buttons_block
    for label in ["Извлечь в итог", "Совместить в итог", "Разделить строку", "Связать материал с работой", "Ручной ввод ВОР", "Добавить работу БД", "Добавить материал БД"]:
        assert label in source_buttons_block
    assert "Применить БД" not in source_buttons_block
    assert "on_target_combo_selected" in source
    assert "apply_first_target_match" in source
    assert "preview_page_slider_widget" in source
    assert "preview_zoom_entry_var" in source
    assert "def edit_selected_source" in source
    assert "ref_search_var" in source
    assert 'target_combo.bind("<<ComboboxSelected>>", on_ref_selected)' in source
    assert "def choose_source_work_for_materials" in source
    assert "linked_work_source_index" in source
    assert "Совмещайте отдельно работы или отдельно материалы" in source
    assert "Для объединенного материала выберите материал по БД" in source
    assert '("Связать материал", self.attach_selected_materials' not in source
    assert "def add_manual_source_row" in source
    assert "def split_selected_source_row" in source
    assert "ttk.Entry(win, textvariable=status_var)" in source
    assert "unit_var" in source
    assert "volume_var" in source
    assert "analytics_tag" in create_work_block
    assert "Группа" not in create_work_block
    assert "def create_object" in source
    assert "object_hidden_from_rd" in source
    assert "STOCKPILE" not in source
    assert "ПК" in source
    assert "source_item_numbers" in core
    assert "_vsm_new_work_types" in core
    assert "_vsm_new_materials" in core
    assert "_vsm_new_objects" in core
    assert "_vsm_aliases" in core
    assert "support-листы" in readme
    assert "support-листы" in portable_builder
    assert "Fallback рамки" not in readme
    assert "PDF Inspector" in readme
    assert "Режим границ" in readme
    assert "Fallback рамки" not in portable_builder
    assert "PDF Inspector" in portable_builder
    assert "Кмат = Vматериала/Vработы" in readme
    assert "Кмат = Vматериала/Vработы" in portable_builder
    assert "Текущая выборка перед добавлением" not in ui_block
