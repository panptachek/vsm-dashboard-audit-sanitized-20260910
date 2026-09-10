"""Safe DB admin endpoints for manual dashboard corrections.

The API intentionally exposes only a small whitelist of project tables and
never accepts raw SQL from the browser.
"""
from __future__ import annotations

from contextlib import contextmanager
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import re
import urllib.parse
from typing import Any, Literal
from uuid import UUID
from zipfile import ZIP_DEFLATED, ZipFile

import psycopg2
import psycopg2.extras
from psycopg2 import sql
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from main import get_conn
from auth import current_user, require_admin
from report_text_parser import parse_pk_value
from temp_roads_catalog import ensure_temp_road_catalog_rows


router = APIRouter(prefix="/api/db-admin", tags=["db-admin"])

INTERNAL_POVSR_WORK_TYPE_SQL = "code NOT LIKE 'POVSR\\_%%' ESCAPE '\\'"


def _attachment_headers(filename_ascii: str, filename_utf8: str) -> dict[str, str]:
    return {
        "Content-Disposition": (
            f"attachment; filename=\"{filename_ascii}\"; "
            f"filename*=UTF-8''{urllib.parse.quote(filename_utf8)}"
        )
    }

EDITABLE_TABLES: dict[str, str] = {
    "construction_sections": "Участки",
    "construction_section_versions": "Границы участков",
    "objects": "Объекты",
    "object_segments": "Сегменты объектов",
    "constructives": "Конструктивы",
    "work_types": "Виды работ",
    "work_type_aliases": "Алиасы работ",
    "materials": "Материалы",
    "contractors": "Подрядчики",
    "daily_reports": "Суточные отчеты",
    "daily_work_items": "Работы за сутки",
    "daily_work_item_segments": "Сегменты работ",
    "material_movements": "Перевозки материалов",
    "equipment_units": "Справочник техники",
    "equipment_unit_identifiers": "Номера техники",
    "report_equipment_units": "Техника отчета",
    "work_item_equipment_usage": "Техника по работам",
    "temporary_roads": "Временные дороги",
    "temporary_road_status_segments": "Статусы временных дорог",
    "mainline_rd_coverage": "РД по ОХ",
    "mainline_work_schedule_ranges": "Окна работ ОХ",
    "pile_delivery_specs": "Спецификации поставки свай",
    "pile_delivery_calendar": "Календарь поставки свай",
    "stockpiles": "Накопители",
    "stockpile_balance_snapshots": "Остатки накопителей",
    "object_types": "Типы объектов",
}

READONLY_TABLES: dict[str, str] = {
    "constructive_work_types": "Связи конструктивы-работы",
    "daily_report_parse_candidates": "Кандидаты парсинга отчетов",
    "daily_report_problems": "Проблемы отчетов",
    "equipment_norms_config": "Настройки норм техники",
    "equipment_productivity_norms": "Нормы производительности техники",
    "material_movement_equipment_usage": "Техника по перевозкам",
    "pile_fields": "Свайные поля",
    "pile_plan_periods": "Проект свай по периодам",
    "project_work_items": "Проектные объемы работ",
    "project_work_item_segments": "Попикетные проектные объемы",
    "route_pickets": "Пикеты основного хода",
    "temp_road_points": "Точки ВАД",
    "temporary_road_import_drafts": "Черновики импорта ВАД",
    "temporary_road_import_runs": "Запуски импорта ВАД",
    "temporary_road_pk_mappings": "Мэппинг пикетов ВАД",
    "temporary_road_state_corrections": "Коррекции статусов ВАД",
}

EXPOSED_TABLES: dict[str, str] = {**EDITABLE_TABLES, **READONLY_TABLES}

RD_INTAKE_SYSTEM_ALIASES: tuple[tuple[str, str, tuple[str, ...]], ...] = (
    (
        "work_type",
        "ZS1",
        (
            "Устройство первого защитного слоя",
            "Устройство первого верхнего защитного слоя",
            "Устройство верхнего защитного слоя",
            "Первый защитный слой",
        ),
    ),
    (
        "work_type",
        "ZS2",
        (
            "Устройство второго защитного слоя",
            "Устройство второго нижнего защитного слоя",
            "Устройство нижнего защитного слоя",
            "Второй защитный слой",
        ),
    ),
    (
        "work_type",
        "WT_D746ECC7",
        ("Погрузка готовой смеси", "Транспортировка готовой смеси", "Дальность возки готовой смеси"),
    ),
    ("work_type", "PAVEMENT_SANDING", ("Досыпка песчаного дренирующего грунта",)),
    (
        "work_type",
        "GEOMAT_PLACEMENT",
        ("Укладка противоэрозионного геомата", "Укладка противоэрозионного геомата по откосу"),
    ),
    (
        "work_type",
        "GEOMAT_ANCHOR_TRENCH",
        ("Анкерная траншея для геоматов", "Разработка анкерной траншеи для геоматов"),
    ),
    (
        "work_type",
        "DITCH_CONCRETE_M150",
        ("Бетонирование откосов и дна канав, М150", "Бетонирование откосов и дна канав М150"),
    ),
    (
        "work_type",
        "EMBANKMENT_CONSTRUCTION",
        (
            "Устройство насыпи из песка",
            "Устройство насыпи из песчаного грунта",
            "Устройство насыпи из песка средней крупности",
            "Устройство насыпи из песка крупного",
            "Устройство дорожных насыпей из песчаного грунта",
            "Досыпка естественной бермы дренирующим грунтом",
        ),
    ),
    (
        "work_type",
        "WT_2667EA39",
        (
            "Устройство насыпи из местного грунта",
            "Устройство досыпки основания насыпи из местного недренирующего грунта",
            "Досыпка бермы из местного недренирующего грунта",
        ),
    ),
    (
        "work_type",
        "EARTH_EXCAVATION",
        (
            "Разработка грунта",
            "Разработка выемки",
            "Выемка грунта",
            "Устройство выемки",
            "Разработка котлована",
        ),
    ),
    (
        "work_type",
        "TOPSOIL_STRIPPING",
        (
            "Снятие растительного грунта",
            "Срезка растительного грунта",
            "Снятие почвенно-растительного слоя",
            "Снятие ПРС",
            "Погрузка растительного грунта",
            "Надвижка растительного грунта",
        ),
    ),
    (
        "work_type",
        "PEAT_REMOVAL",
        ("Выторфовка", "Выемка торфа", "Выемка торфяного грунта"),
    ),
    (
        "work_type",
        "WEAK_SOIL_REPLACEMENT",
        ("Замена слабого грунта", "Замещение слабого грунта", "Замена слабого грунта песком"),
    ),
    (
        "work_type",
        "AREA_GRADING",
        (
            "Планировка поверхности грунта",
            "Планировка верха первого защитного слоя",
            "Планировка верха второго защитного слоя",
            "Планировка верха тела насыпи",
            "Планировка откосов насыпи",
            "Планировка откосов выемки",
            "Планировка полки бермы",
            "Планировка закюветной полки",
        ),
    ),
    (
        "work_type",
        "SLOPE_FORMATION_GG",
        (
            "Укрепление откоса насыпи посевом трав",
            "Укрепление откосов посевом многолетних трав",
            "Укрепление откоса выемки посевом трав",
        ),
    ),
    (
        "work_type",
        "SLOPE_FORMATION_CC",
        (
            "Укрепление откосов щебнем",
            "Укрепление дна кюветов щебнем",
            "Укрепление обочины щебнем",
            "Укрепление откоса выемки щебнем",
        ),
    ),
    (
        "work_type",
        "DITCH_CONSTRUCTION",
        ("Устройство кюветов", "Разработка водоотводных канав", "Устройство водоотводных канав"),
    ),
    (
        "work_type",
        "ASPHALT_PAVEMENT",
        (
            "Устройство асфальтобетонного покрытия",
            "Устройство дорожного покрытия из асфальтобетона",
            "Устройство покрытия толщиной 12 см асфальтоукладчиком",
        ),
    ),
    (
        "work_type",
        "CRUSHED_STONE_PLACEMENT",
        (
            "Устройство участков переменной жесткости материала ЩПГС",
            "Устройство участков переменной жесткости материала ЩПС",
            "Устройство слоя из ЩПГС",
            "Устройство дорожного покрытия из ЩПС",
        ),
    ),
    (
        "work_type",
        "WT_FA506691",
        ("Погрузка песка", "Погрузка песчаного грунта", "Погрузка песка экскаватором"),
    ),
    (
        "work_type",
        "WT_70E50F83",
        ("Погрузка грунта", "Погрузка местного грунта", "Погрузка растительного грунта экскаватором"),
    ),
    (
        "work_type",
        "WT_73AD9101",
        ("Погрузка ЩПГС", "Погрузка ЩПС", "Погрузка щебеночно-песчаной смеси"),
    ),
    (
        "work_type",
        "WT_B505C78",
        (
            "Транспортировка грунта на расстояние",
            "Дальность возки грунта",
            "Транспортировка местного грунта",
            "Транспортировка растительного грунта",
        ),
    ),
    (
        "material",
        "SHPGS",
        (
            "Смесь щебеночно-песчаная",
            "Смесь щебеночно песчаная",
            "Щебеночно-песчаная смесь",
            "Щебеночно песчаная смесь",
            "ЩПС",
        ),
    ),
    (
        "material",
        "SAND",
        (
            "Песок средней крупности",
            "Песок крупный",
            "Песок мелкий",
            "Песчаный грунт",
            "Песок по ГОСТ 25100",
            "Песок средней крупности по ГОСТ 25100",
        ),
    ),
    (
        "material",
        "SOIL",
        ("Местный грунт", "Растительный грунт", "Грунт ИГЭ", "Недренирующий грунт"),
    ),
    ("material", "PGS", ("Песчано-гравийная смесь", "Песчано гравийная смесь", "ПГС")),
    ("material", "ASPHALT_MIX", ("Асфальтобетонная смесь", "Асфальтобетонная смесь А32Нт")),
    ("material", "CEMENT", ("Цемент ПЦ", "Цемент ПЦ 500")),
    ("material", "CRUSHED_STONE", ("Щебень фр", "Щебень фракции", "Щебень 10-20", "Щебень 20-40")),
    ("material", "GEOMAT", ("Противоэрозионный геомат", "Противоэрозионного геомата", "Для геоматов")),
    ("material", "CONCRETE_M150", ("Бетон М150", "Бетона М150", "Бетонирование М150")),
)

SYSTEM_COLUMNS = {"created_at", "updated_at", "approved_at"}
DEPRECATED_EDIT_COLUMNS: dict[str, set[str]] = {
    "work_types": {"work_group"},
}

CASCADE_DELETE_CHILD_REFS: dict[str, dict[tuple[str, str], str]] = {
    "objects": {
        ("object_segments", "object_id"): "пикетажные границы объекта",
        ("object_map_points", "object_id"): "точки объекта на карте",
    },
    "temporary_roads": {
        ("temporary_road_pk_mappings", "road_id"): "пикетажные привязки ВАД",
        ("temp_road_points", "road_id"): "точки схемы ВАД",
    },
}

REFERENCE_BLOCKER_LABELS: dict[tuple[str, str], str] = {
    ("daily_work_items", "object_id"): "факты работ",
    ("material_movements", "from_object_id"): "перевозки, где объект указан источником",
    ("material_movements", "to_object_id"): "перевозки, где объект указан получателем",
    ("temporary_roads", "object_id"): "связанная временная дорога",
    ("temporary_road_status_segments", "road_id"): "статусы отсыпки временной дороги",
    ("temporary_road_state_corrections", "road_id"): "ручные коррекции статусов временной дороги",
    ("planned_work_items", "object_id"): "плановые объемы",
    ("project_work_items", "object_id"): "проектные объемы",
    ("object_segments", "object_id"): "пикетажные границы объекта",
    ("object_map_points", "object_id"): "точки объекта на карте",
}


@contextmanager
def _db_cursor():
    conn = get_conn()
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            yield conn, cur
    finally:
        conn.close()


def _non_pile_control_report_condition(report_alias: str = "dr") -> str:
    return (
        f"NOT (COALESCE({report_alias}.source_type, '') = 'web_pile_control' "
        f"OR COALESCE({report_alias}.review_payload->'review_flags', '[]'::jsonb) ?| ARRAY['isso', 'контроль свай'])"
    )


def _daily_report_filter_parts(table: str, columns: set[str]) -> tuple[sql.SQL, sql.SQL]:
    if table == "daily_reports":
        return sql.SQL(""), sql.SQL("WHERE " + _non_pile_control_report_condition("t"))
    if "daily_report_id" in columns:
        return (
            sql.SQL("LEFT JOIN daily_reports dr_scope ON t.daily_report_id = dr_scope.id"),
            sql.SQL("WHERE " + _non_pile_control_report_condition("dr_scope")),
        )
    return sql.SQL(""), sql.SQL("")


def _database_table_filter_parts(table: str, columns: set[str]) -> tuple[sql.SQL, sql.SQL]:
    joins, where = _daily_report_filter_parts(table, columns)
    if table == "work_types":
        return joins, sql.SQL("WHERE t.code NOT LIKE 'POVSR\\_%%' ESCAPE '\\'")
    if table == "work_type_aliases":
        return joins, sql.SQL("WHERE t.kind <> 'work_type' OR t.canonical_code NOT LIKE 'POVSR\\_%%' ESCAPE '\\'")
    return joins, where


class ChangeRequest(BaseModel):
    action: Literal["insert", "update", "delete"]
    pk: dict[str, Any] | None = None
    values: dict[str, Any] | None = None
    confirmed: bool = False
    changed_by: str | None = Field(default="dashboard", max_length=80)


class ReferenceNameUpdate(BaseModel):
    kind: Literal["object", "work_type", "material", "constructive", "object_type", "temporary_road"]
    code: str = Field(min_length=1, max_length=160)
    name: str = Field(min_length=1, max_length=500)
    changed_by: str | None = Field(default="dashboard", max_length=80)


class ReferenceDetailUpdate(BaseModel):
    kind: Literal["object", "work_type", "material", "constructive", "object_type", "temporary_road"]
    code: str = Field(min_length=1, max_length=160)
    values: dict[str, Any] = Field(default_factory=dict)
    changed_by: str | None = Field(default="dashboard", max_length=80)


REFERENCE_NAME_TABLES: dict[str, dict[str, str]] = {
    "object": {
        "table": "objects",
        "code_column": "object_code",
        "name_column": "name",
        "label": "Объекты",
    },
    "work_type": {
        "table": "work_types",
        "code_column": "code",
        "name_column": "name",
        "label": "Виды работ",
    },
    "material": {
        "table": "materials",
        "code_column": "code",
        "name_column": "name",
        "label": "Материалы",
    },
    "constructive": {
        "table": "constructives",
        "code_column": "code",
        "name_column": "name",
        "label": "Конструктивы объектов",
    },
    "object_type": {
        "table": "object_types",
        "code_column": "code",
        "name_column": "name",
        "label": "Типы объектов",
    },
    "temporary_road": {
        "table": "temporary_roads",
        "code_column": "road_code",
        "name_column": "road_name",
        "label": "Временные дороги",
    },
}

REFERENCE_DETAIL_FIELDS: dict[str, set[str]] = {
    "object": {"name", "comment", "is_active", "object_type_code", "pk_start", "pk_end", "pk_raw_text"},
    "work_type": {"name", "default_unit", "analytics_tag", "is_active", "show_in_timeline", "productivity_enabled", "rate_novgorod", "rate_tver"},
    "material": {"name", "default_unit", "is_active"},
    "constructive": {"name", "comment", "is_active", "object_code"},
    "object_type": {
        "name", "is_active", "map_enabled", "work_accounting_enabled",
        "material_accounting_enabled", "is_linear", "accounting_note",
    },
    "temporary_road": {
        "name", "section_code", "road_type", "ad_start_pk", "ad_end_pk", "rail_start_pk",
        "rail_end_pk", "display_in_dashboard", "can_translate_to_rail", "comment",
    },
}


def _parse_pk_input(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    text = str(value).strip().lower().replace("ё", "е").replace(",", ".")
    if not text:
        return None
    numeric = re.fullmatch(r"\d+(?:\.\d+)?", text)
    if numeric:
        raw = Decimal(text)
        return raw * Decimal(100) if raw < Decimal(10000) else raw
    parsed = parse_pk_value(text)
    if parsed is None:
        raise HTTPException(status_code=400, detail=f"Не удалось разобрать пикетаж: {value}")
    return Decimal(str(parsed))


def _format_pk_human(value: Any) -> str | None:
    if value in (None, ""):
        return None
    raw = Decimal(str(value))
    pk = int(raw // Decimal(100))
    plus = raw - Decimal(pk * 100)
    if plus == plus.to_integral_value():
        plus_text = f"{int(plus):02d}"
    else:
        plus_text = f"{plus:05.2f}".replace(".", ",")
    return f"ПК{pk}+{plus_text}"


def _format_pk_range_human(start: Any, end: Any, raw_text: Any = None) -> str | None:
    """Return one canonical human route-picket label.

    Numeric DB columns (`pk_start`/`pk_end`) are the source of truth.  Older
    catalog rows sometimes keep `pk_raw_text` in mixed forms (`ПК...`, bare
    metres, or stale OCR text), so using raw text first made the Database tab
    show different picketage formats for the same kind of value.
    """
    start_label = _format_pk_human(start)
    end_label = _format_pk_human(end)
    if start_label and end_label and start_label != end_label:
        return f"{start_label} - {end_label}"
    if start_label or end_label:
        return start_label or end_label
    if raw_text:
        return str(raw_text)
    return None


def _reference_pk_pair(values: dict[str, Any], start_key: str, end_key: str, label: str) -> tuple[Decimal | None, Decimal | None] | None:
    if start_key not in values and end_key not in values:
        return None
    start_raw = values.get(start_key)
    end_raw = values.get(end_key)
    if start_raw in (None, "") and end_raw in (None, ""):
        return None, None
    if start_raw in (None, "") or end_raw in (None, ""):
        raise HTTPException(status_code=400, detail=f"Для {label} нужно заполнить начало и конец")
    start = _parse_pk_input(start_raw)
    end = _parse_pk_input(end_raw)
    if start is None or end is None:
        raise HTTPException(status_code=400, detail=f"Для {label} нужно заполнить начало и конец")
    if start > end:
        start, end = end, start
    return start, end


def _json_safe(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, dict):
        return {k: _json_safe(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_json_safe(v) for v in value]
    return value


def _allowed_table(table: str, *, require_editable: bool = False) -> None:
    allowed = EDITABLE_TABLES if require_editable else EXPOSED_TABLES
    if table not in allowed:
        raise HTTPException(status_code=404, detail="Table is not exposed for dashboard editing")


def _fetch_all(cur, query, params=()) -> list[dict[str, Any]]:
    cur.execute(query, params)
    return [dict(row) for row in cur.fetchall()]


def _table_exists(cur, table: str) -> bool:
    cur.execute(
        """
        SELECT EXISTS (
          SELECT 1
          FROM information_schema.tables
          WHERE table_schema = 'public' AND table_name = %s
        ) AS exists
        """,
        (table,),
    )
    row = cur.fetchone()
    return bool(row and row["exists"])


WORK_TYPE_ANALYTICS_REQUIRED_COLUMNS: dict[str, tuple[str, ...]] = {
    "work_types": ("analytics_tag",),
    "daily_work_items": ("analytics_tag",),
    "daily_work_item_segments": ("analytics_tag",),
}


def _work_type_analytics_schema_prepared(cur) -> bool:
    cur.execute(
        """
        SELECT table_name, column_name
        FROM information_schema.columns
        WHERE table_schema = 'public'
          AND table_name = ANY(%s)
        """,
        (list(WORK_TYPE_ANALYTICS_REQUIRED_COLUMNS),),
    )
    have_columns = set()
    for row in cur.fetchall():
        if isinstance(row, dict):
            have_columns.add((row.get("table_name"), row.get("column_name")))
        else:
            have_columns.add((row[0], row[1]))
    for table_name, columns in WORK_TYPE_ANALYTICS_REQUIRED_COLUMNS.items():
        for column in columns:
            if (table_name, column) not in have_columns:
                return False
    cur.execute("SELECT to_regclass('public.work_analytics_tags') IS NOT NULL AS exists")
    row = cur.fetchone()
    if not row:
        return False
    if isinstance(row, dict):
        return bool(row.get("exists"))
    return bool(row[0])


def _ensure_work_type_analytics_schema(cur) -> None:
    if _work_type_analytics_schema_prepared(cur):
        return
    cur.execute("SET LOCAL lock_timeout = '2s'")
    cur.execute("SET LOCAL statement_timeout = '60s'")
    cur.execute("ALTER TABLE work_types ADD COLUMN IF NOT EXISTS analytics_tag text")
    cur.execute("ALTER TABLE daily_work_items ADD COLUMN IF NOT EXISTS analytics_tag text")
    cur.execute("ALTER TABLE daily_work_item_segments ADD COLUMN IF NOT EXISTS analytics_tag text")
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS work_analytics_tags (
          name text PRIMARY KEY,
          is_active boolean NOT NULL DEFAULT true,
          created_at timestamptz NOT NULL DEFAULT now()
        )
        """
    )
    cur.execute(
        """
        WITH ranked AS (
          SELECT work_type_id,
                 BTRIM(analytics_tag) AS analytics_tag,
                 COUNT(*) AS cnt,
                 ROW_NUMBER() OVER (
                   PARTITION BY work_type_id
                   ORDER BY COUNT(*) DESC, BTRIM(analytics_tag)
                 ) AS rn
          FROM daily_work_items
          WHERE NULLIF(BTRIM(analytics_tag), '') IS NOT NULL
          GROUP BY work_type_id, BTRIM(analytics_tag)
        )
        UPDATE work_types wt
        SET analytics_tag = ranked.analytics_tag
        FROM ranked
        WHERE ranked.work_type_id = wt.id
          AND ranked.rn = 1
          AND NULLIF(BTRIM(COALESCE(wt.analytics_tag, '')), '') IS NULL
        """
    )
    cur.execute(
        """
        UPDATE daily_work_items dwi
        SET analytics_tag = NULLIF(BTRIM(wt.analytics_tag), '')
        FROM work_types wt
        WHERE wt.id = dwi.work_type_id
          AND COALESCE(dwi.analytics_tag, '') IS DISTINCT FROM COALESCE(NULLIF(BTRIM(wt.analytics_tag), ''), '')
        """
    )
    cur.execute(
        """
        UPDATE daily_work_item_segments seg
        SET analytics_tag = dwi.analytics_tag
        FROM daily_work_items dwi
        WHERE dwi.id = seg.daily_work_item_id
          AND COALESCE(seg.analytics_tag, '') IS DISTINCT FROM COALESCE(dwi.analytics_tag, '')
        """
    )



WORK_TYPE_RATE_REGIONS = {
    "novgorod": "Новгородская область до ПК3039",
    "tver": "Тверская область после ПК3039",
}
PILE_LENGTH_RATE_WORK_CODES = {"PILE_TRIAL", "PILE_MAIN"}
PILE_LENGTH_RATE_WORK_NAMES = {"забивка пробных свай", "забивка основных свай"}
PILE_RATE_FIELD_RE = re.compile(r"^pile_rate_(\d+(?:[._]\d+)?)_(novgorod|tver)$")


def _work_type_supports_pile_length_rates(row: dict[str, Any]) -> bool:
    code = str(row.get("code") or "").strip().upper()
    name = str(row.get("name") or "").strip().lower()
    return code in PILE_LENGTH_RATE_WORK_CODES or name in PILE_LENGTH_RATE_WORK_NAMES


def _pile_length_from_type(value: Any) -> Decimal | None:
    text = str(value or "").strip().upper().replace("C", "С").replace(",", ".")
    if not text or ";" in text:
        return None
    direct = re.fullmatch(r"\s*(\d{1,2}(?:\.\d+)?)\s*", text)
    if direct:
        return _pile_rate_length_value(direct.group(1))
    matches = [Decimal(m) for m in re.findall(r"С(\d{2,3})", text)]
    if not matches:
        return None
    length = sum(matches) / Decimal(10)
    return length.quantize(Decimal("1")) if length == length.to_integral_value() else length.normalize()


def _pile_rate_length_key(value: Any) -> str | None:
    length = _pile_rate_length_value(value)
    if length is None:
        return None
    return format(length.normalize(), "f").rstrip("0").rstrip(".") or "0"


def _pile_rate_row_length_key(row: dict[str, Any]) -> str | None:
    return _pile_rate_length_key(row.get("pile_length_m"))


def _pile_rate_row_for_length(length: Decimal, saved: dict[str, dict[str, Any]]) -> dict[str, Any]:
    key = _pile_rate_length_key(length)
    existing = saved.get(key or "") if key else None
    return {
        "pile_length_m": length,
        "rate_novgorod": (existing or {}).get("rate_novgorod"),
        "rate_tver": (existing or {}).get("rate_tver"),
    }


def _pile_rate_length_value(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    text = str(value).strip().replace("_", ".").replace(",", ".")
    if not text:
        return None
    try:
        length = Decimal(text)
    except Exception:
        raise HTTPException(status_code=400, detail=f"Некорректная длина сваи: {value}")
    if length <= 0:
        raise HTTPException(status_code=400, detail=f"Некорректная длина сваи: {value}")
    return length


def _pile_rate_field_parts(field_name: str) -> tuple[Decimal, str] | None:
    match = PILE_RATE_FIELD_RE.fullmatch(field_name)
    if not match:
        return None
    return _pile_rate_length_value(match.group(1)), match.group(2)


def _ensure_work_type_unit_rates_table(cur) -> None:
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS work_type_unit_rates (
            id uuid PRIMARY KEY DEFAULT gen_random_uuid(),
            work_type_id uuid NOT NULL REFERENCES work_types(id) ON DELETE CASCADE,
            region_code text NOT NULL,
            region_label text NOT NULL,
            rate numeric,
            currency text NOT NULL DEFAULT 'RUB',
            unit text,
            pile_length_m numeric,
            updated_at timestamptz NOT NULL DEFAULT now(),
            updated_by text,
            CONSTRAINT work_type_unit_rates_region_check CHECK (region_code IN ('novgorod', 'tver'))
        )
        """
    )
    cur.execute("ALTER TABLE work_type_unit_rates ADD COLUMN IF NOT EXISTS pile_length_m numeric")
    cur.execute("ALTER TABLE work_type_unit_rates DROP CONSTRAINT IF EXISTS work_type_unit_rates_unique")
    cur.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS work_type_unit_rates_unique_base_idx
        ON work_type_unit_rates(work_type_id, region_code)
        WHERE pile_length_m IS NULL
        """
    )
    cur.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS work_type_unit_rates_unique_length_idx
        ON work_type_unit_rates(work_type_id, region_code, pile_length_m)
        WHERE pile_length_m IS NOT NULL
        """
    )


def _rate_value(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return Decimal(text)
    except Exception:
        raise HTTPException(status_code=400, detail=f"Некорректная расценка: {value}")


def _upsert_work_type_rate(cur, work_type_id: Any, region_code: str, rate_value: Any, unit: Any, changed_by: str | None, pile_length_m: Any = None) -> None:
    rate = _rate_value(rate_value)
    pile_length = _pile_rate_length_value(pile_length_m)
    if rate is None:
        if pile_length is None:
            cur.execute(
                "DELETE FROM work_type_unit_rates WHERE work_type_id = %s AND region_code = %s AND pile_length_m IS NULL",
                (work_type_id, region_code),
            )
        else:
            cur.execute(
                "DELETE FROM work_type_unit_rates WHERE work_type_id = %s AND region_code = %s AND pile_length_m = %s",
                (work_type_id, region_code, pile_length),
            )
        return
    if pile_length is None:
        cur.execute(
            """
            INSERT INTO work_type_unit_rates (work_type_id, region_code, region_label, rate, unit, pile_length_m, updated_at, updated_by)
            VALUES (%s, %s, %s, %s, %s, NULL, now(), %s)
            ON CONFLICT (work_type_id, region_code) WHERE pile_length_m IS NULL DO UPDATE
            SET region_label = EXCLUDED.region_label,
                rate = EXCLUDED.rate,
                unit = EXCLUDED.unit,
                updated_at = now(),
                updated_by = EXCLUDED.updated_by
            """,
            (work_type_id, region_code, WORK_TYPE_RATE_REGIONS[region_code], rate, unit or None, changed_by or "dashboard"),
        )
        return
    cur.execute(
        """
        INSERT INTO work_type_unit_rates (work_type_id, region_code, region_label, rate, unit, pile_length_m, updated_at, updated_by)
        VALUES (%s, %s, %s, %s, %s, %s, now(), %s)
        ON CONFLICT (work_type_id, region_code, pile_length_m) WHERE pile_length_m IS NOT NULL DO UPDATE
        SET region_label = EXCLUDED.region_label,
            rate = EXCLUDED.rate,
            unit = EXCLUDED.unit,
            updated_at = now(),
            updated_by = EXCLUDED.updated_by
        """,
        (work_type_id, region_code, WORK_TYPE_RATE_REGIONS[region_code], rate, unit or None, pile_length, changed_by or "dashboard"),
    )

def _sync_work_type_fact_analytics_tags(cur, work_type_code: str) -> None:
    cur.execute(
        """
        UPDATE daily_work_items dwi
        SET analytics_tag = NULLIF(BTRIM(wt.analytics_tag), '')
        FROM work_types wt
        WHERE wt.id = dwi.work_type_id
          AND wt.code = %s
          AND COALESCE(dwi.analytics_tag, '') IS DISTINCT FROM COALESCE(NULLIF(BTRIM(wt.analytics_tag), ''), '')
        """,
        (work_type_code,),
    )
    cur.execute(
        """
        UPDATE daily_work_item_segments seg
        SET analytics_tag = dwi.analytics_tag
        FROM daily_work_items dwi
        JOIN work_types wt ON wt.id = dwi.work_type_id
        WHERE dwi.id = seg.daily_work_item_id
          AND wt.code = %s
          AND COALESCE(seg.analytics_tag, '') IS DISTINCT FROM COALESCE(dwi.analytics_tag, '')
        """,
        (work_type_code,),
    )


def _columns(cur, table: str) -> list[dict[str, Any]]:
    return _fetch_all(
        cur,
        """
        SELECT column_name, data_type, udt_name, is_nullable, column_default,
               character_maximum_length, numeric_precision, numeric_scale
        FROM information_schema.columns
        WHERE table_schema = 'public' AND table_name = %s
        ORDER BY ordinal_position
        """,
        (table,),
    )


def _primary_key(cur, table: str) -> list[str]:
    cur.execute(
        """
        SELECT a.attname AS column_name
        FROM pg_index i
        JOIN pg_attribute a ON a.attrelid = i.indrelid AND a.attnum = ANY(i.indkey)
        WHERE i.indrelid = %s::regclass AND i.indisprimary
        ORDER BY array_position(i.indkey, a.attnum)
        """,
        (f"public.{table}",),
    )
    return [row["column_name"] for row in cur.fetchall()]


def _foreign_keys(cur, table: str) -> list[dict[str, Any]]:
    return _fetch_all(
        cur,
        """
        SELECT c.conname AS constraint_name,
               src.relname AS source_table,
               src_col.attname AS source_column,
               dst.relname AS target_table,
               dst_col.attname AS target_column,
               c.confdeltype AS on_delete
        FROM pg_constraint c
        JOIN pg_class src ON src.oid = c.conrelid
        JOIN pg_class dst ON dst.oid = c.confrelid
        JOIN unnest(c.conkey) WITH ORDINALITY AS src_key(attnum, ord) ON true
        JOIN unnest(c.confkey) WITH ORDINALITY AS dst_key(attnum, ord) ON dst_key.ord = src_key.ord
        JOIN pg_attribute src_col ON src_col.attrelid = src.oid AND src_col.attnum = src_key.attnum
        JOIN pg_attribute dst_col ON dst_col.attrelid = dst.oid AND dst_col.attnum = dst_key.attnum
        WHERE c.contype = 'f' AND src.relname = %s
        ORDER BY c.conname, src_key.ord
        """,
        (table,),
    )


def _incoming_refs(cur, table: str) -> list[dict[str, Any]]:
    return _fetch_all(
        cur,
        """
        SELECT c.conname AS constraint_name,
               src.relname AS source_table,
               src_col.attname AS source_column,
               dst.relname AS target_table,
               dst_col.attname AS target_column,
               c.confdeltype AS on_delete
        FROM pg_constraint c
        JOIN pg_class src ON src.oid = c.conrelid
        JOIN pg_class dst ON dst.oid = c.confrelid
        JOIN unnest(c.conkey) WITH ORDINALITY AS src_key(attnum, ord) ON true
        JOIN unnest(c.confkey) WITH ORDINALITY AS dst_key(attnum, ord) ON dst_key.ord = src_key.ord
        JOIN pg_attribute src_col ON src_col.attrelid = src.oid AND src_col.attnum = src_key.attnum
        JOIN pg_attribute dst_col ON dst_col.attrelid = dst.oid AND dst_col.attnum = dst_key.attnum
        WHERE c.contype = 'f' AND dst.relname = %s
        ORDER BY src.relname, c.conname, src_key.ord
        """,
        (table,),
    )


def _indexes(cur, table: str) -> list[dict[str, Any]]:
    return _fetch_all(
        cur,
        """
        SELECT indexname, indexdef
        FROM pg_indexes
        WHERE schemaname = 'public' AND tablename = %s
        ORDER BY indexname
        """,
        (table,),
    )


def _fk_for_column(cur, table: str, column: str) -> dict[str, Any] | None:
    for fk in _foreign_keys(cur, table):
        if fk["source_column"] == column:
            return fk
    return None


def _label_expression(cur, target_table: str, target_alias: str = "d") -> sql.Composed:
    target_columns = {col["column_name"] for col in _columns(cur, target_table)}
    alias = sql.Identifier(target_alias)
    if {"code", "name"} <= target_columns:
        return sql.SQL("CONCAT_WS(' · ', {}.{}, {}.{})").format(
            alias,
            sql.Identifier("code"),
            alias,
            sql.Identifier("name"),
        )
    for column in ("name", "short_name", "object_code", "road_code", "code"):
        if column in target_columns:
            return sql.SQL("{}.{}::text").format(alias, sql.Identifier(column))
    return sql.SQL("NULL")


def _where_pk(pk_cols: list[str], pk: dict[str, Any] | None) -> tuple[sql.SQL, list[Any]]:
    if not pk_cols:
        raise HTTPException(status_code=400, detail="Table has no primary key")
    if not pk:
        raise HTTPException(status_code=400, detail="Primary key is required")
    missing = [col for col in pk_cols if col not in pk]
    if missing:
        raise HTTPException(status_code=400, detail=f"Missing primary key column(s): {', '.join(missing)}")
    parts = [sql.SQL("{} = %s").format(sql.Identifier(col)) for col in pk_cols]
    return sql.SQL(" AND ").join(parts), [pk[col] for col in pk_cols]


def _row_by_pk(cur, table: str, pk_cols: list[str], pk: dict[str, Any] | None) -> dict[str, Any] | None:
    where_sql, params = _where_pk(pk_cols, pk)
    query = sql.SQL("SELECT * FROM {} WHERE {} LIMIT 1").format(sql.Identifier(table), where_sql)
    cur.execute(query, params)
    row = cur.fetchone()
    return dict(row) if row else None


def _cascade_delete_child_label(table: str, ref: dict[str, Any]) -> str | None:
    return CASCADE_DELETE_CHILD_REFS.get(table, {}).get((ref["source_table"], ref["source_column"]))


def _reference_blocker_message(ref: dict[str, Any], count: int) -> str:
    label = REFERENCE_BLOCKER_LABELS.get(
        (ref["source_table"], ref["source_column"]),
        EXPOSED_TABLES.get(ref["source_table"], ref["source_table"]),
    )
    return f"Нельзя удалить: есть связанные записи - {label} ({count})"


def _delete_safe_child_refs(cur, table: str, before: dict[str, Any] | None) -> dict[str, int]:
    if not before:
        return {}
    deleted: dict[str, int] = {}
    for ref in _incoming_refs(cur, table):
        label = _cascade_delete_child_label(table, ref)
        if not label:
            continue
        target_value = before.get(ref["target_column"])
        if target_value in (None, ""):
            continue
        query = sql.SQL("DELETE FROM {} WHERE {} = %s").format(
            sql.Identifier(ref["source_table"]),
            sql.Identifier(ref["source_column"]),
        )
        cur.execute(query, (target_value,))
        if cur.rowcount:
            deleted[label] = cur.rowcount
    return deleted


def _temp_road_linked_object_id(before: dict[str, Any] | None) -> str | None:
    if not before:
        return None
    value = before.get("object_id")
    return str(value).strip() if value not in (None, "") else None


def _count_ref_rows(
    cur,
    table: str,
    column: str,
    value: Any,
    *,
    exclude_id: Any | None = None,
    non_demo_only: bool = False,
) -> int:
    where_parts = [sql.SQL("{} = %s").format(sql.Identifier(column))]
    params: list[Any] = [value]
    if exclude_id not in (None, ""):
        where_parts.append(sql.SQL("id <> %s"))
        params.append(exclude_id)
    if non_demo_only:
        table_columns = {col["column_name"] for col in _columns(cur, table)}
        if "is_demo" in table_columns:
            where_parts.append(sql.SQL("COALESCE(is_demo, false) = false"))
    query = sql.SQL("SELECT COUNT(*) AS cnt FROM {} WHERE {}").format(
        sql.Identifier(table),
        sql.SQL(" AND ").join(where_parts),
    )
    cur.execute(query, params)
    row = cur.fetchone()
    return int(row["cnt"] or 0) if row else 0


def _temp_road_linked_object_delete_errors(cur, before: dict[str, Any] | None) -> tuple[list[str], list[str]]:
    object_id = _temp_road_linked_object_id(before)
    if not object_id:
        return [], []

    cur.execute(
        """
        SELECT o.id::text, o.object_code, o.name, ot.code AS object_type_code
        FROM objects o
        LEFT JOIN object_types ot ON ot.id = o.object_type_id
        WHERE o.id = %s
        LIMIT 1
        """,
        (object_id,),
    )
    object_row = cur.fetchone()
    if not object_row:
        return [], []

    errors: list[str] = []
    warnings: list[str] = [
        f"Вместе с дорогой будет удален связанный объект: {object_row['object_code'] or object_row['name']}"
    ]
    road_id = before.get("id") if before else None
    for ref in _incoming_refs(cur, "objects"):
        cascade_label = _cascade_delete_child_label("objects", ref)
        count = _count_ref_rows(
            cur,
            ref["source_table"],
            ref["source_column"],
            object_id,
            exclude_id=road_id if ref["source_table"] == "temporary_roads" else None,
            non_demo_only=ref["source_table"] in {"daily_work_items", "material_movements"},
        )
        if not count:
            continue
        if cascade_label:
            warnings.append(f"Вместе со связанным объектом будут удалены: {cascade_label} ({count})")
            continue
        errors.append("Связанный объект дороги нельзя удалить автоматически: " + _reference_blocker_message(ref, count))
    return errors, warnings


def _object_linked_temp_road_delete_errors(cur, before: dict[str, Any] | None) -> tuple[list[str], list[str]]:
    if not before:
        return [], []
    object_id = before.get("id")
    if object_id in (None, ""):
        return [], []

    cur.execute(
        """
        SELECT id::text, road_code, road_name
        FROM temporary_roads
        WHERE object_id = %s
        ORDER BY road_code
        """,
        (object_id,),
    )
    roads = [dict(row) for row in cur.fetchall()]
    if not roads:
        return [], []

    errors: list[str] = []
    warnings: list[str] = []
    for road in roads:
        warnings.append(f"Вместе с объектом будет удалена связанная временная дорога: {road.get('road_code') or road.get('road_name')}")
        for ref in _incoming_refs(cur, "temporary_roads"):
            cascade_label = _cascade_delete_child_label("temporary_roads", ref)
            count = _count_ref_rows(
                cur,
                ref["source_table"],
                ref["source_column"],
                road["id"],
                non_demo_only=ref["source_table"] in {"temporary_road_status_segments"},
            )
            if not count:
                continue
            if cascade_label:
                warnings.append(f"Вместе со связанной дорогой будут удалены: {cascade_label} ({count})")
                continue
            errors.append("Связанную временную дорогу нельзя удалить автоматически: " + _reference_blocker_message(ref, count))
    return errors, warnings


def _delete_temp_roads_for_object(cur, before: dict[str, Any] | None) -> dict[str, int]:
    if not before:
        return {}
    object_id = before.get("id")
    if object_id in (None, ""):
        return {}
    cur.execute(
        """
        SELECT *
        FROM temporary_roads
        WHERE object_id = %s
        ORDER BY road_code
        """,
        (object_id,),
    )
    deleted: dict[str, int] = {}
    for road in [dict(row) for row in cur.fetchall()]:
        for label, count in _delete_safe_child_refs(cur, "temporary_roads", road).items():
            deleted[label] = deleted.get(label, 0) + count
        cur.execute("DELETE FROM temporary_roads WHERE id = %s RETURNING id", (road["id"],))
        if cur.fetchone():
            deleted["связанная временная дорога"] = deleted.get("связанная временная дорога", 0) + 1
    return deleted


def _delete_temp_road_linked_object(cur, before: dict[str, Any] | None) -> dict[str, int]:
    object_id = _temp_road_linked_object_id(before)
    if not object_id:
        return {}
    object_before = _row_by_pk(cur, "objects", ["id"], {"id": object_id})
    if not object_before:
        return {}
    deleted = _delete_safe_child_refs(cur, "objects", object_before)
    cur.execute("DELETE FROM objects WHERE id = %s RETURNING id", (object_id,))
    if cur.fetchone():
        deleted["связанный объект временной дороги"] = 1
    return deleted


def _ensure_audit_table(cur) -> None:
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS public.db_change_audit (
            id BIGSERIAL PRIMARY KEY,
            table_name TEXT NOT NULL,
            action TEXT NOT NULL,
            pk_json JSONB,
            before_json JSONB,
            after_json JSONB,
            changed_by TEXT NOT NULL DEFAULT 'dashboard',
            created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
        )
        """
    )


def _validate(cur, table: str, request: ChangeRequest) -> dict[str, Any]:
    _allowed_table(table, require_editable=True)
    if not _table_exists(cur, table):
        raise HTTPException(status_code=404, detail="Table does not exist")

    columns = _columns(cur, table)
    column_names = {col["column_name"] for col in columns}
    pk_cols = _primary_key(cur, table)
    values = request.values or {}
    errors: list[str] = []
    warnings: list[str] = []
    before = None
    after = None

    unknown = sorted(set(values) - column_names)
    if unknown:
        errors.append(f"Unknown column(s): {', '.join(unknown)}")
    deprecated = sorted(set(values) & DEPRECATED_EDIT_COLUMNS.get(table, set()))
    if deprecated:
        errors.append(f"Column(s) are disabled: {', '.join(deprecated)}")

    if request.action == "update":
        if not values:
            errors.append("No changed values provided")
        blocked = sorted((set(values) & set(pk_cols)) | (set(values) & SYSTEM_COLUMNS))
        if blocked:
            errors.append(f"Column(s) are protected: {', '.join(blocked)}")
        before = _row_by_pk(cur, table, pk_cols, request.pk)
        if not before:
            errors.append("Target row was not found")
        else:
            changed = {key: value for key, value in values.items() if _json_safe(before.get(key)) != _json_safe(value)}
            if not changed:
                warnings.append("No effective field changes detected")
            after = {**before, **values}

    elif request.action == "insert":
        missing_required = [
            col["column_name"]
            for col in columns
            if col["is_nullable"] == "NO"
            and col["column_default"] is None
            and col["column_name"] not in values
            and col["column_name"] not in SYSTEM_COLUMNS
        ]
        if missing_required:
            errors.append(f"Required column(s) are missing: {', '.join(missing_required)}")
        after = values

    elif request.action == "delete":
        before = _row_by_pk(cur, table, pk_cols, request.pk)
        if not before:
            errors.append("Target row was not found")
        else:
            for ref in _incoming_refs(cur, table):
                target_col = ref["target_column"]
                if target_col not in before:
                    continue
                count_query = sql.SQL("SELECT COUNT(*) AS cnt FROM {} WHERE {} = %s").format(
                    sql.Identifier(ref["source_table"]),
                    sql.Identifier(ref["source_column"]),
                )
                cur.execute(count_query, (before[target_col],))
                count = int(cur.fetchone()["cnt"])
                if count:
                    cascade_label = _cascade_delete_child_label(table, ref)
                    if cascade_label:
                        warnings.append(f"Вместе с записью будут удалены: {cascade_label} ({count})")
                    elif table == "objects" and (ref["source_table"], ref["source_column"]) == ("temporary_roads", "object_id"):
                        linked_errors, linked_warnings = _object_linked_temp_road_delete_errors(cur, before)
                        errors.extend(linked_errors)
                        warnings.extend(linked_warnings)
                    else:
                        errors.append(_reference_blocker_message(ref, count))
            if table == "temporary_roads":
                linked_errors, linked_warnings = _temp_road_linked_object_delete_errors(cur, before)
                errors.extend(linked_errors)
                warnings.extend(linked_warnings)

    if values:
        fk_by_column = {fk["source_column"]: fk for fk in _foreign_keys(cur, table)}
        for column, value in values.items():
            if value in (None, "") or column not in fk_by_column:
                continue
            fk = fk_by_column[column]
            exists_query = sql.SQL("SELECT 1 FROM {} WHERE {} = %s LIMIT 1").format(
                sql.Identifier(fk["target_table"]),
                sql.Identifier(fk["target_column"]),
            )
            cur.execute(exists_query, (value,))
            if not cur.fetchone():
                errors.append(f"{column} references missing {fk['target_table']}.{fk['target_column']}")

    diff = []
    if before and values:
        for key in sorted(values):
            diff.append({"column": key, "before": _json_safe(before.get(key)), "after": _json_safe(values[key])})

    return {
        "ok": not errors,
        "errors": errors,
        "warnings": warnings,
        "diff": diff,
        "before": _json_safe(before),
        "after": _json_safe(after),
    }


@router.get("/reference-names")
def reference_names(
    kind: Literal["object", "work_type", "material", "constructive", "object_type", "temporary_road"] | None = None,
    q: str | None = Query(default=None, max_length=120),
    _user=Depends(current_user),
):
    specs = {kind: REFERENCE_NAME_TABLES[kind]} if kind else REFERENCE_NAME_TABLES
    rows: list[dict[str, Any]] = []
    with _db_cursor() as (conn, cur):
        for ref_kind, spec in specs.items():
            if not _table_exists(cur, spec["table"]):
                continue
            if ref_kind == "object":
                params: list[Any] = []
                where = "TRUE"
                if q and q.strip():
                    where = "(o.object_code ILIKE %s OR o.name ILIKE %s OR ot.name ILIKE %s)"
                    like = f"%{q.strip()}%"
                    params.extend([like, like, like])
                cur.execute(
                    f"""
                    SELECT o.*,
                           ot.code AS object_type_code,
                           ot.name AS object_type_name,
                           os.pk_start,
                           os.pk_end,
                           os.pk_raw_text,
                           os.id::text AS object_segment_id
                    FROM objects o
                    JOIN object_types ot ON ot.id = o.object_type_id
                    LEFT JOIN LATERAL (
                      SELECT id, pk_start, pk_end, pk_raw_text
                      FROM object_segments
                      WHERE object_id = o.id
                      ORDER BY pk_start NULLS LAST, created_at NULLS LAST
                      LIMIT 1
                    ) os ON true
                    WHERE {where}
                    ORDER BY o.name NULLS LAST, o.object_code
                    LIMIT 500
                    """,
                    params,
                )
                for row in cur.fetchall():
                    data = _json_safe(dict(row))
                    if data.get("pk_start") is not None:
                        data["pk_start_label"] = _format_pk_human(data["pk_start"])
                    if data.get("pk_end") is not None:
                        data["pk_end_label"] = _format_pk_human(data["pk_end"])
                    rows.append({
                        "kind": ref_kind,
                        "kind_label": spec["label"],
                        "code": row["object_code"],
                        "name": row["name"] or row["object_code"],
                        "data": data,
                        "table": spec["table"],
                        "code_column": spec["code_column"],
                        "name_column": spec["name_column"],
                    })
                continue
            where = sql.SQL(INTERNAL_POVSR_WORK_TYPE_SQL if ref_kind == "work_type" else "TRUE")
            params: list[Any] = []
            if q and q.strip():
                search_where = sql.SQL("({code} ILIKE %s OR {name} ILIKE %s)").format(
                    code=sql.Identifier(spec["code_column"]),
                    name=sql.Identifier(spec["name_column"]),
                )
                where = sql.SQL("({}) AND ({})").format(where, search_where)
                like = f"%{q.strip()}%"
                params.extend([like, like])
            query = sql.SQL(
                """
                SELECT *,
                       {code}::text AS __ref_code,
                       {name}::text AS __ref_name
                FROM {table}
                WHERE {where}
                ORDER BY {name} NULLS LAST, {code}
                LIMIT 500
                """
            ).format(
                code=sql.Identifier(spec["code_column"]),
                name=sql.Identifier(spec["name_column"]),
                table=sql.Identifier(spec["table"]),
                where=where,
            )
            cur.execute(query, params)
            for row in cur.fetchall():
                data = _json_safe({k: v for k, v in dict(row).items() if not k.startswith("__ref_")})
                rows.append({
                    "kind": ref_kind,
                    "kind_label": spec["label"],
                    "code": row["__ref_code"],
                    "name": row["__ref_name"] or row["__ref_code"],
                    "data": data,
                    "table": spec["table"],
                    "code_column": spec["code_column"],
                    "name_column": spec["name_column"],
                })
    return {"rows": rows, "kinds": [
        {"kind": k, "label": v["label"]} for k, v in REFERENCE_NAME_TABLES.items()
    ]}


@router.get("/reference-overview")
def reference_overview(_user=Depends(current_user)):
    ensure_temp_road_catalog_rows()
    with _db_cursor() as (conn, cur):
        _ensure_work_type_analytics_schema(cur)
        _ensure_work_type_unit_rates_table(cur)
        cur.execute(
            """
            SELECT o.id::text AS id,
                   o.object_code AS code,
                   o.name,
                   o.is_active,
                   o.comment,
                   o.created_at,
                   ot.code AS object_type_code,
                   ot.name AS object_type_name,
                   ot.sort_order AS object_type_sort_order,
                   ot.map_enabled AS object_type_map_enabled,
                   ot.work_accounting_enabled AS object_type_work_accounting_enabled,
                   ot.material_accounting_enabled AS object_type_material_accounting_enabled,
                   ot.is_linear AS object_type_is_linear,
                   c.code AS constructive_code,
                   c.name AS constructive_name,
                   seg.pk_start,
                   seg.pk_end,
                   seg.pk_raw_text,
                   COALESCE(sec.sections, '[]'::jsonb) AS sections,
                   COALESCE(wrk.project_item_count, 0)::int AS project_item_count,
                   COALESCE(wrk.work_type_count, 0)::int AS work_type_count,
                   COALESCE(wrk.work_tags, '') AS work_tags
            FROM objects o
            JOIN object_types ot ON ot.id = o.object_type_id
            LEFT JOIN constructives c ON c.id = o.constructive_id
            LEFT JOIN LATERAL (
              SELECT os.pk_start, os.pk_end, os.pk_raw_text
              FROM object_segments os
              WHERE os.object_id = o.id
              ORDER BY os.pk_start NULLS LAST, os.created_at NULLS LAST
              LIMIT 1
            ) seg ON true
            LEFT JOIN LATERAL (
              SELECT COALESCE(jsonb_agg(
                       jsonb_build_object(
                         'id', s.id::text,
                         'code', s.code,
                         'name', s.name,
                         'sort_order', s.sort_order,
                         'boundary_kinds', s.boundary_kinds
                       )
                       ORDER BY s.sort_order NULLS LAST, s.code
                     ), '[]'::jsonb) AS sections
              FROM (
                SELECT cs.id,
                       cs.code,
                       cs.name,
                       cs.sort_order,
                       string_agg(DISTINCT csv.boundary_kind, ', ' ORDER BY csv.boundary_kind) AS boundary_kinds
                FROM object_segments os_scope
                JOIN construction_section_versions csv
                  ON csv.is_current IS TRUE
                 AND csv.boundary_kind = CASE
                   WHEN ot.code = 'TEMP_ROAD' THEN 'temp_roads'
                   WHEN ot.code = 'SERVICE_ROAD' THEN 'temp_roads'
                   WHEN ot.code IN ('PILE_DRIVING_PAD', 'PILE_FIELD', 'PIPE', 'ISSO_ACCESS', 'BRIDGE', 'OVERPASS', 'TECH_ROAD') THEN 'piles_pipes'
                   ELSE 'main'
                 END
                 AND (
                   (
                     LEAST(os_scope.pk_start, os_scope.pk_end) < GREATEST(os_scope.pk_start, os_scope.pk_end)
                     AND GREATEST(
                       LEAST(os_scope.pk_start, os_scope.pk_end),
                       LEAST(csv.pk_start, csv.pk_end)
                     ) < LEAST(
                       GREATEST(os_scope.pk_start, os_scope.pk_end),
                       GREATEST(csv.pk_start, csv.pk_end)
                     )
                   )
                   OR (
                     os_scope.pk_start = os_scope.pk_end
                     AND os_scope.pk_start >= LEAST(csv.pk_start, csv.pk_end)
                     AND os_scope.pk_start < GREATEST(csv.pk_start, csv.pk_end)
                   )
                 )
                JOIN construction_sections cs
                  ON cs.id = csv.section_id
                 AND cs.is_active IS NOT FALSE
                WHERE os_scope.object_id = o.id
                  AND o.object_code NOT IN ('ASP_2840', 'BRIDGE_282952')
                GROUP BY cs.id, cs.code, cs.name, cs.sort_order
                UNION
                SELECT cs.id,
                       cs.code,
                       cs.name,
                       cs.sort_order,
                       'piles_pipes' AS boundary_kinds
                FROM construction_sections cs
                WHERE o.object_code IN ('ASP_2840', 'BRIDGE_282952')
                  AND cs.code = 'UCH_3'
                  AND cs.is_active IS NOT FALSE
              ) s
            ) sec ON true
            LEFT JOIN LATERAL (
              SELECT COUNT(*)::int AS project_item_count,
                     COUNT(DISTINCT pwi.work_type_id)::int AS work_type_count,
                     string_agg(DISTINCT COALESCE(NULLIF(BTRIM(wt.analytics_tag), ''), 'Без тега'), ', ') AS work_tags
              FROM project_work_items pwi
              JOIN work_types wt ON wt.id = pwi.work_type_id
              WHERE pwi.object_id = o.id
            ) wrk ON true
            ORDER BY ot.sort_order NULLS LAST, ot.name, o.name, o.object_code
            """
        )
        objects = []
        for row in cur.fetchall():
            item = _json_safe(dict(row))
            item["pk_label"] = _format_pk_range_human(row["pk_start"], row["pk_end"], row["pk_raw_text"])
            objects.append(item)

        cur.execute(
            """
            SELECT ot.id::text AS id,
                   ot.code,
                   ot.name,
                   ot.sort_order,
                   ot.is_active,
                   ot.map_enabled,
                   ot.work_accounting_enabled,
                   ot.material_accounting_enabled,
                   ot.is_linear,
                   ot.accounting_note,
                   ot.created_at,
                   COUNT(o.id)::int AS object_count,
                   COUNT(o.id) FILTER (WHERE o.is_active IS TRUE)::int AS active_object_count
            FROM object_types ot
            LEFT JOIN objects o ON o.object_type_id = ot.id
            GROUP BY ot.id, ot.code, ot.name, ot.sort_order, ot.is_active, ot.map_enabled,
                     ot.work_accounting_enabled, ot.material_accounting_enabled, ot.is_linear,
                     ot.accounting_note, ot.created_at
            ORDER BY ot.sort_order NULLS LAST, ot.name, ot.code
            """
        )
        object_types = [_json_safe(dict(row)) for row in cur.fetchall()]

        cur.execute(
            """
            SELECT cs.id::text AS id,
                   cs.code,
                   cs.name,
                   cs.sort_order,
                   cs.is_active,
                   cs.map_color,
                   cs.created_at,
                   COALESCE(jsonb_agg(
                     jsonb_build_object(
                       'boundary_kind', csv.boundary_kind,
                       'pk_start', csv.pk_start,
                       'pk_end', csv.pk_end,
                       'pk_raw_text', csv.pk_raw_text
                     )
                     ORDER BY csv.boundary_kind
                   ) FILTER (WHERE csv.id IS NOT NULL), '[]'::jsonb) AS boundaries
            FROM construction_sections cs
            LEFT JOIN construction_section_versions csv
              ON csv.section_id = cs.id
             AND csv.is_current IS TRUE
            GROUP BY cs.id, cs.code, cs.name, cs.sort_order, cs.is_active, cs.map_color, cs.created_at
            ORDER BY cs.sort_order NULLS LAST, cs.code
            """
        )
        sections = []
        for row in cur.fetchall():
            item = _json_safe(dict(row))
            for boundary in item.get("boundaries") or []:
                boundary["pk_label"] = _format_pk_range_human(
                    boundary.get("pk_start"),
                    boundary.get("pk_end"),
                    boundary.get("pk_raw_text"),
                )
            sections.append(item)

        cur.execute(
            """
            SELECT wt.id::text AS id,
                   wt.code,
                   wt.name,
                   wt.default_unit,
                   wt.created_at,
                   wt.analytics_tag,
                   wt.is_active,
                   wt.show_in_timeline,
                   wt.productivity_enabled,
                   MAX(wtur.rate) FILTER (WHERE wtur.region_code = 'novgorod' AND wtur.pile_length_m IS NULL)::numeric AS rate_novgorod,
                   MAX(wtur.rate) FILTER (WHERE wtur.region_code = 'tver' AND wtur.pile_length_m IS NULL)::numeric AS rate_tver,
                   COUNT(pwi.id)::int AS project_item_count,
                   (
                     COUNT(DISTINCT pwi.object_id)
                     + CASE WHEN EXISTS (
                         SELECT 1
                         FROM object_type_work_type_defaults d_pf
                         WHERE d_pf.work_type_id = wt.id
                           AND d_pf.source_group = 'PILE_FIELD_MAIN'
                           AND d_pf.is_active IS NOT FALSE
                       ) THEN (
                         SELECT COUNT(DISTINCT pf.field_code)
                         FROM pile_fields pf
                         WHERE pf.is_demo IS NOT TRUE
                           AND pf.field_type = 'main'
                       ) ELSE 0 END
                   )::int AS object_count
            FROM work_types wt
            LEFT JOIN project_work_items pwi ON pwi.work_type_id = wt.id
            LEFT JOIN work_type_unit_rates wtur ON wtur.work_type_id = wt.id AND wtur.pile_length_m IS NULL
            WHERE wt.code NOT LIKE 'POVSR\\_%%' ESCAPE '\\'
            GROUP BY wt.id, wt.code, wt.name, wt.default_unit, wt.analytics_tag, wt.is_active,
                     wt.show_in_timeline, wt.productivity_enabled
            ORDER BY wt.analytics_tag NULLS LAST, wt.name, wt.code
            """
        )
        work_types = [_json_safe(dict(row)) for row in cur.fetchall()]
        cur.execute(
            """
            SELECT work_type_id::text AS work_type_id,
                   pile_length_m::numeric AS pile_length_m,
                   MAX(rate) FILTER (WHERE region_code = 'novgorod')::numeric AS rate_novgorod,
                   MAX(rate) FILTER (WHERE region_code = 'tver')::numeric AS rate_tver
            FROM work_type_unit_rates
            WHERE pile_length_m IS NOT NULL
            GROUP BY work_type_id, pile_length_m
            ORDER BY pile_length_m
            """
        )
        saved_pile_rates_by_work: dict[str, dict[str, dict[str, Any]]] = {}
        for row in cur.fetchall():
            item = _json_safe(dict(row))
            length_key = _pile_rate_row_length_key(item)
            if not length_key:
                continue
            saved_pile_rates_by_work.setdefault(str(item.get("work_type_id")), {})[length_key] = item

        cur.execute(
            """
            SELECT wt.id::text AS work_type_id,
                   wt.code AS work_type_code,
                   pf.pile_type
            FROM work_types wt
            JOIN pile_fields pf
              ON (
                (wt.code = 'PILE_TRIAL' AND lower(COALESCE(pf.field_type, '')) IN ('test', 'trial'))
                OR (wt.code = 'PILE_MAIN' AND lower(COALESCE(pf.field_type, 'main')) NOT IN ('test', 'trial'))
              )
            WHERE wt.code IN ('PILE_TRIAL', 'PILE_MAIN')
              AND COALESCE(pf.is_demo, false) IS FALSE
              AND NULLIF(BTRIM(COALESCE(pf.pile_type, '')), '') IS NOT NULL
            UNION
            SELECT wt.id::text AS work_type_id,
                   wt.code AS work_type_code,
                   pf.pile_type
            FROM daily_work_items dwi
            JOIN work_types wt ON wt.id = dwi.work_type_id
            JOIN daily_reports dr ON dr.id = dwi.daily_report_id
            JOIN daily_work_item_segments seg ON seg.daily_work_item_id = dwi.id AND seg.is_demo IS NOT TRUE
            JOIN pile_fields pf ON pf.id = seg.pile_field_id
            WHERE wt.code IN ('PILE_TRIAL', 'PILE_MAIN')
              AND dwi.is_demo IS NOT TRUE
              AND COALESCE(dr.is_demo, false) IS FALSE
              AND COALESCE(dr.status, 'confirmed') IN ('confirmed', 'approved', 'pending_review')
              AND NULLIF(BTRIM(COALESCE(pf.pile_type, '')), '') IS NOT NULL
            """
        )
        current_pile_lengths_by_work: dict[str, set[Decimal]] = {}
        for row in cur.fetchall():
            length = _pile_length_from_type(row.get("pile_type"))
            if length is None:
                continue
            current_pile_lengths_by_work.setdefault(str(row.get("work_type_id")), set()).add(length)

        for item in work_types:
            work_id = str(item.get("id"))
            saved_rates = saved_pile_rates_by_work.get(work_id, {})
            if not _work_type_supports_pile_length_rates(item):
                item["pile_rates"] = []
                continue
            lengths = current_pile_lengths_by_work.get(work_id) or {
                length
                for length in (_pile_rate_length_value(row.get("pile_length_m")) for row in saved_rates.values())
                if length is not None
            }
            item["pile_rates"] = [
                _json_safe(_pile_rate_row_for_length(length, saved_rates))
                for length in sorted(lengths)
            ]

        cur.execute(
            f"""
            SELECT m.id::text AS id,
                   m.code,
                   m.name,
                   m.default_unit,
                   m.created_at,
                   COALESCE(mm.movement_count, 0)::int AS movement_count,
                   COALESCE(mm.section_count, 0)::int AS section_count,
                   COALESCE(mm.movement_volume, 0)::numeric AS movement_volume,
                   COALESCE(sp.stockpile_count, 0)::int AS stockpile_count
            FROM materials m
            LEFT JOIN LATERAL (
              SELECT COUNT(*)::int AS movement_count,
                     COUNT(DISTINCT mm_src.section_id)::int AS section_count,
                     SUM(mm_src.volume)::numeric AS movement_volume
              FROM material_movements mm_src
              LEFT JOIN daily_reports dr ON dr.id = mm_src.daily_report_id
              WHERE mm_src.material_id = m.id
                AND {_non_pile_control_report_condition('dr')}
            ) mm ON true
            LEFT JOIN LATERAL (
              SELECT COUNT(*)::int AS stockpile_count
              FROM stockpiles
              WHERE material_id = m.id
            ) sp ON true
            ORDER BY m.name, m.code
            """
        )
        materials = [_json_safe(dict(row)) for row in cur.fetchall()]

        cur.execute(
            """
            SELECT tr.id::text AS id,
                   tr.road_code AS code,
                   tr.road_name AS name,
                   tr.road_type,
                   tr.ad_start_pk,
                   tr.ad_end_pk,
                   tr.rail_start_pk,
                   tr.rail_end_pk,
                   tr.display_in_dashboard,
                   tr.can_translate_to_rail,
                   tr.comment,
                   tr.created_at,
                   tr.updated_at,
                   cs.id::text AS section_id,
                   cs.code AS section_code,
                   cs.name AS section_name,
                   COALESCE(status_stats.status_segment_count, 0)::int AS status_segment_count,
                   status_stats.latest_status_date,
                   COALESCE(correction_stats.correction_count, 0)::int AS correction_count,
                   correction_stats.latest_correction_date
            FROM temporary_roads tr
            LEFT JOIN construction_sections cs ON cs.id = tr.section_id
            LEFT JOIN LATERAL (
              SELECT COUNT(*)::int AS status_segment_count,
                     MAX(status_date) AS latest_status_date
              FROM temporary_road_status_segments s
              WHERE s.road_id = tr.id
                AND COALESCE(s.is_demo, false) IS FALSE
                AND COALESCE(s.review_tag, '') = ''
            ) status_stats ON true
            LEFT JOIN LATERAL (
              SELECT COUNT(*)::int AS correction_count,
                     MAX(effective_date) AS latest_correction_date
              FROM temporary_road_state_corrections c
              WHERE c.road_id = tr.id
            ) correction_stats ON true
            ORDER BY tr.created_at DESC NULLS LAST, tr.road_code
            """
        )
        temporary_roads = []
        for row in cur.fetchall():
            item = _json_safe(dict(row))
            item["ad_pk_label"] = _format_pk_range_human(row.get("ad_start_pk"), row.get("ad_end_pk"))
            item["rail_pk_label"] = _format_pk_range_human(row.get("rail_start_pk"), row.get("rail_end_pk"))
            item["section_code"] = item.get("section_code") or ""
            item["section_name"] = item.get("section_name") or ""
            temporary_roads.append(item)

        return {
            "summary": [
                {
                    "key": "objects",
                    "label": "Объекты",
                    "count": len(objects),
                    "active_count": sum(1 for item in objects if item.get("is_active")),
                },
                {
                    "key": "object_types",
                    "label": "Типы объектов",
                    "count": len(object_types),
                    "active_count": sum(1 for item in object_types if item.get("is_active")),
                },
                {
                    "key": "work_types",
                    "label": "Виды работ",
                    "count": len(work_types),
                    "active_count": sum(1 for item in work_types if item.get("is_active")),
                },
                {
                    "key": "materials",
                    "label": "Материалы",
                    "count": len(materials),
                    "active_count": len(materials),
                },
                {
                    "key": "temporary_roads",
                    "label": "Временные дороги",
                    "count": len(temporary_roads),
                    "active_count": sum(1 for item in temporary_roads if item.get("display_in_dashboard") is not False),
                },
            ],
            "objects": objects,
            "object_types": object_types,
            "sections": sections,
            "work_types": work_types,
            "materials": materials,
            "temporary_roads": temporary_roads,
        }


@router.get("/reference-export.xlsx")
def reference_export_xlsx(_user=Depends(current_user)):
    output = _build_reference_export_workbook()
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_attachment_headers("rd-reference-snapshot.xlsx", "Справочник для инструмента ввода РД.xlsx"),
    )


def _append_rd_intake_system_aliases(
    aliases: list[dict[str, Any]],
    *,
    work_types: list[dict[str, Any]],
    materials: list[dict[str, Any]],
) -> None:
    valid_codes = {
        "work_type": {str(row.get("code") or "") for row in work_types},
        "material": {str(row.get("code") or "") for row in materials},
    }
    seen = {(str(row.get("kind") or ""), str(row.get("alias_text") or "").casefold()) for row in aliases}
    for kind, canonical_code, alias_texts in RD_INTAKE_SYSTEM_ALIASES:
        if canonical_code not in valid_codes.get(kind, set()):
            continue
        for alias_text in alias_texts:
            key = (kind, alias_text.casefold())
            if key in seen:
                continue
            aliases.append({"kind": kind, "alias_text": alias_text, "canonical_code": canonical_code})
            seen.add(key)


def _build_reference_export_workbook() -> BytesIO:
    """Export reference snapshot for the offline RD intake Windows tool."""
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise HTTPException(status_code=500, detail="openpyxl is not installed") from exc

    output = BytesIO()
    wb = Workbook()
    wb.remove(wb.active)

    def append_sheet(name: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header) for header in headers])
        ws.freeze_panes = "A2"
        for column_cells in ws.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 60)
            ws.column_dimensions[column_cells[0].column_letter].width = width

    with _db_cursor() as (conn, cur):
        _ensure_work_type_analytics_schema(cur)
        cur.execute(
            """
            SELECT o.object_code AS code,
                   o.name AS name,
                   ot.code AS type,
                   ot.name AS type_name,
                   MIN(LEAST(os.pk_start, os.pk_end))::numeric AS pk_start,
                   MAX(GREATEST(os.pk_start, os.pk_end))::numeric AS pk_end,
                   string_agg(DISTINCT NULLIF(os.pk_raw_text, ''), '; ' ORDER BY NULLIF(os.pk_raw_text, '')) AS pk_raw_text
            FROM objects o
            LEFT JOIN object_types ot ON ot.id = o.object_type_id
            LEFT JOIN object_segments os ON os.object_id = o.id
            WHERE COALESCE(o.is_active, true) IS TRUE
            GROUP BY o.id, o.object_code, o.name, ot.code, ot.name
            ORDER BY o.name NULLS LAST, o.object_code
            """
        )
        objects = [_json_safe(dict(row)) for row in cur.fetchall()]

        cur.execute(
            """
            SELECT EXISTS (
              SELECT 1
              FROM information_schema.columns
              WHERE table_schema = 'public'
                AND table_name = 'work_types'
                AND column_name = 'work_group'
            ) AS exists
            """
        )
        work_group_exists = bool((cur.fetchone() or {}).get("exists"))
        analytics_group_expr = (
            "COALESCE(NULLIF(BTRIM(analytics_tag), ''), NULLIF(BTRIM(work_group), ''))"
            if work_group_exists
            else "NULLIF(BTRIM(analytics_tag), '')"
        )
        cur.execute(
            f"""
            SELECT code,
                   name,
                   default_unit,
                   analytics_tag,
                   {analytics_group_expr} AS analytics_group
            FROM work_types
            WHERE COALESCE(is_active, true) IS TRUE
              AND COALESCE(code, '') NOT ILIKE 'PIPE_ARRANGEMENT%'
              AND NOT (
                COALESCE(name, '') ILIKE '%Труба Вр. пр.%'
                OR COALESCE(name, '') ILIKE '%врем%водопропуск%труб%'
                OR COALESCE(name, '') ILIKE '%врем%труб%водопропуск%'
              )
            ORDER BY name NULLS LAST, code
            """
        )
        work_types = [_json_safe(dict(row)) for row in cur.fetchall()]

        cur.execute(
            """
            SELECT code, name, default_unit
            FROM materials
            ORDER BY name NULLS LAST, code
            """
        )
        materials = [_json_safe(dict(row)) for row in cur.fetchall()]

        aliases: list[dict[str, Any]] = []
        if _table_exists(cur, "work_type_aliases"):
            cur.execute(
                """
                SELECT kind, alias_text, canonical_code
                FROM work_type_aliases
                WHERE COALESCE(alias_text, '') <> ''
                  AND COALESCE(canonical_code, '') <> ''
                ORDER BY kind, alias_text
                """
            )
            aliases = [_json_safe(dict(row)) for row in cur.fetchall()]

    _append_rd_intake_system_aliases(aliases, work_types=work_types, materials=materials)

    append_sheet("objects", ["code", "name", "type", "type_name", "pk_start", "pk_end", "pk_raw_text"], objects)
    append_sheet("work_types", ["code", "name", "default_unit", "analytics_tag", "analytics_group"], work_types)
    append_sheet("materials", ["code", "name", "default_unit"], materials)
    append_sheet("aliases", ["kind", "alias_text", "canonical_code"], aliases)
    wb.save(output)
    output.seek(0)
    return output


def _rd_intake_portable_zip_candidates() -> list[Path]:
    root = Path(__file__).resolve().parents[1]
    return [
        Path("/app/public_templates/rd_intake_windows_portable.zip"),
        root / "frontend" / "public" / "templates" / "rd_intake_windows_portable.zip",
        root / "artifacts" / "rd_intake_windows_portable.zip",
    ]


def _rd_intake_snapshot_archive_name(names: list[str]) -> str:
    for name in names:
        normalized = name.replace("\\", "/")
        if normalized.endswith("/app.py"):
            return f"{normalized.rsplit('/', 1)[0]}/rd_reference_snapshot.xlsx"
    return "rd_reference_snapshot.xlsx"


@router.get("/rd-intake-portable.zip")
def rd_intake_portable_zip(_user=Depends(current_user)):
    base_zip = next((path for path in _rd_intake_portable_zip_candidates() if path.exists()), None)
    if not base_zip:
        raise HTTPException(status_code=500, detail="RD intake portable ZIP is not available on this server")

    snapshot_bytes = _build_reference_export_workbook().getvalue()
    output = BytesIO()
    with ZipFile(base_zip, "r") as source, ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
        names = source.namelist()
        snapshot_name = _rd_intake_snapshot_archive_name(names)
        skip_names = {
            snapshot_name,
            snapshot_name.rsplit("/", 1)[-1],
            "reference_snapshot.xlsx",
            "rd_reference_snapshot.xlsx",
        }
        for item in source.infolist():
            normalized = item.filename.replace("\\", "/")
            if normalized in skip_names or normalized.endswith("/reference_snapshot.xlsx") or normalized.endswith("/rd_reference_snapshot.xlsx"):
                continue
            target.writestr(item, source.read(item.filename))
        target.writestr(snapshot_name, snapshot_bytes)

    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/zip",
        headers=_attachment_headers("rd-intake-windows.zip", "Инструмент ввода РД для Windows.zip"),
    )


@router.get("/reference-relations")
def reference_relations(kind: str, code: str, _user=Depends(current_user)):
    normalized_kind = (kind or "").strip()
    normalized_code = (code or "").strip()
    if normalized_kind not in {"object", "work_type", "material", "object_type", "temporary_road"}:
        raise HTTPException(status_code=400, detail="Неизвестный тип справочника")
    if not normalized_code:
        raise HTTPException(status_code=400, detail="Код обязателен")

    with _db_cursor() as (conn, cur):
        _ensure_work_type_analytics_schema(cur)
        if normalized_kind == "temporary_road":
            cur.execute(
                """
                WITH target_road AS (
                  SELECT id
                  FROM temporary_roads
                  WHERE road_code = %s
                  LIMIT 1
                ),
                status_rows AS (
                  SELECT s.status_type,
                         COUNT(*)::int AS segment_count,
                         MIN(s.status_date) AS first_date,
                         MAX(s.status_date) AS latest_date,
                         SUM(ABS(COALESCE(s.road_pk_end, s.rail_pk_end, 0) - COALESCE(s.road_pk_start, s.rail_pk_start, 0)))::numeric AS length_m,
                         'отчетные статусы'::text AS source
                  FROM temporary_road_status_segments s
                  JOIN target_road tr ON tr.id = s.road_id
                  WHERE COALESCE(s.is_demo, false) IS FALSE
                    AND COALESCE(s.review_tag, '') = ''
                  GROUP BY s.status_type
                ),
                correction_rows AS (
                  SELECT c.status_type,
                         COUNT(*)::int AS segment_count,
                         MIN(c.effective_date) AS first_date,
                         MAX(c.effective_date) AS latest_date,
                         SUM(ABS(COALESCE(c.road_pk_end, c.rail_pk_end, 0) - COALESCE(c.road_pk_start, c.rail_pk_start, 0)))::numeric AS length_m,
                         'ручные корректировки'::text AS source
                  FROM temporary_road_state_corrections c
                  JOIN target_road tr ON tr.id = c.road_id
                  GROUP BY c.status_type
                )
                SELECT *
                FROM (
                  SELECT * FROM status_rows
                  UNION ALL
                  SELECT * FROM correction_rows
                ) rows
                ORDER BY latest_date DESC NULLS LAST, source, status_type
                """,
                (normalized_code,),
            )
            return {"kind": normalized_kind, "code": normalized_code, "rows": _json_safe([dict(r) for r in cur.fetchall()])}

        if normalized_kind == "object":
            cur.execute(
                f"""
                WITH target_object AS (
                  SELECT o.id, o.object_code, o.name, ot.code AS object_type_code, ot.name AS object_type_name
                  FROM objects o
                  JOIN object_types ot ON ot.id = o.object_type_id
                  WHERE o.object_code = %s
                  LIMIT 1
                ),
                project_rows AS (
                  SELECT pwi.work_type_id,
                         COALESCE(NULLIF(MAX(pwi.unit), ''), MAX(wt.default_unit)) AS project_unit,
                         SUM(COALESCE(pwi.project_volume, 0))::numeric AS project_volume,
                         COUNT(*)::int AS project_item_count
                  FROM project_work_items pwi
                  JOIN target_object obj ON obj.id = pwi.object_id
                  JOIN work_types wt ON wt.id = pwi.work_type_id
                  WHERE COALESCE(pwi.project_volume, 0) <> 0
                  GROUP BY pwi.work_type_id
                ),
                fact_rows AS (
                  SELECT dwi.work_type_id,
                         COALESCE(NULLIF(MAX(dwi.unit), ''), MAX(wt.default_unit)) AS fact_unit,
                         SUM(COALESCE(dwi.volume, 0))::numeric AS fact_volume,
                         COUNT(*)::int AS fact_row_count,
                         COUNT(DISTINCT dr.id)::int AS fact_report_count,
                         MIN(dwi.report_date) AS first_fact_date,
                         MAX(dwi.report_date) AS last_fact_date
                  FROM daily_work_items dwi
                  JOIN target_object obj ON obj.id = dwi.object_id
                  JOIN work_types wt ON wt.id = dwi.work_type_id
                  JOIN daily_reports dr ON dr.id = dwi.daily_report_id
                  WHERE dwi.is_demo IS NOT TRUE
                    AND COALESCE(dr.is_demo, false) IS FALSE
                    AND COALESCE(dr.status, 'confirmed') IN ('confirmed', 'approved', 'pending_review')
                    AND {_non_pile_control_report_condition('dr')}
                    AND COALESCE(dwi.volume, 0) <> 0
                  GROUP BY dwi.work_type_id
                ),
                merged AS (
                  SELECT COALESCE(p.work_type_id, f.work_type_id) AS work_type_id,
                         COALESCE(f.fact_unit, p.project_unit, wt.default_unit) AS unit,
                         COALESCE(p.project_volume, 0)::numeric AS project_volume,
                         COALESCE(p.project_item_count, 0)::int AS project_item_count,
                         COALESCE(f.fact_volume, 0)::numeric AS fact_volume,
                         COALESCE(f.fact_row_count, 0)::int AS fact_row_count,
                         COALESCE(f.fact_report_count, 0)::int AS fact_report_count,
                         f.first_fact_date,
                         f.last_fact_date
                  FROM project_rows p
                  FULL OUTER JOIN fact_rows f ON f.work_type_id = p.work_type_id
                  JOIN work_types wt ON wt.id = COALESCE(p.work_type_id, f.work_type_id)
                )
                SELECT obj.object_code,
                       obj.name AS object_name,
                       obj.object_type_code,
                       obj.object_type_name,
                       wt.code AS work_type_code,
                       wt.name AS work_type_name,
                       wt.analytics_tag,
                       m.unit,
                       m.project_volume,
                       m.fact_volume,
                       m.project_item_count,
                       m.fact_row_count,
                       m.fact_report_count,
                       m.first_fact_date,
                       m.last_fact_date,
                       (m.project_item_count + m.fact_row_count)::int AS row_count
                FROM merged m
                JOIN target_object obj ON true
                JOIN work_types wt ON wt.id = m.work_type_id
                WHERE COALESCE(m.project_volume, 0) <> 0
                   OR COALESCE(m.fact_volume, 0) <> 0
                ORDER BY wt.name, wt.code
                """,
                (normalized_code,),
            )
            return {"kind": normalized_kind, "code": normalized_code, "rows": _json_safe([dict(r) for r in cur.fetchall()])}

        if normalized_kind == "work_type":
            cur.execute(
                f"""
                WITH target_work AS (
                  SELECT id, code, name, default_unit
                  FROM work_types
                  WHERE code = %s
                  LIMIT 1
                ),
                project_rows AS (
                  SELECT pwi.object_id,
                         COALESCE(NULLIF(MAX(pwi.unit), ''), MAX(tw.default_unit)) AS project_unit,
                         SUM(COALESCE(pwi.project_volume, 0))::numeric AS project_volume,
                         COUNT(*)::int AS project_item_count
                  FROM project_work_items pwi
                  JOIN target_work tw ON tw.id = pwi.work_type_id
                  WHERE COALESCE(pwi.project_volume, 0) <> 0
                  GROUP BY pwi.object_id
                ),
                fact_rows AS (
                  SELECT dwi.object_id,
                         COALESCE(NULLIF(MAX(dwi.unit), ''), MAX(tw.default_unit)) AS fact_unit,
                         SUM(COALESCE(dwi.volume, 0))::numeric AS fact_volume,
                         COUNT(*)::int AS fact_row_count,
                         COUNT(DISTINCT dr.id)::int AS fact_report_count,
                         MIN(dwi.report_date) AS first_fact_date,
                         MAX(dwi.report_date) AS last_fact_date
                  FROM daily_work_items dwi
                  JOIN target_work tw ON tw.id = dwi.work_type_id
                  JOIN daily_reports dr ON dr.id = dwi.daily_report_id
                  WHERE dwi.is_demo IS NOT TRUE
                    AND COALESCE(dr.is_demo, false) IS FALSE
                    AND COALESCE(dr.status, 'confirmed') IN ('confirmed', 'approved', 'pending_review')
                    AND {_non_pile_control_report_condition('dr')}
                    AND COALESCE(dwi.volume, 0) <> 0
                  GROUP BY dwi.object_id
                ),
                merged AS (
                  SELECT COALESCE(p.object_id, f.object_id) AS object_id,
                         COALESCE(f.fact_unit, p.project_unit, tw.default_unit) AS unit,
                         COALESCE(p.project_volume, 0)::numeric AS project_volume,
                         COALESCE(p.project_item_count, 0)::int AS project_item_count,
                         COALESCE(f.fact_volume, 0)::numeric AS fact_volume,
                         COALESCE(f.fact_row_count, 0)::int AS fact_row_count,
                         COALESCE(f.fact_report_count, 0)::int AS fact_report_count,
                         f.first_fact_date,
                         f.last_fact_date
                  FROM target_work tw
                  LEFT JOIN project_rows p ON true
                  FULL OUTER JOIN fact_rows f ON f.object_id = p.object_id
                )
                SELECT o.object_code,
                       o.name AS object_name,
                       ot.code AS object_type_code,
                       ot.name AS object_type_name,
                       tw.code AS work_type_code,
                       tw.name AS work_type_name,
                       m.unit,
                       m.project_volume,
                       m.fact_volume,
                       m.project_item_count,
                       m.fact_row_count,
                       m.fact_report_count,
                       m.first_fact_date,
                       m.last_fact_date,
                       (m.project_item_count + m.fact_row_count)::int AS row_count
                FROM merged m
                JOIN target_work tw ON true
                JOIN objects o ON o.id = m.object_id
                JOIN object_types ot ON ot.id = o.object_type_id
                WHERE COALESCE(m.project_volume, 0) <> 0
                   OR COALESCE(m.fact_volume, 0) <> 0
                ORDER BY ot.name, o.name, o.object_code
                """,
                (normalized_code,),
            )
            return {"kind": normalized_kind, "code": normalized_code, "rows": _json_safe([dict(r) for r in cur.fetchall()])}

        if normalized_kind == "object_type":
            cur.execute(
                """
                SELECT o.object_code,
                       o.name AS object_name,
                       o.is_active,
                       seg.pk_start,
                       seg.pk_end,
                       seg.pk_raw_text,
                       COUNT(pwi.id)::int AS project_item_count
                FROM object_types ot
                JOIN objects o ON o.object_type_id = ot.id
                LEFT JOIN LATERAL (
                  SELECT os.pk_start, os.pk_end, os.pk_raw_text
                  FROM object_segments os
                  WHERE os.object_id = o.id
                  ORDER BY os.pk_start NULLS LAST, os.created_at NULLS LAST
                  LIMIT 1
                ) seg ON true
                LEFT JOIN project_work_items pwi ON pwi.object_id = o.id AND COALESCE(pwi.project_volume, 0) <> 0
                WHERE ot.code = %s
                GROUP BY o.object_code, o.name, o.is_active, seg.pk_start, seg.pk_end, seg.pk_raw_text
                ORDER BY o.name, o.object_code
                """,
                (normalized_code,),
            )
            rows = []
            for r in cur.fetchall():
                item = _json_safe(dict(r))
                item["pk_label"] = _format_pk_range_human(r["pk_start"], r["pk_end"], r["pk_raw_text"])
                rows.append(item)
            return {"kind": normalized_kind, "code": normalized_code, "rows": rows}

        cur.execute(
            f"""
            WITH movement_rows AS (
              SELECT COALESCE(cs.name, 'Без участка') AS section_name,
                     COALESCE(cs.code, '—') AS section_code,
                     SUM(mm.volume)::numeric AS movement_volume,
                     COUNT(mm.id)::int AS movement_count
              FROM materials m
              LEFT JOIN material_movements mm ON mm.material_id = m.id
              LEFT JOIN daily_reports dr ON dr.id = mm.daily_report_id
              LEFT JOIN construction_sections cs ON cs.id = mm.section_id
              WHERE m.code = %s
                AND {_non_pile_control_report_condition('dr')}
              GROUP BY cs.name, cs.code
            ),
            stockpile_rows AS (
              SELECT COUNT(sp.id)::int AS stockpile_count
              FROM materials m
              LEFT JOIN stockpiles sp ON sp.material_id = m.id
              WHERE m.code = %s
            )
            SELECT movement_rows.section_name,
                   movement_rows.section_code,
                   COALESCE(movement_volume, 0)::numeric AS movement_volume,
                   movement_count,
                   stockpile_rows.stockpile_count
            FROM movement_rows
            CROSS JOIN stockpile_rows
            ORDER BY section_code NULLS LAST
            """,
            (normalized_code, normalized_code),
        )
        return {"kind": normalized_kind, "code": normalized_code, "rows": _json_safe([dict(r) for r in cur.fetchall()])}


@router.get("/reference-facts")
def reference_facts(
    kind: str,
    code: str,
    limit: int = Query(1000, ge=1, le=3000),
    _user=Depends(current_user),
):
    normalized_kind = (kind or "").strip()
    normalized_code = (code or "").strip()
    if normalized_kind not in {"object", "work_type", "material", "object_type", "temporary_road"}:
        raise HTTPException(status_code=400, detail="Неизвестный тип справочника")
    if not normalized_code:
        raise HTTPException(status_code=400, detail="Код обязателен")

    with _db_cursor() as (conn, cur):
        _ensure_work_type_analytics_schema(cur)
        if normalized_kind == "temporary_road":
            return {"kind": normalized_kind, "code": normalized_code, "rows": []}

        work_where = ""
        work_params: list[Any] = []
        if normalized_kind == "object":
            work_where = "AND o.object_code = %s"
            work_params.append(normalized_code)
        elif normalized_kind == "work_type":
            work_where = "AND wt.code = %s"
            work_params.append(normalized_code)
        elif normalized_kind == "object_type":
            work_where = "AND ot.code = %s"
            work_params.append(normalized_code)

        work_rows: list[dict[str, Any]] = []
        if normalized_kind != "material":
            cur.execute(
                f"""
                WITH fact_rows AS (
                  SELECT 'work' AS fact_kind,
                         dwi.id::text AS fact_id,
                         dr.id::text AS report_id,
                         dwi.report_date,
                         dwi.shift,
                         dr.status AS report_status,
                         dr.source_type,
                         dr.source_reference,
                         dr.uploaded_by_username,
                         cs.code AS section_code,
                         cs.name AS section_name,
                         o.object_code,
                         o.name AS object_name,
                         ot.code AS object_type_code,
                         ot.name AS object_type_name,
                         wt.code AS work_type_code,
                         wt.name AS work_type_name,
                         wt.analytics_tag,
                         NULL::text AS material_code,
                         NULL::text AS material_name,
                         COALESCE(NULLIF(dwi.unit, ''), wt.default_unit) AS unit,
                         CASE
                           WHEN seg.id IS NULL THEN COALESCE(dwi.volume, 0)
                           WHEN seg.volume_segment IS NOT NULL THEN seg.volume_segment
                           WHEN seg_count.cnt > 0 THEN COALESCE(dwi.volume, 0) / seg_count.cnt
                           ELSE COALESCE(dwi.volume, 0)
                         END::numeric AS volume,
                         seg.pk_start,
                         seg.pk_end,
                         seg.pk_raw_text,
                         dwi.labor_source_type,
                         dwi.contractor_name,
                         dwi.comment,
                         dwi.created_at
                  FROM daily_work_items dwi
                  JOIN daily_reports dr ON dr.id = dwi.daily_report_id
                  JOIN objects o ON o.id = dwi.object_id
                  JOIN object_types ot ON ot.id = o.object_type_id
                  JOIN work_types wt ON wt.id = dwi.work_type_id
                  LEFT JOIN construction_sections cs ON cs.id = dwi.section_id
                  LEFT JOIN LATERAL (
                    SELECT COUNT(*)::numeric AS cnt
                    FROM daily_work_item_segments s
                    WHERE s.daily_work_item_id = dwi.id
                      AND s.is_demo IS NOT TRUE
                  ) seg_count ON true
                  LEFT JOIN daily_work_item_segments seg
                    ON seg.daily_work_item_id = dwi.id
                   AND seg.is_demo IS NOT TRUE
                  WHERE dwi.is_demo IS NOT TRUE
                    AND COALESCE(dr.is_demo, false) IS FALSE
                    AND COALESCE(dr.status, 'confirmed') IN ('confirmed', 'approved', 'pending_review')
                    AND {_non_pile_control_report_condition('dr')}
                    AND COALESCE(dwi.volume, 0) <> 0
                    {work_where}
                )
                SELECT *
                FROM fact_rows
                WHERE COALESCE(volume, 0) <> 0
                ORDER BY report_date DESC NULLS LAST,
                         shift NULLS LAST,
                         object_name,
                         work_type_name,
                         pk_start NULLS LAST,
                         created_at DESC NULLS LAST
                LIMIT %s
                """,
                tuple(work_params + [limit]),
            )
            for row in cur.fetchall():
                item = _json_safe(dict(row))
                item["pk_label"] = _format_pk_range_human(row.get("pk_start"), row.get("pk_end"), row.get("pk_raw_text"))
                work_rows.append(item)

        movement_where = ""
        movement_params: list[Any] = []
        if normalized_kind == "object":
            movement_where = "AND (from_o.object_code = %s OR to_o.object_code = %s)"
            movement_params.extend([normalized_code, normalized_code])
        elif normalized_kind == "material":
            movement_where = "AND m.code = %s"
            movement_params.append(normalized_code)

        movement_rows: list[dict[str, Any]] = []
        if normalized_kind in {"object", "material"}:
            cur.execute(
                f"""
                SELECT 'material_movement' AS fact_kind,
                       mm.id::text AS fact_id,
                       dr.id::text AS report_id,
                       mm.report_date,
                       mm.shift,
                       dr.status AS report_status,
                       dr.source_type,
                       dr.source_reference,
                       dr.uploaded_by_username,
                       cs.code AS section_code,
                       cs.name AS section_name,
                       COALESCE(to_o.object_code, from_o.object_code) AS object_code,
                       CONCAT_WS(' → ', NULLIF(from_o.name, ''), NULLIF(to_o.name, '')) AS object_name,
                       NULL::text AS object_type_code,
                       NULL::text AS object_type_name,
                       NULL::text AS work_type_code,
                       CASE
                         WHEN mm.movement_type = 'out' THEN 'Вывоз материала'
                         WHEN mm.movement_type = 'in' THEN 'Завоз материала'
                         ELSE 'Перевозка материала'
                       END AS work_type_name,
                       NULL::text AS analytics_tag,
                       m.code AS material_code,
                       m.name AS material_name,
                       COALESCE(NULLIF(mm.unit, ''), m.default_unit) AS unit,
                       mm.volume::numeric AS volume,
                       NULL::numeric AS pk_start,
                       NULL::numeric AS pk_end,
                       NULL::text AS pk_raw_text,
                       mm.labor_source_type,
                       mm.contractor_name,
                       mm.comment,
                       mm.created_at
                FROM material_movements mm
                JOIN materials m ON m.id = mm.material_id
                LEFT JOIN daily_reports dr ON dr.id = mm.daily_report_id
                LEFT JOIN construction_sections cs ON cs.id = mm.section_id
                LEFT JOIN objects from_o ON from_o.id = mm.from_object_id
                LEFT JOIN objects to_o ON to_o.id = mm.to_object_id
                WHERE COALESCE(mm.volume, 0) <> 0
                  AND COALESCE(mm.is_demo, false) IS FALSE
                  AND COALESCE(dr.is_demo, false) IS FALSE
                  AND COALESCE(dr.status, 'confirmed') IN ('confirmed', 'approved', 'pending_review')
                  AND {_non_pile_control_report_condition('dr')}
                  {movement_where}
                ORDER BY mm.report_date DESC NULLS LAST,
                         mm.shift NULLS LAST,
                         m.name,
                         object_name,
                         mm.created_at DESC NULLS LAST
                LIMIT %s
                """,
                tuple(movement_params + [limit]),
            )
            for row in cur.fetchall():
                item = _json_safe(dict(row))
                item["pk_label"] = ""
                movement_rows.append(item)

        rows = sorted(
            work_rows + movement_rows,
            key=lambda item: (
                str(item.get("report_date") or ""),
                str(item.get("shift") or ""),
                str(item.get("object_name") or ""),
                str(item.get("work_type_name") or ""),
                str(item.get("pk_start") or ""),
            ),
            reverse=True,
        )[:limit]
        return {"kind": normalized_kind, "code": normalized_code, "rows": rows}


@router.get("/reference-relations/work-type-object-detail")
def work_type_object_detail(work_code: str, object_code: str, _user=Depends(current_user)):
    normalized_work_code = (work_code or "").strip()
    normalized_object_code = (object_code or "").strip()
    if not normalized_work_code or not normalized_object_code:
        raise HTTPException(status_code=400, detail="Коды работы и объекта обязательны")

    with _db_cursor() as (conn, cur):
        _ensure_work_type_analytics_schema(cur)
        cur.execute(
            f"""
            WITH target_work AS (
              SELECT id, code, name, default_unit, analytics_tag
              FROM work_types
              WHERE code = %s
              LIMIT 1
            ),
            target_object AS (
              SELECT o.id, o.object_code, o.name, ot.name AS object_type_name
              FROM objects o
              JOIN object_types ot ON ot.id = o.object_type_id
              WHERE o.object_code = %s
              LIMIT 1
            ),
            fact_rows AS (
              SELECT dwi.id::text AS daily_work_item_id,
                     dr.id::text AS report_id,
                     dwi.report_date,
                     dwi.shift,
                     dr.status AS report_status,
                     dr.source_type,
                     dr.source_reference,
                     dr.uploaded_by_username,
                     cs.code AS section_code,
                     cs.name AS section_name,
                     obj.object_code,
                     obj.name AS object_name,
                     obj.object_type_name,
                     tw.code AS work_type_code,
                     tw.name AS work_type_name,
                     tw.analytics_tag,
                     dwi.work_name_raw,
                     COALESCE(NULLIF(dwi.unit, ''), tw.default_unit) AS unit,
                     CASE
                       WHEN seg.id IS NULL THEN COALESCE(dwi.volume, 0)
                       WHEN seg.volume_segment IS NOT NULL THEN seg.volume_segment
                       WHEN seg_count.cnt > 0 THEN COALESCE(dwi.volume, 0) / seg_count.cnt
                       ELSE COALESCE(dwi.volume, 0)
                     END::numeric AS volume,
                     seg.pk_start,
                     seg.pk_end,
                     seg.pk_raw_text,
                     dwi.labor_source_type,
                     dwi.contractor_name,
                     dwi.comment,
                     dwi.created_at
              FROM daily_work_items dwi
              JOIN target_work tw ON tw.id = dwi.work_type_id
              JOIN target_object obj ON obj.id = dwi.object_id
              JOIN daily_reports dr ON dr.id = dwi.daily_report_id
              LEFT JOIN construction_sections cs ON cs.id = dwi.section_id
              LEFT JOIN LATERAL (
                SELECT COUNT(*)::numeric AS cnt
                FROM daily_work_item_segments s
                WHERE s.daily_work_item_id = dwi.id
                  AND s.is_demo IS NOT TRUE
              ) seg_count ON true
              LEFT JOIN daily_work_item_segments seg
                ON seg.daily_work_item_id = dwi.id
               AND seg.is_demo IS NOT TRUE
              WHERE dwi.is_demo IS NOT TRUE
                AND COALESCE(dr.is_demo, false) IS FALSE
                AND COALESCE(dr.status, 'confirmed') IN ('confirmed', 'approved', 'pending_review')
                AND {_non_pile_control_report_condition('dr')}
                AND COALESCE(dwi.volume, 0) <> 0
            )
            SELECT *
            FROM fact_rows
            ORDER BY report_date DESC, shift, created_at DESC, pk_start NULLS LAST
            LIMIT 500
            """,
            (normalized_work_code, normalized_object_code),
        )
        rows = []
        for row in cur.fetchall():
            item = _json_safe(dict(row))
            item["pk_label"] = _format_pk_range_human(row.get("pk_start"), row.get("pk_end"), row.get("pk_raw_text"))
            rows.append(item)
        return {
            "work_code": normalized_work_code,
            "object_code": normalized_object_code,
            "rows": rows,
        }


@router.patch("/reference-names")
def update_reference_name(request: ReferenceNameUpdate, _admin=Depends(require_admin)):
    spec = REFERENCE_NAME_TABLES.get(request.kind)
    if not spec:
        raise HTTPException(status_code=400, detail="Unknown reference kind")
    name = request.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Name is required")
    with _db_cursor() as (conn, cur):
        if not _table_exists(cur, spec["table"]):
            raise HTTPException(status_code=404, detail="Reference table does not exist")
        before_query = sql.SQL("SELECT * FROM {table} WHERE {code} = %s LIMIT 1").format(
            table=sql.Identifier(spec["table"]),
            code=sql.Identifier(spec["code_column"]),
        )
        cur.execute(before_query, (request.code,))
        before_row = cur.fetchone()
        if not before_row:
            raise HTTPException(status_code=404, detail="Reference row was not found")
        before = _json_safe(dict(before_row))
        update_query = sql.SQL(
            "UPDATE {table} SET {name} = %s WHERE {code} = %s RETURNING *"
        ).format(
            table=sql.Identifier(spec["table"]),
            name=sql.Identifier(spec["name_column"]),
            code=sql.Identifier(spec["code_column"]),
        )
        cur.execute(update_query, (name, request.code))
        after = _json_safe(dict(cur.fetchone()))
        _ensure_audit_table(cur)
        cur.execute(
            """
            INSERT INTO public.db_change_audit
              (table_name, action, pk_json, before_json, after_json, changed_by)
            VALUES (%s, 'update', %s, %s, %s, %s)
            """,
            (
                spec["table"],
                psycopg2.extras.Json({spec["code_column"]: request.code}),
                psycopg2.extras.Json(before),
                psycopg2.extras.Json(after),
                request.changed_by or "dashboard",
            ),
        )
        conn.commit()
        return {"ok": True, "row": {
            "kind": request.kind,
            "kind_label": spec["label"],
            "code": after[spec["code_column"]],
            "name": after[spec["name_column"]],
            "data": after,
            "table": spec["table"],
            "code_column": spec["code_column"],
            "name_column": spec["name_column"],
        }}


@router.patch("/reference-details")
def update_reference_detail(request: ReferenceDetailUpdate, _admin=Depends(require_admin)):
    spec = REFERENCE_NAME_TABLES.get(request.kind)
    if not spec:
        raise HTTPException(status_code=400, detail="Неизвестный справочник")
    allowed = set(REFERENCE_DETAIL_FIELDS.get(request.kind, set()))
    dynamic_pile_rate_fields = {key for key in (request.values or {}) if request.kind == "work_type" and _pile_rate_field_parts(key)}
    allowed.update(dynamic_pile_rate_fields)
    values = {k: v for k, v in (request.values or {}).items() if k in allowed}
    unknown = sorted(set(request.values or {}) - allowed)
    if unknown:
        raise HTTPException(status_code=400, detail=f"Недоступные поля: {', '.join(unknown)}")
    if not values:
        raise HTTPException(status_code=400, detail="Нет данных для сохранения")

    with _db_cursor() as (conn, cur):
        _ensure_audit_table(cur)
        if request.kind == "work_type":
            _ensure_work_type_analytics_schema(cur)
            _ensure_work_type_unit_rates_table(cur)
        table = spec["table"]
        code_col = spec["code_column"]
        cur.execute(
            sql.SQL("SELECT * FROM {table} WHERE {code_col} = %s LIMIT 1").format(
                table=sql.Identifier(table),
                code_col=sql.Identifier(code_col),
            ),
            (request.code,),
        )
        before_row = cur.fetchone()
        if not before_row:
            raise HTTPException(status_code=404, detail="Строка справочника не найдена")
        before = _json_safe(dict(before_row))
        if request.kind == "work_type" and dynamic_pile_rate_fields and not _work_type_supports_pile_length_rates(before):
            raise HTTPException(
                status_code=400,
                detail="Расценки забивки свай по длине доступны только для работ PILE_TRIAL и PILE_MAIN",
            )

        update_values = dict(values)
        table_columns = {col["column_name"] for col in _columns(cur, table)}
        virtual_fields = {
            "object": {"object_type_code", "pk_start", "pk_end", "pk_raw_text"},
            "constructive": {"object_code"},
            "work_type": {"rate_novgorod", "rate_tver", *dynamic_pile_rate_fields},
            "temporary_road": {"name", "section_code"},
        }.get(request.kind, set())
        for key in list(update_values):
            if key not in table_columns and key not in virtual_fields:
                update_values.pop(key)

        if request.kind == "object":
            object_id = before["id"]
            segment_fields = {k: update_values.pop(k) for k in list(update_values) if k in {"pk_start", "pk_end", "pk_raw_text"}}
            if "object_type_code" in update_values:
                type_code = str(update_values.pop("object_type_code") or "").strip()
                cur.execute("SELECT id FROM object_types WHERE code = %s AND is_active IS NOT FALSE LIMIT 1", (type_code,))
                type_row = cur.fetchone()
                if not type_row:
                    raise HTTPException(status_code=400, detail=f"Тип объекта не найден: {type_code}")
                update_values["object_type_id"] = type_row["id"]
            if update_values:
                set_sql = sql.SQL(", ").join(
                    sql.SQL("{} = %s").format(sql.Identifier(col)) for col in update_values
                )
                cur.execute(
                    sql.SQL("UPDATE objects SET {set_sql} WHERE object_code = %s RETURNING *").format(set_sql=set_sql),
                    list(update_values.values()) + [request.code],
                )
            if segment_fields:
                pk_start = _parse_pk_input(segment_fields.get("pk_start"))
                pk_end = _parse_pk_input(segment_fields.get("pk_end")) if segment_fields.get("pk_end") not in (None, "") else pk_start
                pk_raw = str(segment_fields.get("pk_raw_text") or "").strip() or (
                    f"{_format_pk_human(pk_start)} - {_format_pk_human(pk_end)}" if pk_start is not None and pk_end is not None and pk_start != pk_end
                    else _format_pk_human(pk_start)
                )
                if pk_start is None:
                    raise HTTPException(status_code=400, detail="Для сегмента объекта нужен пикетаж начала")
                cur.execute("SELECT id FROM object_segments WHERE object_id = %s ORDER BY pk_start NULLS LAST LIMIT 1", (object_id,))
                seg = cur.fetchone()
                if seg:
                    cur.execute(
                        """
                        UPDATE object_segments
                        SET pk_start = %s, pk_end = %s, pk_raw_text = %s
                        WHERE id = %s
                        """,
                        (pk_start, pk_end, pk_raw, seg["id"]),
                    )
                else:
                    cur.execute(
                        """
                        INSERT INTO object_segments (object_id, pk_start, pk_end, pk_raw_text, comment)
                        VALUES (%s, %s, %s, %s, %s)
                        """,
                        (object_id, pk_start, pk_end, pk_raw, "created from DB reference card"),
                    )

        elif request.kind == "constructive":
            if "object_code" in update_values:
                object_code = str(update_values.pop("object_code") or "").strip()
                if object_code:
                    cur.execute("SELECT id FROM objects WHERE object_code = %s LIMIT 1", (object_code,))
                    object_row = cur.fetchone()
                    if not object_row:
                        raise HTTPException(status_code=400, detail=f"Объект не найден: {object_code}")
                    update_values["object_id"] = object_row["id"]
                else:
                    update_values["object_id"] = None
            if update_values:
                set_sql = sql.SQL(", ").join(
                    sql.SQL("{} = %s").format(sql.Identifier(col)) for col in update_values
                )
                cur.execute(
                    sql.SQL("UPDATE constructives SET {set_sql} WHERE code = %s RETURNING *").format(set_sql=set_sql),
                    list(update_values.values()) + [request.code],
                )

        elif request.kind == "work_type":
            rate_updates = {k: update_values.pop(k) for k in list(update_values) if k in {"rate_novgorod", "rate_tver"} or k in dynamic_pile_rate_fields}
            rate_unit = update_values.get("default_unit") or before.get("default_unit")
            if "rate_novgorod" in rate_updates:
                _upsert_work_type_rate(cur, before["id"], "novgorod", rate_updates["rate_novgorod"], rate_unit, request.changed_by)
            if "rate_tver" in rate_updates:
                _upsert_work_type_rate(cur, before["id"], "tver", rate_updates["rate_tver"], rate_unit, request.changed_by)
            for field_name in sorted(dynamic_pile_rate_fields):
                if field_name not in rate_updates:
                    continue
                field_parts = _pile_rate_field_parts(field_name)
                if field_parts is None:
                    continue
                pile_length, region_code = field_parts
                _upsert_work_type_rate(cur, before["id"], region_code, rate_updates[field_name], rate_unit, request.changed_by, pile_length)
            if "analytics_tag" in update_values:
                tag = str(update_values.get("analytics_tag") or "").strip()
                update_values["analytics_tag"] = tag or None
                if tag:
                    cur.execute(
                        """
                        INSERT INTO work_analytics_tags (name)
                        VALUES (%s)
                        ON CONFLICT (name) DO UPDATE SET is_active = true
                        """,
                        (tag,),
                    )
            if update_values:
                set_sql = sql.SQL(", ").join(
                    sql.SQL("{} = %s").format(sql.Identifier(col)) for col in update_values
                )
                cur.execute(
                    sql.SQL("UPDATE work_types SET {set_sql} WHERE code = %s RETURNING *").format(set_sql=set_sql),
                    list(update_values.values()) + [request.code],
                )
                if "analytics_tag" in update_values:
                    _sync_work_type_fact_analytics_tags(cur, request.code)

        elif request.kind == "temporary_road":
            if "name" in update_values:
                road_name = str(update_values.pop("name") or "").strip()
                if not road_name:
                    raise HTTPException(status_code=400, detail="Название временной дороги обязательно")
                update_values["road_name"] = road_name
            if "section_code" in update_values:
                section_code = str(update_values.pop("section_code") or "").strip()
                if section_code:
                    cur.execute(
                        "SELECT id FROM construction_sections WHERE code = %s AND is_active IS NOT FALSE LIMIT 1",
                        (section_code,),
                    )
                    section_row = cur.fetchone()
                    if not section_row:
                        raise HTTPException(status_code=400, detail=f"Участок не найден: {section_code}")
                    update_values["section_id"] = section_row["id"]
                else:
                    update_values["section_id"] = None

            ad_pair = _reference_pk_pair(update_values, "ad_start_pk", "ad_end_pk", "АД ПК")
            if ad_pair is not None:
                update_values["ad_start_pk"], update_values["ad_end_pk"] = ad_pair
            rail_pair = _reference_pk_pair(update_values, "rail_start_pk", "rail_end_pk", "ВСЖМ ПК")
            if rail_pair is not None:
                update_values["rail_start_pk"], update_values["rail_end_pk"] = rail_pair
            if "can_translate_to_rail" in update_values and update_values.get("can_translate_to_rail"):
                ad_ready = update_values.get("ad_start_pk", before.get("ad_start_pk")) is not None and update_values.get("ad_end_pk", before.get("ad_end_pk")) is not None
                rail_ready = update_values.get("rail_start_pk", before.get("rail_start_pk")) is not None and update_values.get("rail_end_pk", before.get("rail_end_pk")) is not None
                if not (ad_ready and rail_ready):
                    raise HTTPException(status_code=400, detail="Перевод AD↔ВСЖМ можно включить только при заполненных AD и ВСЖМ пикетажах")
            if update_values:
                set_sql = sql.SQL(", ").join(
                    sql.SQL("{} = %s").format(sql.Identifier(col)) for col in update_values
                )
                cur.execute(
                    sql.SQL("UPDATE temporary_roads SET {set_sql}, updated_at = now() WHERE road_code = %s RETURNING *").format(set_sql=set_sql),
                    list(update_values.values()) + [request.code],
                )

        else:
            if update_values:
                set_sql = sql.SQL(", ").join(
                    sql.SQL("{} = %s").format(sql.Identifier(col)) for col in update_values
                )
                cur.execute(
                    sql.SQL("UPDATE {table} SET {set_sql} WHERE {code_col} = %s RETURNING *").format(
                        table=sql.Identifier(table),
                        set_sql=set_sql,
                        code_col=sql.Identifier(code_col),
                    ),
                    list(update_values.values()) + [request.code],
                )

        cur.execute(
            sql.SQL("SELECT * FROM {table} WHERE {code_col} = %s LIMIT 1").format(
                table=sql.Identifier(table),
                code_col=sql.Identifier(code_col),
            ),
            (request.code,),
        )
        after = _json_safe(dict(cur.fetchone()))
        cur.execute(
            """
            INSERT INTO public.db_change_audit
              (table_name, action, pk_json, before_json, after_json, changed_by)
            VALUES (%s, 'update', %s, %s, %s, %s)
            """,
            (
                table,
                psycopg2.extras.Json({code_col: request.code}),
                psycopg2.extras.Json(before),
                psycopg2.extras.Json(after),
                request.changed_by or "dashboard",
            ),
        )
        conn.commit()
        return {"ok": True, "row": after}


@router.get("/tables")
def list_tables(_user=Depends(current_user)):
    with _db_cursor() as (conn, cur):
        rows = []
        for table, label in EXPOSED_TABLES.items():
            if not _table_exists(cur, table):
                continue
            columns = {col["column_name"] for col in _columns(cur, table)}
            joins, where = _database_table_filter_parts(table, columns)
            cur.execute(
                sql.SQL("SELECT COUNT(*) AS cnt FROM {} t {} {}").format(
                    sql.Identifier(table),
                    joins,
                    where,
                )
            )
            rows.append({
                "name": table,
                "label": label,
                "rows": int(cur.fetchone()["cnt"]),
                "editable": table in EDITABLE_TABLES,
            })
        return {"tables": rows}


@router.get("/tables/{table}/schema")
def table_schema(table: str, _user=Depends(current_user)):
    _allowed_table(table)
    with _db_cursor() as (conn, cur):
        if not _table_exists(cur, table):
            raise HTTPException(status_code=404, detail="Table does not exist")
        return {
            "table": table,
            "label": EXPOSED_TABLES[table],
            "primary_key": _primary_key(cur, table),
            "columns": _columns(cur, table),
            "foreign_keys": _foreign_keys(cur, table),
            "incoming_refs": _incoming_refs(cur, table),
            "indexes": _indexes(cur, table),
        }


@router.get("/tables/{table}/rows")
def table_rows(
    table: str,
    limit: int = Query(default=100, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
    _user=Depends(current_user),
):
    _allowed_table(table)
    if table == "temporary_roads":
        ensure_temp_road_catalog_rows()
    with _db_cursor() as (conn, cur):
        if not _table_exists(cur, table):
            raise HTTPException(status_code=404, detail="Table does not exist")
        pk_cols = _primary_key(cur, table)
        columns = [col["column_name"] for col in _columns(cur, table)]
        order_cols = pk_cols or columns[:1]
        order_sql = sql.SQL(", ").join(sql.Identifier("t", col) for col in order_cols)
        joins, where = _database_table_filter_parts(table, set(columns))
        query = sql.SQL("SELECT t.* FROM {} t {} {} ORDER BY {} LIMIT %s OFFSET %s").format(
            sql.Identifier(table),
            joins,
            where,
            order_sql,
        )
        cur.execute(query, (limit, offset))
        return {"rows": [_json_safe(dict(row)) for row in cur.fetchall()], "limit": limit, "offset": offset}


@router.get("/tables/{table}/columns/{column}/suggestions")
def column_suggestions(
    table: str,
    column: str,
    section_id: str | None = None,
    limit: int = Query(default=25, ge=1, le=100),
    _user=Depends(current_user),
):
    _allowed_table(table)
    with _db_cursor() as (conn, cur):
        if not _table_exists(cur, table):
            raise HTTPException(status_code=404, detail="Table does not exist")

        columns = _columns(cur, table)
        column_names = {col["column_name"] for col in columns}
        if column not in column_names:
            raise HTTPException(status_code=400, detail="Column does not exist")

        joins: list[sql.SQL] = []
        params: list[Any] = []
        where_parts = [
            sql.SQL("t.{} IS NOT NULL").format(sql.Identifier(column)),
            sql.SQL("t.{}::text <> ''").format(sql.Identifier(column)),
        ]
        if table == "work_types":
            where_parts.append(sql.SQL("t.code NOT LIKE 'POVSR\\_%%' ESCAPE '\\'"))
        elif table == "work_type_aliases":
            where_parts.append(sql.SQL("t.kind <> 'work_type' OR t.canonical_code NOT LIKE 'POVSR\\_%%' ESCAPE '\\'"))
        label_expr: sql.SQL | sql.Composed = sql.SQL("NULL")
        joined_report_scope = False

        if table == "daily_reports":
            where_parts.append(sql.SQL(_non_pile_control_report_condition("t")))
        elif "daily_report_id" in column_names:
            joins.append(sql.SQL("LEFT JOIN daily_reports dr_scope ON t.daily_report_id = dr_scope.id"))
            where_parts.append(sql.SQL(_non_pile_control_report_condition("dr_scope")))
            joined_report_scope = True

        fk = _fk_for_column(cur, table, column)
        if fk and _table_exists(cur, fk["target_table"]):
            label_expr = _label_expression(cur, fk["target_table"])
            joins.append(
                sql.SQL("LEFT JOIN {} d ON t.{} = d.{}").format(
                    sql.Identifier(fk["target_table"]),
                    sql.Identifier(column),
                    sql.Identifier(fk["target_column"]),
                )
            )

        if section_id:
            if "section_id" in column_names:
                where_parts.append(sql.SQL("t.section_id = %s"))
                params.append(section_id)
            elif "daily_report_id" in column_names:
                if not joined_report_scope:
                    joins.append(sql.SQL("LEFT JOIN daily_reports dr_scope ON t.daily_report_id = dr_scope.id"))
                where_parts.append(sql.SQL("dr_scope.section_id = %s"))
                params.append(section_id)

        query = sql.SQL(
            """
            SELECT t.{column}::text AS value,
                   {label} AS label,
                   COUNT(*)::int AS count
            FROM {table} t
            {joins}
            WHERE {where}
            GROUP BY 1, 2
            ORDER BY COUNT(*) DESC, label NULLS LAST, value
            LIMIT %s
            """
        ).format(
            column=sql.Identifier(column),
            label=label_expr,
            table=sql.Identifier(table),
            joins=sql.SQL(" ").join(joins),
            where=sql.SQL(" AND ").join(where_parts),
        )
        cur.execute(query, params + [limit])
        return {
            "table": table,
            "column": column,
            "section_id": section_id,
            "suggestions": [_json_safe(dict(row)) for row in cur.fetchall()],
        }


@router.post("/tables/{table}/validate")
def validate_change(table: str, request: ChangeRequest, _admin=Depends(require_admin)):
    with _db_cursor() as (conn, cur):
        result = _validate(cur, table, request)
        conn.rollback()
        return result


@router.post("/tables/{table}/apply")
def apply_change(table: str, request: ChangeRequest, _admin=Depends(require_admin)):
    if not request.confirmed:
        raise HTTPException(status_code=400, detail="confirmed=true is required")
    _allowed_table(table, require_editable=True)
    with _db_cursor() as (conn, cur):
        result = _validate(cur, table, request)
        if not result["ok"]:
            conn.rollback()
            return result

        pk_cols = _primary_key(cur, table)
        values = request.values or {}
        before = result.get("before")
        after = result.get("after")
        deleted_child_refs: dict[str, int] = {}

        if request.action == "update":
            set_sql = sql.SQL(", ").join(
                sql.SQL("{} = %s").format(sql.Identifier(col)) for col in values
            )
            where_sql, pk_params = _where_pk(pk_cols, request.pk)
            query = sql.SQL("UPDATE {} SET {} WHERE {} RETURNING *").format(
                sql.Identifier(table),
                set_sql,
                where_sql,
            )
            cur.execute(query, list(values.values()) + pk_params)
            after = _json_safe(dict(cur.fetchone()))

        elif request.action == "insert":
            columns = list(values.keys())
            query = sql.SQL("INSERT INTO {} ({}) VALUES ({}) RETURNING *").format(
                sql.Identifier(table),
                sql.SQL(", ").join(sql.Identifier(col) for col in columns),
                sql.SQL(", ").join(sql.Placeholder() for _ in columns),
            )
            cur.execute(query, list(values.values()))
            after = _json_safe(dict(cur.fetchone()))

        elif request.action == "delete":
            where_sql, pk_params = _where_pk(pk_cols, request.pk)
            if table == "objects":
                deleted_child_refs = _delete_temp_roads_for_object(cur, before)
            for label, count in _delete_safe_child_refs(cur, table, before).items():
                deleted_child_refs[label] = deleted_child_refs.get(label, 0) + count
            query = sql.SQL("DELETE FROM {} WHERE {} RETURNING *").format(sql.Identifier(table), where_sql)
            cur.execute(query, pk_params)
            before = _json_safe(dict(cur.fetchone()))
            if table == "temporary_roads":
                for label, count in _delete_temp_road_linked_object(cur, before).items():
                    deleted_child_refs[label] = deleted_child_refs.get(label, 0) + count
            after = None

        _ensure_audit_table(cur)
        cur.execute(
            """
            INSERT INTO public.db_change_audit
              (table_name, action, pk_json, before_json, after_json, changed_by)
            VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (
                table,
                request.action,
                psycopg2.extras.Json(_json_safe(request.pk)),
                psycopg2.extras.Json(before),
                psycopg2.extras.Json(after),
                request.changed_by or "dashboard",
            ),
        )
        conn.commit()
        return {**result, "ok": True, "before": before, "after": after, "deleted_child_refs": deleted_child_refs}
