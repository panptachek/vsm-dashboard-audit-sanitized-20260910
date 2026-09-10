#!/usr/bin/env python3
from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import rd_parser


class RdParserTests(unittest.TestCase):
    def test_rows_from_table_extracts_vor_columns(self) -> None:
        rows = rd_parser.rows_from_table(
            [["N", "Наименование", "Ед. изм.", "Кол-во"], ["1", "Устройство насыпи", "м3", "1 234,5"]],
            page_number=2,
            table_index=1,
            table_bbox={"x0": 10.0, "top": 20.0, "x1": 300.0, "bottom": 90.0},
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].raw_name, "Устройство насыпи")
        self.assertEqual(rows[0].unit, "м3")
        self.assertEqual(rows[0].volume, 1234.5)
        self.assertEqual(rows[0].page_number, 2)
        self.assertEqual(rows[0].bbox, {"x0": 10.0, "top": 20.0, "x1": 300.0, "bottom": 90.0})
        self.assertEqual(rows[0].row_bbox, {"x0": 10.0, "top": 55.0, "x1": 300.0, "bottom": 90.0})

    def test_rows_from_table_extracts_picket_range_and_multiline_name(self) -> None:
        rows = rd_parser.rows_from_table(
            [["ПК", "Наименование работ", "Ед", "Объем"], ["", "Устройство", "", ""], ["ПК 10+00 - ПК 11+50", "насыпи", "м3", "100"]]
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].raw_name, "Устройство насыпи")
        self.assertEqual(rows[0].pk_start, 1000.0)
        self.assertEqual(rows[0].pk_end, 1150.0)
        self.assertEqual(rows[0].pk_raw, "ПК 10+00 - ПК 11+50")

    def test_aggregate_rows_sums_by_object_work_unit_and_pk(self) -> None:
        source_rows = [
            {"id": "row-1", "object_id": "obj-1", "object_name": "ОХ", "work_type_id": "work-1", "work_type_name": "Насыпь", "unit": "м3", "volume": 10, "pk_start": 1000, "pk_end": 1100, "confidence": 0.99},
            {"id": "row-2", "object_id": "obj-1", "object_name": "ОХ", "work_type_id": "work-1", "work_type_name": "Насыпь", "unit": "м3", "volume": 15.5, "pk_start": 1000, "pk_end": 1100, "confidence": 0.99},
            {"id": "row-3", "object_id": "obj-1", "work_type_id": None, "unit": "м3", "volume": 99, "confidence": 0},
        ]
        aggregates = rd_parser.aggregate_rows(source_rows)
        self.assertEqual(len(aggregates), 1)
        self.assertEqual(aggregates[0]["volume"], 25.5)
        self.assertEqual(aggregates[0]["source_row_count"], 2)
        self.assertEqual(aggregates[0]["row_ids"], ["row-1", "row-2"])

    def test_layout_parser_extracts_named_vor_rows_without_numeric_header_noise(self) -> None:
        rows = rd_parser.rows_from_layout_text(
            """
            № Наименование Ед. изм. Кол. Примечание
            п/п
            1 2 3 4 5
            Устройство насыпи из песка средней крупности по
            ГОСТ 25100 из карьера. Разработка грунта с
            1.1 перемещением до 30 м бульдозерами мощностью 130 м³ 46981,49
            Песок средней крупности по ГОСТ 25100 для
            насыпи.
            1.2 Объём дан с учётом: м³ 55992,54
            """,
            page_number=1,
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].raw_name, "Устройство насыпи из песка средней крупности по ГОСТ 25100 из карьера. Разработка грунта с перемещением до 30 м бульдозерами мощностью 130")
        self.assertEqual(rows[0].unit, "м3")
        self.assertEqual(rows[0].volume, 46981.49)
        self.assertFalse(any(row.raw_name == "3" for row in rows))


    def test_layout_parser_does_not_shift_lowercase_continuation_to_next_material(self) -> None:
        rows = rd_parser.rows_from_layout_text(
            """
            № Наименование Ед. изм. Кол.
            Устройство второго защитного слоя 1,9м из природной песчано-гравийной смеси
            1 Уплотнение ПГС, с м3 57617
            катком массой 25 т за 12 проходов по 1 слою, толщина уплотняемого слоя до 0,25 м.
            Песчано-гравийная смесь по ГОСТ 25100 для
            2 второго (нижнего) защитного слоя. м3 70292,74
            """,
            page_number=1,
        )
        self.assertEqual(len(rows), 2)
        self.assertIn("Устройство второго защитного слоя", rows[0].raw_name)
        self.assertEqual(rows[0].volume, 57617.0)
        self.assertTrue(rows[1].raw_name.startswith("Песчано-гравийная смесь"), rows[1].raw_name)
        self.assertNotIn("катком массой 25", rows[1].raw_name)
        self.assertEqual(rows[1].volume, 70292.74)

    def test_layout_parser_extracts_earthwork_picket_rows(self) -> None:
        rows = rd_parser.rows_from_layout_text(
            """
            Пикетаж Длина участка, м Насыпь, м3 Выемка, м3 Кюветы, м3 Примечание
            1 2 3 4 5 6
            70+00,00
                         13,50     400,45    0,00      2,71
            """,
            page_number=2,
        )
        by_name = {row.raw_name: row for row in rows}
        self.assertAlmostEqual(by_name["Длина участка"].volume, 13.5)
        self.assertAlmostEqual(by_name["Насыпь"].volume, 400.45)
        self.assertAlmostEqual(by_name["Кюветы"].volume, 2.71)
        self.assertEqual(by_name["Насыпь"].pk_start, 7000.0)
        self.assertEqual(by_name["Насыпь"].pk_end, 7013.5)

    def test_layout_parser_extracts_prs_good_and_bad_totals(self) -> None:
        rows = rd_parser.rows_from_layout_text(
            """
            Пикетаж Длина Площадь снятия пригодного ПРС, м2 Площадь снятия не пригодного ПРС, м2 Объём снятия пригодного ПРС, м3 Объём снятия не пригодного ПРС, м3
            участка, м
            Слева Справа Итого Слева Справа Итого Слева Справа Итого Слева Справа Итого
            ПК90+33,06           0,00   0,00      0,00      2,00   2,04       4,04
                        12,00                                                            0,00    0,00     0,00    24,00  24,48    48,48
            """,
            page_number=1,
        )
        by_name = {row.raw_name: row for row in rows}
        self.assertAlmostEqual(by_name["Площадь снятия непригодного ПРС"].volume, 4.04)
        self.assertAlmostEqual(by_name["Объем снятия непригодного ПРС"].volume, 48.48)
        self.assertEqual(by_name["Объем снятия непригодного ПРС"].unit, "м3")

    def test_table_parser_prefers_quantity_column_over_formula_column(self) -> None:
        rows = rd_parser.rows_from_table(
            [
                [
                    "№ п.п",
                    "№ в ЛСР",
                    "Наименование работ",
                    "Ед. изм.",
                    "Кол.",
                    "Ссылки на чертежи",
                    "Формулы расчета, расчет объемов работ и расхода материалов",
                ],
                ["1", "2", "3", "4", "5", "6", "7"],
                [
                    "1",
                    "",
                    "Устройство второго защитного слоя с послойным уплотнением пневмоколесным катком",
                    "м3",
                    "57 617",
                    "",
                    "V = ((S1 + S2) / 2 * L1-2)",
                ],
                [
                    "2",
                    "",
                    "Песчано-гравийная смесь по ГОСТ 25100 для второго защитного слоя",
                    "м3",
                    "70 292,74",
                    "",
                    "57617*1,22",
                ],
            ]
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].raw_name, "Устройство второго защитного слоя с послойным уплотнением пневмоколесным катком")
        self.assertEqual(rows[0].volume, 57617.0)
        self.assertEqual(rows[1].volume, 70292.74)
        self.assertFalse(any(row.raw_name == "3" for row in rows))


    def test_table_parser_attaches_leading_page_continuation_to_previous_row(self) -> None:
        previous_rows = rd_parser.rows_from_table(
            [
                ["№ п.п", "№ в ЛСР", "Наименование работ", "Ед. изм.", "Кол.", "Ссылки", "Формулы"],
                ["1", "2", "3", "4", "5", "6", "7"],
                [
                    "18",
                    "",
                    "Дальность возки песка средней крупности по ГОСТ 25100 для",
                    "т",
                    "669,33",
                    "",
                    "G = V * gamma",
                ],
            ],
            page_number=5,
            table_index=2,
        )
        self.assertEqual(len(previous_rows), 1)
        continuations: list[tuple[str, list[str]]] = []
        next_rows = rd_parser._rows_from_table(
            [
                ["№ п.п", "№ в ЛСР", "Наименование работ", "Ед. изм.", "Кол.", "Ссылки", "Формулы"],
                ["1", "2", "3", "4", "5", "6", "7"],
                [
                    "",
                    "",
                    "насыпи внутри региона на расстояние 24 км",
                    "",
                    "",
                    "",
                    "V - объём грунта, м3",
                ],
                [
                    "19",
                    "",
                    "Устройство участков переменной жесткости материала ЩПГС",
                    "м3",
                    "2228",
                    "",
                    "Объём материала определяется",
                ],
            ],
            page_number=6,
            table_index=2,
            page_continuations=continuations,
        )
        for text, raw_cells in continuations:
            rd_parser._append_table_continuation(previous_rows[-1], text, raw_cells)
        self.assertIn("насыпи внутри региона", previous_rows[-1].raw_name)
        self.assertEqual(next_rows[0].raw_name, "Устройство участков переменной жесткости материала ЩПГС")
        self.assertNotIn("насыпи внутри региона", next_rows[0].raw_name)

    def test_table_parser_treats_uppercase_leading_page_fragment_as_continuation(self) -> None:
        previous_rows = rd_parser.rows_from_table(
            [
                ["№ п.п", "Наименование работ", "Ед. изм.", "Кол.", "Формулы"],
                ["1", "2", "3", "4", "5"],
                [
                    "141",
                    "Погрузка растительного грунта на временной площадке.",
                    "м3",
                    "1624,95",
                    "W = ...",
                ],
            ],
            page_number=32,
            table_index=2,
        )
        continuations: list[tuple[str, list[str]]] = []
        next_rows = rd_parser._rows_from_table(
            [
                ["№ п.п", "Наименование работ", "Ед. изм.", "Кол.", "Формулы"],
                ["1", "2", "3", "4", "5"],
                [
                    "",
                    "Надвижка растительного грунта бульдозером на расстояние 10 м",
                    "",
                    "",
                    "",
                ],
                [
                    "142",
                    "Укрепление дна канавы с щебневанием (фр. 10-20мм)",
                    "м2",
                    "600",
                    "W = ...",
                ],
            ],
            page_number=33,
            table_index=2,
            page_continuations=continuations,
        )
        for text, raw_cells in continuations:
            rd_parser._append_table_continuation(previous_rows[-1], text, raw_cells)

        self.assertIn("Надвижка растительного грунта", previous_rows[-1].raw_name)
        self.assertEqual(next_rows[0].raw_name, "Укрепление дна канавы с щебневанием (фр. 10-20мм)")
        self.assertNotIn("Надвижка растительного грунта", next_rows[0].raw_name)



    def test_table_parser_ignores_incomplete_noise_before_next_position(self) -> None:
        continuations: list[tuple[str, list[str]]] = []
        rows = rd_parser._rows_from_table(
            [
                ["№ п.п", "Наименование работ", "Ед. изм.", "Кол.", "Формулы"],
                ["1", "2", "3", "4", "5"],
                ["", "Лист 6", "", "", ""],
                ["", "насыпи внутри региона на расстояние 24 км", "", "", "V - объем грунта"],
                ["19", "Устройство участков переменной жесткости материала ЩПГС", "м3", "2228", ""],
            ],
            page_number=6,
            table_index=1,
            page_continuations=continuations,
        )

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].raw_name, "Устройство участков переменной жесткости материала ЩПГС")
        self.assertEqual(rows[0].volume, 2228.0)
        self.assertIn("насыпи внутри региона", continuations[0][0])
        self.assertNotIn("Лист 6", rows[0].raw_name)

    def test_material_names_are_manual_review(self) -> None:
        row = rd_parser.ParsedRdRow(raw_name="Песок средней крупности по ГОСТ 25100", unit="м3", volume=10)
        self.assertEqual(rd_parser.review_status_for_row(row), "материал / проверить")

    def test_work_actions_with_material_words_are_not_materials(self) -> None:
        self.assertFalse(rd_parser.is_material_like("Устройство насыпи из песка средней крупности"))
        self.assertFalse(rd_parser.is_material_like("Транспортировка цемента ПЦ 500 на площадку"))
        self.assertFalse(rd_parser.is_material_like("Разработка грунта ИГЭ 14д I группы"))
        self.assertTrue(rd_parser.is_material_like("Смесь щебеночно-песчаная типа 0/31,5"))

    def test_layout_parser_skips_empty_text_pages(self) -> None:
        self.assertEqual(rd_parser.rows_from_layout_text("\n \n", page_number=1), [])

    def test_parse_document_extracts_xlsx_sheets_as_tables(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "rd_tables.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "ВОР"
            sheet.append(["№ п.п", "Наименование работ", "Ед. изм.", "Кол."])
            sheet.append(["1", "Устройство насыпи", "м3", "1 234,5"])

            second = workbook.create_sheet("Ведомость")
            second.append(["№", "Наименование работ", "Ед. изм.", "Объем"])
            second.append(["1", "Устройство второго защитного слоя", "м3", 57617])
            second.append(["2", "Песчано-гравийная смесь для второго защитного слоя", "м3", "70 292,74"])

            hidden = workbook.create_sheet("Скрытый черновик")
            hidden.sheet_state = "hidden"
            hidden.append(["№", "Наименование работ", "Ед. изм.", "Объем"])
            hidden.append(["1", "Скрытая строка", "м3", 999])

            workbook.save(path)
            workbook.close()

            rows = rd_parser.parse_document(path)

        self.assertEqual([row.raw_name for row in rows], [
            "Устройство насыпи",
            "Устройство второго защитного слоя",
            "Песчано-гравийная смесь для второго защитного слоя",
        ])
        self.assertEqual([row.page_number for row in rows], [1, 2, 2])
        self.assertEqual([row.page_label for row in rows], ["Лист 1 · ВОР", "Лист 2 · Ведомость", "Лист 2 · Ведомость"])
        self.assertEqual(rows[0].volume, 1234.5)
        self.assertEqual(rows[1].volume, 57617.0)
        self.assertEqual(rows[2].volume, 70292.74)


if __name__ == "__main__":
    unittest.main()
