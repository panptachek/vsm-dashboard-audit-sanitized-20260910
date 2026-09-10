#!/usr/bin/env python3
from __future__ import annotations

import asyncio
import sys
import types
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))


class DummyRouter:
    def __init__(self, *args, **kwargs):
        pass

    def get(self, *args, **kwargs):
        return lambda fn: fn

    post = put = patch = delete = get


class DummyHTTPException(Exception):
    def __init__(self, status_code: int, detail: str):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


def passthrough(*args, **kwargs):
    return None


fastapi_stub = sys.modules.get("fastapi") or types.ModuleType("fastapi")
fastapi_stub.APIRouter = getattr(fastapi_stub, "APIRouter", DummyRouter)
fastapi_stub.Depends = getattr(fastapi_stub, "Depends", passthrough)
fastapi_stub.File = getattr(fastapi_stub, "File", passthrough)
fastapi_stub.Header = getattr(fastapi_stub, "Header", passthrough)
fastapi_stub.Query = getattr(fastapi_stub, "Query", passthrough)
fastapi_stub.HTTPException = getattr(fastapi_stub, "HTTPException", DummyHTTPException)
fastapi_stub.UploadFile = getattr(fastapi_stub, "UploadFile", object)
sys.modules.setdefault("fastapi", fastapi_stub)

responses_stub = sys.modules.get("fastapi.responses") or types.ModuleType("fastapi.responses")
responses_stub.FileResponse = getattr(responses_stub, "FileResponse", object)
responses_stub.Response = getattr(responses_stub, "Response", object)
responses_stub.StreamingResponse = getattr(responses_stub, "StreamingResponse", object)
sys.modules.setdefault("fastapi.responses", responses_stub)

pydantic_stub = sys.modules.get("pydantic") or types.ModuleType("pydantic")
pydantic_stub.BaseModel = getattr(pydantic_stub, "BaseModel", object)
pydantic_stub.Field = getattr(pydantic_stub, "Field", lambda default=None, **kwargs: default)
sys.modules.setdefault("pydantic", pydantic_stub)

main_stub = sys.modules.get("main") or types.ModuleType("main")
main_stub.query = getattr(main_stub, "query", lambda *args, **kwargs: [])
main_stub.query_one = getattr(main_stub, "query_one", lambda *args, **kwargs: None)
main_stub.get_conn = getattr(main_stub, "get_conn", lambda: None)
sys.modules.setdefault("main", main_stub)

auth_stub = sys.modules.get("auth") or types.ModuleType("auth")
auth_stub.current_user = getattr(auth_stub, "current_user", lambda: None)
auth_stub.require_admin = getattr(auth_stub, "require_admin", lambda: None)
auth_stub.require_settings_manager = getattr(auth_stub, "require_settings_manager", lambda: None)
auth_stub.require_statements_project_writer = getattr(auth_stub, "require_statements_project_writer", lambda: None)
sys.modules.setdefault("auth", auth_stub)

import wip_routes


PROJECT_WRITER_USER = types.SimpleNamespace(role="user", permissions=["statements:project:write"])


class FakeUpload:
    def __init__(self, data: bytes):
        self._data = data

    async def read(self) -> bytes:
        return self._data


class FakeCursor:
    def __init__(self):
        self.executed: list[tuple[str, list[object] | tuple[object, ...] | None]] = []
        self.rowcount = 0

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False

    def execute(self, sql: str, params=None):
        self.executed.append((sql, params))
        self.rowcount = 1 if sql.lstrip().upper().startswith("UPDATE") else 0


class FakeConn:
    def __init__(self):
        self.cursor_obj = FakeCursor()
        self.closed = False

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False

    def cursor(self):
        return self.cursor_obj

    def close(self):
        self.closed = True


class RowcountCursor:
    def __init__(self, update_rowcounts: list[int] | None = None, insert_rowcounts: list[int] | None = None, fetch_results: list[object] | None = None):
        self.update_rowcounts = list(update_rowcounts or [])
        self.insert_rowcounts = list(insert_rowcounts or [])
        self.fetch_results = list(fetch_results or [])
        self.executed: list[tuple[str, list[object] | tuple[object, ...] | None]] = []
        self.rowcount = 0

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False

    def execute(self, sql: str, params=None):
        self.executed.append((sql, params))
        command = sql.lstrip().upper()
        if command.startswith("UPDATE"):
            self.rowcount = self.update_rowcounts.pop(0) if self.update_rowcounts else 0
        elif command.startswith("INSERT"):
            self.rowcount = self.insert_rowcounts.pop(0) if self.insert_rowcounts else 0
        else:
            self.rowcount = 0

    def fetchone(self):
        if not self.fetch_results:
            return None
        return self.fetch_results.pop(0)


class RowcountConn:
    def __init__(self, update_rowcounts: list[int] | None = None, insert_rowcounts: list[int] | None = None, fetch_results: list[object] | None = None):
        self.cursor_obj = RowcountCursor(update_rowcounts, insert_rowcounts, fetch_results)
        self.closed = False

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False

    def cursor(self):
        return self.cursor_obj

    def close(self):
        self.closed = True


def workbook_bytes(headers: list[str], rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def workbook_bytes_with_statement_meta(headers: list[str], rows: list[list[object]], periods: list[tuple[str, str, str, str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    meta = wb.create_sheet(wip_routes.STATEMENT_EXPORT_META_SHEET)
    meta.sheet_state = "hidden"
    meta.append(["schema", "vsm_statements_export_v1"])
    meta.append(["period_key", "label", "start", "end"])
    for period in periods:
        meta.append(list(period))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class StatementsImportTests(unittest.TestCase):
    def test_rejects_non_xlsx_payload_before_db_work(self):
        with mock.patch.object(wip_routes, "_statement_table_exists", return_value=True), \
             mock.patch.object(wip_routes, "get_conn") as get_conn_mock:
            with self.assertRaises(Exception) as raised:
                asyncio.run(wip_routes.statements_import(file=FakeUpload(b"object,plan\n1,2\n"), _user=None))

        self.assertEqual(getattr(raised.exception, "status_code", None), 400)
        self.assertIn("XLSX", getattr(raised.exception, "detail", str(raised.exception)))
        get_conn_mock.assert_not_called()


    def test_existing_plan_update_preserves_period_and_pk_when_columns_are_absent(self):
        data = workbook_bytes(
            ["Группа", "Плановый объем", "plan_id", "object_id", "work_type_id", "unit", "source_rows", "__is_group"],
            [["Монтаж", 77, "33333333-3333-3333-3333-333333333333", "11111111-1111-1111-1111-111111111111", "22222222-2222-2222-2222-222222222222", "м3", 1, False]],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=None))

        self.assertTrue(result["ok"])
        self.assertEqual(result["plan"], {"inserted": 0, "updated": 1})
        updates = [entry for entry in conn.cursor_obj.executed if "UPDATE planned_work_items" in entry[0]]
        self.assertEqual(len(updates), 1)
        sql, params = updates[0]
        self.assertIn("planned_volume = %s", sql)
        self.assertIn("updated_at = now()", sql)
        self.assertNotIn("period_start =", sql)
        self.assertNotIn("period_end =", sql)
        self.assertNotIn("pk_start =", sql)
        self.assertNotIn("pk_end =", sql)
        self.assertEqual(params[-1], "33333333-3333-3333-3333-333333333333")
        self.assertTrue(conn.closed)

    def test_duplicate_visible_volume_columns_import_rightmost_filled_value(self):
        data = workbook_bytes(
            [
                "Группа",
                "Плановый объем #1",
                "Плановый объем #2",
                "plan_id",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                "Монтаж",
                77,
                88,
                "33333333-3333-3333-3333-333333333333",
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                1,
                False,
            ]],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=None))

        self.assertTrue(result["ok"])
        updates = [entry for entry in conn.cursor_obj.executed if "UPDATE planned_work_items" in entry[0]]
        self.assertEqual(len(updates), 1)
        self.assertEqual(updates[0][1][0], 88.0)

    def test_existing_plan_update_changes_period_only_when_period_columns_are_present(self):
        data = workbook_bytes(
            ["Группа", "Плановый объем", "Дата/период от", "Дата/период до", "plan_id", "object_id", "work_type_id", "unit", "source_rows", "__is_group"],
            [["Монтаж", 12.5, "2026-05-01", "2026-05-10", "33333333-3333-3333-3333-333333333333", "11111111-1111-1111-1111-111111111111", "22222222-2222-2222-2222-222222222222", "м3", 1, False]],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=None))

        self.assertTrue(result["ok"])
        updates = [entry for entry in conn.cursor_obj.executed if "UPDATE planned_work_items" in entry[0]]
        self.assertEqual(len(updates), 1)
        sql, _params = updates[0]
        self.assertIn("period_start = %s", sql)
        self.assertIn("period_end = %s", sql)
        self.assertIn("plan_type = CASE", sql)
        self.assertNotIn("pk_start =", sql)
        self.assertNotIn("pk_end =", sql)

    def test_aggregate_row_with_period_values_does_not_import_total_plan(self):
        periods = wip_routes._statement_period_buckets("month_third", date(2026, 5, 1), date(2026, 5, 20))
        first_period = periods[0]
        second_period = periods[1]
        data = workbook_bytes(
            [
                "Группа",
                "Плановый объем",
                f"{first_period['label']} план",
                f"{second_period['label']} план",
                "plan_id",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                "Монтаж",
                99,
                10,
                25.5,
                "33333333-3333-3333-3333-333333333333",
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                2,
                False,
            ]],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(
                wip_routes.statements_import(
                    file=FakeUpload(data),
                    period_from="2026-05-01",
                    period_to="2026-05-20",
                    period_mode="month_third",
                    _user=None,
                )
            )

        self.assertTrue(result["ok"])
        self.assertEqual(result["plan"], {"inserted": 0, "updated": 2})
        updates = [entry for entry in conn.cursor_obj.executed if "UPDATE planned_work_items" in entry[0]]
        self.assertEqual(len(updates), 2)
        self.assertTrue(all("WHERE id = %s::uuid" not in sql for sql, _params in updates))
        self.assertTrue(all("period_start IS NOT DISTINCT FROM %s" in sql for sql, _params in updates))

    def test_multi_source_row_without_period_values_skips_direct_total_imports(self):
        data = workbook_bytes(
            [
                "Группа",
                "Плановый объем",
                "Проектный объем",
                "plan_id",
                "project_item_id",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                "Монтаж",
                99,
                120,
                "33333333-3333-3333-3333-333333333333",
                "44444444-4444-4444-4444-444444444444",
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                2,
                False,
            ]],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=PROJECT_WRITER_USER))

        self.assertTrue(result["ok"])
        self.assertEqual(result["plan"], {"inserted": 0, "updated": 0})
        self.assertEqual(result["project"], {"inserted": 0, "updated": 0})
        self.assertEqual(result["skipped"], 2)
        self.assertEqual(conn.cursor_obj.executed, [])
        self.assertTrue(conn.closed)


    def test_period_columns_use_exported_workbook_metadata_without_current_filters(self):
        data = workbook_bytes_with_statement_meta(
            [
                "Группа",
                "Плановый объем",
                "1/3 05.2026 план",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                "Монтаж",
                99,
                33.5,
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                2,
                False,
            ]],
            [("2026-05-01_2026-05-10", "1/3 05.2026", "2026-05-01", "2026-05-10")],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=None))

        self.assertTrue(result["ok"])
        self.assertEqual(result["plan"], {"inserted": 0, "updated": 1})
        updates = [entry for entry in conn.cursor_obj.executed if "UPDATE planned_work_items" in entry[0]]
        self.assertEqual(len(updates), 1)
        _sql, params = updates[0]
        self.assertEqual(params[0], 33.5)
        self.assertEqual(params[-6:-4], (date(2026, 5, 1), date(2026, 5, 10)))
        self.assertTrue(conn.closed)

    def test_period_plan_import_matches_pk_range_when_export_has_pk_columns(self):
        data = workbook_bytes_with_statement_meta(
            [
                "Группа",
                "1/3 05.2026 план",
                "ПК от",
                "ПК до",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                "Монтаж",
                17.5,
                "ПК1+00",
                "ПК2+00",
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                1,
                False,
            ]],
            [("2026-05-01_2026-05-10", "1/3 05.2026", "2026-05-01", "2026-05-10")],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=None))

        self.assertTrue(result["ok"])
        self.assertEqual(result["plan"], {"inserted": 0, "updated": 1})
        updates = [entry for entry in conn.cursor_obj.executed if "UPDATE planned_work_items" in entry[0]]
        self.assertEqual(len(updates), 1)
        sql, params = updates[0]
        self.assertIn("pk_start IS NOT DISTINCT FROM", sql)
        self.assertIn("pk_end IS NOT DISTINCT FROM", sql)
        self.assertEqual(params[-4:], (100.0, 100.0, 200.0, 200.0))
        self.assertTrue(conn.closed)


    def test_period_plan_values_skip_total_plan_even_for_single_source_row(self):
        periods = wip_routes._statement_period_buckets("day", date(2026, 5, 1), date(2026, 5, 1))
        period = periods[0]
        data = workbook_bytes(
            [
                "Группа",
                "Плановый объем",
                f"{period['label']} план",
                "plan_id",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                "Монтаж",
                99,
                12,
                "33333333-3333-3333-3333-333333333333",
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                1,
                False,
            ]],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(
                wip_routes.statements_import(
                    file=FakeUpload(data),
                    period_from="2026-05-01",
                    period_to="2026-05-01",
                    period_mode="day",
                    _user=None,
                )
            )

        self.assertTrue(result["ok"])
        self.assertEqual(result["plan"], {"inserted": 0, "updated": 1})
        updates = [entry for entry in conn.cursor_obj.executed if "UPDATE planned_work_items" in entry[0]]
        self.assertEqual(len(updates), 1)
        sql, params = updates[0]
        self.assertNotIn("WHERE id = %s::uuid", sql)
        self.assertIn("period_start IS NOT DISTINCT FROM %s", sql)
        self.assertEqual(params[0], 12.0)
        self.assertEqual(params[-6:-4], (date(2026, 5, 1), date(2026, 5, 1)))
        self.assertTrue(conn.closed)


    def test_project_volume_update_by_project_item_id(self):
        data = workbook_bytes(
            [
                "Группа",
                "Проектный объем",
                "project_item_id",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                "Монтаж",
                55.5,
                "44444444-4444-4444-4444-444444444444",
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                1,
                False,
            ]],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=PROJECT_WRITER_USER))

        self.assertTrue(result["ok"])
        self.assertEqual(result["plan"], {"inserted": 0, "updated": 0})
        self.assertEqual(result["project"], {"inserted": 0, "updated": 1})
        project_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_items" in entry[0]]
        self.assertEqual(len(project_updates), 1)
        self.assertIn("WHERE id = %s::uuid", project_updates[0][0])
        self.assertEqual(project_updates[0][1], (55.5, "м3", None, "44444444-4444-4444-4444-444444444444"))
        plan_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE planned_work_items" in entry[0]]
        self.assertEqual(plan_updates, [])
        self.assertTrue(conn.closed)

    def test_project_volume_import_requires_project_writer_permission(self):
        data = workbook_bytes(
            [
                "Группа",
                "Проектный объем",
                "project_item_id",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                "Монтаж",
                55.5,
                "44444444-4444-4444-4444-444444444444",
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                1,
                False,
            ]],
        )
        conn = FakeConn()
        blocked_user = types.SimpleNamespace(role="user", permissions=[])

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            with self.assertRaises(Exception) as raised:
                asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=blocked_user))

        self.assertEqual(getattr(raised.exception, "status_code", None), 403)
        self.assertIn("проектных данных", getattr(raised.exception, "detail", str(raised.exception)).lower())
        self.assertEqual(conn.cursor_obj.executed, [])
        self.assertTrue(conn.closed)

    def test_project_volume_update_by_object_work_when_project_id_absent(self):
        data = workbook_bytes(
            [
                "Группа",
                "Проектный объем",
                "object_id",
                "work_type_id",
                "unit",
                "Комментарий",
                "source_rows",
                "__is_group",
            ],
            [[
                "Монтаж",
                62,
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "шт",
                "project import",
                1,
                False,
            ]],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=PROJECT_WRITER_USER))

        self.assertTrue(result["ok"])
        self.assertEqual(result["project"], {"inserted": 0, "updated": 1})
        project_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_items" in entry[0]]
        self.assertEqual(len(project_updates), 1)
        self.assertIn("WHERE object_id = %s::uuid", project_updates[0][0])
        self.assertEqual(
            project_updates[0][1],
            (62.0, "шт", "project import", "11111111-1111-1111-1111-111111111111", "22222222-2222-2222-2222-222222222222", "шт"),
        )
        self.assertTrue(conn.closed)

    def test_planned_volume_import_insert_validates_object_and_work(self):
        data = workbook_bytes(
            [
                "Плановый объем",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                15,
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                1,
                False,
            ]],
        )
        conn = RowcountConn(update_rowcounts=[0], insert_rowcounts=[1])

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=None))

        self.assertTrue(result["ok"])
        self.assertEqual(result["plan"], {"inserted": 1, "updated": 0})
        self.assertEqual(result["skipped"], 0)
        plan_inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO planned_work_items" in entry[0]]
        self.assertEqual(len(plan_inserts), 1)
        sql, params = plan_inserts[0]
        self.assertIn("FROM objects o", sql)
        self.assertIn("CROSS JOIN work_types wt", sql)
        self.assertIn("WHERE o.id = %s::uuid", sql)
        self.assertEqual(params[-2:], ("11111111-1111-1111-1111-111111111111", "22222222-2222-2222-2222-222222222222"))
        self.assertTrue(conn.closed)

    def test_planned_volume_import_missing_object_or_work_is_skipped(self):
        data = workbook_bytes(
            [
                "Плановый объем",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                15,
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                1,
                False,
            ]],
        )
        conn = RowcountConn(update_rowcounts=[0], insert_rowcounts=[0])

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=None))

        self.assertTrue(result["ok"])
        self.assertEqual(result["plan"], {"inserted": 0, "updated": 0})
        self.assertEqual(result["skipped"], 1)
        plan_inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO planned_work_items" in entry[0]]
        self.assertEqual(len(plan_inserts), 1)
        self.assertIn("FROM objects o", plan_inserts[0][0])
        self.assertTrue(conn.closed)

    def test_project_volume_import_insert_validates_object_and_work(self):
        data = workbook_bytes(
            [
                "Проектный объем",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                62,
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "шт",
                1,
                False,
            ]],
        )
        conn = RowcountConn(update_rowcounts=[0], insert_rowcounts=[1])

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=PROJECT_WRITER_USER))

        self.assertTrue(result["ok"])
        self.assertEqual(result["project"], {"inserted": 1, "updated": 0})
        self.assertEqual(result["skipped"], 0)
        project_inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO project_work_items" in entry[0]]
        self.assertEqual(len(project_inserts), 1)
        sql, params = project_inserts[0]
        self.assertIn("FROM objects o", sql)
        self.assertIn("CROSS JOIN work_types wt", sql)
        self.assertIn("WHERE o.id = %s::uuid", sql)
        self.assertEqual(params[-2:], ("11111111-1111-1111-1111-111111111111", "22222222-2222-2222-2222-222222222222"))
        self.assertTrue(conn.closed)

    def test_project_volume_import_missing_object_or_work_is_skipped(self):
        data = workbook_bytes(
            [
                "Проектный объем",
                "object_id",
                "work_type_id",
                "unit",
                "source_rows",
                "__is_group",
            ],
            [[
                62,
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "шт",
                1,
                False,
            ]],
        )
        conn = RowcountConn(update_rowcounts=[0], insert_rowcounts=[0])

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=PROJECT_WRITER_USER))

        self.assertTrue(result["ok"])
        self.assertEqual(result["project"], {"inserted": 0, "updated": 0})
        self.assertEqual(result["skipped"], 1)
        project_inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO project_work_items" in entry[0]]
        self.assertEqual(len(project_inserts), 1)
        self.assertIn("FROM objects o", project_inserts[0][0])
        self.assertTrue(conn.closed)

    def test_project_volume_with_exported_object_pk_and_zero_segments_updates_parent_project_item(self):
        data = workbook_bytes(
            [
                "Группа",
                "Проектный объем",
                "project_item_id",
                "object_id",
                "work_type_id",
                "unit",
                "ПК от",
                "ПК до",
                "project_segment_count",
                "source_rows",
                "__is_group",
            ],
            [[
                "Линейный объект",
                88,
                "44444444-4444-4444-4444-444444444444",
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                "ПК1+00",
                "ПК2+00",
                0,
                1,
                False,
            ]],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=PROJECT_WRITER_USER))

        self.assertTrue(result["ok"])
        self.assertEqual(result["project"], {"inserted": 0, "updated": 1})
        segment_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_item_segments" in entry[0]]
        self.assertEqual(segment_updates, [])
        project_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_items" in entry[0] and "WHERE id = %s::uuid" in entry[0]]
        self.assertEqual(len(project_updates), 1)
        self.assertEqual(project_updates[0][1], (88.0, "м3", None, "44444444-4444-4444-4444-444444444444"))
        self.assertTrue(conn.closed)

    def test_project_volume_with_exported_multi_segment_count_is_skipped(self):
        data = workbook_bytes(
            [
                "Проектный объем",
                "project_item_id",
                "object_id",
                "work_type_id",
                "unit",
                "ПК от",
                "ПК до",
                "project_segment_count",
                "source_rows",
                "__is_group",
            ],
            [[
                99,
                "44444444-4444-4444-4444-444444444444",
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                "ПК1+00",
                "ПК2+00",
                2,
                1,
                False,
            ]],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=PROJECT_WRITER_USER))

        self.assertTrue(result["ok"])
        self.assertEqual(result["project"], {"inserted": 0, "updated": 0})
        self.assertEqual(result["skipped"], 1)
        self.assertEqual(conn.cursor_obj.executed, [])
        self.assertTrue(conn.closed)

    def test_project_volume_with_project_id_and_pk_updates_project_segment(self):
        data = workbook_bytes(
            [
                "Проектный объем",
                "project_item_id",
                "object_id",
                "work_type_id",
                "unit",
                "ПК от",
                "ПК до",
                "ПК текст",
                "Комментарий",
                "source_rows",
                "__is_group",
            ],
            [[
                77,
                "44444444-4444-4444-4444-444444444444",
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "м3",
                "ПК1+00",
                "ПК2+00",
                "ПК1+00-ПК2+00",
                "segment project",
                1,
                False,
            ]],
        )
        conn = FakeConn()

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=PROJECT_WRITER_USER))

        self.assertTrue(result["ok"])
        self.assertEqual(result["project"], {"inserted": 0, "updated": 1})
        segment_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_item_segments" in entry[0]]
        self.assertEqual(len(segment_updates), 1)
        sql, params = segment_updates[0]
        self.assertIn("volume_segment = %s", sql)
        self.assertEqual(params, (77.0, "ПК1+00-ПК2+00", "segment project", "44444444-4444-4444-4444-444444444444", 100.0, 200.0))
        project_item_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_items" in entry[0] and "WHERE id = %s::uuid" in entry[0]]
        self.assertEqual(project_item_updates, [])
        self.assertTrue(conn.closed)

    def test_project_volume_with_object_work_and_pk_updates_project_segment(self):
        data = workbook_bytes(
            [
                "Проектный объем",
                "object_id",
                "work_type_id",
                "unit",
                "ПК от",
                "ПК до",
                "Комментарий",
                "source_rows",
                "__is_group",
            ],
            [[
                33,
                "11111111-1111-1111-1111-111111111111",
                "22222222-2222-2222-2222-222222222222",
                "шт",
                150,
                250,
                "object segment project",
                1,
                False,
            ]],
        )
        conn = SequencedConn([("project-id",)])

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = asyncio.run(wip_routes.statements_import(file=FakeUpload(data), _user=PROJECT_WRITER_USER))

        self.assertTrue(result["ok"])
        self.assertEqual(result["project"], {"inserted": 0, "updated": 1})
        ensure_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_items" in entry[0]]
        self.assertEqual(len(ensure_updates), 1)
        self.assertIn("RETURNING id::text", ensure_updates[0][0])
        segment_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_item_segments" in entry[0]]
        self.assertEqual(len(segment_updates), 1)
        self.assertEqual(segment_updates[0][1], (33.0, None, "object segment project", "project-id", 150.0, 250.0))
        self.assertTrue(conn.closed)


class StatementsPeriodAndPkTests(unittest.TestCase):
    def test_parse_statement_pk_accepts_route_formats(self):
        self.assertIsNone(wip_routes._parse_statement_pk(None))
        self.assertIsNone(wip_routes._parse_statement_pk(""))
        self.assertEqual(wip_routes._parse_statement_pk("2702"), 270200.0)
        self.assertEqual(wip_routes._parse_statement_pk("270250"), 270250.0)
        self.assertEqual(wip_routes._parse_statement_pk("PK2702+50"), 270250.0)
        self.assertEqual(wip_routes._parse_statement_pk("ПК 2702+50,5"), 270250.5)

    def test_scope_filter_uses_overlap_for_segment_mode(self):
        sql, params = wip_routes._statement_scope_filter_sql(
            ["UCH_1"],
            "11111111-1111-1111-1111-111111111111",
            270200.0,
            270300.0,
            "MAIN_TRACK",
        )

        self.assertIn("section_code = ANY(%s)", sql)
        self.assertIn("id = ANY(%s::uuid[])", sql)
        self.assertIn("object_type_code = ANY(%s)", sql)
        self.assertIn("pk_end IS NOT NULL AND pk_end >= %s", sql)
        self.assertIn("pk_start IS NOT NULL AND pk_start <= %s", sql)
        self.assertEqual(
            params,
            [["UCH_1"], ["11111111-1111-1111-1111-111111111111"], ["MAIN_TRACK"], 270200.0, 270300.0],
        )

    def test_period_buckets_day_mode(self):
        buckets = wip_routes._statement_period_buckets("day", date(2026, 5, 1), date(2026, 5, 3))

        self.assertEqual([item["key"] for item in buckets], ["2026-05-01", "2026-05-02", "2026-05-03"])
        self.assertEqual([item["label"] for item in buckets], ["01.05", "02.05", "03.05"])

    def test_period_buckets_quarter_mode_clips_partial_month(self):
        buckets = wip_routes._statement_period_buckets("month_quarter", date(2026, 5, 6), date(2026, 5, 25))

        self.assertEqual(
            [(item["label"], item["start_date"], item["end_date"]) for item in buckets],
            [
                ("1/4 05.2026", date(2026, 5, 6), date(2026, 5, 7)),
                ("2/4 05.2026", date(2026, 5, 8), date(2026, 5, 15)),
                ("3/4 05.2026", date(2026, 5, 16), date(2026, 5, 23)),
                ("4/4 05.2026", date(2026, 5, 24), date(2026, 5, 25)),
            ],
        )

    def test_period_buckets_third_mode_clips_partial_month(self):
        buckets = wip_routes._statement_period_buckets("month_third", date(2026, 5, 9), date(2026, 5, 22))

        self.assertEqual(
            [(item["label"], item["start_date"], item["end_date"]) for item in buckets],
            [
                ("1/3 05.2026", date(2026, 5, 9), date(2026, 5, 10)),
                ("2/3 05.2026", date(2026, 5, 11), date(2026, 5, 20)),
                ("3/3 05.2026", date(2026, 5, 21), date(2026, 5, 22)),
            ],
        )

    def test_period_buckets_month_mode_across_months(self):
        buckets = wip_routes._statement_period_buckets("month", date(2026, 5, 10), date(2026, 6, 5))

        self.assertEqual(
            [(item["label"], item["start_date"], item["end_date"]) for item in buckets],
            [
                ("05.2026", date(2026, 5, 10), date(2026, 5, 31)),
                ("06.2026", date(2026, 6, 1), date(2026, 6, 5)),
            ],
        )

    def test_period_buckets_reject_unknown_mode(self):
        with self.assertRaises(Exception) as raised:
            wip_routes._statement_period_buckets("week", date(2026, 5, 1), date(2026, 5, 7))

        self.assertEqual(getattr(raised.exception, "status_code", None), 400)


class StatementsSearchFilterTests(unittest.TestCase):
    def test_statement_filter_rows_by_query_matches_all_terms_across_labels(self):
        rows = [
            {
                "object_code": "AD5",
                "object_name": "АД5 ПК2702+50-ПК2703+00",
                "object_type_name": "Временная дорога",
                "section_name": "Участок 1",
                "work_type_name": "Устройство песчаного слоя",
                "unit": "м3",
                "pk_start": 270250,
                "pk_end": 270300,
                "comment": "левая сторона",
            },
            {
                "object_code": "P12",
                "object_name": "Площадка 12",
                "object_type_name": "Площадка",
                "section_name": "Участок 2",
                "work_type_name": "Щебень",
                "unit": "м3",
                "pk_start": 271000,
                "pk_end": 271100,
            },
        ]

        result = wip_routes._statement_filter_rows_by_query(rows, "АД5 песчаного ПК2702")

        self.assertEqual(result, [rows[0]])

    def test_statement_filter_rows_by_query_blank_keeps_rows(self):
        rows = [{"object_name": "АД5"}, {"object_name": "АД6"}]

        self.assertIs(wip_routes._statement_filter_rows_by_query(rows, "   "), rows)




class StatementsLegacyProjectPayloadTests(unittest.TestCase):
    def test_legacy_payload_filters_by_work_type_and_recomputes_totals(self):
        source_payload = {
            "sections": [],
            "objects": [],
            "rows": [
                {"object_id": "obj-1", "work_type_id": "work-keep", "unit": "м3", "project_volume": 100, "fact_volume": 40, "can_edit_project": True},
                {"object_id": "obj-2", "work_type_id": "work-drop", "unit": "м3", "project_volume": 55, "fact_volume": 10, "can_edit_project": True},
            ],
            "totals": [{"unit": "м3", "project_volume": 155, "fact_volume": 50}],
            "row_count": 2,
        }

        with (
            mock.patch.object(wip_routes, "statements", return_value=source_payload),
            mock.patch.object(wip_routes, "_statement_schema_flags", return_value={"editable": True}),
        ):
            payload = wip_routes._legacy_statement_rows_payload(
                section=None,
                object_id=None,
                pk_from=None,
                pk_to=None,
                reason="requested legacy source",
                _user=None,
                object_type=None,
                work_type_id="work-keep",
            )

        self.assertEqual(payload["row_count"], 1)
        self.assertEqual([row["work_type_id"] for row in payload["rows"]], ["work-keep"])
        self.assertEqual(payload["rows"][0]["planned_volume"], 100)
        self.assertEqual(payload["totals"], [{
            "unit": "м3",
            "project_volume": 100.0,
            "planned_volume": 100.0,
            "fact_volume": 40.0,
            "remaining_volume": 60.0,
            "percent_complete": 40.0,
        }])

    def test_legacy_payload_forwards_period_filters_and_keeps_period_values(self):
        captured = {}
        source_payload = {
            "sections": [],
            "objects": [],
            "periods": [{"key": "2026-05-02", "label": "02.05", "start": "2026-05-02", "end": "2026-05-02"}],
            "rows": [
                {
                    "object_id": "obj-1",
                    "work_type_id": "work-keep",
                    "unit": "м3",
                    "project_volume": 100,
                    "fact_volume": 8,
                    "can_edit_project": True,
                    "period_values": [{"period_key": "2026-05-02", "planned_volume": 0, "fact_volume": 8}],
                },
            ],
            "totals": [],
            "row_count": 1,
        }

        def fake_statements(**kwargs):
            captured.update(kwargs)
            return source_payload

        with (
            mock.patch.object(wip_routes, "statements", side_effect=fake_statements),
            mock.patch.object(wip_routes, "_statement_schema_flags", return_value={"editable": True}),
        ):
            payload = wip_routes._legacy_statement_rows_payload(
                section="UCH_1",
                object_id="obj-1",
                pk_from="ПК1+00",
                pk_to="ПК2+00",
                reason="requested legacy source",
                _user=None,
                object_type="PIPE",
                work_type_id="work-keep",
                period_from="2026-05-01",
                period_to="2026-05-03",
                period_mode="day",
            )

        self.assertEqual(captured["period_from"], "2026-05-01")
        self.assertEqual(captured["period_to"], "2026-05-03")
        self.assertEqual(captured["period_mode"], "day")
        self.assertEqual(captured["work_type_id"], "work-keep")
        self.assertEqual(payload["periods"], source_payload["periods"])
        self.assertEqual(payload["rows"][0]["period_values"][0]["fact_volume"], 8)

    def test_legacy_statements_filters_fact_by_period_and_returns_period_values(self):
        work_type_id = "22222222-2222-2222-2222-222222222222"
        calls = []
        main_row = {
            "object_id": "obj-1",
            "object_code": "OBJ",
            "object_name": "Объект",
            "object_type_code": "PIPE",
            "object_type_name": "Трубы",
            "section_code": "UCH_1",
            "section_name": "Участок 1",
            "work_type_id": work_type_id,
            "work_type_code": "FILL",
            "work_type_name": "Отсыпка",
            "unit": "м3",
            "pk_start": None,
            "pk_end": None,
            "project_volume": 100,
            "fact_volume": 5,
            "fact_pk_ranges": None,
            "project_item_id": "project-1",
            "project_item_count": 1,
        }
        period_fact_row = {
            "object_id": "obj-1",
            "work_type_id": work_type_id,
            "unit": "м3",
            "report_date": date(2026, 5, 2),
            "fact_volume": 5,
        }

        def fake_query(sql, params=None):
            calls.append((sql, params or []))
            if len(calls) == 1:
                return []
            if len(calls) == 2:
                return [main_row]
            return [period_fact_row]

        with (
            mock.patch.object(wip_routes, "query", side_effect=fake_query),
            mock.patch.object(wip_routes, "_statement_sections", return_value=[]),
            mock.patch.object(wip_routes, "_statement_object_types_payload", return_value=[]),
            mock.patch.object(wip_routes, "_statement_work_types_payload", return_value=[]),
        ):
            payload = wip_routes.statements(
                period_from="2026-05-01",
                period_to="2026-05-03",
                work_type_id=work_type_id,
                period_mode="day",
                _user=None,
            )

        self.assertEqual([period["key"] for period in payload["periods"]], ["2026-05-01", "2026-05-02", "2026-05-03"])
        self.assertEqual(payload["filters"]["period_from"], "2026-05-01")
        self.assertEqual(payload["filters"]["period_to"], "2026-05-03")
        self.assertEqual(payload["filters"]["work_type_id"], work_type_id)
        self.assertIn("dwi.work_type_id = ANY(%s::uuid[])", calls[1][0])
        self.assertIn("dr.report_date >= %s::date", calls[1][0])
        self.assertIn("COALESCE(dr.status, 'confirmed') IN ('confirmed', 'approved', 'pending_review', 'review')", calls[1][0])
        self.assertIn([work_type_id], calls[1][1])
        self.assertIn("fact_items AS", calls[2][0])
        self.assertIn("COALESCE(dr.status, 'confirmed') IN ('confirmed', 'approved', 'pending_review', 'review')", calls[2][0])
        self.assertIn("SUM(COALESCE(ds.volume_segment, 0) * row_scope.row_pk_factor)", calls[2][0])
        self.assertIn("ELSE MAX(dwi.volume)", calls[2][0])
        values_by_key = {item["period_key"]: item for item in payload["rows"][0]["period_values"]}
        self.assertEqual(values_by_key["2026-05-02"]["fact_volume"], 5.0)
        self.assertEqual(values_by_key["2026-05-01"]["fact_volume"], 0.0)



class StatementsFactSqlTests(unittest.TestCase):
    def test_general_workbook_fact_queries_include_reports_on_review(self):
        calls = []

        def fake_query(sql, params=None):
            calls.append(sql)
            return []

        with (
            mock.patch.object(wip_routes, "query", side_effect=fake_query),
            mock.patch.object(wip_routes, "_general_report_work_rules", return_value=[]),
            mock.patch.object(wip_routes, "_statement_add_quarry_cumulative_totals", return_value=None),
        ):
            wip_routes._statement_sand_by_section(date(2026, 5, 1), date(2026, 5, 31), None)
            wip_routes._statement_general_add_fact_rows(
                {},
                days=[],
                month_start=date(2026, 5, 1),
                month_end=date(2026, 5, 31),
                cumulative_end=date(2026, 5, 31),
                section_codes=None,
                object_type=None,
                object_id=None,
                q=None,
            )

        self.assertEqual(len(calls), 2)
        for sql in calls:
            self.assertIn("COALESCE(dr_scope.status, 'confirmed') IN ('confirmed', 'approved', 'pending_review', 'review')", sql)
            self.assertIn("daily_reports dr_scope", sql)

    def test_statement_quarry_cumulative_query_includes_reports_on_review(self):
        calls = []

        def fake_query(sql, params=None):
            calls.append((sql, params or []))
            return []

        with mock.patch.object(wip_routes, "query", side_effect=fake_query):
            rows = wip_routes._statement_quarry_cumulative_rows(date(2026, 5, 31), ["UCH_1"])

        self.assertEqual(rows, [])
        self.assertEqual(len(calls), 1)
        self.assertIn("COALESCE(dr_scope.status, 'confirmed') IN ('confirmed', 'approved', 'pending_review', 'review')", calls[0][0])
        self.assertIn("cs.code = ANY(%s)", calls[0][0])

    def test_statement_quarry_cumulative_totals_do_not_use_dashboard_material_flow(self):
        with (
            mock.patch.object(wip_routes, "_statement_quarry_cumulative_rows", return_value=[]),
            mock.patch.object(wip_routes, "dashboard_material_flow", side_effect=AssertionError("unexpected material-flow call")),
        ):
            wip_routes._statement_add_quarry_cumulative_totals({}, cumulative_end=date(2026, 5, 31), section_codes=None)

    def test_rows_payload_uses_segment_volumes_for_fact_totals_and_periods(self):
        calls = []

        def fake_query(sql, params=None):
            calls.append((sql, params or []))
            return []

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "query", side_effect=fake_query),
            mock.patch.object(wip_routes, "_statement_metadata_payload", return_value={"schema": {}, "sections": [], "object_types": [], "objects": [], "work_types": []}),
        ):
            payload = wip_routes._statement_rows_payload(
                section=None,
                object_id=None,
                pk_from="ПК1+00",
                pk_to="ПК2+00",
                period_from="2026-05-01",
                period_to="2026-05-03",
                work_type_id="22222222-2222-2222-2222-222222222222",
                source="auto",
                _user=None,
                period_mode="day",
                object_type=None,
            )

        self.assertEqual(payload["rows"], [])
        self.assertGreaterEqual(len(calls), 3)
        row_sql = calls[0][0]
        period_fact_sql = calls[2][0]
        for sql in (row_sql, period_fact_sql):
            self.assertIn("fact_items AS", sql)
            self.assertIn("COALESCE(dr.status, 'confirmed') IN ('confirmed', 'approved', 'pending_review', 'review')", sql)
            self.assertIn("SUM(COALESCE(ds.volume_segment, 0) * row_scope.row_pk_factor)", sql)
            self.assertIn("ELSE MAX(dwi.volume)", sql)
        self.assertNotIn("SUM(COALESCE(dwi.volume, 0))::numeric AS fact_volume", row_sql)
        self.assertNotIn("SUM(COALESCE(dwi.volume, 0))::numeric AS fact_volume", period_fact_sql)

    def test_rows_payload_uses_project_segments_for_project_totals(self):
        calls = []

        def fake_query(sql, params=None):
            calls.append((sql, params or []))
            return []

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "query", side_effect=fake_query),
            mock.patch.object(wip_routes, "_statement_metadata_payload", return_value={"schema": {}, "sections": [], "object_types": [], "objects": [], "work_types": []}),
        ):
            payload = wip_routes._statement_rows_payload(
                section=None,
                object_id=None,
                pk_from="ПК1+00",
                pk_to="ПК2+00",
                period_from=None,
                period_to=None,
                work_type_id="22222222-2222-2222-2222-222222222222",
                source="auto",
                _user=None,
                period_mode="day",
                object_type=None,
            )

        self.assertEqual(payload["rows"], [])
        self.assertGreaterEqual(len(calls), 1)
        row_sql = calls[0][0]
        self.assertIn("project_items AS", row_sql)
        self.assertIn("LEFT JOIN project_work_item_segments ps ON ps.project_work_item_id = pwi.id", row_sql)
        self.assertIn("LEFT JOIN pile_fields pf ON pf.id = pwi.source_pile_field_id", row_sql)
        self.assertIn("COALESCE(ps.pk_end, pf.pk_end, so.pk_end) IS NOT NULL AND COALESCE(ps.pk_end, pf.pk_end, so.pk_end) >= %s", row_sql)
        self.assertIn("COALESCE(ps.pk_start, pf.pk_start, so.pk_start) IS NOT NULL AND COALESCE(ps.pk_start, pf.pk_start, so.pk_start) <= %s", row_sql)
        self.assertIn("SUM(COALESCE(ps.volume_segment, 0) * row_scope.row_pk_factor)", row_sql)
        self.assertIn("ELSE MAX(pwi.project_volume)", row_sql)
        self.assertIn("COUNT(ps.id) FILTER (WHERE ps.id IS NOT NULL) AS project_segment_count", row_sql)
        self.assertIn("SUM(COALESCE(project_segment_count, 0))::int AS project_segment_count", row_sql)
        self.assertIn("COALESCE(pa.project_segment_count, 0) AS project_segment_count", row_sql)
        self.assertIn("COALESCE(pwi.project_volume, 0) <> 0 OR COALESCE(ps.volume_segment, 0) <> 0", row_sql)
        self.assertNotIn("SUM(COALESCE(pwi.project_volume, 0))::numeric AS project_volume", row_sql)

    def test_rows_payload_keeps_period_values_separate_by_pk_segment(self):
        calls = []
        base_row = {
            "plan_id": "plan-id",
            "object_id": "obj-1",
            "object_code": "OBJ",
            "object_name": "Объект",
            "object_type_code": "PIPE",
            "object_type_name": "Трубы",
            "section_code": "UCH_1",
            "section_name": "Участок 1",
            "work_type_id": "work-1",
            "work_type_code": "FILL",
            "work_type_name": "Отсыпка",
            "unit": "м3",
            "pk_raw_text": None,
            "project_volume": 0,
            "planned_volume": 0,
            "fact_volume": 0,
            "fact_pk_ranges": None,
            "project_item_id": None,
            "project_item_count": 0,
            "project_segment_count": 0,
            "period_start": None,
            "period_end": None,
            "plan_type": "manual_period",
            "source_reference": "manual:test",
            "comment": None,
            "updated_at": None,
        }
        row_rows = [
            {**base_row, "plan_id": "plan-a", "pk_start": 100.0, "pk_end": 200.0},
            {**base_row, "plan_id": "plan-b", "pk_start": 200.0, "pk_end": 300.0},
        ]
        plan_period_rows = [
            {"object_id": "obj-1", "work_type_id": "work-1", "unit": "м3", "pk_start": 100.0, "pk_end": 200.0, "period_start": date(2026, 5, 1), "period_end": date(2026, 5, 1), "planned_volume": 3},
            {"object_id": "obj-1", "work_type_id": "work-1", "unit": "м3", "pk_start": 200.0, "pk_end": 300.0, "period_start": date(2026, 5, 1), "period_end": date(2026, 5, 1), "planned_volume": 7},
        ]
        fact_period_rows = [
            {"object_id": "obj-1", "work_type_id": "work-1", "unit": "м3", "pk_start": 100.0, "pk_end": 200.0, "report_date": date(2026, 5, 1), "fact_volume": 2},
            {"object_id": "obj-1", "work_type_id": "work-1", "unit": "м3", "pk_start": 200.0, "pk_end": 300.0, "report_date": date(2026, 5, 1), "fact_volume": 5},
        ]

        def fake_query(sql, params=None):
            calls.append((sql, params or []))
            if len(calls) == 1:
                return row_rows
            if len(calls) == 2:
                return plan_period_rows
            if len(calls) == 3:
                return fact_period_rows
            return []

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "query", side_effect=fake_query),
            mock.patch.object(wip_routes, "_statement_metadata_payload", return_value={"schema": {}, "sections": [], "object_types": [], "objects": [], "work_types": []}),
        ):
            payload = wip_routes._statement_rows_payload(
                section=None,
                object_id=None,
                pk_from=None,
                pk_to=None,
                period_from="2026-05-01",
                period_to="2026-05-01",
                work_type_id=None,
                source="auto",
                _user=None,
                period_mode="day",
                object_type=None,
            )

        self.assertEqual(len(payload["rows"]), 2)
        by_pk = {row["pk_start"]: row for row in payload["rows"]}
        first_period = {item["period_key"]: item for item in by_pk[100.0]["period_values"]}["2026-05-01"]
        second_period = {item["period_key"]: item for item in by_pk[200.0]["period_values"]}["2026-05-01"]
        self.assertEqual(first_period["planned_volume"], 3.0)
        self.assertEqual(first_period["fact_volume"], 2.0)
        self.assertEqual(second_period["planned_volume"], 7.0)
        self.assertEqual(second_period["fact_volume"], 5.0)
        self.assertIn("row_scope.row_pk_start AS pk_start", calls[1][0])
        self.assertIn("row_scope.row_pk_end AS pk_end", calls[1][0])
        self.assertIn("row_scope.row_pk_start AS pk_start", calls[2][0])
        self.assertIn("LEFT JOIN pile_fields pf ON pf.id = ds.pile_field_id", calls[2][0])
        self.assertIn("GROUP BY object_id, work_type_id, unit, pk_start, pk_end, report_date", calls[2][0])

    def test_rows_payload_totals_do_not_duplicate_project_fact_across_pk_plan_rows(self):
        calls = []
        base_row = {
            "plan_id": "plan-id",
            "object_id": "obj-1",
            "object_code": "OBJ",
            "object_name": "Объект",
            "object_type_code": "PIPE",
            "object_type_name": "Трубы",
            "section_code": "UCH_1",
            "section_name": "Участок 1",
            "work_type_id": "work-1",
            "work_type_code": "FILL",
            "work_type_name": "Отсыпка",
            "unit": "м3",
            "pk_raw_text": None,
            "project_volume": 10,
            "fact_volume": 4,
            "fact_pk_ranges": "ПК1+00-ПК3+00",
            "project_item_id": "project-id",
            "project_item_count": 1,
            "project_segment_count": 2,
            "period_start": None,
            "period_end": None,
            "plan_type": "manual_period",
            "source_reference": "manual:test",
            "comment": None,
            "updated_at": None,
        }
        row_rows = [
            {**base_row, "plan_id": "plan-a", "pk_start": 100.0, "pk_end": 200.0, "planned_volume": 3},
            {**base_row, "plan_id": "plan-b", "pk_start": 200.0, "pk_end": 300.0, "planned_volume": 7},
        ]

        def fake_query(sql, params=None):
            calls.append((sql, params or []))
            return row_rows if len(calls) == 1 else []

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "query", side_effect=fake_query),
            mock.patch.object(wip_routes, "_statement_metadata_payload", return_value={"schema": {}, "sections": [], "object_types": [], "objects": [], "work_types": []}),
        ):
            payload = wip_routes._statement_rows_payload(
                section=None,
                object_id=None,
                pk_from=None,
                pk_to=None,
                period_from=None,
                period_to=None,
                work_type_id=None,
                source="auto",
                _user=None,
                period_mode="day",
                object_type=None,
            )

        self.assertEqual(len(payload["rows"]), 2)
        self.assertEqual(payload["totals"], [{
            "unit": "м3",
            "project_volume": 10.0,
            "planned_volume": 10.0,
            "fact_volume": 4.0,
            "remaining_volume": 6.0,
            "percent_complete": 40.0,
        }])

    def test_rows_payload_exposes_project_segment_count(self):
        base = {
            "plan_id": None,
            "object_id": "obj-1",
            "object_code": "OBJ",
            "object_name": "Объект",
            "object_type_code": "PIPE",
            "object_type_name": "Трубы",
            "section_code": "UCH_1",
            "section_name": "Участок 1",
            "work_type_id": "work-1",
            "work_type_code": "FILL",
            "work_type_name": "Отсыпка",
            "unit": "м3",
            "pk_start": 100.0,
            "pk_end": 200.0,
            "pk_raw_text": "ПК1+00-ПК2+00",
            "project_volume": 10,
            "planned_volume": 0,
            "fact_volume": 0,
            "fact_pk_ranges": None,
            "project_item_id": "project-id",
            "project_item_count": 1,
            "project_segment_count": 2,
            "period_start": None,
            "period_end": None,
            "plan_type": "project_fact_without_plan",
            "source_reference": "auto:project_or_fact",
            "comment": None,
            "updated_at": None,
        }

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "query", return_value=[base]),
            mock.patch.object(wip_routes, "_statement_metadata_payload", return_value={"schema": {}, "sections": [], "object_types": [], "objects": [], "work_types": []}),
        ):
            payload = wip_routes._statement_rows_payload(
                section=None,
                object_id=None,
                pk_from=None,
                pk_to=None,
                period_from=None,
                period_to=None,
                work_type_id=None,
                source="auto",
                _user=None,
                period_mode="day",
                object_type=None,
            )

        self.assertEqual(payload["rows"][0]["project_segment_count"], 2)
        self.assertTrue(payload["rows"][0]["can_edit_project"])

    def test_rows_payload_fact_source_returns_only_fact_rows_and_keeps_unplanned_union(self):
        calls = []
        base = {
            "plan_id": None,
            "object_type_code": "PIPE",
            "object_type_name": "Трубы",
            "section_code": "UCH_1",
            "section_name": "Участок 1",
            "work_type_id": "work-1",
            "work_type_code": "FILL",
            "work_type_name": "Отсыпка",
            "unit": "м3",
            "pk_start": 100.0,
            "pk_end": 200.0,
            "pk_raw_text": None,
            "project_volume": 10,
            "planned_volume": 0,
            "fact_pk_ranges": None,
            "project_item_id": None,
            "project_item_count": 0,
            "period_start": None,
            "period_end": None,
            "plan_type": "project_fact_without_plan",
            "source_reference": "auto:project_or_fact",
            "comment": None,
            "updated_at": None,
        }

        def fake_query(sql, params=None):
            calls.append((sql, params or []))
            return [
                {**base, "object_id": "obj-plan", "object_code": "PLAN", "object_name": "План без факта", "fact_volume": 0},
                {**base, "object_id": "obj-fact", "object_code": "FACT", "object_name": "Факт", "fact_volume": 7},
            ]

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "query", side_effect=fake_query),
            mock.patch.object(wip_routes, "_statement_metadata_payload", return_value={"schema": {}, "sections": [], "object_types": [], "objects": [], "work_types": []}),
        ):
            payload = wip_routes._statement_rows_payload(
                section=None,
                object_id=None,
                pk_from=None,
                pk_to=None,
                period_from=None,
                period_to=None,
                work_type_id=None,
                source="fact",
                _user=None,
                period_mode="day",
                object_type=None,
            )

        self.assertEqual(payload["source"], "fact_daily_work_items")
        self.assertEqual(payload["row_count"], 1)
        self.assertEqual(payload["rows"][0]["object_code"], "FACT")
        self.assertEqual(payload["totals"][0]["fact_volume"], 7.0)
        self.assertIn("UNION ALL", calls[0][0])


class StatementsRowsPayloadMetadataTests(unittest.TestCase):
    def test_rows_payload_includes_dictionaries_even_when_rows_are_empty(self):
        metadata = {
            "schema": {"editable": True},
            "sections": [{"code": "UCH_1", "name": "Участок 1"}],
            "object_types": [{"code": "PIPE", "name": "Трубы"}],
            "objects": [{"id": "obj-1", "object_code": "OBJ", "object_name": "Объект"}],
            "work_types": [{"id": "work-1", "code": "FILL", "name": "Отсыпка", "default_unit": "м3"}],
        }

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "query", return_value=[]),
            mock.patch.object(wip_routes, "_statement_metadata_payload", return_value=metadata),
        ):
            payload = wip_routes._statement_rows_payload(
                section=None,
                object_id=None,
                pk_from=None,
                pk_to=None,
                period_from=None,
                period_to=None,
                work_type_id=None,
                source="auto",
                _user=None,
                period_mode="month_third",
                object_type=None,
            )

        self.assertEqual(payload["rows"], [])
        self.assertEqual(payload["sections"], metadata["sections"])
        self.assertEqual(payload["object_types"], metadata["object_types"])
        self.assertEqual(payload["objects"], metadata["objects"])
        self.assertEqual(payload["work_types"], metadata["work_types"])

class SequencedCursor:
    def __init__(self, fetch_results: list[object]):
        self.fetch_results = list(fetch_results)
        self.executed: list[tuple[str, object]] = []
        self.rowcount = 0

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False

    def execute(self, sql: str, params=None):
        self.executed.append((sql, params))
        self.rowcount = 1

    def fetchone(self):
        if not self.fetch_results:
            return None
        return self.fetch_results.pop(0)

    def fetchall(self):
        if not self.fetch_results:
            return []
        result = self.fetch_results.pop(0)
        if result is None:
            return []
        if isinstance(result, list):
            return result
        return [result]


class SequencedConn:
    def __init__(self, fetch_results: list[object]):
        self.cursor_obj = SequencedCursor(fetch_results)
        self.closed = False

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False

    def cursor(self):
        return self.cursor_obj

    def close(self):
        self.closed = True


class StatementsObjectCreateTests(unittest.TestCase):
    def test_create_object_uses_selected_section_boundaries_when_pk_is_absent(self):
        conn = SequencedConn([
            ("type-id",),
            ("section-id", "UCH_1", "Участок 1", 100.0, 200.0, "ПК1+00-ПК2+00"),
            ("object-id", "OBJ_1", "Объект 1"),
            None,
        ])
        body = types.SimpleNamespace(
            object_code="obj-1",
            object_name="Объект 1",
            object_type_code="MAIN_TRACK",
            section_code="UCH_1",
            pk_start=None,
            pk_end=None,
            pk_raw_text=None,
            comment=None,
        )

        with mock.patch.object(wip_routes, "get_conn", return_value=conn):
            result = wip_routes.statements_create_object(body, _user=None)

        self.assertEqual(result["section_code"], "UCH_1")
        self.assertEqual(result["section_name"], "Участок 1")
        self.assertEqual(result["pk_start"], 100.0)
        self.assertEqual(result["pk_end"], 200.0)
        segment_inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO object_segments" in entry[0]]
        self.assertEqual(len(segment_inserts), 1)
        self.assertEqual(segment_inserts[0][1][1:4], (100.0, 200.0, "ПК1+00-ПК2+00"))
        self.assertTrue(conn.closed)

    def test_create_object_rejects_pk_outside_selected_section(self):
        conn = SequencedConn([
            ("type-id",),
            ("section-id", "UCH_1", "Участок 1", 100.0, 200.0, "ПК1+00-ПК2+00"),
        ])
        body = types.SimpleNamespace(
            object_code="obj-2",
            object_name="Объект 2",
            object_type_code="MAIN_TRACK",
            section_code="UCH_1",
            pk_start=250.0,
            pk_end=260.0,
            pk_raw_text=None,
            comment=None,
        )

        with mock.patch.object(wip_routes, "get_conn", return_value=conn):
            with self.assertRaises(Exception) as raised:
                wip_routes.statements_create_object(body, _user=None)

        self.assertEqual(getattr(raised.exception, "status_code", None), 400)
        object_inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO objects" in entry[0]]
        self.assertEqual(object_inserts, [])
        self.assertTrue(conn.closed)

    def test_create_optional_pk_object_keeps_empty_pk_even_with_section(self):
        conn = SequencedConn([
            ("type-id", "BORROW_PIT", False),
            ("section-id", "UCH_1", "Участок 1", 100.0, 200.0, "ПК1+00-ПК2+00"),
            ("object-id", "BORROW", "Карьер"),
        ])
        body = types.SimpleNamespace(
            object_code="borrow",
            object_name="Карьер",
            object_type_code="BORROW_PIT",
            section_code="UCH_1",
            pk_start=None,
            pk_end=None,
            pk_raw_text=None,
            comment=None,
        )

        with mock.patch.object(wip_routes, "get_conn", return_value=conn):
            result = wip_routes.statements_create_object(body, _user=None)

        self.assertEqual(result["section_code"], "UCH_1")
        self.assertIsNone(result["pk_start"])
        self.assertIsNone(result["pk_end"])
        segment_inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO object_segments" in entry[0]]
        self.assertEqual(segment_inserts, [])
        self.assertTrue(conn.closed)

    def test_clear_optional_object_picketage_deletes_object_segments(self):
        conn = SequencedConn([
            ("object-id", "BORROW", "Карьер", "BORROW_PIT", "Карьер", False),
            ("object-id", "BORROW", "Карьер", "BORROW_PIT", "Карьер", None, None, None, None),
        ])
        body = types.SimpleNamespace(clear=True, pk_start=None, pk_end=None, pk_raw_text=None, section_code=None, comment=None)

        with mock.patch.object(wip_routes, "get_conn", return_value=conn):
            result = wip_routes.statements_update_object_picketage("object-id", body, _user=None)

        self.assertIsNone(result["pk_start"])
        self.assertIsNone(result["pk_end"])
        deletes = [entry for entry in conn.cursor_obj.executed if "DELETE FROM object_segments" in entry[0]]
        self.assertEqual(len(deletes), 1)
        segment_inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO object_segments" in entry[0]]
        self.assertEqual(segment_inserts, [])
        self.assertTrue(conn.closed)

    def test_clear_route_object_picketage_is_rejected(self):
        conn = SequencedConn([
            ("object-id", "PIPE", "Труба", "PIPE", "Труба", False),
        ])
        body = types.SimpleNamespace(clear=True, pk_start=None, pk_end=None, pk_raw_text=None, section_code=None, comment=None)

        with mock.patch.object(wip_routes, "get_conn", return_value=conn):
            with self.assertRaises(Exception) as raised:
                wip_routes.statements_update_object_picketage("object-id", body, _user=None)

        self.assertEqual(getattr(raised.exception, "status_code", None), 400)
        deletes = [entry for entry in conn.cursor_obj.executed if "DELETE FROM object_segments" in entry[0]]
        self.assertEqual(deletes, [])
        self.assertTrue(conn.closed)


class StatementsWorkTypeCreateTests(unittest.TestCase):
    def test_create_work_type_normalizes_code_and_returns_unit(self):
        conn = SequencedConn([
            ("work-id", "MANUAL_FILL", "Ручная отсыпка", "м3", "manual"),
        ])
        body = types.SimpleNamespace(
            code="manual fill",
            name="  Ручная отсыпка  ",
            default_unit="м3",
            work_group="",
            analytics_tag=" earth ",
            productivity_enabled=True,
        )

        with mock.patch.object(wip_routes, "get_conn", return_value=conn):
            result = wip_routes.statements_create_work_type(body, _user=None)

        self.assertEqual(result, {
            "id": "work-id",
            "code": "MANUAL_FILL",
            "name": "Ручная отсыпка",
            "default_unit": "м3",
            "unit": "м3",
            "work_group": "manual",
        })
        inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO work_types" in entry[0]]
        self.assertEqual(len(inserts), 1)
        sql, params = inserts[0]
        self.assertIn("ON CONFLICT (code) DO UPDATE", sql)
        self.assertEqual(params, ("MANUAL_FILL", "Ручная отсыпка", "м3", "manual", "earth", True, "pending_admin_review"))
        self.assertTrue(conn.closed)

    def test_create_work_type_rejects_empty_normalized_code_before_db_work(self):
        body = types.SimpleNamespace(
            code="!!!",
            name="Работа без кода",
            default_unit="м3",
            work_group="manual",
            analytics_tag=None,
            productivity_enabled=True,
        )

        with mock.patch.object(wip_routes, "get_conn") as get_conn:
            with self.assertRaises(Exception) as raised:
                wip_routes.statements_create_work_type(body, _user=None)

        self.assertEqual(getattr(raised.exception, "status_code", None), 400)
        self.assertEqual(getattr(raised.exception, "detail", None), "code is required")
        get_conn.assert_not_called()


class StatementsPlannedRowCreateTests(unittest.TestCase):
    def test_create_planned_row_validates_object_and_work_before_insert(self):
        conn = SequencedConn([("plan-id",)])
        body = types.SimpleNamespace(
            object_id="11111111-1111-1111-1111-111111111111",
            work_type_id="22222222-2222-2222-2222-222222222222",
            planned_volume=12.5,
            unit="м3",
            period_start="2026-05-01",
            period_end="2026-05-10",
            pk_start=100.0,
            pk_end=200.0,
            pk_raw_text="ПК1+00-ПК2+00",
            comment="manual planned",
        )

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = wip_routes.statements_create_row(body, _user=None)

        self.assertEqual(result, {"ok": True, "plan_id": "plan-id"})
        inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO planned_work_items" in entry[0]]
        self.assertEqual(len(inserts), 1)
        sql, params = inserts[0]
        self.assertIn("FROM objects o", sql)
        self.assertIn("CROSS JOIN work_types wt", sql)
        self.assertIn("WHERE o.id = %s::uuid", sql)
        self.assertEqual(params[0], "м3")
        self.assertEqual(params[-2:], (body.object_id, body.work_type_id))
        self.assertTrue(conn.closed)

    def test_create_planned_row_returns_404_when_object_or_work_is_missing(self):
        conn = SequencedConn([None])
        body = types.SimpleNamespace(
            object_id="11111111-1111-1111-1111-111111111111",
            work_type_id="22222222-2222-2222-2222-222222222222",
            planned_volume=12.5,
            unit="м3",
            period_start=None,
            period_end=None,
            pk_start=None,
            pk_end=None,
            pk_raw_text=None,
            comment=None,
        )

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            with self.assertRaises(Exception) as raised:
                wip_routes.statements_create_row(body, _user=None)

        self.assertEqual(getattr(raised.exception, "status_code", None), 404)
        self.assertEqual(getattr(raised.exception, "detail", None), "object_id or work_type_id not found")
        inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO planned_work_items" in entry[0]]
        self.assertEqual(len(inserts), 1)
        self.assertTrue(conn.closed)


class StatementsProjectRowCreateTests(unittest.TestCase):
    def test_create_project_row_updates_existing_object_work_unit(self):
        conn = SequencedConn([("project-id",)])
        body = types.SimpleNamespace(
            object_id="11111111-1111-1111-1111-111111111111",
            work_type_id="22222222-2222-2222-2222-222222222222",
            project_volume=42.5,
            unit="м3",
            comment="manual project",
        )

        with mock.patch.object(wip_routes, "get_conn", return_value=conn):
            result = wip_routes.statements_create_project_row(body, _user=None)

        self.assertEqual(result, {"ok": True, "created": False, "project_item_id": "project-id"})
        updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_items" in entry[0]]
        inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO project_work_items" in entry[0]]
        self.assertEqual(len(updates), 1)
        self.assertEqual(inserts, [])
        self.assertEqual(updates[0][1], (42.5, "м3", "manual project", body.object_id, body.work_type_id, "м3"))
        self.assertTrue(conn.closed)

    def test_create_project_row_inserts_when_no_existing_row(self):
        conn = SequencedConn([None, ("new-project-id",)])
        body = types.SimpleNamespace(
            object_id="11111111-1111-1111-1111-111111111111",
            work_type_id="22222222-2222-2222-2222-222222222222",
            project_volume=18,
            unit="шт",
            comment=None,
        )

        with mock.patch.object(wip_routes, "get_conn", return_value=conn):
            result = wip_routes.statements_create_project_row(body, _user=None)

        self.assertEqual(result, {"ok": True, "created": True, "project_item_id": "new-project-id"})
        inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO project_work_items" in entry[0]]
        self.assertEqual(len(inserts), 1)
        self.assertEqual(inserts[0][1], (18.0, "шт", None, body.object_id, body.work_type_id))
        self.assertTrue(conn.closed)

    def test_create_project_row_rejects_negative_volume_before_db_work(self):
        body = types.SimpleNamespace(
            object_id="11111111-1111-1111-1111-111111111111",
            work_type_id="22222222-2222-2222-2222-222222222222",
            project_volume=-1,
            unit="м3",
            comment=None,
        )

        with mock.patch.object(wip_routes, "get_conn") as get_conn_mock:
            with self.assertRaises(Exception) as raised:
                wip_routes.statements_create_project_row(body, _user=None)

        self.assertEqual(getattr(raised.exception, "status_code", None), 400)
        get_conn_mock.assert_not_called()

    def test_create_project_row_with_pk_updates_project_segment(self):
        conn = SequencedConn([("project-id",)])
        body = types.SimpleNamespace(
            object_id="11111111-1111-1111-1111-111111111111",
            work_type_id="22222222-2222-2222-2222-222222222222",
            project_volume=42.5,
            unit="м3",
            pk_start=100.0,
            pk_end=200.0,
            pk_raw_text="ПК1+00-ПК2+00",
            comment="manual segment project",
        )

        with mock.patch.object(wip_routes, "get_conn", return_value=conn):
            result = wip_routes.statements_create_project_row(body, _user=None)

        self.assertEqual(result, {"ok": True, "created": False, "project_item_id": "project-id", "segment": {"created": False}})
        project_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_items" in entry[0]]
        self.assertEqual(len(project_updates), 1)
        self.assertNotIn("SET project_volume = %s", project_updates[0][0])
        segment_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_item_segments" in entry[0]]
        self.assertEqual(len(segment_updates), 1)
        self.assertEqual(segment_updates[0][1], (42.5, "ПК1+00-ПК2+00", "manual segment project", "project-id", 100.0, 200.0))
        self.assertTrue(conn.closed)

    def test_create_project_row_rejects_inverted_pk_before_db_work(self):
        body = types.SimpleNamespace(
            object_id="11111111-1111-1111-1111-111111111111",
            work_type_id="22222222-2222-2222-2222-222222222222",
            project_volume=1,
            unit="м3",
            pk_start=200.0,
            pk_end=100.0,
            pk_raw_text=None,
            comment=None,
        )

        with mock.patch.object(wip_routes, "get_conn") as get_conn_mock:
            with self.assertRaises(Exception) as raised:
                wip_routes.statements_create_project_row(body, _user=None)

        self.assertEqual(getattr(raised.exception, "status_code", None), 400)
        get_conn_mock.assert_not_called()


class StatementsProjectRowUpdateTests(unittest.TestCase):
    def test_update_project_row_without_pk_updates_parent_volume(self):
        conn = SequencedConn([("project-id",)])
        body = types.SimpleNamespace(
            project_volume=42.5,
            unit="м3",
            comment="manual project",
        )

        with mock.patch.object(wip_routes, "get_conn", return_value=conn):
            result = wip_routes.statements_update_project_row("project-id", body, _user=None)

        self.assertEqual(result, {"ok": True, "changed": 1, "project_item_id": "project-id"})
        project_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_items" in entry[0]]
        segment_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_item_segments" in entry[0]]
        self.assertEqual(len(project_updates), 1)
        self.assertIn("project_volume = %s", project_updates[0][0])
        self.assertEqual(project_updates[0][1], [42.5, "м3", "manual project", "project-id"])
        self.assertEqual(segment_updates, [])
        self.assertTrue(conn.closed)

    def test_update_project_row_with_pk_updates_segment_not_parent_volume(self):
        conn = SequencedConn([("project-id",)])
        body = types.SimpleNamespace(
            project_volume=55,
            unit="м3",
            pk_start=100.0,
            pk_end=200.0,
            pk_raw_text="ПК1+00-ПК2+00",
            comment="segment update",
        )

        with mock.patch.object(wip_routes, "get_conn", return_value=conn):
            result = wip_routes.statements_update_project_row("project-id", body, _user=None)

        self.assertEqual(result, {"ok": True, "changed": 1, "project_item_id": "project-id", "segment": {"changed": 1, "created": False}})
        project_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_items" in entry[0]]
        self.assertEqual(len(project_updates), 1)
        self.assertNotIn("project_volume = %s", project_updates[0][0])
        segment_updates = [entry for entry in conn.cursor_obj.executed if "UPDATE project_work_item_segments" in entry[0]]
        self.assertEqual(len(segment_updates), 1)
        self.assertEqual(segment_updates[0][1], (55.0, "ПК1+00-ПК2+00", "segment update", "project-id", 100.0, 200.0))
        self.assertTrue(conn.closed)

    def test_update_project_row_rejects_inverted_pk_before_db_work(self):
        body = types.SimpleNamespace(
            project_volume=1,
            unit="м3",
            pk_start=200.0,
            pk_end=100.0,
            pk_raw_text=None,
            comment=None,
        )

        with mock.patch.object(wip_routes, "get_conn") as get_conn_mock:
            with self.assertRaises(Exception) as raised:
                wip_routes.statements_update_project_row("project-id", body, _user=None)

        self.assertEqual(getattr(raised.exception, "status_code", None), 400)
        get_conn_mock.assert_not_called()


class StatementsPeriodPlanUpsertTests(unittest.TestCase):
    def test_period_plan_upsert_matches_pk_range_when_supplied(self):
        conn = SequencedConn([[('plan-id',)]])
        body = types.SimpleNamespace(
            object_id="11111111-1111-1111-1111-111111111111",
            work_type_id="22222222-2222-2222-2222-222222222222",
            planned_volume=9.5,
            unit="м3",
            period_start="2026-05-01",
            period_end="2026-05-10",
            pk_start=100.0,
            pk_end=200.0,
            pk_raw_text="ПК1+00-ПК2+00",
            comment="segment plan",
        )

        with (
            mock.patch.object(wip_routes, "_statement_table_exists", return_value=True),
            mock.patch.object(wip_routes, "get_conn", return_value=conn),
        ):
            result = wip_routes.statements_upsert_period_plan(body, _user=None)

        self.assertEqual(result, {"ok": True, "updated": 1, "inserted": 0, "plan_ids": ["plan-id"]})
        updates = [entry for entry in conn.cursor_obj.executed if "UPDATE planned_work_items" in entry[0]]
        inserts = [entry for entry in conn.cursor_obj.executed if "INSERT INTO planned_work_items" in entry[0]]
        self.assertEqual(len(updates), 1)
        self.assertEqual(inserts, [])
        sql, params = updates[0]
        self.assertIn("pk_start IS NOT DISTINCT FROM", sql)
        self.assertIn("pk_end IS NOT DISTINCT FROM", sql)
        self.assertEqual(params[-4:], (100.0, 100.0, 200.0, 200.0))
        self.assertTrue(conn.closed)


class StatementsExportHeaderTests(unittest.TestCase):
    def test_export_metadata_sheet_roundtrips_period_definitions(self):
        wb = Workbook()
        payload = {
            "filters": {"period_from": "2026-05-01", "period_to": "2026-05-10", "q": "АД5 песок"},
            "periods": [{"key": "2026-05-01_2026-05-10", "label": "1/3 05.2026", "start": "2026-05-01", "end": "2026-05-10"}],
        }

        wip_routes._statement_write_workbook_metadata(
            wb,
            payload,
            period_mode="month_third",
            source="auto",
            display_mode="section_tree",
            group_by="section",
            columns="unit,planned_volume",
        )
        self.assertEqual(wb[wip_routes.STATEMENT_EXPORT_META_SHEET].sheet_state, "hidden")
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        loaded = load_workbook(buf, data_only=True)
        meta_ws = loaded[wip_routes.STATEMENT_EXPORT_META_SHEET]
        meta_values = {str(row[0].value): row[1].value for row in meta_ws.iter_rows(min_row=1, max_col=2) if row[0].value}
        self.assertEqual(meta_values["q"], "АД5 песок")

        periods = wip_routes._statement_workbook_periods(loaded)

        self.assertEqual(periods, [{
            "key": "2026-05-01_2026-05-10",
            "label": "1/3 05.2026",
            "start_date": date(2026, 5, 1),
            "end_date": date(2026, 5, 10),
        }])

    def test_excel_headers_include_all_visible_statement_columns_and_duplicates(self):
        headers = wip_routes._statement_excel_headers(
            "pk_raw_text,plan_type,source_reference,updated_at,unit,unit"
        )

        self.assertEqual(
            headers,
            [
                ("group_label", "Группа"),
                ("pk_raw_text", "ПК текст"),
                ("plan_type", "Тип плана"),
                ("source_reference", "Источник"),
                ("updated_at", "Обновлено"),
                ("unit", "Ед. изм. #1"),
                ("unit", "Ед. изм. #2"),
            ],
        )

    def test_display_headers_keep_service_columns_for_hierarchical_export(self):
        headers = wip_routes._statement_display_excel_headers(
            "section_name,object_type_name,object_name,work_type_name,pk_raw_text,source_reference,updated_at",
            "section_tree",
        )

        self.assertEqual(
            headers,
            [
                ("__structure", "Структура"),
                ("pk_raw_text", "ПК текст"),
                ("source_reference", "Источник"),
                ("updated_at", "Обновлено"),
            ],
        )

    def test_import_identity_headers_add_hidden_pk_identity_when_not_visible(self):
        visible = [
            ("__structure", "Структура"),
            ("planned_volume", "Плановый объем"),
            ("period:2026-05-01:planned_volume", "01.05 план"),
        ]

        hidden = wip_routes._statement_import_identity_headers(visible)

        self.assertIn(("pk_start", "pk_start"), hidden)
        self.assertIn(("pk_end", "pk_end"), hidden)
        self.assertIn(("pk_raw_text", "pk_raw_text"), hidden)
        self.assertIn(("project_segment_count", "project_segment_count"), hidden)
        self.assertIn(("object_id", "object_id"), hidden)
        self.assertIn(("work_type_id", "work_type_id"), hidden)

    def test_import_identity_headers_do_not_duplicate_visible_pk_columns(self):
        visible = [
            ("__structure", "Структура"),
            ("pk_start", "ПК от"),
            ("pk_end", "ПК до"),
            ("pk_raw_text", "ПК текст"),
        ]

        hidden = wip_routes._statement_import_identity_headers(visible)
        hidden_keys = [key for key, _label in hidden]

        self.assertNotIn("pk_start", hidden_keys)
        self.assertNotIn("pk_end", hidden_keys)
        self.assertNotIn("pk_raw_text", hidden_keys)

    def test_export_sheet_styles_unlock_only_editable_leaf_cells(self):
        wb = Workbook()
        ws = wb.active
        headers = [
            ("__structure", "Структура"),
            ("planned_volume", "Плановый объем"),
            ("period:2026-05-01:planned_volume", "01.05 план"),
            ("period:2026-05-01:fact_volume", "01.05 факт"),
            ("fact_volume", "Фактический объем"),
            ("comment", "Комментарий"),
            ("object_id", "object_id"),
            ("work_type_id", "work_type_id"),
            ("source_rows", "source_rows"),
            ("__is_group", "__is_group"),
        ]
        ws.append([label for _key, label in headers])
        ws.append(["Работа", 10, 4, 2, 2, "leaf", "obj-1", "work-1", 1, False])
        ws.append(["Сводная работа", 30, 12, 6, 6, "aggregate", "obj-1", "work-1", 2, False])
        ws.append(["Группа", 30, 12, 6, 6, "group", "", "", 2, True])

        wip_routes._style_statement_export_sheet(
            ws,
            headers,
            {"object_id", "work_type_id", "source_rows", "__is_group"},
            "section_tree",
        )

        self.assertTrue(ws.protection.sheet)
        self.assertEqual(ws.freeze_panes, "B2")
        self.assertEqual(ws.auto_filter.ref, ws.dimensions)
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertTrue(ws.column_dimensions["G"].hidden)
        self.assertTrue(ws.column_dimensions["H"].hidden)
        self.assertTrue(ws.column_dimensions["I"].hidden)
        self.assertTrue(ws.column_dimensions["J"].hidden)
        self.assertFalse(ws.cell(row=2, column=2).protection.locked)
        self.assertFalse(ws.cell(row=2, column=3).protection.locked)
        self.assertFalse(ws.cell(row=2, column=6).protection.locked)
        self.assertTrue(ws.cell(row=2, column=4).protection.locked)
        self.assertTrue(ws.cell(row=2, column=5).protection.locked)
        self.assertTrue(ws.cell(row=2, column=7).protection.locked)
        self.assertTrue(ws.cell(row=3, column=2).protection.locked)
        self.assertFalse(ws.cell(row=3, column=3).protection.locked)
        self.assertTrue(ws.cell(row=3, column=4).protection.locked)
        self.assertTrue(ws.cell(row=3, column=6).protection.locked)
        self.assertTrue(ws.cell(row=4, column=2).protection.locked)
        self.assertTrue(ws.cell(row=4, column=3).protection.locked)
        self.assertEqual(ws.cell(row=2, column=2).fill.fgColor.rgb, "00FFF3CD")
        self.assertEqual(ws.cell(row=2, column=4).fill.fgColor.rgb, "00FAFBFC")

    def test_display_export_keeps_period_values_separate_by_pk_without_double_counting_totals(self):
        base = {
            "object_id": "obj-1",
            "object_code": "OBJ",
            "object_name": "Объект",
            "object_type_code": "PIPE",
            "object_type_name": "Трубы",
            "section_code": "UCH_1",
            "section_name": "Участок 1",
            "work_type_id": "work-fill",
            "work_type_code": "FILL",
            "work_type_name": "Отсыпка",
            "unit": "м3",
            "project_volume": 10,
            "fact_volume": 4,
            "remaining_volume": 6,
            "period_values": [{"period_key": "2026-05-01", "planned_volume": 5, "fact_volume": 1}],
        }
        rows = [
            {**base, "plan_id": "plan-a", "planned_volume": 3, "pk_start": 100.0, "pk_end": 200.0},
            {**base, "plan_id": "plan-b", "planned_volume": 7, "pk_start": 200.0, "pk_end": 300.0},
        ]

        display_rows = wip_routes._statement_display_rows_for_export(rows, "section_tree")
        leaf = next(row for row in display_rows if row.get("__group_kind") == "work")
        period = {item["period_key"]: item for item in leaf["period_values"]}["2026-05-01"]

        self.assertEqual(leaf["project_volume"], 10.0)
        self.assertEqual(leaf["planned_volume"], 10.0)
        self.assertEqual(leaf["fact_volume"], 4.0)
        self.assertEqual(period["planned_volume"], 10.0)
        self.assertEqual(period["fact_volume"], 2.0)

    def test_export_sheet_locks_project_volume_for_multi_segment_project_leaf(self):
        wb = Workbook()
        ws = wb.active
        headers = [
            ("__structure", "Структура"),
            ("project_volume", "Проектный объем"),
            ("object_id", "object_id"),
            ("work_type_id", "work_type_id"),
            ("project_segment_count", "project_segment_count"),
            ("source_rows", "source_rows"),
            ("__is_group", "__is_group"),
        ]
        ws.append([label for _key, label in headers])
        ws.append(["Отсыпка", 10, "obj-1", "work-1", 2, 1, False])

        wip_routes._style_statement_export_sheet(
            ws,
            headers,
            {"object_id", "work_type_id", "project_segment_count", "source_rows", "__is_group"},
            "picketage",
        )

        self.assertTrue(ws.cell(row=2, column=2).protection.locked)
        self.assertEqual(ws.cell(row=2, column=2).fill.fgColor.rgb, "00FAFBFC")

    def test_picketage_display_export_splits_same_work_by_pk_segment(self):
        base = {
            "object_id": "obj-1",
            "object_code": "OBJ",
            "object_name": "Линейный объект",
            "object_type_code": "PIPE",
            "object_type_name": "Трубы",
            "section_code": "UCH_1",
            "section_name": "Участок 1",
            "work_type_id": "work-fill",
            "work_type_code": "FILL",
            "work_type_name": "Отсыпка",
            "unit": "м3",
            "project_volume": 10,
            "planned_volume": 5,
            "fact_volume": 0,
            "remaining_volume": 5,
            "period_values": [],
        }
        rows = [
            {**base, "plan_id": "plan-a", "pk_start": 100.0, "pk_end": 200.0},
            {**base, "plan_id": "plan-b", "pk_start": 200.0, "pk_end": 300.0},
        ]

        display_rows = wip_routes._statement_display_rows_for_export(rows, "picketage")
        leaf_rows = [row for row in display_rows if row.get("__group_kind") == "work"]

        self.assertEqual(len(leaf_rows), 2)
        self.assertEqual([row.get("source_rows") for row in leaf_rows], [1, 1])
        self.assertEqual([row.get("plan_id") for row in leaf_rows], ["plan-a", "plan-b"])
        self.assertEqual([row.get("pk_start") for row in leaf_rows], [100.0, 200.0])
        self.assertIn("ПК1+00", leaf_rows[0]["__structure"])
        self.assertIn("Отсыпка", leaf_rows[0]["__structure"])

    def test_route_display_export_object_group_label_spans_all_pk_segments(self):
        base = {
            "object_id": "obj-1",
            "object_code": "OBJ",
            "object_name": "Линейный объект",
            "object_type_code": "PIPE",
            "object_type_name": "Трубы",
            "section_code": "UCH_1",
            "section_name": "Участок 1",
            "work_type_id": "work-fill",
            "work_type_code": "FILL",
            "work_type_name": "Отсыпка",
            "unit": "м3",
            "project_volume": 10,
            "planned_volume": 5,
            "fact_volume": 0,
            "remaining_volume": 5,
            "period_values": [],
        }
        rows = [
            {**base, "plan_id": "plan-a", "pk_start": 100.0, "pk_end": 200.0},
            {**base, "plan_id": "plan-b", "pk_start": 250.0, "pk_end": 300.0},
        ]

        for mode in ("picketage", "segment"):
            with self.subTest(mode=mode):
                display_rows = wip_routes._statement_display_rows_for_export(rows, mode)
                object_row = next(row for row in display_rows if row.get("__group_kind") == "object")

                self.assertIn("ПК1+00 - ПК3+00", object_row["__structure"])
                self.assertIn("Линейный объект", object_row["__structure"])

    def test_picketage_export_sorts_same_start_by_object_code_before_end_pk(self):
        base = {
            "object_type_code": "PIPE",
            "object_type_name": "Трубы",
            "section_code": "UCH_1",
            "section_name": "Участок 1",
            "work_type_id": "work-fill",
            "work_type_code": "FILL",
            "work_type_name": "Отсыпка",
            "unit": "м3",
            "project_volume": 1,
            "planned_volume": 1,
            "fact_volume": 0,
            "remaining_volume": 1,
            "period_values": [],
        }
        rows = [
            {**base, "object_id": "obj-b", "object_code": "B_OBJ", "object_name": "Б", "pk_start": 100.0, "pk_end": 110.0},
            {**base, "object_id": "obj-a", "object_code": "A_OBJ", "object_name": "А", "pk_start": 100.0, "pk_end": 200.0},
        ]

        display_rows = wip_routes._statement_display_rows_for_export(rows, "picketage")
        object_rows = [row for row in display_rows if row.get("__group_kind") == "object"]

        self.assertEqual([row.get("object_code") for row in object_rows], ["A_OBJ", "B_OBJ"])


class StatementsGeneralWorkbookHelpersTests(unittest.TestCase):
    def test_statement_scope_has_no_actual_object_exceptions(self):
        sql = wip_routes._statement_scope_cte()

        self.assertNotIn("VPD_012", sql)

    def test_actual_plan_query_uses_saved_row_section_without_object_exceptions(self):
        captured = []

        def capture(sql, params=None):
            captured.append((sql, params))
            return []

        with mock.patch.object(wip_routes, "_statement_table_exists", return_value=True), \
             mock.patch.object(wip_routes, "_general_report_work_rules", return_value=[]), \
             mock.patch.object(wip_routes, "query", side_effect=capture):
            wip_routes._statement_general_add_plan_rows(
                {},
                days=[],
                period_start=date(2026, 9, 1),
                period_end=date(2026, 9, 30),
                section_codes=["UCH_3"],
                object_type=None,
                object_id=None,
                q=None,
                actual_section_binding=True,
            )

        sql = captured[-1][0]
        self.assertIn("assigned_cs.id = pwi.assigned_section_id", sql)
        self.assertIn("COALESCE(assigned_cs.code, so.section_code)", sql)
        self.assertIn("AND assigned_cs.id IS NULL", sql)
        self.assertNotIn("VPD_012", sql)

    def test_actual_project_query_keeps_project_geography_without_object_exceptions(self):
        captured = []

        def capture(sql, params=None):
            captured.append((sql, params))
            return []

        with mock.patch.object(wip_routes, "_general_report_work_rules", return_value=[]), \
             mock.patch.object(wip_routes, "query", side_effect=capture):
            wip_routes._statement_general_add_project_rows(
                {},
                days=[],
                section_codes=["UCH_3"],
                object_type=None,
                object_id=None,
                q=None,
                actual_section_binding=True,
            )

        sql = captured[-1][0]
        self.assertIn("JOIN object_sections so ON so.id = COALESCE(pwi.object_id, con.object_id)", sql)
        self.assertNotIn("assigned_section_id", sql)
        self.assertNotIn("VPD_012", sql)

    def test_actual_fact_query_keeps_explicit_report_section_out_of_geographic_clip(self):
        captured = []

        def capture(sql, params=None):
            captured.append((sql, params))
            return []

        with mock.patch.object(wip_routes, "_general_report_work_rules", return_value=[]), \
             mock.patch.object(wip_routes, "query", side_effect=capture):
            wip_routes._statement_general_add_fact_rows(
                {},
                days=[],
                month_start=date(2026, 9, 1),
                month_end=date(2026, 9, 30),
                cumulative_end=date(2026, 9, 30),
                section_codes=["UCH_3"],
                object_type=None,
                object_id=None,
                q=None,
                actual_section_binding=True,
            )

        sql = captured[-1][0]
        self.assertIn("AND cs_report.id IS NULL", sql)
        self.assertIn("cs_report.id,", sql)
        self.assertNotIn("VPD_012", sql)

    def test_nominal_fact_query_uses_report_section_only_after_object_scope(self):
        captured = []

        def capture(sql, params=None):
            captured.append((sql, params))
            return []

        with mock.patch.object(wip_routes, "_general_report_work_rules", return_value=[]), \
             mock.patch.object(wip_routes, "query", side_effect=capture):
            wip_routes._statement_general_add_fact_rows(
                {},
                days=[],
                month_start=date(2026, 9, 1),
                month_end=date(2026, 9, 30),
                cumulative_end=date(2026, 9, 30),
                section_codes=["UCH_3"],
                object_type=None,
                object_id=None,
                q=None,
                actual_section_binding=False,
            )

        sql = captured[-1][0]
        coalesce_sql = sql[sql.index("ON cs_eff.id = COALESCE("):]
        self.assertLess(coalesce_sql.index("effective_section.section_id"), coalesce_sql.index("cs_report.id"))
        self.assertNotIn("assigned_section_id", sql)

    def test_general_workbook_cache_key_separates_binding_modes(self):
        base = {
            "month_key": "2026-09",
            "section": "UCH_3",
            "object_type": None,
            "object_id": None,
            "q": None,
            "include_metadata": False,
        }
        nominal = wip_routes._statement_general_workbook_cache_params(**base, actual_section_binding=False)
        actual = wip_routes._statement_general_workbook_cache_params(**base, actual_section_binding=True)

        self.assertNotEqual(
            wip_routes._statement_general_workbook_short_key(nominal, "fingerprint"),
            wip_routes._statement_general_workbook_short_key(actual, "fingerprint"),
        )

    def test_month_plan_persists_and_deletes_within_selected_actual_section(self):
        conn = RowcountConn(insert_rowcounts=[1], fetch_results=[("section-3-id",)])
        body = types.SimpleNamespace(
            object_id="00000000-0000-0000-0000-000000000001",
            work_type_id="00000000-0000-0000-0000-000000000002",
            planned_volume=31,
            unit="м3",
            month="2026-09",
            section_code="UCH_3",
            pk_start=317640,
            pk_end=317900,
            pk_raw_text=None,
            comment=None,
            actual_section_binding=True,
        )

        with mock.patch.object(wip_routes, "_statement_table_exists", return_value=True), \
             mock.patch.object(wip_routes, "_statement_signed_pile_plan_for_period", return_value=None), \
             mock.patch.object(wip_routes, "_statement_cur_column_exists", return_value=True), \
             mock.patch.object(wip_routes, "_statement_general_day_response", return_value=[{"date": "2026-09-01"}]), \
             mock.patch.object(wip_routes, "get_conn", return_value=conn):
            result = wip_routes.statements_general_workbook_month_plan(body, _user=None)

        delete_sql, delete_params = next((sql, params) for sql, params in conn.cursor_obj.executed if sql.lstrip().upper().startswith("DELETE"))
        insert_sql, insert_params = next((sql, params) for sql, params in conn.cursor_obj.executed if sql.lstrip().upper().startswith("INSERT"))
        self.assertIn("assigned_section_id IS NOT DISTINCT FROM", delete_sql)
        self.assertIn(True, delete_params)
        self.assertIn("section-3-id", delete_params)
        self.assertIn("object_id, work_type_id, assigned_section_id", insert_sql)
        self.assertEqual(insert_params[1], "section-3-id")
        self.assertEqual(result["assigned_section_code"], "UCH_3")

    def test_statement_pk_segments_merge_small_gaps(self):
        merged = wip_routes._statement_merge_pk_segments([
            {"pk_start": 100.0, "pk_end": 120.0},
            {"pk_start": 124.5, "pk_end": 130.0},
            {"pk_start": 140.5, "pk_end": 150.0},
        ])

        self.assertEqual(len(merged), 2)
        self.assertEqual((merged[0]["pk_start"], merged[0]["pk_end"]), (100.0, 130.0))
        self.assertEqual((merged[1]["pk_start"], merged[1]["pk_end"]), (140.5, 150.0))
        self.assertIn("ПК1+00", merged[0]["label"])

    def test_statement_search_pk_range_extracts_pk_and_keeps_text_terms(self):
        start, end, cleaned = wip_routes._statement_search_pk_range("АД9 песок ПК100+00 - ПК101+50")

        self.assertEqual(start, 10000.0)
        self.assertEqual(end, 10150.0)
        self.assertEqual(cleaned, "АД9 песок")

    def test_statement_compaction_matches_pipeline_specificity_and_scales_work_row(self):
        rules = [
            {
                "id": "generic",
                "object_type_codes": [],
                "object_codes": [],
                "work_codes": ["FILL"],
                "tags": [],
                "coefficient": 1.1,
            },
            {
                "id": "specific",
                "object_type_codes": ["TEMP_ROAD"],
                "object_codes": [],
                "work_codes": ["FILL"],
                "tags": [],
                "coefficient": 1.25,
            },
        ]
        row = {
            "object_type_code": "TEMP_ROAD",
            "object_code": "AD9",
            "work_type_code": "FILL",
            "work_type_analytics_tag": "Песок",
            "project_volume": 125.0,
            "month_plan": 50.0,
            "month_fact": 25.0,
            "cumulative_plan": 100.0,
            "cumulative_fact": 80.0,
            "month_details": [{"volume": 25.0}],
            "cumulative_details": [{"volume": 80.0}],
            "day_values": {"2026-05-01": {"plan_volume": 10.0, "fact_volume": 5.0, "details": []}},
        }

        wip_routes._statement_apply_compaction(row, rules)

        self.assertEqual(row["compaction_rule_id"], "specific")
        self.assertAlmostEqual(row["compaction_factor"], 0.8)
        self.assertEqual(row["material_cumulative_fact"], 80.0)
        self.assertEqual(row["cumulative_fact"], 64.0)
        self.assertEqual(row["day_values"]["2026-05-01"]["material_fact_volume"], 5.0)
        self.assertEqual(row["day_values"]["2026-05-01"]["fact_volume"], 4.0)

    def test_work_amount_without_pk_uses_identical_regional_rate(self):
        amount, missing, regions = wip_routes._statement_work_type_amount(
            {"month_plan": 27, "pk_start": None, "pk_end": None},
            {"novgorod": 91750, "tver": 91750},
            "month_plan",
        )

        self.assertFalse(missing)
        self.assertEqual(amount, 2477250)
        self.assertEqual(len(regions), 1)
        self.assertEqual(regions[0]["rate"], 91750)

    def test_work_amount_without_pk_keeps_ambiguous_different_rates_missing(self):
        amount, missing, regions = wip_routes._statement_work_type_amount(
            {"month_plan": 27, "pk_start": None, "pk_end": None},
            {"novgorod": 91750, "tver": 92000},
            "month_plan",
        )

        self.assertTrue(missing)
        self.assertIsNone(amount)
        self.assertEqual(regions, [])

    def test_work_amount_rounds_after_rows_are_aggregated(self):
        row_amount, missing, _regions = wip_routes._statement_work_type_amount(
            {"month_plan": 1, "pk_start": 100, "pk_end": 100},
            {"novgorod": 0.335, "tver": None},
            "month_plan",
        )
        summary = wip_routes._statement_amount_summary(
            [
                {"month_plan": 1, "month_plan_amount": row_amount, "month_plan_amount_missing_rate": missing},
                {"month_plan": 1, "month_plan_amount": row_amount, "month_plan_amount_missing_rate": missing},
            ],
            "month_plan_amount",
            "month_plan_amount_missing_rate",
            "month_plan",
        )

        self.assertAlmostEqual(row_amount, 0.335)
        self.assertEqual(summary["amount"], 0.67)


if __name__ == "__main__":
    unittest.main()
