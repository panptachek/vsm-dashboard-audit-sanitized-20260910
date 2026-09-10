"""Pure RD/VOR parsing helpers.

This module intentionally has no FastAPI, PostgreSQL, or dashboard state
dependencies. It is shared by local tooling and tests; the legacy RD web route
must not be used as the parser boundary.
"""
from __future__ import annotations

import gc
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable, Optional


LOW_CONFIDENCE_THRESHOLD = 0.72
PDF_ENGINE_STABLE = "stable"
PDF_ENGINE_PDF_INSPECTOR = "pdf_inspector"
PDF_ENGINES = {PDF_ENGINE_STABLE, PDF_ENGINE_PDF_INSPECTOR}


class RdParserError(RuntimeError):
    """Raised when an RD document cannot be parsed in the local parser."""


@dataclass
class ParsedRdRow:
    raw_name: str
    unit: Optional[str]
    volume: Optional[float]
    item_no: Optional[str] = None
    pk_start: Optional[float] = None
    pk_end: Optional[float] = None
    pk_raw: Optional[str] = None
    page_number: Optional[int] = None
    page_label: Optional[str] = None
    table_index: Optional[int] = None
    row_number: Optional[int] = None
    bbox: Optional[dict[str, float]] = None
    row_bbox: Optional[dict[str, float]] = None
    raw_cells: Optional[list[str]] = None


@dataclass
class _InspectorText:
    text: str
    page: int
    x: float
    y: float
    width: float
    height: float


@dataclass
class _InspectorLine:
    page: int
    y: float
    items: list[_InspectorText]

    @property
    def text(self) -> str:
        return _clean_text(" ".join(item.text for item in sorted(self.items, key=lambda item: item.x)))


def _clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip()


def _normalize_text(value: Any) -> str:
    text = _clean_text(value).lower().replace("ё", "е")
    return re.sub(r"[^0-9a-zа-я]+", " ", text).strip()


def _compact_text(value: Any) -> str:
    return re.sub(r"[^0-9a-zа-я]+", "", _normalize_text(value))


def _norm_contains(norm: str, token: str) -> bool:
    token_norm = _normalize_text(token)
    if not token_norm:
        return False
    if " " in token_norm:
        return token_norm in norm
    return bool(re.search(rf"(?:^|\s){re.escape(token_norm)}(?:\s|$)", norm))


def _parse_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _clean_text(value)
    if not text:
        return None
    text = text.replace(" ", "").replace("\xa0", "").replace(",", ".")
    text = re.sub(r"(?<=\d)[^0-9.+-]+(?=\d|$)", "", text)
    match = re.search(r"[-+]?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def _parse_pk_value(value: Any) -> Optional[float]:
    if value is None:
        return None
    text = _clean_text(value)
    if not text:
        return None
    normalized = text.replace(",", ".")
    match = re.search(
        r"(?:п\s*\.?\s*к\s*\.?|pk)?\s*(\d{1,5})(?:\s*\+\s*(\d{1,3}(?:\.\d+)?))?",
        normalized,
        re.IGNORECASE,
    )
    if match:
        pk = float(match.group(1))
        plus = float(match.group(2) or 0)
        return pk * 100.0 + plus
    number = _parse_float(normalized)
    if number is None:
        return None
    return number * 100.0 if abs(number) < 10000 else number


def _parse_pk_range(*values: Any) -> tuple[Optional[float], Optional[float], Optional[str]]:
    raw = " ".join(_clean_text(value) for value in values if _clean_text(value))
    if not raw:
        return None, None, None
    matches = re.findall(
        r"(?:п\s*\.?\s*к\s*\.?|pk)?\s*(\d{1,5})(?:\s*\+\s*(\d{1,3}(?:[\.,]\d+)?))?",
        raw,
        flags=re.IGNORECASE,
    )
    parsed: list[float] = []
    for pk, plus in matches:
        parsed.append(float(pk) * 100.0 + float((plus or "0").replace(",", ".")))
    if len(parsed) >= 2:
        start, end = parsed[0], parsed[1]
    elif len(parsed) == 1:
        start = end = parsed[0]
    else:
        start = end = _parse_pk_value(raw)
    if start is not None and end is not None and start > end:
        start, end = end, start
    return start, end, raw


_NUMBER_RE = r"[-+]?(?:\d{1,3}(?:[\s\xa0]\d{3})+|\d+)(?:[\.,]\d+)?"
_UNIT_RE = r"(?:км|м\s*[³3]|м\s*[²2]|м[³3]|м[²2]|м3|м2|м|т|шт\.?|га|компл\.?|ед\.?)"
_UNIT_VOLUME_RE = re.compile(
    rf"(?<![0-9A-Za-zА-Яа-я])(?P<unit>{_UNIT_RE})\s+(?P<volume>{_NUMBER_RE})(?![0-9])",
    re.IGNORECASE,
)
_PK_RE = re.compile(r"(?:п\s*\.?\s*к\s*\.?\s*)?\d{1,5}\s*\+\s*\d{1,3}(?:[\.,]\d+)?", re.IGNORECASE)


def _strip_formula_columns(value: Any) -> str:
    text = _clean_text(value)
    markers = [
        r"\b[vgslfs]\s*=",
        r"\b[vgslfs]\s+-",
        r"\b[vgslfs]\s*[–—]",
        r"\bγ\s*[-=]",
        r"\bк\s*=",
        r"\bкпотери\b",
        r"=\s*\(\(",
        r"\bL\d+(?:[-–—]|\b)",
        r"\bS\d+\b",
        r"\bSn\b",
        r"\bплотность\s+(?:грунта|материала)",
        r"\bрасстояние\s+между\s+поперечниками",
        r"\bплощадь\s+грунта\s+на\s+поперечнике",
        r"\bвес\s+грунта",
        r"\bобъ[её]м\s+грунта",
        r"\bдлина\s+элемента\s+укрепления",
    ]
    cut_at: Optional[int] = None
    for marker in markers:
        match = re.search(marker, text, flags=re.IGNORECASE)
        if match:
            if match.start() == 0:
                return ""
            cut_at = match.start() if cut_at is None else min(cut_at, match.start())
    if cut_at is not None:
        text = text[:cut_at]
    text = re.sub(r"\b(?:У|П)\s*-\s*", " ", text)
    text = re.sub(r"\b(?:усовершенс|твованные|переходные)\b", " ", text, flags=re.IGNORECASE)
    return _clean_text(text)


def _normalize_unit(value: Any) -> str:
    unit = _clean_text(value).lower().replace("³", "3").replace("²", "2").replace(".", "")
    unit = re.sub(r"\s+", "", unit)
    if unit in {"м3", "м2", "м", "км", "т", "шт", "га"}:
        return unit
    if unit.startswith("компл"):
        return "компл"
    if unit.startswith("ед"):
        return "ед"
    return unit or _clean_text(value)


def _strip_name_prefix(value: Any) -> str:
    text = _clean_text(value)
    text = re.sub(r"^\s*(?:№\s*)?(?:\d+(?:[\.,]\d+)*|[ivxlcdm]+)\s+", "", text, flags=re.IGNORECASE)
    text = re.sub(r"^\s*(?:п\s*/\s*п|п\.?\s*п\.?)\s+", "", text, flags=re.IGNORECASE)
    return _clean_text(text.strip(" -–—:;"))


def _looks_like_numeric_legend(line: str) -> bool:
    text = _clean_text(line)
    return bool(re.fullmatch(r"(?:\d+\s+){2,}\d+", text))


def _is_layout_noise(line: str) -> bool:
    text = _clean_text(line)
    norm = _normalize_text(text)
    if not norm:
        return True
    if _looks_like_numeric_legend(text):
        return True
    noise_tokens = (
        "создание высокоскоростной железнодорожной",
        "алабушево",
        "обухово",
        "выползово",
        "разработал",
        "проверил",
        "н контроль",
        "нконтроль",
        "н контр",
        "нконтр",
        "гип",
        "стадия",
        "лист",
        "листов",
        "формат",
        "инв",
        "подпись и дата",
        "взам инв",
        "согласовано",
        "продолжение таблицы",
        "окончание таблицы",
    )
    if any(token in norm for token in noise_tokens):
        return True
    header_tokens = (
        "наименование работ",
        "пикетаж",
        "ед изм",
        "кол",
        "номер",
        "формула расчета",
        "расчет объемов",
        "расхода материалов",
        "ссылка на чертежи",
        "ссылки на чертежи",
        "спецификации",
        "примечание",
    )
    if any(token in norm for token in header_tokens) and not _UNIT_VOLUME_RE.search(text):
        return True
    return False


def _is_name_candidate(line: str) -> bool:
    text = _strip_name_prefix(_strip_formula_columns(line))
    if not text or _is_layout_noise(text):
        return False
    if _PK_RE.search(text):
        return False
    if re.fullmatch(r"[\d\s,.\-+]+", text):
        return False
    letters = re.findall(r"[A-Za-zА-Яа-я]", text)
    if len(letters) < 4:
        return False
    if _normalize_text(text).startswith(("итого", "всего")):
        return False
    return True


def _valid_parsed_name(name: str) -> bool:
    text = _strip_name_prefix(_strip_formula_columns(name))
    norm = _normalize_text(text)
    if len(text) < 4 or _is_layout_noise(text):
        return False
    if norm.startswith(("итого", "всего", "раздел")) or "пикетаж" in norm:
        return False
    if len(re.findall(r"[A-Za-zА-Яа-я]", text)) < 4:
        return False
    return True


def _valid_table_name(name: str) -> bool:
    text = _strip_name_prefix(_strip_formula_columns(name))
    norm = _normalize_text(text)
    if len(text) < 4:
        return False
    if _looks_like_numeric_legend(text) or re.fullmatch(r"[\d\s,.\-+]+", text):
        return False
    if norm.startswith(("итого", "всего", "раздел")) or _norm_contains(norm, "пикетаж"):
        return False
    table_noise = (
        "наименование работ",
        "ед изм",
        "формула расчета",
        "формулы расчета",
        "расчет объемов",
        "ссылки на чертежи",
        "спецификации",
        "примечание",
        "лист",
        "листов",
        "согласовано",
        "подпись и дата",
        "взам инв",
        "продолжение таблицы",
        "окончание таблицы",
    )
    if any(_norm_contains(norm, token) for token in table_noise):
        return False
    if len(re.findall(r"[A-Za-zА-Яа-я]", text)) < 4:
        return False
    return True


def _best_unit_volume_match(line: str) -> Optional[re.Match[str]]:
    matches = list(_UNIT_VOLUME_RE.finditer(line))
    if not matches:
        return None
    return matches[-1]


def _can_start_layout_name_fragment(value: str) -> bool:
    text = _strip_name_prefix(_strip_formula_columns(value))
    if not text:
        return False
    match = re.search(r"[A-Za-zА-Яа-я]", text)
    if not match:
        return False
    first = match.group(0)
    return first.isupper() or text.startswith(('"', "'", "В том числе", "в том числе"))


def _append_pending(pending: list[str], line: str) -> None:
    name = _strip_name_prefix(_strip_formula_columns(line))
    if not name:
        return
    if not pending and not _can_start_layout_name_fragment(name):
        return
    pending.append(name)
    if len(pending) > 8:
        del pending[:-8]


def _rows_from_named_layout_lines(lines: list[str], *, page_number: int) -> list[ParsedRdRow]:
    parsed: list[ParsedRdRow] = []
    pending: list[str] = []
    for row_number, line in enumerate(lines, start=1):
        clean = _clean_text(line)
        named_clean = _strip_formula_columns(clean)
        if not named_clean:
            continue
        match = _best_unit_volume_match(named_clean)
        if match:
            unit = _normalize_unit(match.group("unit"))
            if unit == "м":
                pending = []
                continue
            volume = _parse_float(match.group("volume"))
            prefix = _strip_name_prefix(named_clean[: match.start()])
            name_parts = [part for part in [*pending, prefix] if _valid_parsed_name(part)]
            name = _clean_text(" ".join(name_parts))
            pending = []
            if volume is None or volume <= 0 or not _valid_parsed_name(name):
                continue
            parsed.append(
                ParsedRdRow(
                    raw_name=name,
                    unit=unit,
                    volume=volume,
                    page_number=page_number,
                    row_number=row_number,
                    raw_cells=[clean],
                )
            )
            continue
        if _is_name_candidate(named_clean):
            _append_pending(pending, named_clean)
        elif _is_layout_noise(named_clean):
            pending = []
    return parsed


def _extract_numbers(value: str) -> list[float]:
    result: list[float] = []
    for match in re.finditer(_NUMBER_RE, value):
        parsed = _parse_float(match.group(0))
        if parsed is not None:
            result.append(parsed)
    return result


def _picket_raw(match_text: str) -> str:
    raw = _clean_text(match_text)
    if not raw.lower().replace(" ", "").startswith(("пк", "п.к")):
        raw = f"ПК {raw}"
    return raw


def _metric_row(
    *,
    name: str,
    unit: str,
    volume: Optional[float],
    pk_start: Optional[float],
    pk_end: Optional[float],
    pk_raw: Optional[str],
    page_number: int,
    row_number: int,
    raw_cells: list[str],
) -> Optional[ParsedRdRow]:
    if volume is None or volume <= 0:
        return None
    if pk_start is not None and pk_end is None:
        pk_end = pk_start
    if pk_end is not None and pk_start is None:
        pk_start = pk_end
    if pk_start is not None and pk_end is not None and pk_start > pk_end:
        pk_start, pk_end = pk_end, pk_start
    return ParsedRdRow(
        raw_name=name,
        unit=unit,
        volume=round(float(volume), 6),
        pk_start=pk_start,
        pk_end=pk_end,
        pk_raw=pk_raw,
        page_number=page_number,
        row_number=row_number,
        raw_cells=raw_cells,
    )


def _rows_from_earthwork_picket(
    *,
    pk_start: Optional[float],
    pk_raw: str,
    current_values: list[float],
    next_values: list[float],
    page_number: int,
    row_number: int,
    raw_cells: list[str],
    header_text: str,
) -> list[ParsedRdRow]:
    values = next_values if next_values else current_values
    if not values:
        return []
    lower = _normalize_text(header_text)
    labels: list[tuple[str, str]] = []
    if "длина участка" in lower:
        labels.append(("Длина участка", "м"))
    if "насып" in lower:
        labels.append(("Насыпь", "м3"))
    if "выем" in lower:
        labels.append(("Выемка", "м3"))
    if "кювет" in lower:
        labels.append(("Кюветы", "м3"))
    if not labels:
        return []
    if len(values) == len(labels) - 1 and labels and labels[0][0] == "Длина участка":
        labels = labels[1:]
    if len(values) < len(labels):
        return []
    pk_end = None
    if labels and labels[0][0] == "Длина участка":
        pk_end = (pk_start + values[0]) if pk_start is not None else None
    result: list[ParsedRdRow] = []
    for idx, (label, unit) in enumerate(labels):
        row = _metric_row(
            name=label,
            unit=unit,
            volume=values[idx] if idx < len(values) else None,
            pk_start=pk_start,
            pk_end=pk_end,
            pk_raw=pk_raw,
            page_number=page_number,
            row_number=row_number,
            raw_cells=raw_cells,
        )
        if row:
            result.append(row)
    return result


def _rows_from_prs_picket(
    *,
    pk_start: Optional[float],
    pk_raw: str,
    current_values: list[float],
    next_values: list[float],
    page_number: int,
    row_number: int,
    raw_cells: list[str],
    header_text: str,
) -> list[ParsedRdRow]:
    lower = _normalize_text(header_text)
    has_bad_prs = "не пригод" in lower or "непригод" in lower
    length = next_values[0] if next_values else None
    pk_end = (pk_start + length) if pk_start is not None and length is not None else None
    if has_bad_prs:
        if len(current_values) >= 6:
            area_good = current_values[2]
            area_bad = current_values[5]
        else:
            area_good = area_bad = None
        volume_values = next_values[1:] if len(next_values) >= 7 else current_values[-6:]
        volume_good = volume_values[2] if len(volume_values) >= 3 else None
        volume_bad = volume_values[5] if len(volume_values) >= 6 else None
        metrics = [
            ("Площадь снятия пригодного ПРС", "м2", area_good),
            ("Площадь снятия непригодного ПРС", "м2", area_bad),
            ("Объем снятия пригодного ПРС", "м3", volume_good),
            ("Объем снятия непригодного ПРС", "м3", volume_bad),
        ]
    else:
        values = next_values[1:] if len(next_values) >= 7 else current_values[-6:]
        area = values[2] if len(values) >= 3 else None
        volume = values[5] if len(values) >= 6 else None
        metrics = [
            ("Площадь снятия пригодного ПРС", "м2", area),
            ("Объем снятия пригодного ПРС", "м3", volume),
        ]
    result: list[ParsedRdRow] = []
    for label, unit, volume in metrics:
        row = _metric_row(
            name=label,
            unit=unit,
            volume=volume,
            pk_start=pk_start,
            pk_end=pk_end,
            pk_raw=pk_raw,
            page_number=page_number,
            row_number=row_number,
            raw_cells=raw_cells,
        )
        if row:
            result.append(row)
    return result


def _rows_from_picket_layout_lines(lines: list[str], *, page_number: int) -> list[ParsedRdRow]:
    header_text = _normalize_text(" ".join(lines[:12]))
    if "пикетаж" not in header_text:
        return []
    table_kind = ""
    if "прс" in header_text and ("снятия" in header_text or "снятие" in header_text):
        table_kind = "prs"
    elif any(token in header_text for token in ("насып", "выем", "кювет")):
        table_kind = "earthwork"
    if not table_kind:
        return []

    parsed: list[ParsedRdRow] = []
    for idx, line in enumerate(lines):
        clean = _clean_text(line)
        match = _PK_RE.search(clean)
        if not match:
            continue
        pk_raw = _picket_raw(match.group(0))
        pk_start = _parse_pk_value(pk_raw)
        current_values = _extract_numbers(clean[match.end() :])
        next_values: list[float] = []
        raw_cells = [clean]
        for follow in lines[idx + 1 : idx + 3]:
            follow_clean = _clean_text(follow)
            if not follow_clean or _PK_RE.search(follow_clean) or _is_layout_noise(follow_clean):
                continue
            candidate_values = _extract_numbers(follow_clean)
            if candidate_values:
                next_values = candidate_values
                raw_cells.append(follow_clean)
                break
        if table_kind == "prs":
            parsed.extend(
                _rows_from_prs_picket(
                    pk_start=pk_start,
                    pk_raw=pk_raw,
                    current_values=current_values,
                    next_values=next_values,
                    page_number=page_number,
                    row_number=idx + 1,
                    raw_cells=raw_cells,
                    header_text=header_text,
                )
            )
        else:
            parsed.extend(
                _rows_from_earthwork_picket(
                    pk_start=pk_start,
                    pk_raw=pk_raw,
                    current_values=current_values,
                    next_values=next_values,
                    page_number=page_number,
                    row_number=idx + 1,
                    raw_cells=raw_cells,
                    header_text=header_text,
                )
            )
    return parsed


def _rows_from_layout_text(text: str, *, page_number: int) -> list[ParsedRdRow]:
    lines = [_clean_text(line) for line in text.splitlines() if _clean_text(line)]
    if not lines:
        return []
    header_text = _normalize_text(" ".join(lines[:12]))
    picket_rows = _rows_from_picket_layout_lines(lines, page_number=page_number)
    named_rows: list[ParsedRdRow] = []
    if "пикетаж" not in header_text or "наименование работ" in header_text:
        named_rows = _rows_from_named_layout_lines(lines, page_number=page_number)
    return [*named_rows, *picket_rows]


def _dedupe_parsed_rows(rows: list[ParsedRdRow]) -> list[ParsedRdRow]:
    result: list[ParsedRdRow] = []
    seen: set[tuple[str, str, str, str, str, int, int]] = set()
    for row in rows:
        key = (
            _normalize_text(row.raw_name),
            _normalize_unit(row.unit),
            "" if row.volume is None else f"{float(row.volume):.6f}",
            "" if row.pk_start is None else f"{float(row.pk_start):.6f}",
            "" if row.pk_end is None else f"{float(row.pk_end):.6f}",
            int(row.page_number or 0),
            int(row.row_number or 0),
        )
        if key in seen:
            continue
        seen.add(key)
        result.append(row)
    return result


def _header_key(value: Any) -> str:
    return _compact_text(value)


def _classify_headers(row: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, value in enumerate(row):
        key = _header_key(value)
        if not key:
            continue
        if "наимен" in key or "видработ" in key or key in {"работы", "работа", "материал"}:
            mapping.setdefault("name", idx)
        if "едизм" in key or key.startswith("ед") or key in {"изм", "единицаизмерения"}:
            mapping.setdefault("unit", idx)
        is_formula_or_note = any(
            token in key
            for token in (
                "формул",
                "расчет",
                "расход",
                "чертеж",
                "спецификац",
                "примеч",
            )
        )
        if not is_formula_or_note and (
            key in {"кол", "колво", "количество", "объем", "обьем", "volume", "qty"}
            or key.startswith(("колво", "колич", "объем", "обьем"))
        ):
            mapping.setdefault("volume", idx)
        if "пкнач" in key or "pkstart" in key:
            mapping.setdefault("pk_start", idx)
        elif "пккон" in key or "pkend" in key:
            mapping.setdefault("pk_end", idx)
        elif "пикет" in key or key == "пк" or key.startswith("pk"):
            mapping.setdefault("pk", idx)
    return mapping


def _find_header(rows: list[list[str]]) -> tuple[Optional[int], dict[str, int]]:
    for row_idx, row in enumerate(rows[:10]):
        mapping = _classify_headers(row)
        if "name" in mapping and "unit" in mapping and "volume" in mapping:
            return row_idx, mapping
    return None, {}


def _row_cell(row: list[str], idx: Optional[int]) -> str:
    if idx is None or idx < 0 or idx >= len(row):
        return ""
    return _clean_text(row[idx])


def _table_row_ordinal(row: list[str], mapping: dict[str, int]) -> Optional[str]:
    name_idx = mapping.get("name", len(row))
    prefix_cells = row[:name_idx] if name_idx > 0 else row[:1]
    for cell in prefix_cells:
        text = _clean_text(cell)
        if re.fullmatch(r"(?:№\s*)?\d+(?:[\.,]\d+)*", text, flags=re.IGNORECASE):
            return re.sub(r"^№\s*", "", text, flags=re.IGNORECASE).replace(",", ".")
    return None


def _table_row_has_ordinal(row: list[str], mapping: dict[str, int]) -> bool:
    return _table_row_ordinal(row, mapping) is not None


def _is_likely_page_continuation_fragment(value: str) -> bool:
    text = _strip_name_prefix(_strip_formula_columns(value))
    if not text:
        return False
    match = re.search(r"[A-Za-zА-Яа-я]", text)
    if not match:
        return False
    if match.group(0).islower():
        return True
    norm = _normalize_text(text)
    return norm.startswith((
        "и ",
        "а ",
        "с ",
        "со ",
        "для ",
        "до ",
        "на ",
        "в ",
        "во ",
        "объем дан",
        "обьем дан",
        "кпотери",
    ))


def _append_table_continuation(row: ParsedRdRow, text: str, raw_cells: list[str]) -> None:
    fragment = _strip_name_prefix(_strip_formula_columns(text))
    if not fragment:
        return
    row.raw_name = _clean_text(f"{row.raw_name} {fragment}")
    row.raw_cells = [*(row.raw_cells or []), *raw_cells]


def _bbox_dict_from_values(values: Any) -> Optional[dict[str, float]]:
    if not values or len(values) != 4:
        return None
    try:
        x0, top, x1, bottom = (float(value) for value in values)
    except (TypeError, ValueError):
        return None
    if x1 <= x0 or bottom <= top:
        return None
    return {"x0": x0, "top": top, "x1": x1, "bottom": bottom}


def _table_row_bboxes(found_table: Any) -> list[Optional[dict[str, float]]]:
    result: list[Optional[dict[str, float]]] = []
    for table_row in getattr(found_table, "rows", []) or []:
        row_bbox = _bbox_dict_from_values(getattr(table_row, "bbox", None))
        if row_bbox is None:
            cell_bboxes = []
            for cell in getattr(table_row, "cells", []) or []:
                cell_bbox = _bbox_dict_from_values(cell)
                if cell_bbox is not None:
                    cell_bboxes.append(cell_bbox)
            if cell_bboxes:
                row_bbox = {
                    "x0": min(cell["x0"] for cell in cell_bboxes),
                    "top": min(cell["top"] for cell in cell_bboxes),
                    "x1": max(cell["x1"] for cell in cell_bboxes),
                    "bottom": max(cell["bottom"] for cell in cell_bboxes),
                }
        result.append(row_bbox)
    return result


def _row_bbox_from_table_bbox(
    table_bbox: Optional[dict[str, float]],
    *,
    row_number: int,
    row_count: int,
) -> Optional[dict[str, float]]:
    if not isinstance(table_bbox, dict) or row_count <= 0:
        return None
    try:
        x0 = float(table_bbox.get("x0"))
        x1 = float(table_bbox.get("x1"))
        top = float(table_bbox.get("top", table_bbox.get("y0")))
        bottom = float(table_bbox.get("bottom", table_bbox.get("y1")))
    except (TypeError, ValueError):
        return None
    if x1 <= x0 or bottom <= top:
        return None
    bounded_row = max(1, min(row_count, int(row_number or 1)))
    row_height = (bottom - top) / row_count
    row_top = top + (bounded_row - 1) * row_height
    row_bottom = top + bounded_row * row_height
    return {"x0": x0, "top": row_top, "x1": x1, "bottom": row_bottom}


def _merge_bboxes(bboxes: Iterable[Optional[dict[str, float]]]) -> Optional[dict[str, float]]:
    valid: list[dict[str, float]] = []
    for bbox in bboxes:
        if not isinstance(bbox, dict):
            continue
        try:
            x0 = float(bbox.get("x0"))
            x1 = float(bbox.get("x1"))
            top = float(bbox.get("top"))
            bottom = float(bbox.get("bottom"))
        except (TypeError, ValueError):
            continue
        if x1 <= x0 or bottom <= top:
            continue
        valid.append({"x0": x0, "top": top, "x1": x1, "bottom": bottom})
    if not valid:
        return None
    return {
        "x0": min(bbox["x0"] for bbox in valid),
        "top": min(bbox["top"] for bbox in valid),
        "x1": max(bbox["x1"] for bbox in valid),
        "bottom": max(bbox["bottom"] for bbox in valid),
    }


def _rows_from_table(
    table_rows: list[list[Any]],
    *,
    page_number: Optional[int] = None,
    table_index: Optional[int] = None,
    table_bbox: Optional[dict[str, float]] = None,
    table_row_bboxes: Optional[list[Optional[dict[str, float]]]] = None,
    page_continuations: Optional[list[tuple[str, list[str]]]] = None,
    require_numbered_positions: bool = False,
    allow_estimated_row_bboxes: bool = False,
) -> list[ParsedRdRow]:
    rows: list[list[str]] = []
    row_bboxes: list[Optional[dict[str, float]]] = []
    for original_idx, row in enumerate(table_rows):
        cleaned = [_clean_text(cell) for cell in row]
        if not any(cleaned):
            continue
        rows.append(cleaned)
        row_bboxes.append(table_row_bboxes[original_idx] if table_row_bboxes and original_idx < len(table_row_bboxes) else None)
    header_idx, mapping = _find_header(rows)
    if header_idx is None:
        return []

    parsed: list[ParsedRdRow] = []
    pending_name = ""
    pending_item_no: Optional[str] = None
    pending_pk_start: Optional[float] = None
    pending_pk_end: Optional[float] = None
    pending_pk_raw: Optional[str] = None
    pending_bboxes: list[dict[str, float]] = []
    pending_raw_cells: list[str] = []
    completed_item_numbers: set[str] = set()

    def row_bbox_for(raw_idx: int) -> Optional[dict[str, float]]:
        bbox = row_bboxes[raw_idx - 1] if raw_idx - 1 < len(row_bboxes) else None
        if bbox is None and allow_estimated_row_bboxes:
            bbox = _row_bbox_from_table_bbox(table_bbox, row_number=raw_idx, row_count=len(rows))
        return bbox

    def reset_pending() -> None:
        nonlocal pending_name, pending_item_no, pending_pk_start, pending_pk_end, pending_pk_raw, pending_bboxes, pending_raw_cells
        pending_name = ""
        pending_item_no = None
        pending_pk_start = None
        pending_pk_end = None
        pending_pk_raw = None
        pending_bboxes = []
        pending_raw_cells = []

    for raw_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        row_bbox = row_bbox_for(raw_idx)
        item_no = _table_row_ordinal(row, mapping)
        name = _strip_name_prefix(_strip_formula_columns(_row_cell(row, mapping.get("name"))))
        unit = _normalize_unit(_row_cell(row, mapping.get("unit"))) or None
        volume = _parse_float(_row_cell(row, mapping.get("volume")))
        if unit and not re.fullmatch(_UNIT_RE, unit, flags=re.IGNORECASE):
            unit = None
        pk_start = pk_end = None
        pk_raw = None
        if "pk_start" in mapping or "pk_end" in mapping:
            pk_start = _parse_pk_value(_row_cell(row, mapping.get("pk_start")))
            pk_end = _parse_pk_value(_row_cell(row, mapping.get("pk_end")))
            pk_raw = " - ".join(
                part for part in [_row_cell(row, mapping.get("pk_start")), _row_cell(row, mapping.get("pk_end"))] if part
            ) or None
        elif "pk" in mapping:
            pk_start, pk_end, pk_raw = _parse_pk_range(_row_cell(row, mapping.get("pk")))

        if item_no and item_no in completed_item_numbers:
            # Some VOR tables repeat one position with a lower material/alternate unit row.
            # The import workflow uses the upper quantity row for that source position.
            continue

        has_quantity = unit is not None and volume is not None
        has_name_fragment = bool(name and _valid_table_name(name))

        if not has_quantity:
            if not has_name_fragment:
                if item_no:
                    reset_pending()
                continue
            if item_no:
                pending_name = name
                pending_item_no = item_no
                pending_pk_start = pk_start
                pending_pk_end = pk_end
                pending_pk_raw = pk_raw
                pending_bboxes = [row_bbox] if row_bbox is not None else []
                pending_raw_cells = list(row)
                continue
            if pending_item_no:
                pending_name = f"{pending_name} {name}".strip()
                if row_bbox is not None:
                    pending_bboxes.append(row_bbox)
                pending_raw_cells.extend(cell for cell in row if cell)
                continue
            is_leading_unnumbered_fragment = page_continuations is not None and not parsed
            if is_leading_unnumbered_fragment and (
                _is_likely_page_continuation_fragment(name) or int(page_number or 0) > 1
            ):
                page_continuations.append((name, row))
                continue
            if not require_numbered_positions and not parsed:
                pending_name = f"{pending_name} {name}".strip() if pending_name else name
                if row_bbox is not None:
                    pending_bboxes.append(row_bbox)
                pending_raw_cells.extend(cell for cell in row if cell)
                pending_pk_start = pk_start if pk_start is not None else pending_pk_start
                pending_pk_end = pk_end if pk_end is not None else pending_pk_end
                pending_pk_raw = pk_raw or pending_pk_raw
                continue
            if parsed:
                _append_table_continuation(parsed[-1], name, row)
                parsed[-1].row_bbox = _merge_bboxes([parsed[-1].row_bbox, row_bbox]) or parsed[-1].row_bbox
            continue

        if require_numbered_positions and not item_no and not pending_item_no:
            # A valid VOR position must have a point number. Unnumbered quantity rows
            # are usually calculation/material tails of the previous item; ignore them.
            continue

        if pending_item_no or (pending_name and not parsed and not require_numbered_positions):
            item_no = pending_item_no
            name = f"{pending_name} {name}".strip() if name else pending_name
            row_bbox = _merge_bboxes([*pending_bboxes, row_bbox]) or row_bbox
            raw_cells = [*pending_raw_cells, *[cell for cell in row if cell]]
            pk_start = pk_start if pk_start is not None else pending_pk_start
            pk_end = pk_end if pk_end is not None else pending_pk_end
            pk_raw = pk_raw or pending_pk_raw
            reset_pending()
        else:
            raw_cells = row

        if item_no and item_no in completed_item_numbers:
            continue
        if not name or unit is None or volume is None or not _valid_table_name(name):
            if item_no:
                reset_pending()
            continue
        if pk_start is not None and pk_end is None:
            pk_end = pk_start
        if pk_end is not None and pk_start is None:
            pk_start = pk_end
        if pk_start is not None and pk_end is not None and pk_start > pk_end:
            pk_start, pk_end = pk_end, pk_start
        parsed.append(
            ParsedRdRow(
                raw_name=name,
                unit=unit,
                volume=volume,
                item_no=item_no,
                pk_start=pk_start,
                pk_end=pk_end,
                pk_raw=pk_raw,
                page_number=page_number,
                table_index=table_index,
                row_number=raw_idx,
                bbox=table_bbox,
                row_bbox=row_bbox,
                raw_cells=raw_cells,
            )
        )
        if item_no:
            completed_item_numbers.add(item_no)
    return parsed


def _inspector_text_from_item(item: Any) -> Optional[_InspectorText]:
    text = _clean_text(getattr(item, "text", ""))
    if not text:
        return None
    try:
        return _InspectorText(
            text=text,
            page=int(getattr(item, "page", 0) or 0),
            x=float(getattr(item, "x", 0.0) or 0.0),
            y=float(getattr(item, "y", 0.0) or 0.0),
            width=float(getattr(item, "width", 0.0) or 0.0),
            height=float(getattr(item, "height", 0.0) or 0.0),
        )
    except (TypeError, ValueError):
        return None


def _inspector_lines(items: Iterable[_InspectorText], *, y_tolerance: float = 5.5) -> list[_InspectorLine]:
    lines: list[_InspectorLine] = []
    for item in sorted(items, key=lambda value: (value.page, -value.y, value.x)):
        if not lines or lines[-1].page != item.page or abs(lines[-1].y - item.y) > y_tolerance:
            lines.append(_InspectorLine(page=item.page, y=item.y, items=[item]))
            continue
        lines[-1].items.append(item)
        lines[-1].y = max(lines[-1].y, item.y)
    for line in lines:
        line.items.sort(key=lambda value: value.x)
    return lines


def _inspector_item_bbox(item: _InspectorText) -> dict[str, float]:
    return {"x0": item.x, "y0": item.y, "x1": item.x + max(0.1, item.width), "y1": item.y + max(0.1, item.height)}


def _merge_pdf_point_bboxes(bboxes: Iterable[Optional[dict[str, float]]]) -> Optional[dict[str, float]]:
    valid: list[dict[str, float]] = []
    for bbox in bboxes:
        if not isinstance(bbox, dict):
            continue
        try:
            x0 = float(bbox.get("x0"))
            x1 = float(bbox.get("x1"))
            y0 = float(bbox.get("y0"))
            y1 = float(bbox.get("y1"))
        except (TypeError, ValueError):
            continue
        if x1 <= x0 or y1 <= y0:
            continue
        valid.append({"x0": x0, "y0": y0, "x1": x1, "y1": y1})
    if not valid:
        return None
    return {
        "x0": min(bbox["x0"] for bbox in valid),
        "y0": min(bbox["y0"] for bbox in valid),
        "x1": max(bbox["x1"] for bbox in valid),
        "y1": max(bbox["y1"] for bbox in valid),
    }


def _inspector_line_bbox(line: _InspectorLine) -> Optional[dict[str, float]]:
    return _merge_pdf_point_bboxes(_inspector_item_bbox(item) for item in line.items)


def _inspector_line_name(line: _InspectorLine) -> str:
    parts = [
        item.text
        for item in line.items
        if 90.0 <= item.x < 390.0 and not re.fullmatch(r"\d+(?:[\.,]\d+)*", item.text)
    ]
    return _clean_text(" ".join(parts))


def _inspector_line_unit(line: _InspectorLine) -> Optional[str]:
    parts = [
        item.text
        for item in line.items
        if 365.0 <= item.x <= 455.0 and (re.fullmatch(_UNIT_RE, _normalize_unit(item.text), flags=re.IGNORECASE) or item.text in {"2", "3"})
    ]
    if not parts:
        return None
    unit = _normalize_unit(" ".join(parts))
    if re.fullmatch(_UNIT_RE, unit, flags=re.IGNORECASE):
        return unit
    return None


def _inspector_line_volume(line: _InspectorLine) -> Optional[float]:
    values: list[float] = []
    for item in line.items:
        if 405.0 <= item.x <= 530.0:
            parsed = _parse_float(item.text)
            if parsed is not None:
                values.append(parsed)
    if not values:
        return None
    return values[-1]


def _inspector_line_is_numeric_header(line: _InspectorLine) -> bool:
    digit_items = [item for item in line.items if re.fullmatch(r"\d+", item.text)]
    return len(digit_items) >= 4 and len(digit_items) == len([item for item in line.items if item.text])


def _inspector_line_item_no(line: _InspectorLine) -> Optional[str]:
    if _inspector_line_is_numeric_header(line):
        return None
    for item in line.items:
        if item.x > 95.0:
            continue
        text = _clean_text(item.text)
        if re.fullmatch(r"(?:№\s*)?\d+(?:[\.,]\d+)*", text, flags=re.IGNORECASE):
            return re.sub(r"^№\s*", "", text, flags=re.IGNORECASE).replace(",", ".")
    return None


def _inspector_line_has_quantity(line: _InspectorLine) -> bool:
    return _inspector_line_unit(line) is not None and _inspector_line_volume(line) is not None


def _inspector_line_is_noise(line: _InspectorLine) -> bool:
    text = line.text
    if _is_layout_noise(text) or _inspector_line_is_numeric_header(line):
        return True
    norm = _normalize_text(text)
    return norm.startswith(("раздел", "сигнальный экземпляр"))


def _inspector_line_is_relevant(line: _InspectorLine) -> bool:
    if line.y < 90.0 or _inspector_line_is_noise(line):
        return False
    if _inspector_line_item_no(line) is not None or _inspector_line_has_quantity(line):
        return True
    name = _inspector_line_name(line)
    return bool(name and _valid_table_name(name))


def _inspector_line_starts_position(line: _InspectorLine) -> bool:
    name = _strip_name_prefix(_strip_formula_columns(_inspector_line_name(line)))
    if not _valid_table_name(name):
        return False
    norm = _normalize_text(name)
    if norm.startswith("основные работы"):
        return False
    if _starts_with_material(norm) or norm.startswith("в том числе"):
        return True
    if not _can_start_layout_name_fragment(name):
        return False
    return any(token in norm for token in WORK_ACTION_KEYWORDS)


def _inspector_next_item_no_soon(lines: list[_InspectorLine], start_idx: int, *, max_lines: int = 8) -> bool:
    for line in lines[start_idx + 1 : start_idx + 1 + max_lines]:
        if _inspector_line_item_no(line) is not None:
            return True
        if _inspector_line_has_quantity(line):
            return False
    return False


def _inspector_split_leading_prefix(prefix: list[_InspectorLine]) -> tuple[list[_InspectorLine], list[_InspectorLine]]:
    for idx, line in enumerate(prefix):
        if _inspector_line_starts_position(line):
            return prefix[:idx], prefix[idx:]
    return prefix, []


def _inspector_table_bbox(lines: list[_InspectorLine]) -> Optional[dict[str, float]]:
    bbox = _merge_pdf_point_bboxes(_inspector_line_bbox(line) for line in lines)
    if bbox is None:
        return None
    return {
        "x0": max(0.0, bbox["x0"] - 12.0),
        "y0": max(0.0, bbox["y0"] - 4.0),
        "x1": bbox["x1"] + 4.0,
        "y1": bbox["y1"] + 4.0,
    }


def _inspector_row_bbox(group: list[_InspectorLine], table_bbox: Optional[dict[str, float]]) -> Optional[dict[str, float]]:
    bbox = _merge_pdf_point_bboxes(_inspector_line_bbox(line) for line in group)
    if bbox is None:
        return None
    x0 = table_bbox["x0"] if isinstance(table_bbox, dict) and "x0" in table_bbox else max(0.0, bbox["x0"] - 12.0)
    x1 = table_bbox["x1"] if isinstance(table_bbox, dict) and "x1" in table_bbox else bbox["x1"] + 4.0
    return {"x0": x0, "y0": max(0.0, bbox["y0"] - 3.0), "x1": x1, "y1": bbox["y1"] + 3.0}


def _inspector_group_to_row(
    group: list[_InspectorLine],
    *,
    page_number: int,
    table_index: int,
    row_number: int,
    table_bbox: Optional[dict[str, float]],
) -> Optional[ParsedRdRow]:
    item_no = next((_inspector_line_item_no(line) for line in group if _inspector_line_item_no(line)), None)
    quantity_indexes = [idx for idx, line in enumerate(group) if _inspector_line_has_quantity(line)]
    if not item_no or not quantity_indexes:
        return None
    quantity_line = group[quantity_indexes[0]]
    unit = _inspector_line_unit(quantity_line)
    volume = _inspector_line_volume(quantity_line)
    if unit == "м":
        exponent = next(
            (
                item.text
                for line in group
                for item in line.items
                if item.text in {"2", "3"}
                and 386.0 <= item.x <= 406.0
                and abs(float(item.y) - float(quantity_line.y)) <= 9.0
            ),
            "",
        )
        if exponent:
            unit = f"м{exponent}"
    name_cutoff = quantity_indexes[1] if len(quantity_indexes) > 1 else len(group)
    name_parts: list[str] = []
    for line in group[:name_cutoff]:
        name = _strip_name_prefix(_strip_formula_columns(_inspector_line_name(line)))
        if _valid_table_name(name):
            name_parts.append(name)
    name = _clean_text(" ".join(name_parts))
    if not unit or volume is None or volume <= 0 or not _valid_table_name(name):
        return None
    return ParsedRdRow(
        raw_name=name,
        unit=unit,
        volume=volume,
        item_no=item_no,
        page_number=page_number,
        table_index=table_index,
        row_number=row_number,
        bbox=table_bbox,
        row_bbox=_inspector_row_bbox(group, table_bbox),
        raw_cells=[item_no, name, unit, format(volume, "f").rstrip("0").rstrip(".")],
    )


def _rows_from_pdf_inspector_page(
    page_lines: list[_InspectorLine],
    *,
    page_number: int,
    previous_row: Optional[ParsedRdRow],
) -> list[ParsedRdRow]:
    lines = [line for line in page_lines if _inspector_line_is_relevant(line)]
    if not lines:
        return []
    first_item_idx = next((idx for idx, line in enumerate(lines) if _inspector_line_item_no(line) is not None), None)
    if first_item_idx is None:
        return []

    leading, pending_pre = _inspector_split_leading_prefix(lines[:first_item_idx])
    if leading and previous_row is not None:
        continuation = _clean_text(" ".join(_inspector_line_name(line) for line in leading if _inspector_line_name(line)))
        if continuation:
            raw_cells = [line.text for line in leading if line.text]
            _append_table_continuation(previous_row, continuation, raw_cells)

    table_bbox = _inspector_table_bbox(lines)
    groups: list[list[_InspectorLine]] = []
    current: list[_InspectorLine] = []
    for idx, line in enumerate(lines[first_item_idx:], start=first_item_idx):
        if _inspector_line_item_no(line) is not None:
            if current:
                if any(_inspector_line_has_quantity(candidate) for candidate in current):
                    groups.append(current)
                else:
                    _leading, carried = _inspector_split_leading_prefix(current)
                    pending_pre = [*pending_pre, *carried]
            current = [*pending_pre, line]
            pending_pre = []
            continue
        if current:
            if (
                any(_inspector_line_has_quantity(candidate) for candidate in current)
                and _inspector_line_starts_position(line)
                and _inspector_next_item_no_soon(lines, idx)
            ):
                groups.append(current)
                current = []
                pending_pre = [line]
            else:
                current.append(line)
        else:
            pending_pre.append(line)
    if current:
        groups.append(current)

    parsed: list[ParsedRdRow] = []
    for group_idx, group in enumerate(groups, start=1):
        row = _inspector_group_to_row(
            group,
            page_number=page_number,
            table_index=1,
            row_number=group_idx,
            table_bbox=table_bbox,
        )
        if row:
            parsed.append(row)
    return parsed


def _parse_pdf_with_pdf_inspector(path: Path) -> list[ParsedRdRow]:
    try:
        from pdf_inspector import detect_pdf, extract_text_with_positions  # type: ignore
    except ImportError as exc:
        raise RdParserError("pdf-inspector is not installed") from exc

    try:
        detection = detect_pdf(str(path))
        pdf_type = _clean_text(getattr(detection, "pdf_type", ""))
        if pdf_type and pdf_type not in {"text_based", "mixed", "text"}:
            return []
    except Exception:
        pass

    try:
        raw_items = extract_text_with_positions(str(path)) or []
    except Exception as exc:
        raise RdParserError(f"pdf-inspector не смог прочитать PDF: {exc}") from exc
    inspector_items = [item for item in (_inspector_text_from_item(raw_item) for raw_item in raw_items) if item is not None]
    if not inspector_items:
        return []

    parsed: list[ParsedRdRow] = []
    lines_by_page: dict[int, list[_InspectorLine]] = {}
    for line in _inspector_lines(inspector_items):
        if line.page <= 0:
            continue
        lines_by_page.setdefault(line.page, []).append(line)
    for page_number in sorted(lines_by_page):
        page_rows = _rows_from_pdf_inspector_page(
            lines_by_page[page_number],
            page_number=page_number,
            previous_row=parsed[-1] if parsed else None,
        )
        parsed.extend(page_rows)
    return _dedupe_parsed_rows(parsed)

def _parse_pdf(path: Path, *, engine: str = PDF_ENGINE_STABLE) -> list[ParsedRdRow]:
    if engine == PDF_ENGINE_PDF_INSPECTOR:
        return _parse_pdf_with_pdf_inspector(path)
    if engine != PDF_ENGINE_STABLE:
        raise RdParserError(f"Unknown PDF parser engine: {engine}")
    try:
        import pdfplumber  # type: ignore
    except ImportError as exc:
        raise RdParserError("pdfplumber is not installed") from exc

    parsed: list[ParsedRdRow] = []
    text_chars = 0
    with pdfplumber.open(str(path)) as pdf:
        for page_idx, page in enumerate(pdf.pages, start=1):
            page_rows: list[ParsedRdRow] = []
            try:
                seen_table_keys: set[tuple[str, str]] = set()

                def collect_table_sources(
                    table_settings_candidates: list[Optional[dict[str, Any]]],
                ) -> list[tuple[list[list[Any]], Optional[dict[str, float]], Optional[list[Optional[dict[str, float]]]]]]:
                    table_sources: list[tuple[list[list[Any]], Optional[dict[str, float]], Optional[list[Optional[dict[str, float]]]]]] = []
                    for table_settings in table_settings_candidates:
                        try:
                            if table_settings is None:
                                found_tables = page.find_tables() or []
                            else:
                                found_tables = page.find_tables(table_settings=table_settings) or []
                        except Exception:
                            found_tables = []
                        for found_table in found_tables:
                            try:
                                table_data = found_table.extract() or []
                            except Exception:
                                table_data = []
                            if not table_data:
                                continue
                            table_bbox = _bbox_dict_from_values(getattr(found_table, "bbox", None))
                            table_row_bboxes = _table_row_bboxes(found_table)
                            bbox_key = "" if table_bbox is None else ",".join(
                                f"{float(table_bbox.get(key) or 0):.1f}" for key in ("x0", "top", "x1", "bottom")
                            )
                            text_key = "|".join("/".join(_clean_text(cell) for cell in row) for row in table_data[:4])
                            table_key = (bbox_key, _normalize_text(text_key)[:240])
                            if table_key in seen_table_keys:
                                continue
                            seen_table_keys.add(table_key)
                            table_sources.append((table_data, table_bbox, table_row_bboxes))
                    return table_sources

                table_index_sequence = {"next": 1}

                def append_table_rows(
                    table_sources: list[tuple[list[list[Any]], Optional[dict[str, float]], Optional[list[Optional[dict[str, float]]]]]],
                ) -> None:
                    for table, table_bbox, table_row_bboxes in table_sources:
                        table_idx = table_index_sequence["next"]
                        table_index_sequence["next"] += 1
                        continuations: list[tuple[str, list[str]]] = []
                        table_rows = _rows_from_table(
                            table,
                            page_number=page_idx,
                            table_index=table_idx,
                            table_bbox=table_bbox,
                            table_row_bboxes=table_row_bboxes,
                            page_continuations=continuations,
                        )
                        if continuations:
                            previous = page_rows[-1] if page_rows else (parsed[-1] if parsed else None)
                            if previous is not None:
                                for text, raw_cells in continuations:
                                    _append_table_continuation(previous, text, raw_cells)
                        page_rows.extend(table_rows)

                line_table_sources = collect_table_sources(
                    [
                        None,
                        {
                            "vertical_strategy": "lines",
                            "horizontal_strategy": "lines",
                            "snap_tolerance": 3,
                            "join_tolerance": 3,
                            "intersection_tolerance": 5,
                        },
                    ]
                )
                if not line_table_sources:
                    line_table_sources = [(table, None, None) for table in (page.extract_tables() or [])]
                append_table_rows(line_table_sources)
                if not page_rows:
                    text_table_sources = collect_table_sources(
                        [
                            {
                                "vertical_strategy": "text",
                                "horizontal_strategy": "text",
                                "snap_tolerance": 3,
                                "join_tolerance": 3,
                                "intersection_tolerance": 5,
                                "min_words_vertical": 2,
                                "min_words_horizontal": 1,
                            }
                        ]
                    )
                    append_table_rows(text_table_sources)
            except Exception:
                page_rows = []

            if not page_rows:
                try:
                    text = page.extract_text(x_tolerance=2, y_tolerance=3, layout=True) or ""
                except Exception:
                    text = ""
                text_chars += len(text.strip())
                if text.strip():
                    page_rows = _rows_from_layout_text(text, page_number=page_idx)

            parsed.extend(page_rows)
            try:
                page.flush_cache()
            except Exception:
                pass
            if page_idx % 5 == 0:
                gc.collect()
    if text_chars < 20 and not parsed:
        return []
    return _dedupe_parsed_rows(parsed)


def _parse_docx(path: Path) -> list[ParsedRdRow]:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise RdParserError("python-docx is not installed") from exc

    document = Document(str(path))
    parsed: list[ParsedRdRow] = []
    for table_idx, table in enumerate(document.tables, start=1):
        table_rows = [[cell.text for cell in row.cells] for row in table.rows]
        parsed.extend(_rows_from_table(table_rows, page_number=1, table_index=table_idx))
    return _dedupe_parsed_rows(parsed)


def _parse_xlsx(path: Path) -> list[ParsedRdRow]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise RdParserError("openpyxl is not installed") from exc

    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    parsed: list[ParsedRdRow] = []
    try:
        visible_sheet_idx = 0
        for worksheet in workbook.worksheets:
            if getattr(worksheet, "sheet_state", "visible") != "visible":
                continue
            visible_sheet_idx += 1
            table_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            if not any(any(_clean_text(cell) for cell in row) for row in table_rows):
                continue
            sheet_title = _clean_text(getattr(worksheet, "title", ""))
            page_label = f"Лист {visible_sheet_idx}" + (f" · {sheet_title}" if sheet_title else "")
            sheet_rows = _rows_from_table(table_rows, page_number=visible_sheet_idx, table_index=1)
            for parsed_row in sheet_rows:
                parsed_row.page_label = page_label
            parsed.extend(sheet_rows)
    finally:
        workbook.close()
    return _dedupe_parsed_rows(parsed)


def _parse_document(path: Path, filename: str, *, pdf_engine: str = PDF_ENGINE_STABLE) -> list[ParsedRdRow]:
    if pdf_engine not in PDF_ENGINES:
        raise RdParserError(f"Unknown PDF parser engine: {pdf_engine}")
    name = filename.lower()
    if name.endswith(".pdf"):
        return _parse_pdf(path, engine=pdf_engine)
    if name.endswith(".docx"):
        return _parse_docx(path)
    if name.endswith((".xlsx", ".xlsm")):
        return _parse_xlsx(path)
    raise RdParserError("Поддерживаются только PDF, DOCX и XLSX")



def _row_match_status(row: dict[str, Any]) -> str:
    if not row.get("object_id") or not row.get("work_type_id") or not row.get("unit") or row.get("volume") in (None, ""):
        return "unresolved"
    try:
        volume = float(row.get("volume") or 0)
    except (TypeError, ValueError):
        return "unresolved"
    if volume <= 0:
        return "unresolved"
    if float(row.get("confidence") or 0) < LOW_CONFIDENCE_THRESHOLD:
        return "low_confidence"
    return "resolved"


def _row_has_publish_fields(row: dict[str, Any]) -> bool:
    if not row.get("object_id") or not row.get("work_type_id") or not row.get("unit") or row.get("volume") in (None, ""):
        return False
    try:
        return float(row.get("volume") or 0) > 0
    except (TypeError, ValueError):
        return False



def _as_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _aggregate_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    buckets: dict[tuple, dict[str, Any]] = {}
    for row in rows:
        if _row_match_status(row) != "resolved":
            continue
        pk_start = _as_float(row.get("pk_start"))
        pk_end = _as_float(row.get("pk_end"))
        if pk_start is not None and pk_end is None:
            pk_end = pk_start
        if pk_end is not None and pk_start is None:
            pk_start = pk_end
        if pk_start is not None and pk_end is not None and pk_start > pk_end:
            pk_start, pk_end = pk_end, pk_start
        key = (
            str(row.get("object_id") or ""),
            str(row.get("work_type_id") or ""),
            str(row.get("material_id") or ""),
            str(row.get("unit") or ""),
            "" if pk_start is None else f"{pk_start:.6f}",
            "" if pk_end is None else f"{pk_end:.6f}",
            str(row.get("pk_raw") or ""),
        )
        bucket = buckets.setdefault(
            key,
            {
                "object_id": row.get("object_id"),
                "object_code": row.get("object_code"),
                "object_name": row.get("object_name"),
                "work_type_id": row.get("work_type_id"),
                "work_type_code": row.get("work_type_code"),
                "work_type_name": row.get("work_type_name"),
                "material_id": row.get("material_id"),
                "material_code": row.get("material_code"),
                "material_name": row.get("material_name"),
                "unit": row.get("unit"),
                "pk_start": pk_start,
                "pk_end": pk_end,
                "pk_raw": row.get("pk_raw"),
                "volume": 0.0,
                "row_ids": [],
                "source_row_count": 0,
            },
        )
        bucket["volume"] += float(row.get("volume") or 0)
        bucket["row_ids"].append(row.get("id"))
        bucket["source_row_count"] += 1
    result = list(buckets.values())
    for item in result:
        item["volume"] = round(float(item["volume"]), 6)
    result.sort(
        key=lambda item: (
            str(item.get("object_name") or item.get("object_code") or ""),
            str(item.get("work_type_name") or item.get("work_type_code") or ""),
            item.get("pk_start") is None,
            float(item.get("pk_start") or 0),
        )
    )
    return result



MATERIAL_KEYWORDS = (
    "песок",
    "песчано гравийная смесь",
    "пгс",
    "щпгс",
    "смесь щебеночно песчаная",
    "щебень",
    "щеб",
    "грунт",
    "геотекстиль",
    "геотекст",
    "геомат",
    "бетон",
    "асфальтобетонная смесь",
    "битумная эмульсия",
    "цемент",
    "арматура",
    "анкер",
)

WORK_ACTION_KEYWORDS = (
    "устройств",
    "транспортировк",
    "дальность возк",
    "возк",
    "перевозк",
    "погрузк",
    "разработк",
    "выборк",
    "выторфовк",
    "забивк",
    "досыпк",
    "планировк",
    "укреплен",
    "укреплени",
    "укладк",
    "розлив",
    "бетонирован",
    "надвижк",
    "обратная засыпк",
    "уплотнен",
    "уплотнени",
    "распределен",
    "сняти",
)

MATERIAL_START_PATTERNS = (
    "песок",
    "песчано гравийная смесь",
    "пгс",
    "щпгс",
    "смесь щебеночно песчаная",
    "щебень",
    "грунт",
    "геотекстиль",
    "геомат",
    "бетон",
    "асфальтобетонная смесь",
    "битумная эмульсия",
    "цемент",
    "арматура",
    "анкер",
)


def parsed_row_to_dict(row: ParsedRdRow) -> dict[str, Any]:
    return asdict(row)


def parsed_rows_to_dicts(rows: list[ParsedRdRow]) -> list[dict[str, Any]]:
    return [parsed_row_to_dict(row) for row in rows]


def rows_from_table(
    table_rows: list[list[Any]],
    *,
    page_number: Optional[int] = None,
    table_index: Optional[int] = None,
    table_bbox: Optional[dict[str, float]] = None,
) -> list[ParsedRdRow]:
    return _rows_from_table(
        table_rows,
        page_number=page_number,
        table_index=table_index,
        table_bbox=table_bbox,
        allow_estimated_row_bboxes=True,
    )


def rows_from_layout_text(text: str, *, page_number: int) -> list[ParsedRdRow]:
    return _rows_from_layout_text(text, page_number=page_number)


def parse_document(
    path: str | Path,
    filename: Optional[str] = None,
    *,
    pdf_engine: str = PDF_ENGINE_STABLE,
) -> list[ParsedRdRow]:
    doc_path = Path(path)
    return _parse_document(doc_path, filename or doc_path.name, pdf_engine=pdf_engine)


def parse_vor_document(
    path: str | Path,
    filename: Optional[str] = None,
    *,
    pdf_engine: str = PDF_ENGINE_STABLE,
) -> list[ParsedRdRow]:
    doc_path = Path(path)
    rows = _parse_document(doc_path, filename or doc_path.name, pdf_engine=pdf_engine)
    return [
        row
        for row in rows
        if _clean_text(row.item_no) and row.table_index is not None and row.unit and row.volume is not None
    ]


def _first_norm_index(norm: str, tokens: tuple[str, ...]) -> Optional[int]:
    positions = [idx for token in tokens if (idx := norm.find(_normalize_text(token))) >= 0]
    return min(positions) if positions else None


def _starts_with_material(norm: str) -> bool:
    for token in MATERIAL_START_PATTERNS:
        token_norm = _normalize_text(token)
        if norm == token_norm or norm.startswith(f"{token_norm} "):
            return True
    return False


def is_material_like(name: Any) -> bool:
    norm = _normalize_text(name)
    if not norm:
        return False
    first_material = _first_norm_index(norm, MATERIAL_KEYWORDS)
    if first_material is None:
        return False
    first_action = _first_norm_index(norm, WORK_ACTION_KEYWORDS)
    if first_action is not None and first_action <= first_material:
        return False
    return True


def review_status_for_row(row: ParsedRdRow | dict[str, Any]) -> str:
    name = row.raw_name if isinstance(row, ParsedRdRow) else row.get("raw_name")
    return "материал / проверить" if is_material_like(name) else "работа / проверить"


def row_match_status(row: dict[str, Any]) -> str:
    return _row_match_status(row)


def aggregate_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return _aggregate_rows(rows)
