#!/usr/bin/env python3
from __future__ import annotations

from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

import wip_routes


def _merged_ranges(ws) -> set[str]:
    return {str(merged) for merged in ws.merged_cells.ranges}


def _row_output_values(ws, row: int) -> list[object]:
    return [
        ws.cell(row=row, column=col).value
        for col in range(
            wip_routes.GENERAL_ANALYTICS_VALUE_START_COL,
            wip_routes.GENERAL_ANALYTICS_VALUE_START_COL + 1 + len(wip_routes.GENERAL_ANALYTICS_SECTION_CODES),
        )
    ]


def _ready_general_template_paths() -> list[Path]:
    repo_root = Path(wip_routes.__file__).resolve().parent.parent
    return [
        repo_root / "frontend" / "public" / "templates" / wip_routes.GENERAL_ANALYTICS_TEMPLATE_FILENAME,
        repo_root / "frontend" / "dist" / "templates" / wip_routes.GENERAL_ANALYTICS_TEMPLATE_FILENAME,
    ]


def _conditional_format_sqrefs(ws) -> set[str]:
    return {str(item.sqref) for item in ws.conditional_formatting}


def _completion_metric_rows(ws) -> list[int]:
    labels = {
        "% выполнения объема плана за день",
        "% выполнения объема плана за неделю",
        "% выполнения объема плана за месяц",
        "% выполнения объема плана за все время",
        "% выполнения проектного объема",
    }
    rows: list[int] = []
    for row_idx in range(1, ws.max_row + 1):
        if str(ws.cell(row=row_idx, column=3).value or "").strip() in labels:
            rows.append(row_idx)
    return rows


def test_resolve_general_analytics_template_prefers_ready_candidates_before_legacy(monkeypatch, tmp_path):
    ready = tmp_path / wip_routes.GENERAL_ANALYTICS_TEMPLATE_FILENAME
    legacy = tmp_path / wip_routes.GENERAL_ANALYTICS_LEGACY_TEMPLATE_FILENAME
    ready.write_text("ready")
    legacy.write_text("legacy")

    def fake_candidates(filename: str):
        return [ready] if filename == wip_routes.GENERAL_ANALYTICS_TEMPLATE_FILENAME else [legacy]

    monkeypatch.setattr(wip_routes, "GENERAL_ANALYTICS_TEMPLATE_ENV", None)
    monkeypatch.setattr(wip_routes, "_general_analytics_template_candidates", fake_candidates)

    assert wip_routes._resolve_general_analytics_template() == ready



def test_resolve_general_analytics_template_falls_back_to_legacy_when_ready_missing(monkeypatch, tmp_path):
    legacy = tmp_path / wip_routes.GENERAL_ANALYTICS_LEGACY_TEMPLATE_FILENAME
    legacy.write_text("legacy")

    def fake_candidates(filename: str):
        return [] if filename == wip_routes.GENERAL_ANALYTICS_TEMPLATE_FILENAME else [legacy]

    monkeypatch.setattr(wip_routes, "GENERAL_ANALYTICS_TEMPLATE_ENV", None)
    monkeypatch.setattr(wip_routes, "_general_analytics_template_candidates", fake_candidates)

    assert wip_routes._resolve_general_analytics_template() == legacy


def test_ready_general_analytics_templates_keep_completion_bars_on_metric_rows():
    stale_sqrefs = {
        "E236:M236",
        "E295:M295",
        "E354:M354",
        "E413:M413",
        "E492:M492",
        "E571:M571",
        "E650:M650",
    }
    expected_layout: set[str] | None = None

    for template_path in _ready_general_template_paths():
        assert template_path.exists(), f"missing ready template: {template_path}"
        wb = load_workbook(template_path)
        ws = wb["Аналитика"]
        metric_rows = _completion_metric_rows(ws)
        sqrefs = _conditional_format_sqrefs(ws)
        expected_sqrefs = {f"E{row_idx}:M{row_idx}" for row_idx in metric_rows}

        assert len(metric_rows) == 8
        assert expected_sqrefs.issubset(sqrefs)
        assert stale_sqrefs.isdisjoint(sqrefs)

        if expected_layout is None:
            expected_layout = expected_sqrefs
        else:
            assert expected_sqrefs == expected_layout



def test_general_equipment_class_detects_supported_slash_descriptors_only():
    cases = [
        ({"equipment_type": "номер", "unit_number": "242/340м3/экскаватор/ЖДС"}, "excavator"),
        ({"equipment_type": "номер", "unit_number": "A123/самосвал/ЖДС"}, "dump_truck"),
        ({"equipment_type": "номер", "unit_number": "T-77/тонар/ЖДС"}, "dump_truck"),
        ({"equipment_type": "номер", "unit_number": "BD-9/бульдозер/ЖДС"}, "bulldozer"),
    ]

    for row, expected in cases:
        assert wip_routes._general_equipment_class_from_row(row) == expected


def test_general_equipment_class_excludes_cranes_and_pile_drivers_for_now():
    rows = [
        {"equipment_type": "кран", "unit_number": "КС-557/ЖДС"},
        {"equipment_type": "номер", "unit_number": "КС-557/кран/ЖДС"},
        {"equipment_type": "сваебойная установка", "unit_number": "СП-49/ЖДС"},
        {"equipment_type": "номер", "unit_number": "СП-49/сваебойная установка/ЖДС"},
        {"equipment_type": "копер", "unit_number": "КОП-1/ЖДС"},
    ]

    for row in rows:
        assert wip_routes._general_equipment_class_from_row(row) is None


def test_general_equipment_work_count_slots_use_descriptor_when_type_is_generic():
    rows = [
        {
            "section_code": "UCH_2",
            "shift": "day",
            "equipment_type": "номер",
            "unit_number": "242/340м3/экскаватор/ЖДС",
            "category": "SAND",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
        },
        {
            "section_code": "UCH_2",
            "shift": "night",
            "equipment_type": "номер",
            "unit_number": "A123/самосвал/ЖДС",
            "category": "SAND",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
        },
        {
            "section_code": "UCH_3",
            "shift": "day",
            "equipment_type": "номер",
            "unit_number": "BD-9/бульдозер/ЖДС",
            "category": "PRS",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
        },
    ]

    slots = wip_routes._general_equipment_work_count_slots_from_rows(rows, rules=[])

    assert slots["excavator"]["SAND"]["day"]["UCH_2"] == 1.0
    assert slots["excavator"]["SAND"]["total"]["UCH_2"] == 1.0
    assert slots["dump_truck"]["SAND"]["night"]["UCH_2"] == 1.0
    assert slots["dump_truck"]["SAND"]["total"]["UCH_2"] == 1.0
    assert slots["bulldozer"]["PRS"]["day"]["UCH_3"] == 1.0


def test_general_vyemka_category_includes_vytorfovka_tag_and_work_type():
    assert wip_routes._general_work_category("Выторфовка", None, "TEMP_ROAD") == "VYEMKA"
    assert wip_routes._general_work_category("Выторфовка", None, "MAIN_TRACK") == "VYEMKA_OH"
    assert wip_routes._general_work_category(None, "PEAT_REMOVAL", "TEMP_ROAD") == "VYEMKA"
    assert wip_routes._general_work_category(None, "PEAT_REMOVAL", "MAIN_TRACK") == "VYEMKA_OH"


def test_general_default_vyemka_rules_also_route_kanavy_and_kyuvety_into_vyemka_rows():
    rules = wip_routes._default_general_report_work_rows()

    assert wip_routes._general_effective_work_category(rules, "Канавы", None, "TEMP_ROAD", "AD_1") == "VYEMKA"
    assert wip_routes._general_effective_work_category(rules, "Канавы", None, "MAIN_TRACK", "MAIN_1") == "VYEMKA_OH"
    assert wip_routes._general_effective_work_category(rules, "Кюветы", None, "ISSO_ACCESS", "ISSO_1") == "VYEMKA_OTHER"
    assert wip_routes._general_effective_work_category(rules, "Кюветы", None, "STATION_TRACK", "ST_1") == "VYEMKA_OH"


def test_general_effective_work_category_routes_pile_access_into_pile_buckets() -> None:
    rules = wip_routes._default_general_report_work_rows()

    assert wip_routes._general_effective_work_category(rules, "Выемка", "EARTH_EXCAVATION", "PILE_DRIVING_PAD", "PFP_1") == "VYEMKA_PILE"
    assert wip_routes._general_effective_work_category(rules, "Выемка", "EARTH_EXCAVATION", "TECH_ROAD", "PFT_10N") == "VYEMKA_PILE"
    assert wip_routes._general_effective_work_category(rules, "Песок (Насыпь)", "EMBANKMENT_CONSTRUCTION", "PILE_DRIVING_PAD", "PFP_1") == "SAND_PILE"
    assert wip_routes._general_effective_work_category(rules, "Песок (Насыпь)", "EMBANKMENT_CONSTRUCTION", "TECH_ROAD", "PFT_10N") == "SAND_PILE"
    assert wip_routes._general_effective_work_category(rules, "Канавы", None, "PILE_DRIVING_PAD", "PFP_1") == "VYEMKA_PILE"
    assert wip_routes._general_effective_work_category(rules, "Кюветы", None, "TECH_ROAD", "PFT_10N") == "VYEMKA_PILE"


def test_general_station_tracks_route_earthwork_and_sand_to_oh_bucket():
    rules = wip_routes._default_general_report_work_rows()

    assert wip_routes._general_work_category("Выторфовка", None, "STATION_TRACK") == "VYEMKA_OH"
    assert wip_routes._general_work_category(None, "PEAT_REMOVAL", "STATION_TRACK") == "VYEMKA_OH"
    assert wip_routes._general_work_category("Песок (Насыпь)", None, "STATION_TRACK") == "SAND_OH"
    assert wip_routes._general_work_category(None, "EMBANKMENT_CONSTRUCTION", "STATION_TRACK") == "SAND_OH"

    assert wip_routes._general_effective_work_category(rules, "Выторфовка", None, "STATION_TRACK", None) == "VYEMKA_OH"
    assert wip_routes._general_effective_work_category(rules, "Песок (Насыпь)", None, "STATION_TRACK", None) == "SAND_OH"


def test_general_additional_split_categories_cover_tp_and_station_scopes():
    assert wip_routes._general_additional_split_categories("PRS", "TEMP_ROAD", "VAD_1") == ["PRS_VAD"]
    assert wip_routes._general_additional_split_categories("PRS", "MAIN_TRACK", "MAIN_1") == ["PRS_OH"]
    assert wip_routes._general_additional_split_categories("PRS", "PTP", "PTP_4") == ["PRS_VALDAI"]
    assert wip_routes._general_additional_split_categories("PRS", "PTP", "PTP_8") == ["PRS_VYPOLZOVO"]
    assert wip_routes._general_additional_split_categories("PRS", "TRACTION_SUBSTATION", "TRACTION_SUBSTATION_2905") == ["PRS_VALDAI"]
    assert wip_routes._general_additional_split_categories("PRS", "TRACTION_SUBSTATION", "STATION_TRACK_3294") == ["PRS_VYPOLZOVO"]
    assert wip_routes._general_additional_split_categories("SAND", "PTP", "PTP_4") == ["SAND_VALDAI"]
    assert wip_routes._general_additional_split_categories("SAND", "TRACTION_SUBSTATION", "STATION_TRACK_3294") == ["SAND_VYPOLZOVO"]
    assert wip_routes._general_additional_split_categories("VYEMKA_OTHER", "PTP", "PTP_8") == ["VYEMKA_VYPOLZOVO"]
    assert wip_routes._general_additional_split_categories("VYEMKA_OTHER", "TRACTION_SUBSTATION", "TRACTION_SUBSTATION_2905") == ["VYEMKA_VALDAI"]



def test_query_general_work_plans_matches_zero_length_pk_points_to_sections():
    captured_sql: list[str] = []

    def fake_query_one(*_args, **_kwargs):
        return {"exists": True}

    def fake_query(sql, params):
        captured_sql.append(sql)
        return []

    with mock.patch.object(wip_routes, "query_one", side_effect=fake_query_one), mock.patch.object(wip_routes, "query", side_effect=fake_query):
        result = wip_routes._query_general_work_plans(date(2026, 6, 1), date(2026, 6, 29), None)

    assert result == {}
    assert len(captured_sql) == 1
    assert "pwi.pk_start = pwi.pk_end" in captured_sql[0]
    assert "os.pk_start = os.pk_end" in captured_sql[0]
    assert "pwi.pk_start < GREATEST(csv.pk_start, csv.pk_end)" in captured_sql[0]



def test_general_vyemka_category_routes_other_objects_to_misc_but_equipment_keeps_broad_bucket():
    assert wip_routes._general_work_category("Выторфовка", None, "ISSO_ACCESS") == "VYEMKA_OTHER"
    assert wip_routes._general_work_category(None, "PEAT_REMOVAL", "ISSO_ACCESS") == "VYEMKA_OTHER"
    assert wip_routes._general_equipment_work_category([], None, "PEAT_REMOVAL", "ISSO_ACCESS", None) == "VYEMKA"


def test_general_pile_field_related_objects_route_into_pile_bucket():
    assert wip_routes._general_additional_split_categories("PRS", "PILE_FIELD", "PF_1") == ["PRS_PILE"]
    assert wip_routes._general_additional_split_categories("PRS", "PILE_DRIVING_PAD", "PFP_1") == ["PRS_PILE"]
    assert wip_routes._general_additional_split_categories("PRS", "TECH_ROAD", "PFT_10N") == ["PRS_PILE"]
    assert wip_routes._general_work_category("Выторфовка", None, "PILE_FIELD") == "VYEMKA_PILE"
    assert wip_routes._general_work_category(None, "PEAT_REMOVAL", "PILE_DRIVING_PAD") == "VYEMKA_PILE"
    assert wip_routes._general_work_category(None, "PEAT_REMOVAL", "TECH_ROAD", "PFT_10N") == "VYEMKA_PILE"
    assert wip_routes._general_work_category("Песок (Насыпь)", None, "PILE_FIELD") == "SAND_PILE"
    assert wip_routes._general_work_category(None, "EMBANKMENT_CONSTRUCTION", "PILE_DRIVING_PAD") == "SAND_PILE"
    assert wip_routes._general_work_category(None, "EMBANKMENT_CONSTRUCTION", "TECH_ROAD", "PFT_10N") == "SAND_PILE"
    assert wip_routes._general_equipment_work_category([], None, "PEAT_REMOVAL", "PILE_FIELD", None) == "VYEMKA"


def test_general_equipment_vyemka_includes_pads_and_isso_access():
    assert wip_routes._general_work_category(None, "PEAT_REMOVAL", "PILE_DRIVING_PAD") == "VYEMKA_PILE"
    assert wip_routes._general_equipment_work_category([], None, "PEAT_REMOVAL", "PILE_DRIVING_PAD", None) == "VYEMKA"
    assert wip_routes._general_equipment_work_category([], "Выторфовка", None, "ISSO_ACCESS", None) == "VYEMKA"
    assert wip_routes._general_equipment_work_category([], "Выемка", None, "BRIDGE", None) == "VYEMKA"


def test_general_equipment_kanavy_and_kyuvety_stay_in_ditch_rows_even_when_work_rows_merge_into_vyemka():
    rules = wip_routes._default_general_report_work_rows()

    assert wip_routes._general_equipment_work_category(rules, "Канавы", None, "TEMP_ROAD", "AD_1") == "DITCH"
    assert wip_routes._general_equipment_work_category(rules, "Кюветы", None, "MAIN_TRACK", "MAIN_1") == "DITCH"
    assert wip_routes._general_equipment_work_category(rules, None, "DITCH_CONSTRUCTION", "ISSO_ACCESS", "ISSO_1") == "DITCH"


def test_general_equipment_work_counts_put_pad_vytorfovka_into_vyemka():
    rows = [
        {
            "section_code": "UCH_7",
            "shift": "day",
            "equipment_type": "экскаватор",
            "unit_number": "EX-1",
            "wt_code": "PEAT_REMOVAL",
            "object_type_code": "PILE_DRIVING_PAD",
            "analytics_tag": "Выторфовка",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
        },
        {
            "section_code": "UCH_8",
            "shift": "night",
            "equipment_type": "бульдозер",
            "unit_number": "BD-1",
            "wt_code": "EARTH_EXCAVATION",
            "object_type_code": "ISSO_ACCESS",
            "analytics_tag": "Выемка",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
        },
    ]

    slots = wip_routes._general_equipment_work_count_slots_from_rows(rows, rules=[])

    assert slots["excavator"]["VYEMKA"]["day"]["UCH_7"] == 1.0
    assert slots["bulldozer"]["VYEMKA"]["night"]["UCH_8"] == 1.0


def test_general_equipment_counts_exclude_mechanization_report_m_snapshot_source():
    captured_sql: list[str] = []

    def fake_query(sql, params):
        captured_sql.append(sql)
        return []

    with mock.patch.object(wip_routes, "query", side_effect=fake_query):
        result = wip_routes._query_general_equipment_counts(date(2026, 6, 14), None)

    assert result == {}
    assert len(captured_sql) == 1
    assert "COALESCE(dr_scope.source_type, '') <> 'mechanization_report_m'" in captured_sql[0]


def test_general_dump_truck_plan_counts_exclude_mechanization_report_m_snapshot_source():
    captured_sql: list[str] = []

    def fake_query(sql, params):
        captured_sql.append(sql)
        return []

    with mock.patch.object(wip_routes, "query", side_effect=fake_query):
        result = wip_routes._query_general_dump_truck_sand_productivity_plan_counts(date(2026, 6, 14), None)

    assert result["total"] == wip_routes._section_values()
    assert len(captured_sql) == 1
    assert "COALESCE(dr_scope.source_type, '') <> 'mechanization_report_m'" in captured_sql[0]


def test_general_dump_truck_plan_counts_use_linked_transport_units_once_per_day(monkeypatch):
    rows = [
        {
            "section_code": "UCH_1",
            "report_date": date(2026, 7, 6),
            "shift": "day",
            "movement_type": "pit_to_constructive",
            "equipment_type": "самосвал",
            "unit_key": "sam-1",
        },
        {
            "section_code": "UCH_1",
            "report_date": date(2026, 7, 6),
            "shift": "night",
            "movement_type": "pit_to_constructive",
            "equipment_type": "самосвал",
            "unit_key": "sam-1",
        },
        {
            "section_code": "UCH_1",
            "report_date": date(2026, 7, 6),
            "shift": "night",
            "movement_type": "pit_to_constructive",
            "equipment_type": "самосвал",
            "unit_key": "sam-2",
        },
    ]

    monkeypatch.setattr(wip_routes, "query", lambda *_args, **_kwargs: rows)
    monkeypatch.setattr(
        wip_routes,
        "DUMP_TRUCK_NORMS",
        {(1, "SAND"): {"trips": 9, "m3_per_trip": 16, "include_pit_to_stockpile": True}},
    )

    result = wip_routes._query_general_dump_truck_sand_productivity_plan_counts(date(2026, 7, 6), ["UCH_1"])

    assert result["day"]["UCH_1"] == 1
    assert result["night"]["UCH_1"] == 2
    assert result["total"]["UCH_1"] == 2
    assert result["non_hired"]["UCH_1"] == 2



def _general_fact_slot(**buckets):
    slot = wip_routes._new_work_fact_slot()
    for bucket, values in buckets.items():
        slot[bucket].update(values)
    return slot


def _write_ready_general_volume_period(ws, start_row: int, period_label: str) -> dict[str, int]:
    ws.cell(row=start_row, column=2, value=period_label)
    ws.cell(row=start_row + 2, column=3, value="% выполнения объема плана")

    row = start_row + 3
    anchors: dict[str, int] = {"percent": start_row + 2}

    def add_row(
        label_b: str,
        label_c: str,
        d_label: str,
        anchor_key: str | None = None,
        note: str | None = None,
    ) -> None:
        nonlocal row
        if label_b:
            ws.cell(row=row, column=2, value=label_b)
        if label_c:
            ws.cell(row=row, column=3, value=label_c)
        ws.cell(row=row, column=4, value=d_label)
        if note:
            ws.cell(row=row, column=15, value=note)
        if anchor_key:
            anchors[anchor_key] = row
        row += 1

    add_row("ПРС", "ВСЕГО", "План", "prs_plan")
    add_row("", "", "Факт ВСЕГО", "prs_fact")
    add_row("", "ВАД", "План", "prs_vad_plan")
    add_row("", "", "Факт", "prs_vad_fact")
    add_row("", "ОХ", "План", "prs_oh_plan")
    add_row("", "", "Факт", "prs_oh_fact")
    add_row("", "Свайные поля", "План", "prs_pile_plan")
    add_row("", "", "Факт", "prs_pile_fact")
    add_row("", "ТП", "План", "prs_tp_plan", "Тяговая подстанция Валдай + Тяговая подстанция Выползово")
    add_row("", "", "Факт ТП", "prs_tp_fact", "Тяговая подстанция Валдай + Тяговая подстанция Выползово")
    add_row("", "", "План", "prs_valdai_plan", "Тяговая подстанция Валдай")
    add_row("", "", "Факт Валдай", "prs_valdai_fact", "Тяговая подстанция Валдай")
    add_row("", "", "План", "prs_vypolzovo_plan", "Тяговая подстанция Выползово")
    add_row("", "", "Факт Выползово", "prs_vypolzovo_fact", "Тяговая подстанция Выползово")
    add_row("", "Прочее", "План", "prs_other_plan", "Все объекты, кроме временных дорог, основного хода и тяговых подстанций")
    add_row("", "", "Факт", "prs_other_fact", "Все объекты, кроме временных дорог, основного хода и тяговых подстанций")

    add_row("Завоз песка", "ВСЕГО", "План", "sand_haul_plan")
    add_row("", "", "Факт ВСЕГО", "sand_haul_total_fact")
    add_row("", "", "Факт ЖДС", "sand_haul_non_hired_fact")
    add_row("", "", "Факт НАЕМ", "sand_haul_hired_fact")
    add_row("", "", "Факт АЛМАЗ", "sand_haul_almaz_fact")
    add_row("", "", "Факт местный грунт", "sand_haul_local_soil_fact")

    add_row("Укладка песка", "ВСЕГО", "План", "sand_total_plan")
    add_row("", "", "Факт ВСЕГО", "sand_total_fact")
    add_row("", "ВАД", "План", "sand_vad_plan")
    add_row("", "", "Факт", "sand_vad_fact")
    add_row("", "ОХ", "План", "sand_oh_plan")
    add_row("", "", "Факт", "sand_oh_fact")
    add_row("", "Свайные поля", "План", "sand_pile_plan")
    add_row("", "", "Факт", "sand_pile_fact")
    add_row("", "ТП", "План", "sand_tp_plan", "Тяговая подстанция Валдай + Тяговая подстанция Выползово")
    add_row("", "", "Факт ТП", "sand_tp_fact", "Тяговая подстанция Валдай + Тяговая подстанция Выползово")
    add_row("", "", "План", "sand_valdai_plan", "Тяговая подстанция Валдай")
    add_row("", "", "Факт Валдай", "sand_valdai_fact", "Тяговая подстанция Валдай")
    add_row("", "", "План", "sand_vypolzovo_plan", "Тяговая подстанция Выползово")
    add_row("", "", "Факт Выползово", "sand_vypolzovo_fact", "Тяговая подстанция Выползово")
    add_row("", "Прочее", "План", "sand_other_plan", "Все объекты, кроме временных дорог, основного хода и тяговых подстанций")
    add_row("", "", "Факт", "sand_other_fact", "Все объекты, кроме временных дорог, основного хода и тяговых подстанций")

    add_row("Выемка", "ВСЕГО", "План", "vyemka_total_plan")
    add_row("", "", "Факт ВСЕГО", "vyemka_total_fact")
    add_row("", "ВАД", "План", "vyemka_vad_plan")
    add_row("", "", "Факт", "vyemka_vad_fact")
    add_row("", "ОХ", "План", "vyemka_oh_plan")
    add_row("", "", "Факт", "vyemka_oh_fact")
    add_row("", "Свайные поля", "План", "vyemka_pile_plan")
    add_row("", "", "Факт", "vyemka_pile_fact")
    add_row("", "ТП", "План", "vyemka_tp_plan", "Тяговая подстанция Валдай + Тяговая подстанция Выползово")
    add_row("", "", "Факт ТП", "vyemka_tp_fact", "Тяговая подстанция Валдай + Тяговая подстанция Выползово")
    add_row("", "", "План", "vyemka_valdai_plan", "Тяговая подстанция Валдай")
    add_row("", "", "Факт Валдай", "vyemka_valdai_fact", "Тяговая подстанция Валдай")
    add_row("", "", "План", "vyemka_vypolzovo_plan", "Тяговая подстанция Выползово")
    add_row("", "", "Факт Выползово", "vyemka_vypolzovo_fact", "Тяговая подстанция Выползово")
    add_row("", "Прочее", "План", "vyemka_other_plan", "Все объекты, кроме временных дорог, основного хода и тяговых подстанций")
    add_row("", "", "Факт", "vyemka_other_fact", "Все объекты, кроме временных дорог, основного хода и тяговых подстанций")

    add_row("Щебень", "ВСЕГО", "План", "scheben_plan")
    add_row("", "", "Факт ВСЕГО", "scheben_fact")
    add_row("", "", "Факт завоз в накопитель", "scheben_stockpile_fact")

    add_row("ЩПС", "ВСЕГО", "План", "shps_plan")
    add_row("", "", "Факт ВСЕГО", "shps_fact")
    add_row("", "", "Факт завоз в накопитель", "shps_stockpile_fact")
    return anchors


def test_ready_general_volume_summary_does_not_put_unclassified_sand_into_vad():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = _write_ready_general_volume_period(ws, 2, "Текущий день")

    work_periods = {
        "day": {
            "SAND": _general_fact_slot(total={"UCH_2": 100.0}),
            "SAND_VAD": _general_fact_slot(total={"UCH_2": 30.0}),
        }
    }

    wip_routes._fill_ready_general_volume_summary(ws, work_periods, {}, {}, wip_routes._section_values())

    sand_vad_fact_row = anchors["sand_vad_fact"]
    sand_other_fact_row = anchors["sand_other_fact"]
    assert wip_routes._read_section_values_at_row(ws, sand_vad_fact_row)["UCH_2"] == 30.0
    assert ws.cell(row=sand_vad_fact_row, column=wip_routes.GENERAL_ANALYTICS_VALUE_START_COL).value == 30
    assert wip_routes._read_section_values_at_row(ws, sand_other_fact_row)["UCH_2"] == 100.0


def test_ready_general_volume_summary_zero_fills_empty_plan_rows():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = _write_ready_general_volume_period(ws, 2, "Текущий день")

    wip_routes._fill_ready_general_volume_summary(ws, {}, {}, {}, wip_routes._section_values())

    assert _row_output_values(ws, anchors["shps_plan"]) == [0] * (1 + len(wip_routes.GENERAL_ANALYTICS_SECTION_CODES))
    assert _row_output_values(ws, anchors["sand_haul_plan"]) == [0] * (1 + len(wip_routes.GENERAL_ANALYTICS_SECTION_CODES))


def test_write_ready_plan_from_counts_zero_fills_empty_related_work_plan_rows():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"

    row_map = {"plan_total": 10, "plan_day": 11, "plan_night": 12}
    wip_routes._write_ready_plan_from_counts(ws, row_map, {}, {}, wip_routes._constant_norms(1.0))

    for row in row_map.values():
        assert _row_output_values(ws, row) == [0] * (1 + len(wip_routes.GENERAL_ANALYTICS_SECTION_CODES))


def test_write_optional_section_values_at_row_zero_fills_missing_manual_plan_counts():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"

    wip_routes._write_optional_section_values_at_row(ws, 10, {})

    assert _row_output_values(ws, 10) == [0] * (1 + len(wip_routes.GENERAL_ANALYTICS_SECTION_CODES))


def test_write_ready_general_manual_equipment_count_rows_zero_fill_when_payload_missing():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    _seed_ready_bulldozer_block(ws)

    wip_routes._write_ready_general_manual_equipment_count_rows(ws, None)

    assert _row_output_values(ws, 120) == [0] * (1 + len(wip_routes.GENERAL_ANALYTICS_SECTION_CODES))
    assert _row_output_values(ws, 121) == [0] * (1 + len(wip_routes.GENERAL_ANALYTICS_SECTION_CODES))


def test_ready_general_volume_category_recognizes_excavation_plus_peat_total_labels():
    assert wip_routes._ready_general_volume_category_from_label("Выемка+выторфовка ВСЕГО") == "VYEMKA_TOTAL"
    assert wip_routes._ready_general_volume_category_from_label("Выемка/выторфовка ВСЕГО") == "VYEMKA_TOTAL"


def test_find_ready_general_volume_rows_detects_subtotals_in_all_period_blocks():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = {
        "day": _write_ready_general_volume_period(ws, 2, "Текущий день"),
        "week": _write_ready_general_volume_period(ws, 80, "Текущая неделя"),
        "month": _write_ready_general_volume_period(ws, 158, "Текущий месяц"),
        "total": _write_ready_general_volume_period(ws, 236, "Накопительным итогом"),
    }

    rows = wip_routes._find_ready_general_volume_rows(ws)

    for period_key, period_anchors in anchors.items():
        assert rows[period_key]["categories"]["PRS_PILE"]["plan"] == period_anchors["prs_pile_plan"]
        assert rows[period_key]["categories"]["PRS_TP"]["plan"] == period_anchors["prs_tp_plan"]
        assert rows[period_key]["categories"]["PRS_VALDAI"]["plan"] == period_anchors["prs_valdai_plan"]
        assert rows[period_key]["categories"]["PRS_VYPOLZOVO"]["plan"] == period_anchors["prs_vypolzovo_plan"]
        assert rows[period_key]["categories"]["SAND_TOTAL"]["plan"] == period_anchors["sand_total_plan"]
        assert rows[period_key]["categories"]["SAND_PILE"]["plan"] == period_anchors["sand_pile_plan"]
        assert rows[period_key]["categories"]["SAND_TP"]["plan"] == period_anchors["sand_tp_plan"]
        assert rows[period_key]["categories"]["SAND_VALDAI"]["plan"] == period_anchors["sand_valdai_plan"]
        assert rows[period_key]["categories"]["SAND_VYPOLZOVO"]["plan"] == period_anchors["sand_vypolzovo_plan"]
        assert rows[period_key]["categories"]["SAND"]["plan"] == period_anchors["sand_other_plan"]
        assert rows[period_key]["categories"]["VYEMKA_TOTAL"]["plan"] == period_anchors["vyemka_total_plan"]
        assert rows[period_key]["categories"]["VYEMKA_PILE"]["plan"] == period_anchors["vyemka_pile_plan"]
        assert rows[period_key]["categories"]["VYEMKA_TP"]["plan"] == period_anchors["vyemka_tp_plan"]
        assert rows[period_key]["categories"]["VYEMKA_VALDAI"]["plan"] == period_anchors["vyemka_valdai_plan"]
        assert rows[period_key]["categories"]["VYEMKA_VYPOLZOVO"]["plan"] == period_anchors["vyemka_vypolzovo_plan"]
        assert rows[period_key]["categories"]["VYEMKA_OTHER"]["plan"] == period_anchors["vyemka_other_plan"]
        assert rows[period_key]["categories"]["SCHEBEN_STOCKPILE"]["total"] == period_anchors["scheben_stockpile_fact"]
        assert rows[period_key]["categories"]["SHPS_STOCKPILE"]["total"] == period_anchors["shps_stockpile_fact"]
        assert rows[period_key]["sand_haul"]["local_soil"] == period_anchors["sand_haul_local_soil_fact"]


def test_ready_general_volume_summary_splits_tp_rows_without_double_counting_misc():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = _write_ready_general_volume_period(ws, 2, "Текущий день")

    work_periods = {
        "day": {
            "PRS": _general_fact_slot(total={"UCH_1": 130.0}),
            "PRS_VAD": _general_fact_slot(total={"UCH_1": 10.0}),
            "PRS_OH": _general_fact_slot(total={"UCH_1": 20.0}),
            "PRS_PILE": _general_fact_slot(total={"UCH_1": 30.0}),
            "PRS_VALDAI": _general_fact_slot(total={"UCH_1": 15.0}),
            "PRS_VYPOLZOVO": _general_fact_slot(total={"UCH_1": 5.0}),
            "SAND": _general_fact_slot(total={"UCH_1": 60.0}),
            "SAND_VAD": _general_fact_slot(total={"UCH_1": 10.0}),
            "SAND_OH": _general_fact_slot(total={"UCH_1": 20.0}),
            "SAND_PILE": _general_fact_slot(total={"UCH_1": 12.0}),
            "SAND_VALDAI": _general_fact_slot(total={"UCH_1": 7.0}),
            "SAND_VYPOLZOVO": _general_fact_slot(total={"UCH_1": 8.0}),
            "VYEMKA": _general_fact_slot(total={"UCH_1": 12.0}),
            "VYEMKA_OH": _general_fact_slot(total={"UCH_1": 3.0}),
            "VYEMKA_PILE": _general_fact_slot(total={"UCH_1": 9.0}),
            "VYEMKA_OTHER": _general_fact_slot(total={"UCH_1": 25.0}),
            "VYEMKA_VALDAI": _general_fact_slot(total={"UCH_1": 4.0}),
            "VYEMKA_VYPOLZOVO": _general_fact_slot(total={"UCH_1": 6.0}),
        }
    }

    wip_routes._fill_ready_general_volume_summary(ws, work_periods, {}, {}, wip_routes._section_values())

    assert wip_routes._read_section_values_at_row(ws, anchors["prs_pile_fact"])["UCH_1"] == 30.0
    assert wip_routes._read_section_values_at_row(ws, anchors["prs_tp_fact"])["UCH_1"] == 20.0
    assert wip_routes._read_section_values_at_row(ws, anchors["prs_valdai_fact"])["UCH_1"] == 15.0
    assert wip_routes._read_section_values_at_row(ws, anchors["prs_vypolzovo_fact"])["UCH_1"] == 5.0
    assert wip_routes._read_section_values_at_row(ws, anchors["prs_other_fact"])["UCH_1"] == 50.0
    assert wip_routes._read_section_values_at_row(ws, anchors["sand_pile_fact"])["UCH_1"] == 12.0
    assert wip_routes._read_section_values_at_row(ws, anchors["sand_total_fact"])["UCH_1"] == 102.0
    assert wip_routes._read_section_values_at_row(ws, anchors["sand_tp_fact"])["UCH_1"] == 15.0
    assert wip_routes._read_section_values_at_row(ws, anchors["sand_other_fact"])["UCH_1"] == 45.0
    assert wip_routes._read_section_values_at_row(ws, anchors["vyemka_pile_fact"])["UCH_1"] == 9.0
    assert wip_routes._read_section_values_at_row(ws, anchors["vyemka_total_fact"])["UCH_1"] == 49.0
    assert wip_routes._read_section_values_at_row(ws, anchors["vyemka_tp_fact"])["UCH_1"] == 10.0
    assert wip_routes._read_section_values_at_row(ws, anchors["vyemka_other_fact"])["UCH_1"] == 15.0


def test_ready_general_volume_summary_writes_metadata_columns_for_tp_and_local_soil_rows():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = _write_ready_general_volume_period(ws, 2, "Текущий день")

    wip_routes._fill_ready_general_volume_summary(ws, {}, {}, {}, wip_routes._section_values())

    assert ws.cell(row=1, column=wip_routes.GENERAL_ANALYTICS_META_OBJECT_COL).value == "Объекты / source"
    assert ws.cell(row=1, column=wip_routes.GENERAL_ANALYTICS_META_TAG_COL).value == "Теги / work types"
    assert ws.cell(row=1, column=wip_routes.GENERAL_ANALYTICS_META_LOGIC_COL).value == "Логика / исключения"
    assert "TRACTION_SUBSTATION_2905" in str(ws.cell(row=anchors["prs_tp_plan"], column=wip_routes.GENERAL_ANALYTICS_META_OBJECT_COL).value)
    assert "PILE_FIELD" in str(ws.cell(row=anchors["prs_pile_plan"], column=wip_routes.GENERAL_ANALYTICS_META_OBJECT_COL).value)
    assert "PFT_*" in str(ws.cell(row=anchors["prs_pile_plan"], column=wip_routes.GENERAL_ANALYTICS_META_OBJECT_COL).value)
    assert "TOPSOIL_STRIPPING" in str(ws.cell(row=anchors["prs_tp_plan"], column=wip_routes.GENERAL_ANALYTICS_META_TAG_COL).value)
    assert "ТП = Валдай + Выползово" in str(ws.cell(row=anchors["prs_tp_plan"], column=wip_routes.GENERAL_ANALYTICS_META_LOGIC_COL).value)
    assert "material=SAND" in str(ws.cell(row=anchors["sand_haul_total_fact"], column=wip_routes.GENERAL_ANALYTICS_META_TAG_COL).value)
    assert "SOIL_LOCAL" in str(ws.cell(row=anchors["sand_haul_local_soil_fact"], column=wip_routes.GENERAL_ANALYTICS_META_TAG_COL).value)
    assert "SAND_TOTAL fact − sand haul delivered" in str(ws.cell(row=anchors["sand_haul_local_soil_fact"], column=wip_routes.GENERAL_ANALYTICS_META_LOGIC_COL).value)
    assert ws.cell(row=1, column=wip_routes.GENERAL_ANALYTICS_META_OBJECT_COL).alignment.wrap_text is False
    assert ws.cell(row=anchors["prs_tp_plan"], column=wip_routes.GENERAL_ANALYTICS_META_OBJECT_COL).alignment.wrap_text is False
    assert ws.cell(row=anchors["prs_tp_plan"], column=wip_routes.GENERAL_ANALYTICS_META_TAG_COL).alignment.wrap_text is False
    assert ws.cell(row=anchors["prs_tp_plan"], column=wip_routes.GENERAL_ANALYTICS_META_LOGIC_COL).alignment.wrap_text is False


def test_ready_general_volume_summary_uses_subtotal_rows_for_percent_in_all_period_blocks():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = {
        "day": _write_ready_general_volume_period(ws, 2, "Текущий день"),
        "week": _write_ready_general_volume_period(ws, 80, "Текущая неделя"),
        "month": _write_ready_general_volume_period(ws, 158, "Текущий месяц"),
        "total": _write_ready_general_volume_period(ws, 236, "Накопительным итогом"),
    }

    work_periods = {}
    plan_periods = {}
    movement_periods = {}
    for period_key in anchors:
        work_periods[period_key] = {
            "SAND": _general_fact_slot(total={"UCH_1": 50.0}),
            "SAND_VAD": _general_fact_slot(total={"UCH_1": 10.0}),
            "SAND_OH": _general_fact_slot(total={"UCH_1": 0.0}),
            "VYEMKA": _general_fact_slot(total={"UCH_1": 10.0}),
            "VYEMKA_OH": _general_fact_slot(total={"UCH_1": 5.0}),
        }
        plan_periods[period_key] = {
            "SAND": _general_fact_slot(total={"UCH_1": 100.0}),
            "SAND_VAD": _general_fact_slot(total={"UCH_1": 10.0}),
            "SAND_OH": _general_fact_slot(total={"UCH_1": 5.0}),
            "VYEMKA": _general_fact_slot(total={"UCH_1": 20.0}),
            "VYEMKA_OH": _general_fact_slot(total={"UCH_1": 5.0}),
        }
        movement_periods[period_key] = {}

    wip_routes._fill_ready_general_volume_summary(ws, work_periods, plan_periods, movement_periods, wip_routes._section_values())

    expected_ratio = 0.536
    for period_anchors in anchors.values():
        percent_row = period_anchors["percent"]
        assert ws.cell(row=percent_row, column=5).value == expected_ratio
        assert ws.cell(row=percent_row, column=6).value == expected_ratio
        assert ws.cell(row=period_anchors["scheben_stockpile_fact"], column=5).value == 0
        assert ws.cell(row=period_anchors["shps_stockpile_fact"], column=5).value == 0


def test_ready_general_volume_summary_fills_new_17_06_local_soil_other_and_dual_stockpile_rows():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = _write_ready_general_volume_period(ws, 2, "Текущий день")

    work_periods = {
        "day": {
            "SOIL_LOCAL": _general_fact_slot(total={"UCH_2": 15.0}),
            "SAND": _general_fact_slot(total={"UCH_2": 40.0}),
            "VYEMKA_OTHER": _general_fact_slot(total={"UCH_2": 12.0}),
        }
    }
    plan_periods = {
        "day": {
            "SAND": _general_fact_slot(total={"UCH_2": 50.0}),
            "VYEMKA_OTHER": _general_fact_slot(total={"UCH_2": 20.0}),
        }
    }
    movement_periods = {
        "day": {
            "SAND_SITE_DELIVERED": _general_fact_slot(total={"UCH_2": 30.0}, non_hired={"UCH_2": 10.0}, hired={"UCH_2": 12.0}, almaz={"UCH_2": 8.0}),
            "SCHEBEN_PIT_TO_STOCKPILE": _general_fact_slot(total={"UCH_2": 6.0}),
            "SHPS_PIT_TO_STOCKPILE": _general_fact_slot(total={"UCH_2": 7.0}),
        }
    }

    wip_routes._fill_ready_general_volume_summary(ws, work_periods, plan_periods, movement_periods, wip_routes._section_values())

    assert wip_routes._read_section_values_at_row(ws, anchors["sand_haul_total_fact"])["UCH_2"] == 30.0
    assert wip_routes._read_section_values_at_row(ws, anchors["sand_haul_local_soil_fact"])["UCH_2"] == 15.0
    assert wip_routes._read_section_values_at_row(ws, anchors["sand_other_fact"])["UCH_2"] == 40.0
    assert wip_routes._read_section_values_at_row(ws, anchors["vyemka_other_fact"])["UCH_2"] == 12.0
    assert wip_routes._read_section_values_at_row(ws, anchors["scheben_stockpile_fact"])["UCH_2"] == 6.0
    assert wip_routes._read_section_values_at_row(ws, anchors["shps_stockpile_fact"])["UCH_2"] == 7.0


def test_ready_general_volume_summary_total_local_soil_row_uses_sand_fact_minus_haul_fact():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = _write_ready_general_volume_period(ws, 2, "Накопительным итогом")

    work_periods = {
        "total": {
            "SOIL_LOCAL": _general_fact_slot(total={"UCH_2": 15.0}),
            "SAND": _general_fact_slot(total={"UCH_2": 100.0}),
        }
    }
    movement_periods = {
        "total": {
            "SAND_DASHBOARD_SITE_DELIVERED": _general_fact_slot(total={"UCH_2": 40.0}),
        }
    }

    wip_routes._fill_ready_general_volume_summary(ws, work_periods, {}, movement_periods, wip_routes._section_values())

    assert wip_routes._read_section_values_at_row(ws, anchors["sand_total_fact"])["UCH_2"] == 100.0
    assert wip_routes._read_section_values_at_row(ws, anchors["sand_haul_total_fact"])["UCH_2"] == 40.0
    assert ws.cell(row=anchors["sand_haul_local_soil_fact"], column=wip_routes.GENERAL_ANALYTICS_VALUE_START_COL).value == 60.0
    assert wip_routes._read_section_values_at_row(ws, anchors["sand_haul_local_soil_fact"])["UCH_2"] == 60.0


def test_ready_general_volume_summary_total_local_soil_total_cell_subtracts_hidden_negative_sections():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = _write_ready_general_volume_period(ws, 2, "Накопительным итогом")

    work_periods = {
        "total": {
            "SAND": _general_fact_slot(total={"UCH_1": 100.0, "UCH_2": 20.0}),
        }
    }
    movement_periods = {
        "total": {
            "SAND_DASHBOARD_SITE_DELIVERED": _general_fact_slot(total={"UCH_1": 40.0, "UCH_2": 55.0}),
        }
    }

    wip_routes._fill_ready_general_volume_summary(ws, work_periods, {}, movement_periods, wip_routes._section_values())

    row = anchors["sand_haul_local_soil_fact"]
    assert ws.cell(row=row, column=wip_routes.GENERAL_ANALYTICS_VALUE_START_COL).value == 25.0
    assert wip_routes._read_section_values_at_row(ws, row)["UCH_1"] == 60.0
    assert wip_routes._read_section_values_at_row(ws, row)["UCH_2"] == 0.0


def test_ready_general_volume_summary_total_local_soil_section_cells_never_go_negative():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = _write_ready_general_volume_period(ws, 2, "Накопительным итогом")

    work_periods = {
        "total": {
            "SAND": _general_fact_slot(total={"UCH_2": 20.0}),
        }
    }
    movement_periods = {
        "total": {
            "SAND_DASHBOARD_SITE_DELIVERED": _general_fact_slot(total={"UCH_2": 55.0}),
        }
    }

    wip_routes._fill_ready_general_volume_summary(ws, work_periods, {}, movement_periods, wip_routes._section_values())

    assert wip_routes._read_section_values_at_row(ws, anchors["sand_haul_local_soil_fact"])["UCH_2"] == 0.0


def test_ready_general_volume_summary_month_local_soil_row_never_goes_negative():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = _write_ready_general_volume_period(ws, 2, "Текущий месяц")

    work_periods = {
        "month": {
            "SOIL_LOCAL": _general_fact_slot(total={"UCH_7": -181454.489}),
        }
    }

    wip_routes._fill_ready_general_volume_summary(ws, work_periods, {}, {}, wip_routes._section_values())

    assert wip_routes._read_section_values_at_row(ws, anchors["sand_haul_local_soil_fact"])["UCH_7"] == 0.0


def test_ready_general_volume_summary_keeps_shps_stockpile_row_equal_to_quarry_delivery_only_outside_cumulative_total():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = _write_ready_general_volume_period(ws, 2, "Текущий день")

    work_periods = {
        "day": {
            "SHPS": _general_fact_slot(total={"UCH_2": 11.0}),
        }
    }
    movement_periods = {
        "day": {
            "SHPS_PIT_TO_STOCKPILE": _general_fact_slot(total={"UCH_2": 7.0}),
        }
    }

    wip_routes._fill_ready_general_volume_summary(ws, work_periods, {}, movement_periods, wip_routes._section_values())

    assert wip_routes._read_section_values_at_row(ws, anchors["shps_fact"])["UCH_2"] == 11.0
    assert wip_routes._read_section_values_at_row(ws, anchors["shps_stockpile_fact"])["UCH_2"] == 7.0


def test_ready_general_volume_summary_prefers_dashboard_shpgs_delivery_total_for_cumulative_stockpile_row():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    anchors = _write_ready_general_volume_period(ws, 2, "Накопительным итогом")

    work_periods = {
        "total": {
            "SHPS": _general_fact_slot(total={"UCH_2": 11.0}),
        }
    }
    movement_periods = {
        "total": {
            "SHPS_PIT_TO_STOCKPILE": _general_fact_slot(total={"UCH_2": 7.0}),
            "SHPGS_DASHBOARD_SITE_DELIVERED": _general_fact_slot(total={"UCH_2": 13.0}),
        }
    }

    wip_routes._fill_ready_general_volume_summary(ws, work_periods, {}, movement_periods, wip_routes._section_values())

    assert wip_routes._read_section_values_at_row(ws, anchors["shps_fact"])["UCH_2"] == 11.0
    assert wip_routes._read_section_values_at_row(ws, anchors["shps_stockpile_fact"])["UCH_2"] == 13.0


def test_legacy_general_volume_summary_keeps_shps_stockpile_row_equal_to_quarry_delivery_only_outside_cumulative_total():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"

    work_periods = {
        "day": {
            "SHPS": _general_fact_slot(total={"UCH_2": 11.0}),
        }
    }
    movement_periods = {
        "day": {
            "SHPS_PIT_TO_STOCKPILE": _general_fact_slot(total={"UCH_2": 7.0}),
        }
    }

    wip_routes._fill_general_volume_summary(ws, work_periods, {}, movement_periods, wip_routes._section_values())

    shps_fact_row = wip_routes._general_target_row(252)
    shps_stockpile_row = wip_routes._general_target_row(253)
    assert shps_fact_row is not None and shps_stockpile_row is not None
    assert wip_routes._read_section_values_at_row(ws, shps_fact_row)["UCH_2"] == 11.0
    assert wip_routes._read_section_values_at_row(ws, shps_stockpile_row)["UCH_2"] == 7.0


def test_legacy_general_volume_summary_prefers_dashboard_shpgs_delivery_total_for_cumulative_stockpile_row():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"

    work_periods = {
        "total": {
            "SHPS": _general_fact_slot(total={"UCH_2": 11.0}),
        }
    }
    movement_periods = {
        "total": {
            "SHPS_PIT_TO_STOCKPILE": _general_fact_slot(total={"UCH_2": 7.0}),
            "SHPGS_DASHBOARD_SITE_DELIVERED": _general_fact_slot(total={"UCH_2": 13.0}),
        }
    }

    wip_routes._fill_general_volume_summary(ws, work_periods, {}, movement_periods, wip_routes._section_values())

    shps_fact_row = wip_routes._general_target_row(321)
    shps_stockpile_row = wip_routes._general_target_row(322)
    assert shps_fact_row is not None and shps_stockpile_row is not None
    assert wip_routes._read_section_values_at_row(ws, shps_fact_row)["UCH_2"] == 11.0
    assert wip_routes._read_section_values_at_row(ws, shps_stockpile_row)["UCH_2"] == 13.0


def test_general_soil_cut_to_embankment_counts_any_source_to_non_stockpile_object():
    base = {
        "material_code": "SOIL",
        "material_name": "Грунт",
        "from_object_type_code": "BORROW_PIT",
        "to_object_type_code": "TEMP_ROAD",
    }

    assert wip_routes._general_is_soil_cut_to_embankment_movement(base) is True
    assert wip_routes._general_is_soil_cut_to_embankment_movement({**base, "from_object_type_code": "MAIN_TRACK"}) is True
    assert wip_routes._general_is_soil_cut_to_embankment_movement({**base, "from_object_type_code": "STOCKPILE"}) is True
    assert wip_routes._general_is_soil_cut_to_embankment_movement({**base, "to_object_type_code": "STOCKPILE"}) is False
    assert wip_routes._general_is_soil_cut_to_embankment_movement({**base, "material_code": "SAND", "material_name": "Песок"}) is False


def test_general_pile_unmerge_does_not_touch_non_pile_rotated_left_cells():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"

    ws["B45"] = "Текущий день"
    ws["B53"] = "Самосвал"
    ws["B53"].alignment = Alignment(textRotation=90, vertical="center")
    ws.merge_cells("B53:B89")

    ws["B304"] = "Текущий день"
    ws["B305"] = "Погружено свай"
    ws["B306"] = "Фактическое количество Кран"
    ws["B307"] = "Фактическое количество Сваебойная установка"
    ws.merge_cells("B305:B307")
    ws.merge_cells("C305:C307")

    ranges = wip_routes._unmerge_general_pile_block_label_cells(ws)

    assert "B53:B89" in _merged_ranges(ws)
    assert "B305:B307" not in _merged_ranges(ws)
    assert "C305:C307" not in _merged_ranges(ws)
    assert (305, 2, 307, 2) in ranges


def test_general_excluded_pile_rows_are_hidden_without_shifting_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B10"] = "Текущая неделя"
    ws["B11"] = "Погружено свай"
    ws["C11"] = "Объем"
    ws["D11"] = "План"
    ws["C12"] = "Фактическое количество Кран"
    ws["C13"] = "Фактическое количество Сваебойная установка"
    ws["D14"] = "Факт, в т.ч."
    ws["D15"] = "факт 11 м"
    ws.merge_cells("B11:B15")
    ws.merge_cells("C12:D12")
    ws.merge_cells("C13:D13")

    hidden = wip_routes._delete_general_excluded_pile_rows(ws)

    assert hidden == [12, 13]
    assert ws.max_row == 15
    assert ws.row_dimensions[12].hidden is True
    assert ws.row_dimensions[13].hidden is True
    assert ws["C12"].value == "Фактическое количество Кран"
    assert ws["C13"].value == "Фактическое количество Сваебойная установка"
    assert "B11:B15" in _merged_ranges(ws)
    assert ws["D15"].value == "факт 11 м"


def test_general_zero_pile_length_rows_are_hidden_without_changing_template_merges():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B10"] = "Текущая неделя"
    ws["B11"] = "Погружено свай"
    ws["C11"] = "Объем"
    ws["D11"] = "План"
    ws["D12"] = "Факт, в т.ч."
    ws["D13"] = "факт 10 м"
    ws["E13"] = 0
    ws["D14"] = "факт 11 м"
    ws["E14"] = 5
    ws["D15"] = "факт 11 м"
    ws["E15"] = 2
    ws.merge_cells("B11:B15")
    ws.merge_cells("C11:C15")

    hidden = wip_routes._hide_zero_or_duplicate_general_pile_length_rows(ws)

    assert hidden == [13]
    assert ws.max_row == 15
    # Hidden rows stay in the sheet and original left-side template merges stay
    # intact; the owner wants full template repetition, not split B/C labels.
    # Duplicate non-zero length rows are not hidden: only zero-fact length rows
    # may be hidden in the ready Генеральный отчет template.
    assert "B11:B15" in _merged_ranges(ws)
    assert "C11:C15" in _merged_ranges(ws)
    assert ws["B11"].value == "Погружено свай"
    assert ws["C11"].value == "Объем"
    assert ws.row_dimensions[13].hidden is True
    assert ws.row_dimensions[15].hidden is not True
    assert ws.row_dimensions[14].hidden is not True


def test_general_pile_left_merges_can_cross_hidden_rows_to_preserve_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B10"] = "Текущий день"
    ws["B11"] = "Погружено пробных свай"
    ws["C11"] = "Объем"
    ws["D11"] = "План"
    ws["D12"] = "Факт, в т.ч."
    for offset, length in enumerate(range(7, 12), start=13):
        ws.cell(offset, 4, f"факт {length} м")
        ws.cell(offset, 5, 0)
    ws.cell(18, 4, "факт 12 м")
    ws.cell(18, 5, 4)
    ws.merge_cells("B11:B18")
    ws.merge_cells("C11:C18")

    hidden = wip_routes._hide_zero_or_duplicate_general_pile_length_rows(ws)

    assert hidden == [13, 14, 15, 16, 17]
    assert "B11:B18" in _merged_ranges(ws)
    assert "C11:C18" in _merged_ranges(ws)
    assert ws["B11"].value == "Погружено пробных свай"
    assert ws["C11"].value == "Объем"


def test_general_zero_pile_length_rows_are_hidden_even_without_left_block_label():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B270"] = "Текущий день"
    ws["D276"] = "факт 7 м"
    ws["E276"] = 0
    ws["D277"] = "факт 8 м"
    ws["E277"] = 3

    hidden = wip_routes._hide_zero_or_duplicate_general_pile_length_rows(ws)

    assert hidden == [276]
    assert ws.row_dimensions[276].hidden is True
    assert ws.row_dimensions[277].hidden is not True


def test_ready_general_shps_plan_rows_are_not_filled_from_quarry_fact_without_engaged_trucks(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["C45"] = "№ Участка"
    ws["E39"] = 10
    ws["F39"] = 3
    ws["M39"] = 7
    ws["E40"] = 20
    ws["F40"] = 4
    ws["M40"] = 16

    monkeypatch.setattr(wip_routes, "_query_general_equipment_counts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_facts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_counts", lambda *args, **kwargs: {}, raising=False)
    monkeypatch.setattr(wip_routes, "_query_general_dump_truck_sand_productivity_plan_counts", lambda *args, **kwargs: wip_routes._new_equipment_fact_slot(), raising=False)
    monkeypatch.setattr(wip_routes, "_fill_general_equipment_plan_rows", lambda *args, **kwargs: None)

    wip_routes._fill_ready_general_equipment(ws, wip_routes.date_cls(2026, 6, 2), None, {}, {})

    assert ws.row_dimensions[74].hidden is not True
    assert ws["E75"].value in (None, 0)
    assert ws["F75"].value in (None, 0)
    assert ws["M75"].value in (None, 0)


def test_ready_general_pdf_layout_preserves_template_print_area_row():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["A586"] = "last report row"
    ws["O10"] = "xlsx-only note"
    ws.print_area = "A1:M351"

    wip_routes._configure_general_pdf_layout(ws, ws.max_row, "M")

    assert str(ws.print_area) == "'Аналитика'!$A$1:$M$586"
    assert ws.max_row == 586
    assert ws["O10"].value == "xlsx-only note"
    assert ws.column_dimensions["O"].hidden is not True


def _seed_ready_bulldozer_block(ws):
    ws["B3"] = "ДЕНЬ"
    ws["B45"] = "Текущий день"
    ws["C45"] = "№ Участка"
    labels = {
        120: ("Бульдозер", "Плановое количество бульдозеров ДЕНЬ", None),
        121: (None, "Плановое количество бульдозеров НОЧЬ", None),
        122: (None, "Фактическое количество бульдозеров ДЕНЬ", None),
        123: (None, "Фактическое количество бульдозеров НОЧЬ", None),
        124: (None, "% от норматива в сутки", None),
        125: (None, "ПРС", "Плановая выработка"),
        126: (None, None, "План за день"),
        127: (None, None, "План за ночь"),
        128: (None, None, "Факт за сутки"),
        129: (None, None, "Факт за день"),
        130: (None, None, "Факт за ночь"),
        131: (None, "ЩПС", "План за сутки по ПП"),
        132: (None, None, "Плановая выработка"),
        133: (None, None, "План за день"),
        134: (None, None, "План за ночь"),
        135: (None, None, "Факт за сутки"),
        136: (None, None, "Факт за день"),
        137: (None, None, "Факт за ночь"),
        138: (None, "Песок", "План за сутки по ПП"),
        139: (None, None, "Плановая выработка"),
        140: (None, None, "План за день"),
        141: (None, None, "План за ночь"),
        142: (None, None, "Факт за сутки"),
        143: (None, None, "Факт за день"),
        144: (None, None, "Факт за ночь"),
        145: (None, "Выемка", "План за сутки по ПП"),
        146: (None, None, "Плановая выработка"),
        147: (None, None, "План за день"),
        148: (None, None, "План за ночь"),
        149: (None, None, "Факт за сутки"),
        150: (None, None, "Факт за день"),
        151: (None, None, "Факт за ночь"),
        152: (None, "Планировка", "План за сутки по ПП"),
        153: (None, None, "Плановая выработка"),
        154: (None, None, "План за день"),
        155: (None, None, "План за ночь"),
        156: (None, None, "Факт за сутки"),
        157: (None, None, "Факт за день"),
        158: (None, None, "Факт за ночь"),
        159: (None, "Сопутствующие работы", "Плановая выработка"),
        160: (None, None, "План за день"),
        161: (None, None, "План за ночь"),
        162: (None, None, "Факт за сутки"),
        163: (None, None, "Факт за день"),
        164: (None, None, "Факт за ночь"),
        167: ("Текущий день", "№ Участка", None),
    }
    for row, (b, c, d) in labels.items():
        ws.cell(row, 2, b)
        ws.cell(row, 3, c)
        ws.cell(row, 4, d)


def test_ready_general_equipment_uses_visible_template_geometry_not_old_offsets(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    _seed_ready_bulldozer_block(ws)

    slot = wip_routes._new_equipment_fact_slot()
    slot["day"]["UCH_1"] = 3
    slot["night"]["UCH_1"] = 4
    related = wip_routes._new_equipment_fact_slot()
    related["total"]["UCH_1"] = 9
    related["day"]["UCH_1"] = 5
    related["night"]["UCH_1"] = 4

    monkeypatch.setattr(wip_routes, "_query_general_equipment_counts", lambda *args, **kwargs: {"bulldozer": slot})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_facts", lambda *args, **kwargs: {"bulldozer": {"RELATED": related}})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_counts", lambda *args, **kwargs: {}, raising=False)
    monkeypatch.setattr(wip_routes, "_query_general_dump_truck_sand_productivity_plan_counts", lambda *args, **kwargs: wip_routes._new_equipment_fact_slot(), raising=False)
    monkeypatch.setattr(wip_routes, "_fill_general_equipment_plan_rows", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_general_equipment_rule_for_class", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_sand_pit_norms", lambda: wip_routes._constant_norms(1))
    monkeypatch.setattr(wip_routes, "_norm_from_settings", lambda _eq, _codes, default: default)

    wip_routes._fill_ready_general_equipment(ws, wip_routes.date_cls(2026, 6, 2), None, {}, {})

    assert ws["F122"].value == 3
    assert ws["F123"].value == 4
    assert ws["F162"].value == 9
    assert ws["F163"].value == 5
    assert ws["F164"].value == 4
    assert ws["F174"].value is None


def test_ready_general_excavator_prs_rows_are_filled_from_prs_tagged_work(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B3"] = "ДЕНЬ"
    ws["B45"] = "Текущий день"
    ws["C45"] = "№ Участка"
    labels = {
        90: ("Экскаватор", "Плановое количество экскаваторов ДЕНЬ", None),
        91: (None, "Плановое количество экскаваторов НОЧЬ", None),
        92: (None, "Фактическое количество экскаваторов ДЕНЬ", None),
        93: (None, "Фактическое количество экскаваторов НОЧЬ", None),
        94: (None, "% от норматива в сутки", None),
        101: (None, "ПРС", "План за сутки по ПП"),
        102: (None, None, "Плановая выработка"),
        103: (None, None, "План за день"),
        104: (None, None, "План за ночь"),
        105: (None, None, "Факт за сутки"),
        106: (None, None, "Факт за день"),
        107: (None, None, "Факт за ночь"),
        108: (None, "Выемка", "План за сутки по ПП"),
        109: (None, None, "Плановая выработка"),
        110: (None, None, "План за день"),
        111: (None, None, "План за ночь"),
        112: (None, None, "Факт за сутки"),
        113: (None, None, "Факт за день"),
        114: (None, None, "Факт за ночь"),
        120: ("Текущий день", "№ Участка", None),
    }
    for row, (b, c, d) in labels.items():
        ws.cell(row, 2, b)
        ws.cell(row, 3, c)
        ws.cell(row, 4, d)

    counts = wip_routes._new_equipment_fact_slot()
    counts["day"]["UCH_1"] = 2
    counts["night"]["UCH_2"] = 1
    facts = wip_routes._new_equipment_fact_slot()
    facts["total"]["UCH_1"] = 11
    facts["day"]["UCH_1"] = 7
    facts["night"]["UCH_2"] = 4

    monkeypatch.setattr(wip_routes, "_query_general_equipment_counts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_facts", lambda *args, **kwargs: {"excavator": {"PRS": facts}})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_counts", lambda *args, **kwargs: {"excavator": {"PRS": counts}}, raising=False)
    monkeypatch.setattr(wip_routes, "_query_general_dump_truck_sand_productivity_plan_counts", lambda *args, **kwargs: wip_routes._new_equipment_fact_slot(), raising=False)
    monkeypatch.setattr(wip_routes, "_fill_general_equipment_plan_rows", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_general_equipment_rule_for_class", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_sand_pit_norms", lambda: wip_routes._constant_norms(1))
    monkeypatch.setattr(wip_routes, "_norm_from_settings", lambda _eq, _codes, default: default)

    wip_routes._fill_ready_general_equipment(ws, wip_routes.date_cls(2026, 6, 10), None, {}, {})

    assert ws["F103"].value == 1400
    assert ws["G104"].value == 700
    assert ws["F105"].value == 11
    assert ws["F106"].value == 7
    assert ws["G107"].value == 4


def test_ready_general_dump_truck_plain_daily_total_not_overwritten_by_source_rows(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B3"] = "ДЕНЬ"
    ws["B45"] = "Текущий день"
    ws["C45"] = "№ Участка"
    ws["B53"] = "Самосвал"
    ws["C53"] = "Плановое количество самосвалов ДЕНЬ"
    ws["C54"] = "Плановое количество самосвалов НОЧЬ"
    ws["C55"] = "Фактическое количество самосвалов ДЕНЬ"
    ws["C56"] = "Фактическое количество самосвалов НОЧЬ"
    ws["C57"] = "% от норматива в сутки"
    ws["C60"] = "Песок"
    ws["D60"] = "План за сутки по ПП"
    ws["D61"] = "Плановая выработка"
    ws["D62"] = "План за день"
    ws["D63"] = "Факт за сутки"
    ws["D64"] = "План за ночь"
    ws["D65"] = "Факт за день"
    ws["D66"] = "Факт за ночь"
    ws["D67"] = "Факт за сутки ЖДС"
    ws["D68"] = "Факт за сутки НАЕМ+АЛМАЗ"
    ws["B70"] = "Экскаватор"

    slot = wip_routes._new_equipment_fact_slot()
    slot["total"]["UCH_1"] = 1402
    slot["day"]["UCH_1"] = 849
    slot["night"]["UCH_1"] = 553
    slot["non_hired"]["UCH_1"] = 1402
    slot["day_non_hired"]["UCH_1"] = 849
    slot["night_non_hired"]["UCH_1"] = 553
    slot["hired"]["UCH_1"] = 100
    slot["almaz"]["UCH_1"] = 20

    monkeypatch.setattr(wip_routes, "_query_general_equipment_counts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_facts", lambda *args, **kwargs: {"dump_truck": {"SAND": slot}})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_counts", lambda *args, **kwargs: {}, raising=False)
    monkeypatch.setattr(wip_routes, "_query_general_dump_truck_sand_productivity_plan_counts", lambda *args, **kwargs: wip_routes._new_equipment_fact_slot(), raising=False)
    monkeypatch.setattr(wip_routes, "_fill_general_equipment_plan_rows", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_general_equipment_rule_for_class", lambda _eq: {"enabled": True, "filters": {"source_buckets": ["non_hired"]}})
    monkeypatch.setattr(wip_routes, "_sand_pit_norms", lambda: wip_routes._constant_norms(1))
    monkeypatch.setattr(wip_routes, "_norm_from_settings", lambda _eq, _codes, default: default)

    wip_routes._fill_ready_general_equipment(ws, wip_routes.date_cls(2026, 6, 8), None, {}, {})

    assert ws["F63"].value == 1402
    assert ws["F65"].value == 849
    assert ws["F66"].value == 553
    assert ws["F67"].value == 1402
    assert ws["F68"].value == 120


def test_general_pile_week_plan_horizon_is_full_iso_week():
    monday = wip_routes.date_cls(2026, 6, 8)
    report_day = wip_routes.date_cls(2026, 6, 8)

    assert wip_routes._general_pile_plan_period("day", report_day, report_day) == (report_day, report_day)
    assert wip_routes._general_pile_plan_period("week", monday, report_day) == (
        monday,
        wip_routes.date_cls(2026, 6, 14),
    )



def _seed_ready_volume_period(ws, start_row: int, period_label: str) -> None:
    ws.cell(start_row, 2, period_label)
    ws.cell(start_row, 3, "№ Участка")
    rows = [
        (start_row + 2, None, "% выполнения объема плана за день", None),
        (start_row + 3, "ПРС", "Объем", "План"),
        (start_row + 4, None, None, "Факт"),
        (start_row + 5, "Песок ВАД", "Объем", "План"),
        (start_row + 6, None, None, "Факт ЖДС"),
        (start_row + 7, None, None, "Факт НАЕМ"),
        (start_row + 8, "Песок ОХ", "Объем", "План"),
        (start_row + 9, None, None, "Факт ЖДС"),
        (start_row + 10, None, None, "Факт НАЕМ"),
        (start_row + 11, "Выемка ВАД", "Объем", "План"),
        (start_row + 12, None, None, "Факт ЖДС"),
        (start_row + 13, None, None, "Факт НАЕМ"),
        (start_row + 14, "Выемка ОХ", "Объем", "План"),
        (start_row + 15, None, None, "Факт ЖДС"),
        (start_row + 16, None, None, "Факт НАЕМ"),
        (start_row + 17, "Щебень", "Объем", "План"),
        (start_row + 18, None, None, "Факт ЖДС"),
        (start_row + 19, None, None, "Факт НАЕМ"),
        (start_row + 20, "ЩПС", "Объем", "План"),
        (start_row + 21, None, None, "Факт ЖДС"),
        (start_row + 22, None, None, "Факт НАЕМ"),
        (start_row + 23, None, None, "Факт завоз в накопитель"),
    ]
    for row, b, c, d in rows:
        ws.cell(row, 2, b)
        ws.cell(row, 3, c)
        ws.cell(row, 4, d)


def _seed_ready_volume_period_with_totals(ws, start_row: int, period_label: str) -> None:
    ws.cell(start_row, 2, period_label)
    ws.cell(start_row, 3, "№ Участка")
    rows = [
        (start_row + 2, None, "% выполнения объема плана за день", None),
        (start_row + 3, "Завоз песка", "Объем", "Факт ЖДС"),
        (start_row + 4, None, None, "Факт НАЕМ"),
        (start_row + 5, None, None, "Факт АЛМАЗ"),
        (start_row + 6, "ПРС", "Объем", "План"),
        (start_row + 7, None, None, "Факт"),
        (start_row + 8, "Укладка песка\nВСЕГО", "Объем", "План"),
        (start_row + 9, None, None, "Факт ЖДС"),
        (start_row + 10, None, None, "Факт НАЕМ"),
        (start_row + 11, "Песок ВАД", "Объем", "План"),
        (start_row + 12, None, None, "Факт ЖДС"),
        (start_row + 13, None, None, "Факт НАЕМ"),
        (start_row + 14, "Песок ОХ", "Объем", "План"),
        (start_row + 15, None, None, "Факт ЖДС"),
        (start_row + 16, None, None, "Факт НАЕМ"),
        (start_row + 17, "Выемка\nВСЕГО", "Объем", "План"),
        (start_row + 18, None, None, "Факт ЖДС"),
        (start_row + 19, None, None, "Факт НАЕМ"),
        (start_row + 20, "Выемка ВАД", "Объем", "План"),
        (start_row + 21, None, None, "Факт ЖДС"),
        (start_row + 22, None, None, "Факт НАЕМ"),
        (start_row + 23, "Выемка ОХ", "Объем", "План"),
        (start_row + 24, None, None, "Факт ЖДС"),
        (start_row + 25, None, None, "Факт НАЕМ"),
        (start_row + 26, "Щебень", "Объем", "План"),
        (start_row + 27, None, None, "Факт ЖДС"),
        (start_row + 28, None, None, "Факт НАЕМ"),
        (start_row + 29, "ЩПС", "Объем", "План"),
        (start_row + 30, None, None, "Факт ЖДС"),
        (start_row + 31, None, None, "Факт НАЕМ"),
        (start_row + 32, None, None, "Факт завоз в накопитель"),
    ]
    for row, b, c, d in rows:
        ws.cell(row, 2, b)
        ws.cell(row, 3, c)
        ws.cell(row, 4, d)


def _seed_ready_volume_period_with_sand_haul_plan(ws, start_row: int, period_label: str, percent_label: str) -> None:
    ws.cell(start_row, 2, period_label)
    ws.cell(start_row + 1, 3, "№ Участка")
    rows = [
        (start_row + 2, None, percent_label, None),
        (start_row + 3, "ПРС", "ВСЕГО", "План"),
        (start_row + 4, None, None, "Факт"),
        (start_row + 5, "Завоз песка", "ВСЕГО", "План"),
        (start_row + 6, None, None, "Факт ВСЕГО"),
        (start_row + 7, None, None, "Факт ЖДС"),
        (start_row + 8, None, None, "Факт НАЕМ"),
        (start_row + 9, None, None, "Факт АЛМАЗ"),
        (start_row + 10, None, None, "Факт местный грунт"),
        (start_row + 11, "Укладка песка", "ВСЕГО", "План"),
        (start_row + 12, None, None, "Факт ВСЕГО"),
    ]
    for row, b, c, d in rows:
        ws.cell(row, 2, b)
        ws.cell(row, 3, c)
        ws.cell(row, 4, d)


def test_ready_general_volume_summary_fills_new_compact_geometry_by_labels():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B3"] = "ДЕНЬ"
    _seed_ready_volume_period(ws, 167, "Текущий день")
    _seed_ready_volume_period(ws, 193, "Текущая неделя")

    plan_day = {"PRS": {"total": {"UCH_1": 10}}, "SAND_VAD": {"total": {"UCH_1": 20}}}
    work_day = {"PRS": {"total": {"UCH_1": 7}}, "SAND_VAD": {"non_hired": {"UCH_1": 8}, "hired": {"UCH_1": 2}}}
    plan_week = {"SAND_VAD": {"total": {"UCH_1": 30}}}
    work_week = {"SAND_VAD": {"non_hired": {"UCH_1": 11}, "almaz": {"UCH_1": 3}, "hired": {"UCH_1": 4}}}

    wip_routes._fill_ready_general_volume_summary(
        ws,
        {"day": work_day, "week": work_week},
        {"day": plan_day, "week": plan_week},
        {},
        {},
    )

    assert ws["F170"].value == 10
    assert ws["F171"].value == 7
    assert ws["F172"].value == 20
    assert ws["F173"].value == 8
    assert ws["F174"].value == 2
    assert ws["F198"].value == 30
    assert ws["F199"].value == 11
    assert ws["F200"].value == 7


def test_ready_general_volume_summary_handles_inserted_total_rows_without_double_counting():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B3"] = "ДЕНЬ"
    _seed_ready_volume_period_with_totals(ws, 183, "Текущий день")

    plan_day = {
        "PRS": {"total": {"UCH_1": 10}},
        "SAND": {"total": {"UCH_1": 3}},
        "SAND_VAD": {"total": {"UCH_1": 20}},
        "SAND_OH": {"total": {"UCH_1": 5}},
        "VYEMKA": {"total": {"UCH_1": 4}},
        "VYEMKA_OH": {"total": {"UCH_1": 6}},
        "SCHEBEN": {"total": {"UCH_1": 4}},
        "SHPS": {"total": {"UCH_1": 6}},
    }
    work_day = {
        "PRS": _general_fact_slot(total={"UCH_1": 7}),
        "SAND": _general_fact_slot(total={"UCH_1": 3}, non_hired={"UCH_1": 3}),
        "SAND_VAD": _general_fact_slot(total={"UCH_1": 10}, non_hired={"UCH_1": 8}, hired={"UCH_1": 2}),
        "SAND_OH": _general_fact_slot(total={"UCH_1": 6}, non_hired={"UCH_1": 5}, almaz={"UCH_1": 1}),
        "VYEMKA": _general_fact_slot(total={"UCH_1": 4}, non_hired={"UCH_1": 4}),
        "VYEMKA_OH": _general_fact_slot(total={"UCH_1": 8}, non_hired={"UCH_1": 6}, hired={"UCH_1": 2}),
        "SCHEBEN": _general_fact_slot(total={"UCH_1": 4}, non_hired={"UCH_1": 4}),
        "SHPS": _general_fact_slot(total={"UCH_1": 6}, non_hired={"UCH_1": 6}),
    }
    movement_day = {
        "SHPS_PIT_TO_STOCKPILE": _general_fact_slot(total={"UCH_1": 9}),
    }

    wip_routes._fill_ready_general_volume_summary(
        ws,
        {"day": work_day},
        {"day": plan_day},
        {"day": movement_day},
        {},
    )

    assert ws["F191"].value == 28
    assert ws["F192"].value == 16
    assert ws["F193"].value == 3
    assert ws["F194"].value == 20
    assert ws["F195"].value == 8
    assert ws["F196"].value == 2
    assert ws["F197"].value == 5
    assert ws["F198"].value == 5
    assert ws["F199"].value == 1
    assert ws["F200"].value == 10
    assert ws["F201"].value == 10
    assert ws["F202"].value == 2
    assert ws["F215"].value == 9
    assert ws["F185"].value == 0.828


def test_ready_general_volume_summary_fills_sand_haul_rows_and_uses_dashboard_total_source():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B3"] = "ДЕНЬ"
    for start_row, label in ((167, "Текущий день"), (245, "Накопительным итогом")):
        ws.cell(start_row, 2, label)
        ws.cell(start_row, 3, "№ Участка")
        ws.cell(start_row + 2, 3, "% выполнения объема плана")
        ws.cell(start_row + 3, 2, "Завоз песка")
        ws.cell(start_row + 3, 3, "Объем")
        ws.cell(start_row + 3, 4, "Факт ЖДС")
        ws.cell(start_row + 4, 4, "Факт НАЕМ")
        ws.cell(start_row + 5, 4, "Факт АЛМАЗ")
        ws.cell(start_row + 6, 2, "ПРС")
        ws.cell(start_row + 6, 3, "Объем")
        ws.cell(start_row + 6, 4, "План")
        ws.cell(start_row + 7, 4, "Факт")

    day_slot = wip_routes._new_work_fact_slot()
    day_slot["non_hired"]["UCH_1"] = 11
    day_slot["hired"]["UCH_1"] = 2
    day_slot["almaz"]["UCH_1"] = 3
    dashboard_total_slot = wip_routes._new_work_fact_slot()
    dashboard_total_slot["non_hired"]["UCH_1"] = 101
    dashboard_total_slot["hired"]["UCH_1"] = 22
    dashboard_total_slot["almaz"]["UCH_1"] = 33
    raw_total_slot = wip_routes._new_work_fact_slot()
    raw_total_slot["non_hired"]["UCH_1"] = 1

    wip_routes._fill_ready_general_volume_summary(
        ws,
        {},
        {},
        {
            "day": {"SAND_SITE_DELIVERED": day_slot},
            "total": {
                "SAND_SITE_DELIVERED": raw_total_slot,
                "SAND_DASHBOARD_SITE_DELIVERED": dashboard_total_slot,
            },
        },
        {},
    )

    assert ws["F170"].value == 11
    assert ws["F171"].value == 2
    assert ws["F172"].value == 3
    assert ws["F248"].value == 101
    assert ws["F249"].value == 22
    assert ws["F250"].value == 33


def test_ready_general_volume_summary_fills_sand_haul_plan_rows_from_all_sand_plans():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B3"] = "ДЕНЬ"
    for start_row, label, percent_label in (
        (175, "Текущий день", "% выполнения объема плана за день"),
        (210, "Текущая неделя", "% выполнения объема плана за неделю"),
        (245, "Текущий месяц", "% выполнения объема плана за месяц"),
        (280, "Накопительным итогом", "% выполнения объема плана за все время"),
    ):
        _seed_ready_volume_period_with_sand_haul_plan(ws, start_row, label, percent_label)

    plan_periods = {
        "day": {
            "SAND": {"total": {"UCH_1": 1}},
            "SAND_VAD": {"total": {"UCH_1": 2}},
            "SAND_OH": {"total": {"UCH_1": 3}},
        },
        "week": {
            "SAND": {"total": {"UCH_1": 4}},
            "SAND_VAD": {"total": {"UCH_1": 5}},
            "SAND_OH": {"total": {"UCH_1": 6}},
        },
        "month": {
            "SAND": {"total": {"UCH_1": 7}},
            "SAND_VAD": {"total": {"UCH_1": 8}},
            "SAND_OH": {"total": {"UCH_1": 9}},
        },
        "total": {
            "SAND": {"total": {"UCH_1": 10}},
            "SAND_VAD": {"total": {"UCH_1": 11}},
            "SAND_OH": {"total": {"UCH_1": 12}},
        },
    }

    wip_routes._fill_ready_general_volume_summary(ws, {}, plan_periods, {}, {})

    assert ws["F180"].value == 6
    assert ws["F215"].value == 15
    assert ws["F250"].value == 24
    assert ws["F285"].value == 33
    assert ws["F186"].value == 6
    assert ws["F221"].value == 15
    assert ws["F256"].value == 24
    assert ws["F291"].value == 33


def test_ready_general_volume_summary_ignores_duplicate_pile_period_blocks():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    _seed_ready_volume_period(ws, 167, "Текущий день")
    _seed_ready_volume_period(ws, 193, "Текущая неделя")
    _seed_ready_volume_period(ws, 219, "Текущий месяц")
    _seed_ready_volume_period(ws, 245, "Накопительным итогом")
    # Same period label/percent/header geometry repeats later in the pile area.
    # The volume fill must keep the first period block and must not overwrite
    # pile rows like "Погружено пробных свай" with volume-summary labels.
    ws["B271"] = "Текущий день"
    ws["C273"] = "% выполнения объема плана за день"
    ws["B274"] = "Погружено пробных свай "
    ws["C274"] = "Объем"
    ws["D274"] = "План"
    ws["D275"] = "Факт, в т.ч."

    wip_routes._fill_ready_general_volume_summary(
        ws,
        {"day": {"PRS": {"total": {"UCH_1": 7}}}},
        {"day": {"PRS": {"total": {"UCH_1": 10}}}},
        {},
        {},
    )

    assert ws["F170"].value == 10
    assert ws["F171"].value == 7
    assert ws["B274"].value == "Погружено пробных свай "
    assert ws["F274"].value is None
    assert ws["F275"].value is None


def test_general_top_quarry_rows_group_migoloshi_spelling_variants(monkeypatch):
    assert wip_routes._general_quarry_label_key("Карьер Миголоши (ЩПГС)") == wip_routes._general_quarry_label_key("Миголощи ЩПГС")
    assert wip_routes._general_quarry_label_key("Карьер Милогоши (ЩПГС)") == wip_routes._general_quarry_label_key("Миголощи ЩПГС")

    monkeypatch.setattr(wip_routes, "query", lambda *args, **kwargs: [
        {
            "section_code": "UCH_3",
            "shift": "day",
            "quarry_name": "Карьер Милогоши",
            "material_code": "SHPGS",
            "material_name": "ЩПГС",
            "labor_source_type": "subcontractor",
            "contractor_name": "Наемники",
            "amount": 5,
        },
        {
            "section_code": "UCH_3",
            "shift": "night",
            "quarry_name": "Карьер Миголощи",
            "material_code": "SHPGS",
            "material_name": "ЩПГС",
            "labor_source_type": "subcontractor",
            "contractor_name": "Наемники",
            "amount": 7,
        },
    ])

    detail_rows, _, _ = wip_routes._query_general_top_quarry_rows(date(2026, 7, 20), None)

    assert len(detail_rows) == 1
    assert detail_rows[0]["label"] == "Миголощи (ЩПГС)"
    assert detail_rows[0]["day"]["UCH_3"] == 5
    assert detail_rows[0]["night"]["UCH_3"] == 7
    assert detail_rows[0]["total"]["UCH_3"] == 12


def test_general_top_quarry_total_adds_visible_almaz_and_hired_rows_even_when_not_marked_for_totals(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"

    def slot(day=0, night=0):
        item = wip_routes._new_general_quarry_slot()
        item["day"]["UCH_1"] = day
        item["night"]["UCH_1"] = night
        item["total"]["UCH_1"] = day + night
        return item

    rules = [
        {
            "label": "ЖДС карьер",
            "enabled": True,
            "include_in_totals": True,
            "filters": {"quarry_labels": ["ЖДС карьер"]},
        },
        {
            "label": "Завоз силами АЛМАЗ",
            "enabled": True,
            "include_in_totals": False,
            "filters": {"quarry_labels": ["Завоз силами АЛМАЗ"]},
        },
        {
            "label": "Завоз песка наемниками",
            "enabled": True,
            "include_in_totals": False,
            "filters": {"quarry_labels": ["Завоз песка наемниками"]},
        },
    ]
    facts = {
        wip_routes._general_quarry_label_key("ЖДС карьер"): slot(day=10),
        wip_routes._general_quarry_label_key("Завоз силами АЛМАЗ"): slot(day=3),
        wip_routes._general_quarry_label_key("Завоз песка наемниками"): slot(night=4),
    }

    monkeypatch.setattr(wip_routes, "_general_report_top_quarry_rules", lambda: rules)
    monkeypatch.setattr(wip_routes, "_query_general_quarry_period_facts", lambda *args, **kwargs: facts)

    wip_routes._write_general_top_quarry_table(ws, wip_routes.date_cls(2026, 6, 5), None, {}, {})

    assert ws["F7"].value == 10
    assert ws["F9"].value == 3
    assert ws["F12"].value == 4
    assert ws["F4"].value == 17
    assert ws["F5"].value == 13
    assert ws["F6"].value == 4


def test_general_top_quarry_uses_latest_template_visible_rows_with_korovkino(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B3"] = "ДЕНЬ"
    ws["B47"] = "Текущий день"
    ws["C47"] = "№ Участка"
    for idx, label in enumerate(wip_routes.GENERAL_REPORT_DEFAULT_TOP_QUARRY_LABELS, start=0):
        ws.cell(7 + idx * 2, 4, label)

    def slot(day=0, night=0):
        item = wip_routes._new_general_quarry_slot()
        item["day"]["UCH_1"] = day
        item["night"]["UCH_1"] = night
        item["total"]["UCH_1"] = day + night
        return item

    facts = {
        wip_routes._general_quarry_label_key("Коровкино (ЩПГС)"): slot(day=6),
        wip_routes._general_quarry_label_key("Завоз песка наемниками"): slot(night=4),
    }

    monkeypatch.setattr(wip_routes, "_general_report_top_quarry_rules", wip_routes._default_general_report_quarry_rows)
    monkeypatch.setattr(wip_routes, "_query_general_quarry_period_facts", lambda *args, **kwargs: facts)

    wip_routes._write_general_top_quarry_table(ws, wip_routes.date_cls(2026, 6, 5), None, {}, {})

    assert ws["D33"].value == "Коровкино (ЩПГС)"
    assert ws["F33"].value == 6
    assert ws["D35"].value == "Южные Маяки (ЩПГС)"
    assert ws["D43"].value == "Завоз песка наемниками"
    assert ws["F44"].value == 4
    assert ws["F4"].value == 4
    assert ws["F5"].value == 0
    assert ws["F6"].value == 4


def test_general_dump_truck_metadata_uses_visible_rows_after_latest_shift(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B3"] = "ДЕНЬ"
    ws["B47"] = "Текущий день"
    ws["C47"] = "№ Участка"
    ws["E47"] = "Всего за день"
    metadata_rows = {
        48: "Карьер",
        49: "Плечо возки,км",
        50: "Кол-во рейсов на 1ед.",
        51: "День",
        52: "Ночь",
        53: "День (накопитель-конструктив)",
        54: "Ночь (накопитель-конструктив)",
    }
    for row, label in metadata_rows.items():
        ws.cell(row, 5, label)
    ws["F48"] = "Шаблонный карьер"
    ws["F49"] = "12"

    defaults = wip_routes._read_general_dump_truck_defaults(ws)
    wip_routes._blank_general_value_cells(ws)
    monkeypatch.setattr(wip_routes, "_query_general_sand_quarry_meta", lambda *args, **kwargs: {code: [] for code in wip_routes.GENERAL_ANALYTICS_SECTION_CODES})

    wip_routes._fill_general_dump_truck_metadata(ws, wip_routes.date_cls(2026, 6, 5), None, defaults)

    assert ws["E47"].value == "Всего за день"
    assert ws["E48"].value == "Карьер"
    assert ws["F48"].value == "Шаблонный карьер"
    assert ws["E49"].value == "Плечо возки,км"
    assert ws["F49"].value == "12"
    assert ws["E50"].value == "Кол-во рейсов на 1ед."


def test_general_soil_cut_to_embankment_counts_soil_into_constructives_regardless_of_movement_type(monkeypatch):
    rows = [
        {
            "section_code": "UCH_1",
            "shift": "day",
            "material_code": "SOIL",
            "material_name": "Грунт",
            "movement_type": "stockpile_to_constructive",
            "from_object_type_code": "STOCKPILE",
            "to_object_type_code": "SERVICE_ROAD",
            "labor_source_type": None,
            "contractor_name": None,
            "amount": 10,
        },
        {
            "section_code": "UCH_1",
            "shift": "night",
            "material_code": "SOIL",
            "material_name": "Грунт",
            "movement_type": "pit_to_constructive",
            "from_object_type_code": "MAIN_TRACK",
            "to_object_type_code": "TECH_ROAD",
            "labor_source_type": None,
            "contractor_name": None,
            "amount": 5,
        },
        {
            "section_code": "UCH_1",
            "shift": "day",
            "material_code": "SOIL",
            "material_name": "Грунт",
            "movement_type": "stockpile_to_stockpile",
            "from_object_type_code": "STOCKPILE",
            "to_object_type_code": "STOCKPILE",
            "labor_source_type": None,
            "contractor_name": None,
            "amount": 99,
        },
        {
            "section_code": "UCH_1",
            "shift": "day",
            "material_code": "SAND",
            "material_name": "Песок",
            "movement_type": "stockpile_to_constructive",
            "from_object_type_code": "STOCKPILE",
            "to_object_type_code": "SERVICE_ROAD",
            "labor_source_type": None,
            "contractor_name": None,
            "amount": 42,
        },
    ]

    def fake_query(sql, params):
        assert "src_ot.code AS from_object_type_code" in sql
        return rows

    monkeypatch.setattr(wip_routes, "query", fake_query)

    facts = wip_routes._query_general_movement_facts(
        wip_routes.date_cls(2026, 6, 1),
        wip_routes.date_cls(2026, 6, 1),
        None,
    )

    slot = facts["SOIL_STOCKPILE_TO_CONSTRUCTIVE"]
    assert slot["day"]["UCH_1"] == 10
    assert slot["night"]["UCH_1"] == 5
    assert slot["total"]["UCH_1"] == 15



def _seed_ready_dump_truck_block(ws):
    ws["B3"] = "ДЕНЬ"
    labels = {
        45: ("Текущий день", "№ Участка", None),
        53: ("Самосвал", "Плановое количество самосвалов ДЕНЬ", None),
        54: (None, "Плановое количество самосвалов НОЧЬ", None),
        55: (None, "Фактическое количество самосвалов ЖДС ДЕНЬ", None),
        56: (None, "Фактическое количество самосвалов ЖДС НОЧЬ", None),
        57: (None, "% от норматива в сутки", None),
        58: (None, "Песок", "Плановая выработка"),
        59: (None, None, "План за день"),
        60: (None, None, "План за ночь"),
        61: (None, None, "Факт за сутки"),
        62: (None, None, "Факт за день"),
        63: (None, None, "Факт за ночь"),
        64: (None, None, "Факт завоз в накопитель"),
        65: (None, None, "Факт вывоз с накопителя"),
        66: (None, None, "Факт за сутки НАЕМ+АЛМАЗ"),
        67: (None, "Щебень", "План за сутки по ПП"),
        68: (None, None, "Плановая выработка"),
        69: (None, None, "План за день"),
        70: (None, None, "План за ночь"),
        71: (None, None, "Факт за сутки"),
        72: (None, None, "Факт за день"),
        73: (None, None, "Факт за ночь"),
        74: (None, "ЩПС", "Плановая выработка"),
        75: (None, None, "План за день"),
        76: (None, None, "План за ночь"),
        77: (None, None, "Факт за сутки"),
        78: (None, None, "Факт за день"),
        79: (None, None, "Факт за ночь"),
        80: (None, None, "Факт завоз в накопитель"),
        81: (None, "Перевозка", "Плановая выработка"),
        82: (None, None, "План за день"),
        83: (None, None, "План за ночь"),
        84: (None, None, "Факт за сутки"),
        85: (None, None, "Факт за день"),
        86: (None, None, "Факт за ночь"),
        88: ("Экскаватор", "Плановое количество экскаваторов ДЕНЬ", None),
    }
    for row, (b, c, d) in labels.items():
        ws.cell(row, 2, b)
        ws.cell(row, 3, c)
        ws.cell(row, 4, d)


def test_ready_general_dump_truck_sand_plan_uses_transport_linked_counts(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    _seed_ready_dump_truck_block(ws)

    linked_counts = wip_routes._new_equipment_fact_slot()
    linked_counts["day"]["UCH_1"] = 99
    transport_counts = wip_routes._new_equipment_fact_slot()
    transport_counts["day"]["UCH_1"] = 2
    transport_counts["night"]["UCH_1"] = 1
    transport_counts["total"]["UCH_1"] = 2

    monkeypatch.setattr(wip_routes, "_query_general_equipment_counts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_facts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_counts", lambda *args, **kwargs: {"dump_truck": {"SAND": linked_counts}}, raising=False)
    monkeypatch.setattr(wip_routes, "_query_general_dump_truck_sand_productivity_plan_counts", lambda *args, **kwargs: transport_counts, raising=False)
    monkeypatch.setattr(wip_routes, "_fill_general_equipment_plan_rows", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_general_equipment_rule_for_class", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_sand_pit_norms", lambda: wip_routes._constant_norms(10))

    wip_routes._fill_ready_general_equipment(ws, wip_routes.date_cls(2026, 6, 3), None, {}, {})

    assert ws["F58"].value == 20
    assert ws["F59"].value == 20
    assert ws["F60"].value == 10


def test_ready_general_dump_truck_sand_keeps_jds_and_external_fact_rows_separate(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    _seed_ready_dump_truck_block(ws)

    sand_facts = wip_routes._new_equipment_fact_slot()
    sand_facts["non_hired"]["UCH_1"] = 10
    sand_facts["almaz"]["UCH_1"] = 3
    sand_facts["hired"]["UCH_1"] = 4
    sand_facts["total"]["UCH_1"] = 17

    rule = {"filters": {"source_buckets": ["non_hired"]}}
    monkeypatch.setattr(wip_routes, "_query_general_equipment_counts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_facts", lambda *args, **kwargs: {"dump_truck": {"SAND": sand_facts}})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_counts", lambda *args, **kwargs: {}, raising=False)
    monkeypatch.setattr(wip_routes, "_query_general_dump_truck_sand_productivity_plan_counts", lambda *args, **kwargs: wip_routes._new_equipment_fact_slot(), raising=False)
    monkeypatch.setattr(wip_routes, "_fill_general_equipment_plan_rows", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_general_equipment_rule_for_class", lambda *args, **kwargs: rule)

    wip_routes._fill_ready_general_equipment(ws, wip_routes.date_cls(2026, 6, 3), None, {}, {})

    assert ws["F61"].value == 10
    assert ws["F66"].value == 7


def test_ready_general_dump_truck_sand_jds_row_uses_only_pit_delivered_material(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    _seed_ready_dump_truck_block(ws)
    ws.cell(67, 2).value = None
    ws.cell(67, 3).value = None
    ws.cell(67, 4).value = "Факт за сутки ЖДС"
    ws.cell(68, 2).value = None
    ws.cell(68, 3).value = "Щебень"
    ws.cell(68, 4).value = "План за сутки по ПП"

    sand_facts = wip_routes._new_equipment_fact_slot()
    sand_facts["total"]["UCH_1"] = 999
    sand_facts["non_hired"]["UCH_1"] = 888
    sand_facts["almaz"]["UCH_1"] = 111

    delivered = wip_routes._new_work_fact_slot()
    delivered["non_hired"]["UCH_1"] = 25  # ЖДС карьер→конструктив + карьер→накопитель
    delivered["almaz"]["UCH_1"] = 7
    delivered["hired"]["UCH_1"] = 3
    stockpile_out = wip_routes._new_work_fact_slot()
    stockpile_out["non_hired"]["UCH_1"] = 500  # must not leak into the ЖДС source row

    monkeypatch.setattr(wip_routes, "_query_general_equipment_counts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_facts", lambda *args, **kwargs: {"dump_truck": {"SAND": sand_facts}})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_counts", lambda *args, **kwargs: {}, raising=False)
    monkeypatch.setattr(wip_routes, "_query_general_dump_truck_sand_productivity_plan_counts", lambda *args, **kwargs: wip_routes._new_equipment_fact_slot(), raising=False)
    monkeypatch.setattr(wip_routes, "_fill_general_equipment_plan_rows", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_general_equipment_rule_for_class", lambda *args, **kwargs: None)

    wip_routes._fill_ready_general_equipment(
        ws,
        wip_routes.date_cls(2026, 6, 8),
        None,
        {"SAND_SITE_DELIVERED": delivered, "SAND_STOCKPILE_TO_CONSTRUCTIVE": stockpile_out},
        {},
    )

    assert ws["F61"].value == 999
    assert ws["F66"].value == 10
    assert ws["F67"].value == 25


def test_ready_general_dump_truck_shps_plan_night_uses_engaged_trucks(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    _seed_ready_dump_truck_block(ws)
    # Old code mirrored quarry day+night facts into only the day row, leaving
    # SHPS night plan empty. The plan must come from trucks engaged in SHPS.
    ws["F39"] = 0
    ws["F40"] = 0

    shps_counts = wip_routes._new_equipment_fact_slot()
    shps_counts["day"]["UCH_1"] = 1
    shps_counts["night"]["UCH_1"] = 2

    monkeypatch.setattr(wip_routes, "_query_general_equipment_counts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_facts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_counts", lambda *args, **kwargs: {"dump_truck": {"SHPS": shps_counts}}, raising=False)
    monkeypatch.setattr(wip_routes, "_query_general_dump_truck_sand_productivity_plan_counts", lambda *args, **kwargs: wip_routes._new_equipment_fact_slot(), raising=False)
    monkeypatch.setattr(wip_routes, "_fill_general_equipment_plan_rows", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_general_equipment_rule_for_class", lambda *args, **kwargs: None)

    wip_routes._fill_ready_general_equipment(ws, wip_routes.date_cls(2026, 6, 3), None, {}, {})

    assert ws["F74"].value == 498
    assert ws["F75"].value == 166
    assert ws["F76"].value == 332
    assert ws.row_dimensions[74].hidden is not True


def test_ready_general_dump_truck_transport_block_uses_move_category(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    _seed_ready_dump_truck_block(ws)

    move_counts = wip_routes._new_equipment_fact_slot()
    move_counts["day"]["UCH_1"] = 2
    move_counts["night"]["UCH_1"] = 1
    move_facts = wip_routes._new_equipment_fact_slot()
    move_facts["total"]["UCH_1"] = 123
    move_facts["day"]["UCH_1"] = 80
    move_facts["night"]["UCH_1"] = 43

    monkeypatch.setattr(wip_routes, "_query_general_equipment_counts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_facts", lambda *args, **kwargs: {"dump_truck": {"MOVE": move_facts}})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_counts", lambda *args, **kwargs: {"dump_truck": {"MOVE": move_counts}}, raising=False)
    monkeypatch.setattr(wip_routes, "_query_general_dump_truck_sand_productivity_plan_counts", lambda *args, **kwargs: wip_routes._new_equipment_fact_slot(), raising=False)
    monkeypatch.setattr(wip_routes, "_fill_general_equipment_plan_rows", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_general_equipment_rule_for_class", lambda *args, **kwargs: None)

    assert wip_routes._ready_general_equipment_category_from_label("Перевозка") == "MOVE"

    wip_routes._fill_ready_general_equipment(ws, wip_routes.date_cls(2026, 6, 3), None, {}, {})

    assert ws["F81"].value == 498
    assert ws["F82"].value == 332
    assert ws["F83"].value == 166
    assert ws["F84"].value == 123
    assert ws["F85"].value == 80
    assert ws["F86"].value == 43
    assert ws.row_dimensions[81].hidden is not True


def test_ready_general_bulldozer_plan_uses_only_equipment_engaged_in_each_work_tag(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    _seed_ready_bulldozer_block(ws)
    # Template planned equipment counts are not work-tag specific and must not
    # drive per-work planned output.
    ws["F120"] = 9
    ws["F121"] = 9

    sand_counts = wip_routes._new_equipment_fact_slot()
    sand_counts["day"]["UCH_1"] = 2
    shps_counts = wip_routes._new_equipment_fact_slot()
    shps_counts["night"]["UCH_1"] = 1

    monkeypatch.setattr(wip_routes, "_query_general_equipment_counts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_facts", lambda *args, **kwargs: {})
    monkeypatch.setattr(wip_routes, "_query_general_equipment_work_counts", lambda *args, **kwargs: {"bulldozer": {"SAND": sand_counts, "SHPS": shps_counts}}, raising=False)
    monkeypatch.setattr(wip_routes, "_query_general_dump_truck_sand_productivity_plan_counts", lambda *args, **kwargs: wip_routes._new_equipment_fact_slot(), raising=False)
    monkeypatch.setattr(wip_routes, "_fill_general_equipment_plan_rows", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_general_equipment_rule_for_class", lambda *args, **kwargs: None)
    monkeypatch.setattr(wip_routes, "_norm_from_settings", lambda _eq, _codes, default: default)

    wip_routes._fill_ready_general_equipment(ws, wip_routes.date_cls(2026, 6, 3), None, {}, {})

    assert ws["F132"].value == 1400
    assert ws["F133"].value == 0
    assert ws["F134"].value == 1400
    assert ws["F139"].value == 2800
    assert ws["F140"].value == 2800
    assert ws["F141"].value == 0


def test_general_equipment_work_count_slots_deduplicate_units_by_tag_and_shift():
    rows = [
        {"section_code": "UCH_1", "shift": "day", "equipment_type": "Бульдозер", "ownership_type": "own", "contractor_name": "", "unit_key": "гн:A123", "analytics_tag": "ЩПС"},
        {"section_code": "UCH_1", "shift": "day", "equipment_type": "Бульдозер", "ownership_type": "own", "contractor_name": "", "unit_key": "гн:A123", "analytics_tag": "ЩПС"},
        {"section_code": "UCH_1", "shift": "night", "equipment_type": "Бульдозер", "ownership_type": "own", "contractor_name": "", "unit_key": "гн:A123", "analytics_tag": "ЩПС"},
        {"section_code": "UCH_1", "shift": "night", "equipment_type": "Бульдозер", "ownership_type": "own", "contractor_name": "", "unit_key": "бн:17", "analytics_tag": "ЩПС"},
    ]

    counts = wip_routes._general_equipment_work_count_slots_from_rows(rows)
    slot = counts["bulldozer"]["SHPS"]

    assert slot["day"]["UCH_1"] == 1
    assert slot["night"]["UCH_1"] == 2
    # Total is unique by unit across both day and night for the same work tag.
    assert slot["total"]["UCH_1"] == 2



def test_general_work_facts_fall_back_to_builtin_planировка_and_канава_categories(monkeypatch):
    rows = [
        {
            "section_code": "UCH_2",
            "shift": "day",
            "wt_code": "AREA_GRADING",
            "object_type_code": "TEMP_ROAD",
            "object_code": "AD_1",
            "analytics_tag": "Планировка",
            "labor_source_type": "own",
            "contractor_name": "ЖДС",
            "amount": 24000,
        },
        {
            "section_code": "UCH_2",
            "shift": "night",
            "wt_code": "WT_BAE95AF6",
            "object_type_code": "TEMP_ROAD",
            "object_code": "AD_1",
            "analytics_tag": "Канава",
            "labor_source_type": "own",
            "contractor_name": "ЖДС",
            "amount": 180,
        },
    ]

    monkeypatch.setattr(wip_routes, "query", lambda *args, **kwargs: rows)
    monkeypatch.setattr(wip_routes, "_general_report_work_rules", lambda: [])

    facts = wip_routes._query_general_work_facts(
        wip_routes.date_cls(2026, 6, 8),
        wip_routes.date_cls(2026, 6, 8),
        None,
    )

    assert facts["GRADING"]["day"]["UCH_2"] == 24000
    assert facts["DITCH"]["night"]["UCH_2"] == 180



def test_general_equipment_work_count_slots_fall_back_to_builtin_sidework_categories():
    rows = [
        {
            "section_code": "UCH_2",
            "shift": "day",
            "equipment_type": "Экскаватор",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "eq:1",
            "wt_code": "AREA_GRADING",
            "analytics_tag": "Планировка",
            "object_type_code": "TEMP_ROAD",
        },
        {
            "section_code": "UCH_4",
            "shift": "night",
            "equipment_type": "Экскаватор",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "eq:2",
            "wt_code": "WT_BAE95AF6",
            "analytics_tag": "Канава",
            "object_type_code": "MAIN_TRACK",
        },
    ]

    counts = wip_routes._general_equipment_work_count_slots_from_rows(rows, rules=[])

    assert counts["excavator"]["GRADING"]["day"]["UCH_2"] == 1
    assert counts["excavator"]["DITCH"]["night"]["UCH_4"] == 1



def test_general_material_movement_internal_routes_fill_dump_truck_move_not_sand():
    rows = [
        {
            "section_code": "UCH_3",
            "shift": "day",
            "equipment_type": "Самосвал",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "truck:1",
            "material_code": "SAND",
            "material_name": "Песок",
            "movement_type": "stockpile_to_stockpile",
        },
        {
            "section_code": "UCH_1",
            "shift": "night",
            "equipment_type": "Самосвал",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "truck:2",
            "material_code": "SAND",
            "material_name": "Песок",
            "movement_type": "constructive_to_stockpile",
        },
    ]

    counts = wip_routes._general_equipment_work_count_slots_from_rows(rows, rules=[])

    assert counts["dump_truck"]["MOVE"]["day"]["UCH_3"] == 1
    assert counts["dump_truck"]["MOVE"]["night"]["UCH_1"] == 1
    assert "SAND" not in counts["dump_truck"]


def test_general_material_movement_internal_routes_count_soil_and_peat_as_transport():
    rows = [
        {
            "section_code": "UCH_2",
            "shift": "day",
            "equipment_type": "Самосвал",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "truck:soil",
            "material_code": "SOIL",
            "material_name": "Грунт",
            "movement_type": "constructive_to_constructive",
        },
        {
            "section_code": "UCH_4",
            "shift": "night",
            "equipment_type": "Самосвал",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "truck:peat",
            "material_code": "PEAT",
            "material_name": "Торф",
            "movement_type": "constructive_to_stockpile",
        },
        {
            "section_code": "UCH_5",
            "shift": "day",
            "equipment_type": "Самосвал",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "truck:soil-stockpile-constructive",
            "material_code": "SOIL",
            "material_name": "Грунт",
            "movement_type": "stockpile_to_constructive",
        },
        {
            "section_code": "UCH_6",
            "shift": "night",
            "equipment_type": "Самосвал",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "truck:peat-stockpile-constructive",
            "material_code": "PEAT",
            "material_name": "Торф",
            "movement_type": "stockpile_to_constructive",
        },
        {
            "section_code": "UCH_4",
            "shift": "day",
            "equipment_type": "Самосвал",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "truck:ad15-tech-8-1",
            "material_code": "SOIL",
            "material_name": "Грунт",
            "movement_type": "pit_to_constructive",
            "from_object_type_code": "TEMP_ROAD",
            "to_object_type_code": "TECH_ROAD",
        },
    ]

    counts = wip_routes._general_equipment_work_count_slots_from_rows(rows, rules=[])

    assert counts["dump_truck"]["MOVE"]["day"]["UCH_2"] == 1
    assert counts["dump_truck"]["MOVE"]["night"]["UCH_4"] == 1
    assert counts["dump_truck"]["MOVE"]["day"]["UCH_5"] == 1
    assert counts["dump_truck"]["MOVE"]["night"]["UCH_6"] == 1
    assert counts["dump_truck"]["MOVE"]["day"]["UCH_4"] == 1
    assert "SOIL" not in counts["dump_truck"]
    assert "PEAT" not in counts["dump_truck"]
