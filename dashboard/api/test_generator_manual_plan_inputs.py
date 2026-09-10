from __future__ import annotations

from openpyxl import Workbook

import wip_routes


START_COL = wip_routes.GENERAL_ANALYTICS_VALUE_START_COL


def _ready_general_count_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "Аналитика"
    ws["B3"] = "ДЕНЬ"

    ws.cell(row=10, column=2, value="Самосвал")
    ws.cell(row=11, column=3, value="Плановое количество самосвалов ДЕНЬ")
    ws.cell(row=12, column=3, value="Плановое количество самосвалов НОЧЬ")

    ws.cell(row=20, column=2, value="Экскаватор")
    ws.cell(row=21, column=3, value="Плановое количество экскаваторов ДЕНЬ")
    ws.cell(row=22, column=3, value="Плановое количество экскаваторов НОЧЬ")

    ws.cell(row=30, column=2, value="Бульдозер")
    ws.cell(row=31, column=3, value="Плановое количество бульдозеров ДЕНЬ")
    ws.cell(row=32, column=3, value="Плановое количество бульдозеров НОЧЬ")
    return ws


def test_dim_pipeline_report_supports_cumulative_as_of_end_checkbox():
    assert wip_routes.GENERATOR_PIPELINE_REPORTS["dim"]["supports_cumulative_as_of_end_date"] is True
    assert wip_routes.GENERATOR_PIPELINE_REPORTS["dim"]["supports_daily_pk_details"] is True


def test_write_ready_general_manual_plan_counts_uses_section_inputs():
    ws = _ready_general_count_sheet()

    wip_routes._write_ready_general_manual_equipment_count_rows(
        ws,
        {
            "dump_truck": {
                "day": {"UCH_1": 2, "UCH_3": 4},
                "night": {"UCH_2": 1},
            },
            "excavator": {
                "day": {"UCH_4": 5},
            },
            "bulldozer": {
                "night": {"UCH_1": 1, "UCH_8": 2},
            },
        },
    )

    assert ws.cell(row=11, column=START_COL).value == 6
    assert ws.cell(row=11, column=START_COL + 1).value == 2
    assert ws.cell(row=11, column=START_COL + 2).value is None
    assert ws.cell(row=11, column=START_COL + 3).value == 4

    assert ws.cell(row=12, column=START_COL).value == 1
    assert ws.cell(row=12, column=START_COL + 1).value is None
    assert ws.cell(row=12, column=START_COL + 2).value == 1

    assert ws.cell(row=21, column=START_COL).value == 5
    assert ws.cell(row=21, column=START_COL + 4).value == 5
    assert ws.cell(row=22, column=START_COL).value is None

    assert ws.cell(row=32, column=START_COL).value == 3
    assert ws.cell(row=32, column=START_COL + 1).value == 1
    assert ws.cell(row=32, column=START_COL + 8).value == 2
