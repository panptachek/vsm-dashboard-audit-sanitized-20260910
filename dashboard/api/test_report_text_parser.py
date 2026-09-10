#!/usr/bin/env python3
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from report_text_parser import infer_material_code, parse_equipment_line, parse_pk_range, parse_pk_value, parse_report_text, split_report_texts
from xlsx_report_adapter import xlsx_bytes_to_report_text


SAMPLE = """
%%Комментарий к заполнению
Дата - 04.05.2026
Смена - день
Участок - участок №1
Направление - ЗП
Конструктивы - АД9, ОХ, стартовая площадка участка усиления 5.2
|
Информация по завозу: Карьер Васильки - песок - 666 м3; Карьер Миголощи - ЩПГС - 999 м3
|
===Перевозка===
|
ФИО: Степанов А. А.
Самосвал FAW (б/н 42; г/н 913); ЖДС
/песок/
карьер Васильки → АД5 ПК51+20-ПК53+20
111,1 м3 / 6 рейс
|
===Основные работы===
-ОХ-
/Разработка выемки/
ПК ВСЖМ: ПК2654+50-ПК2655+20
ПК АД:
ФИО: Грошин А. В.
Экскаватор CAT 330 (б/н ; г/н 7439); ЖДС
V = 620 м3
|
-АД5-
/Устройство насыпи из песка/
ПК ВСЖМ: ПК2694+00
ПК АД: ПК51+20-ПК53+20
ФИО: Анипченко В. Е.
Бульдозер SEM (б/н ; г/н 6947); ЖДС
V = 1804,7 м3
===Сопутствующие работы===
|
-Работа на ИССО-
|
/Работа на ИССО/
ПК ВСЖМ: ПК2658
ПК АД: ПК12+20
ФИО: Пулатов Ю. З.
Экскаватор HITACHI330 (б/н ; г/н 0337); ЖДС
===Парк техники===
|
Самосвал FAW (б/н ; г/н 256); ООО "Алмаз"
Самосвал FAW (б/н ; г/н 256); ООО "Алмаз"
Бульдозер SEM (б/н ; г/н 6947); Анипченко В. Е.; в работе;
|
===Персонал===
ДР: 3
ИТР: 4
|
===Накопители===
Накопитель песка АД13 - ПК2715+00 -1000 м3
Накопитель ЩПГС - ПК2687+00 - 2000 м3
|
===Забивка свай===
5_1 Н - ПК2655+73.50 - основные - 4 шт - составная готова
"""


def assert_equipment_number_variants() -> None:
    samples = [
        "Самосвал FAW (бортовой 42; госномер А123ВС178); ЖДС",
        "Самосвал FAW б/н 42 г/н А123ВС178; ЖДС",
        "Самосвал FAW бортовой 42 госномер А123ВС178; ЖДС",
        "Самосвал FAW; бортовой 42; госномер А123ВС178; ЖДС",
    ]
    for line in samples:
        parsed = parse_equipment_line(line)
        assert parsed is not None, line
        assert parsed["unit_number"] == "42", line
        assert parsed["plate_number"] == "А123ВС178", line
        assert parsed["owner"] == "ЖДС", line
    unlabeled = parse_equipment_line("Самосвал FAW (913); ЖДС")
    assert unlabeled is not None
    assert unlabeled["unit_number"] == "913"
    assert unlabeled["plate_number"] is None
    plate_only = parse_equipment_line("Самосвал FAW (С870 ХВ 27); ЖДС")
    assert plate_only is not None
    assert plate_only["unit_number"] is None
    assert plate_only["plate_number"] == "С870ХВ27"
    no_semicolon = parse_equipment_line("Коток грунтовый SANY (3581) ЖДС")
    assert no_semicolon is not None
    assert no_semicolon["equipment_type"] == "каток"
    assert no_semicolon["unit_number"] == "3581"
    assert no_semicolon["owner"] == "ЖДС"


def assert_pk_variants() -> None:
    cases = [
        ("ПК2657+20-ПК2658+00", 265720.0, 265800.0),
        ("ПК 2657+20 - 2658 +00", 265720.0, 265800.0),
        ("ПК2657+20-+80", 265720.0, 265780.0),
        ("ПК2657+20-80", 265720.0, 265780.0),
        ("3230-3231", 323000.0, 323100.0),
        ("ПК2714", 271400.0, 271400.0),
        ("ПК51+20-ПК53+20", 5120.0, 5320.0),
        ("ПК2657+59,63", 265759.63, 265759.63),
        ("П.К. ВСЖМ: 3230-3231", 323000.0, 323100.0),
        # DB stores route picketage as numeric metres from the route origin:
        # ПК2702 => 270200.00, ПК2702+50 => 270250.00.
        ("ПК2702", 270200.0, 270200.0),
        ("ПК2702+50", 270250.0, 270250.0),
        ("ПК2702-ПК2703", 270200.0, 270300.0),
        ("ПК ВСЖМ: 2702-2703", 270200.0, 270300.0),
        ("264022", 264022.0, 264022.0),
        ("264022-264050", 264022.0, 264050.0),
        ("ПК264022", 264022.0, 264022.0),
        ("ПК 2640 + 22", 264022.0, 264022.0),
        ("ПК2640.22", 264022.0, 264022.0),
        ("2640,22-2640,50", 264022.0, 264050.0),
        ("ПК2703+00-ПК2702+50", 270250.0, 270300.0),
        ("ПК 2640,22 - ПК 2639.90", 263990.0, 264022.0),
        ("ПК ВСЖМ: ПК 2799+80 - 2800+40 (49+80-50+50)", 279980.0, 280040.0),
        ("ПК ВСЖМ: ПК 2799+80 - 2800+40 (ПК АД: ПК49+80-ПК50+50)", 279980.0, 280040.0),
    ]
    for raw, start, end in cases:
        parsed = parse_pk_range(raw)
        assert parsed["start"] == start, raw
        assert parsed["end"] == end, raw
    assert parse_pk_value("2655+73.50") == 265573.5
    assert parse_pk_value("2655.73") == 265573.0
    assert parse_pk_value("2655,73") == 265573.0
    assert parse_pk_value("264022") == 264022.0
    assert parse_pk_value("АД13") is None


def assert_unlabeled_parenthetical_pk_range_does_not_override_rail_range() -> None:
    text = """
Дата - 27.05.2026
Смена - день
Участок - участок №5
===Основные работы===
-АД14-
/Устройство насыпи из песка/
ПК ВСЖМ: ПК 2799+80 – 2800+40 (49+80-50+50)
Номер б/н 230 / 718 м3 / бульдозер / ЖДС
"""
    parsed = parse_report_text("pk-parenthetical.txt", text)
    assert len(parsed["main_works"]) == 1
    work = parsed["main_works"][0]
    assert work["pk_rail_start"] == 279980.0
    assert work["pk_rail_end"] == 280040.0
    assert work["pk_start"] == 279980.0
    assert work["pk_end"] == 280040.0
    assert work["volume"] == 718.0
    assert work["unit_number"] == "230"


def assert_work_split_by_performer_volumes() -> None:
    text = """
Дата - 07.05.2026
Смена - день
Участок - участок №1
===Основные работы===
-ОХ-
/Устройство насыпи из песка/
ПК ВСЖМ: ПК2654+50-ПК2655+20
ФИО: Иванов И. И.
V = 40 м3
ФИО: Петров П. П.
V = 60 м3
"""
    parsed = parse_report_text("performers.txt", text)
    rows = parsed["main_works"]
    assert len(rows) == 2
    assert all("operator_name" not in row for row in rows)
    assert [row["volume"] for row in rows] == [40.0, 60.0]
    assert "ФИО" not in parsed["raw_text"]
    assert not parsed["warnings"]


def assert_work_split_by_equipment_volume_variants() -> None:
    text = """
Дата - 07.05.2026
Смена - день
Участок - участок №1
===Основные работы===
-ОХ-
/Устройство насыпи из песка/
ПК ВСЖМ: ПК2654+50-ПК2655+20
ФИО: Иванов И. И.
Бульдозер CAT бортовой D12 госномер А123ВС178; ЖДС
V = 40 м3
ФИО: Петров П. П.
Бульдозер SEM (бортовой D13; госномер В456ОР178); Алмаз
V = 60 м3
"""
    parsed = parse_report_text("equipment-performers.txt", text)
    rows = parsed["main_works"]
    assert len(rows) == 2
    assert all("operator_name" not in row for row in rows)
    assert [row["unit_number"] for row in rows] == ["D12", "D13"]
    assert [row["plate_number"] for row in rows] == ["А123ВС178", "В456ОР178"]
    assert [row["owner"] for row in rows] == ["ЖДС", "Алмаз"]
    assert [row["volume"] for row in rows] == [40.0, 60.0]


def assert_geotextile_without_equipment_variants() -> None:
    text = """
Дата - 10.07.2026
Смена - день
Участок - участок №1
===Работы===
-АД5-
Устройство геотекстиля 300 кН/м 1250 м2
ПК ВСЖМ: ПК2694+00-ПК2695+00
Устройство слоя из геотекстиля 5 кН/м 500 м2
ПК ВСЖМ: ПК2695+00-ПК2695+50
"""
    parsed = parse_report_text("geotextile-no-equipment.txt", text)
    rows = parsed["main_works"]
    assert len(rows) == 2
    assert [row["work_type_code"] for row in rows] == ["GEOTEXTILE_LAYER", "GEOTEXTILE_LAYER_DO"]
    assert [row["volume"] for row in rows] == [1250.0, 500.0]
    assert [row["unit"] for row in rows] == ["м2", "м2"]
    assert all(row["equipment"] == [] for row in rows)
    assert all(row["productivity_enabled"] is False for row in rows)


def assert_one_sheet_xlsx_geotextile_without_equipment() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчет"
    ws["A1"] = "VSM_DAILY_REPORT_ONE_SHEET_V2"
    ws["B4"] = "10.07.2026"
    ws["B5"] = "день"
    ws["B6"] = "участок №1"
    ws["A10"] = "## РАБОТЫ"
    headers = ["Объект/конструктив", "Работа", "ПК ВСЖМ", "ПК АД", "Номер техники", "Объем", "Ед.", "Тип техники", "Организация", "Производительность", "Комментарий"]
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(11, col_idx).value = header
    values = ["АД5", "Укладка геотекстиля", "ПК2694+00-ПК2695+00", "", "", 640, "м2", "", "", "да", ""]
    for col_idx, value in enumerate(values, start=1):
        ws.cell(12, col_idx).value = value
    bio = BytesIO()
    wb.save(bio)

    report_text = xlsx_bytes_to_report_text(bio.getvalue())
    assert "/Укладка геотекстиля/" in report_text
    assert "V = 640 м2" in report_text
    assert "Номер Н/Д" not in report_text

    parsed = parse_report_text("geotextile-no-equipment.xlsx", report_text)
    assert len(parsed["main_works"]) == 1
    work = parsed["main_works"][0]
    assert work["work_type_code"] == "GEOTEXTILE_LAYER"
    assert work["volume"] == 640.0
    assert work["unit"] == "м2"
    assert work["equipment"] == []
    assert work["productivity_enabled"] is False


def assert_real_report_volume_and_operator_variants() -> None:
    text = """
Дата - 06.05.2026
Смена - день
Участок - участок №1
===Основные работы===
-ОХ-
/Разработка выемки грунта с 2-й перекидкой/
ПК ВСЖМ: ПК2654+55-ПК2655+20
V = 1150 м3
ФИО: Вернигоров Р. В.
Бульдозер SHANTUY (7032); ЖДС
V = 720 м3
ФИО: Вишерский О. В.
Бульдозер SEM (3944); ЖДС
ФИО: Матвеев Д. Е.
Экскаватор CAT 330 (3588); ЖДС
|
/Уплотнение насыпи из песка/
S = 1000 м2
Ф.И.О Бадаев Р. А.
Коток грунтовый SANY (3581) ЖДС
"""
    parsed = parse_report_text("real-variants.txt", text)
    rows = parsed["main_works"]
    assert len(rows) == 4
    assert [rows[0]["volume"], rows[1]["volume"]] == [1150.0, 720.0]
    assert all("operator_name" not in row for row in rows)
    assert rows[2]["volume"] is None
    assert rows[3]["vehicle"] == "каток грунтовый SANY"
    assert any("Разработка выемки грунта" in warning and "без объема" in warning for warning in parsed["warnings"])


def assert_section_7_8_report_variants() -> None:
    text = """
===Шапка===
Дата - 06.05.2026
Смена - ночь
Участок - участок №8
Направление - ЗП
Конструктивы - АД 4.8, АД 12
===Перевозка===
-Исаков А.А.-
Самосвал FAW 806 (с527хв27); ЖДС
/песок/
Боковой резерв → АД 4.8
560 м3 / 35 рейсов
===Основные работы===
-АД 4.9-
/Устройство основания земляного полотна/
ПК ВСЖМ: ПК3286+70 – ПК3287+00
Фрелих М.
Экскаватор LonKing 336 845 (27ха6735); ЖДС
1200 м3 (непригодный грунт)
"""
    parsed = parse_report_text("section-8-night.txt", text)
    assert parsed["header"]["report_date"] == "2026-05-06"
    assert parsed["header"]["shift"] == "night"
    assert parsed["header"]["section_code"] == "UCH_8"
    assert len(parsed["transport"]) == 1
    assert "driver" not in parsed["transport"][0]
    assert parsed["transport"][0]["unit_number"] == "806"
    assert parsed["transport"][0]["plate_number"] == "с527хв27"
    assert parsed["transport"][0]["trips"][0]["trips"] == 35
    work = parsed["main_works"][0]
    assert "operator_name" not in work
    assert work["volume"] == 1200.0
    assert work["unit_number"] == "845"
    assert work["plate_number"] == "27ха6735"

    text_7 = """
Дата- 06.05.2026 г
Смена - ночь
Участок-участок №7
Направление – ЗП
Конструктивы – АД4.8
Информация по завозу: Карьер Великий – песок – 1199,9 м3
===Перевозка===
ФИО: Голиков
Самосвал FAW (б/н 709); ЖДС
/песок/
Карьер-> свайная площадка ПК ВСЖМ: 3230-3231
51,2 м3 / 3 рейс
===Основные работы===
-АД4.8-
/Устройство ДСО/
ФИО: Чирков Е.
БульдозерSEM 826 (г/н 3448ха27); ЖДС
V = 837,2 м3
"""
    parsed_7 = parse_report_text("section-7-night.txt", text_7)
    assert parsed_7["header"]["direction"] == "ЗП"
    assert parsed_7["human_summary"]["delivery_rows"][0]["volume"] == 1199.9
    assert parsed_7["transport"][0]["trips"][0]["trips"] == 3
    assert parsed_7["main_works"][0]["vehicle"] == "бульдозер SEM 826"
    assert parsed_7["main_works"][0]["plate_number"] == "3448ха27"


def assert_compact_report_format() -> None:
    text = """
Дата - 04.05.2026
Смена - день
Участок - участок №1
Направление - ЗП
Конструктивы - АД9, АД5, стартовая площадка участка усиления №5.2
|
Информация по завозу: Карьер Васильки - песок - 666666 м3; Карьер Миголощи - ЩПГС - 9999999 м3
|
===Перевозка===
%%%Направления перевозки: карьер-конструктив
=карьер Васильки → АД5 ПК51+20-ПК53+20=
/песок/
Номер 913 / 111,1 м3 / 6 рейс / ЖДС
Номер 256 / 173,8 м3 / 6 рейс / АЛМАЗ
=АД5 ПК2698+00 → тех. проезд к АД13=
/грунт/
Номер 932 / 45 м3 / 5 рейс / ЖДС
===Работы===
-Основной ход-
/Разработка выемки/
ПК ВСЖМ: ПК2654+50-ПК2655+20
Номер 7439 / 620 м3 / экскаватор / ЖДС
Номер 7032 / 1100 м3 / бульдозер / ЖДС
|
-АД5-
/Устройство насыпи/
ПК ВСЖМ: ПК2694+00 (ПК АД: ПК51+20-ПК53+20)
Номер 6947 / 940 м3 / бульдозер / ЖДС
/Приемка песка в накопитель/
ПК ВСЖМ: ПК2715+00
Номер 6041 / 228,8 / экскаватор / ЖДС
/Устройство насыпи с ошибкой объема/
ПК ВСЖМ: ПК 2657+20-ПК2658 +00
Номер 3944 / 124,№ м3 / бульдозер / ЖДС|
===Парк техники===
Самосвал FAW (б/н ; г/н 913); Степанов А. А.; в работе;
===Персонал===
ДР: 3
ИТР: 4
Водители: 18
Механизаторы: 13
|
===Накопители===
Накопитель песка АД13 - ПК2715+00 -1000 м3
"""
    parsed = parse_report_text("compact.txt", text)
    assert parsed["header"]["report_date"] == "2026-05-04"
    assert parsed["header"]["section_code"] == "UCH_1"
    assert parsed["human_summary"]["delivery_rows"][0]["volume"] == 666666.0
    assert len(parsed["transport"]) == 3
    assert parsed["transport"][0]["unit_number"] == "913"
    assert parsed["transport"][0]["plate_number"] is None
    assert parsed["transport"][0]["trips"][0]["from"] == "карьер Васильки"
    assert parsed["transport"][0]["trips"][0]["to"] == "АД5 ПК51+20-ПК53+20"
    assert parsed["transport"][1]["owner"] == "АЛМАЗ"
    assert parsed["transport"][2]["trips"][0]["material_code"] == "SOIL"
    assert len(parsed["main_works"]) == 5
    assert [row["unit_number"] for row in parsed["main_works"][:4]] == ["7439", "7032", "6947", "6041"]
    assert parsed["main_works"][2]["pk_ad_raw"] == "ПК51+20-ПК53+20"
    assert parsed["main_works"][3]["volume"] == 228.8
    assert parsed["main_works"][4]["volume"] is None
    assert any("Устройство насыпи с ошибкой объема" in warning for warning in parsed["warnings"])
    assert parsed["staff_counts"] == []
    assert parsed["park"] == []
    assert "Степанов" not in parsed["raw_text"]
    assert not any("Степанов" in str(row) for row in parsed["park"])
    assert parsed["stockpiles"][0]["rounded_pk"] == 2715


def assert_fill_status_block_and_number_markers() -> None:
    text = """
Дата - 08.05.2026
Смена - день
Участок - участок №1
===Перевозка===
-карьер Васильки -> АД5 ПК51+20-ПК53+20-
/песок/
Номер б/н 913 / 111,1 м3 / 6 рейс / ЖДС
Номер У824АС27 / 173,8 м3 / 6 рейс / АЛМАЗ
===Работы===
-АД5-
/Устройство насыпи/
ПК ВСЖМ: ПК2694+00 (ПК АД: ПК51+20-ПК53+20)
Номер б/н 6947 / 940 м3 / бульдозер / ЖДС
===Данные по отсыпке===
АД5; ПК АД 0+00-2+15; не в отметку
АД5; ПК АД 2+15-3+00; ДСО
АД5; ПК АД 3+00-4+00; под аб
"""
    parsed = parse_report_text("fill-statuses.txt", text)
    assert len(parsed["transport"]) == 2
    assert parsed["transport"][0]["unit_number"] == "913"
    assert parsed["transport"][0]["plate_number"] is None
    assert parsed["transport"][1]["unit_number"] is None
    assert parsed["transport"][1]["plate_number"] == "У824АС27"
    assert parsed["transport"][0]["trips"][0]["from"] == "карьер Васильки"
    assert parsed["transport"][0]["trips"][0]["to"] == "АД5 ПК51+20-ПК53+20"
    assert parsed["main_works"][0]["unit_number"] == "6947"
    assert parsed["main_works"][0]["plate_number"] is None

    fill = parsed["fill_statuses"]
    assert len(fill) == 3
    assert fill[0]["road_code"] == "АД5"
    assert fill[0]["subgrade_not_to_grade"] == "ПК АД 0+00-2+15"
    assert fill[1]["dso"] == "ПК АД 2+15-3+00"
    assert fill[2]["shpgs_done"] == "ПК АД 3+00-4+00"
    assert not parsed["warnings"]


def assert_section_6_transport_variants() -> None:
    text = """
Дата - 18.05.2026
Смена - день
Участок - участок №6
Направление - ЗП
Конструктивы - АД12, АД2.6, ПК3063 технологические площадки
===Перевозка===
-Накопитель ПК3089 - АД12 ПК3111+50-ПК3112+60-
/песок/
Номер С887ХВ27/ 280 м3 / 28 рейс / ЖДС
Номер К512ВХ27/ 180 м3 / 18 рейс / ЖДС

-Накопитель ПК3089 -> ПК3063 технологические площадки-
/песок/
Номер С141СН27/ 80 м3 / 8 рейс / ЖДС
Номер С870ХВ27/ 100 м3 / 10 рейс / ЖДС
Номер С179ТВ27/ 80 м3 / 8 рейс / ЖДС
Номер С302ТВ27/ 80 м3 / 8 рейс / ЖДС
Номер С334ХЕ27/ 80 м3 / 8 рейс / ЖДС
Номер С170ХМ27/ 70 м3 / 7 рейс / ЖДС

-Накопитель ПК3089 -> АД12 ПК3118+70-ПК3117+60-
/песок/
Номер С412ХС27/ 204 м3 / 17 рейс / ЖДС
Номер У615ОХ27/ 180 м3 / 15 рейс / ЖДС
Номер С861ХА27/ 216 м3 / 18 рейс / ЖДС
Номер Т927ЕА27/ 192 м3 / 16 рейс / ЖДС

-карьер Великий -> накопитель АД2.6 ПК3153+00-
/песок/
6 а/с / 262,5 м3 / 15 рейс / Алмаз

-карьер Великий -> АД12 ПК3118+70-ПК3117+80-
/песок/
8 а/с / 436,9 м3 / 24 рейс / Алмаз
"""
    parsed = parse_report_text("section-6-transport.txt", text)
    assert len(parsed["transport"]) == 14
    assert parsed["transport"][0]["plate_number"] == "С887ХВ27"
    assert parsed["transport"][0]["trips"][0]["from"] == "Накопитель ПК3089"
    assert parsed["transport"][0]["trips"][0]["to"] == "АД12 ПК3111+50-ПК3112+60"
    assert parsed["transport"][0]["trips"][0]["trips"] == 28
    assert parsed["transport"][12]["equipment_count"] == 6
    assert parsed["transport"][12]["unit_number"] is None
    assert parsed["transport"][12]["plate_number"] is None
    assert parsed["transport"][12]["trips"][0]["volume"] == 262.5
    assert parsed["transport"][13]["equipment_count"] == 8
    assert parsed["transport"][13]["trips"][0]["to"] == "АД12 ПК3118+70-ПК3117+80"


def assert_tolerant_max_headings() -> None:
    text = """
Дата - 19.05.2026
Смена - день
Участок - участок №6
===Возка песка===
Маршрут: карьер Великий -> накопитель ПК3089
Материал: песок
Номер С887ХВ27 / 280 м3 / 28 рейс / ЖДС

### Выполненные работы
Основной ход:
Устройство насыпи из песка
ПК ВСЖМ: ПК2702-ПК2703
Номер 6947 / 100 м3 / бульдозер / ЖДС
"""
    parsed = parse_report_text("tolerant-max-headings.txt", text)
    assert len(parsed["transport"]) == 1
    assert parsed["transport"][0]["plate_number"] == "С887ХВ27"
    assert parsed["transport"][0]["trips"][0]["material_code"] == "SAND"
    assert parsed["transport"][0]["trips"][0]["from"] == "карьер Великий"
    assert parsed["transport"][0]["trips"][0]["to"] == "накопитель ПК3089"
    assert len(parsed["main_works"]) == 1
    assert parsed["main_works"][0]["constructive"] == "Основной ход"
    assert parsed["main_works"][0]["volume"] == 100.0


def assert_core_blocks_inferred_without_explicit_headings() -> None:
    text = """
Дата - 19.05.2026
Смена - ночь
Участок - участок №6
карьер Великий -> накопитель ПК3089 = песок
Номер С887ХВ27 / 280 м3 / 28 рейс / ЖДС

Основной ход:
Устройство насыпи из песка
ПК ВСЖМ: ПК2702-ПК2703
Номер 6947 / 100 м3 / бульдозер / ЖДС
"""
    parsed = parse_report_text("core-cues-without-headings.txt", text)
    assert len(parsed["transport"]) == 1
    assert parsed["transport"][0]["plate_number"] == "С887ХВ27"
    assert parsed["transport"][0]["trips"][0]["from"] == "карьер Великий"
    assert parsed["transport"][0]["trips"][0]["to"] == "накопитель ПК3089"
    assert len(parsed["main_works"]) == 1
    assert parsed["main_works"][0]["constructive"] == "Основной ход"
    assert parsed["main_works"][0]["volume"] == 100.0
    assert any("Перевозка распознана без явного заголовка" in warning for warning in parsed["warnings"])
    assert any("Работы распознаны без явного заголовка" in warning for warning in parsed["warnings"])


def assert_park_exception_statuses() -> None:
    text = """
Дата - 07.05.2026
Смена - день
Участок - участок №1
===Парк техники===
Самосвал FAW (913); ЖДС; ремонт Боровичи;
Самосвал FAW (838); простой по причине ремонта топливозаправщика;
Автогрейдер SEM 922 (б/н 6945); ЖДС; в ремонте - сцепление
Самосвал FAW (256); ЖДС; в работе;
Самосвал FAW (000); что угодно без статуса;
"""
    parsed = parse_report_text("park-statuses.txt", text)
    assert parsed["park"] == []
    assert parsed["staff_counts"] == []
    assert parsed["warnings"] == []


def assert_multi_report_warning() -> None:
    text = """
Дата - 07.05.2026
Смена - день
Участок - участок №1
===Перевозка===
Дата - 07.05.2026
Смена - ночь
Участок - участок №1
===Перевозка===
"""
    parsed = parse_report_text("two-reports.txt", text)
    assert any("несколько отчетов" in warning for warning in parsed["warnings"])
    reports = split_report_texts(text)
    assert len(reports) == 2
    assert parse_report_text("two-reports.txt #1", reports[0])["header"]["shift"] == "day"
    assert parse_report_text("two-reports.txt #2", reports[1])["header"]["shift"] == "night"


def assert_compact_xlsx_template() -> None:
    template = Path(__file__).resolve().parents[1] / "templates" / "daily_report_compact_template.xlsx"
    text = xlsx_bytes_to_report_text(template.read_bytes())
    parsed = parse_report_text("daily_report_compact_template.xlsx", text)
    assert parsed["header"]["report_date"] == "2026-05-04"
    assert len(parsed["transport"]) == 2
    assert parsed["transport"][0]["unit_number"] == "913"
    assert len(parsed["main_works"]) == 2
    assert parsed["main_works"][1]["pk_ad_raw"] == "ПК51+20-ПК53+20"


def assert_transport_marker_not_used_as_endpoint() -> None:
    text = """
Дата - 20.05.2026
Смена - день
Участок - участок №2
===Перевозка===
%% комментарий между заголовком и маршрутом
-карьер Васильки -> АД 14, ВСЖМ ПК 2768+70-2770+40
/песок/
Номер б/н 209 / 18.4 м3 / 1 рейс / ЖДС

- Накопитель 2776 -> АД 14 ПК 2796+75-2797+40
/песок/
Номер б/н 201 / 180 м3 / 18 рейс / ЖДС
"""
    parsed = parse_report_text("transport-marker-endpoint.txt", text)
    trips = [trip for unit in parsed["transport"] for trip in unit["trips"]]
    assert len(trips) == 2
    assert trips[0]["from"] == "карьер Васильки"
    assert "Перевоз" not in trips[0]["from"]
    assert trips[0]["to"].startswith("АД 14")
    assert trips[1]["from"] == "Накопитель 2776"


def assert_loose_arrow_route_without_opening_dash() -> None:
    text = """
Дата - 20.05.2026
Смена - день
Участок - участок №4
===Перевозка===
Карьер Васильки -> накопитель ПК 2890
/песок/
Номер В499НК977 / 118.9 м3 / 4 рейс / Автомолит
"""
    parsed = parse_report_text("transport-no-opening-dash.txt", text)
    trip = parsed["transport"][0]["trips"][0]
    assert trip["from"] == "Карьер Васильки"
    assert trip["to"] == "накопитель ПК 2890"


def assert_opening_sign_route_without_closing_marker() -> None:
    text = """
Дата - 20.05.2026
Смена - день
Участок - участок №2
===Перевозка===
-карьер Васильки - АД 14 ПК 2768+70-2770+40
/песок/
Номер б/н 209 / 18.4 м3 / 1 рейс / ЖДС
-накопитель песка-АД4.8 ПК3280+50-ПК3280+30
/песок/
Номер 802 / 560 м3 / 35 рейс / ЖДС
===Работы===
-АД14-
/Устройство насыпи из песка/
ПК ВСЖМ: ПК2768+70-ПК2770+40
Номер 231 / 303,4 м3 / бульдозер / ЖДС
===Проблемные вопросы===
   -
   проблем нет
"""
    parsed = parse_report_text("opening-sign-route.txt", text)
    trips = [trip for unit in parsed["transport"] for trip in unit["trips"]]
    assert len(trips) == 2
    assert trips[0]["from"] == "карьер Васильки"
    assert trips[0]["to"] == "АД 14 ПК 2768+70-2770+40"
    assert trips[1]["from"] == "накопитель песка"
    assert trips[1]["to"] == "АД4.8 ПК3280+50-ПК3280+30"
    assert parsed["main_works"][0]["constructive"] == "АД14"
    assert parsed["problems"] == ""


def assert_units_stockpiles_and_own_forces_variants() -> None:
    text = """
Дата - 20.05.2026
Смена - ночь
Участок - участок №4
===Перевозка===
Карьер Васильки -> накопитель ПК 2890
/песок/
Номер В499НК977 / 2 ед. / 1 рейс / Собственные силы
===Работы===
-ОХ-
/Планировка площадки/
ПК ВСЖМ: ПК3039+00-ПК3040+00
Номер 231 / 3 м³ / бульдозер / собственными силами
===Накопители===
Накопитель песка ПК2890 - 1000 м³
Накопитель ЩПГС - ПК2891+50 - 12 ед.
"""
    parsed = parse_report_text("units-stockpiles-own.txt", text)
    vehicle = parsed["transport"][0]
    assert vehicle["owner"] == "Собственные силы"
    trip = vehicle["trips"][0]
    assert trip["unit"] == "шт"
    work = parsed["main_works"][0]
    assert work["unit"] == "м3"
    assert work["equipment"][0]["owner"] == "собственными силами"
    assert len(parsed["stockpiles"]) == 2
    assert parsed["stockpiles"][0]["material_code"] == "SAND"
    assert parsed["stockpiles"][0]["rounded_pk"] == 2890
    assert parsed["stockpiles"][0]["unit"] == "м3"
    assert parsed["stockpiles"][1]["material_code"] == "SHPGS"
    assert parsed["stockpiles"][1]["rounded_pk"] == 2892
    assert parsed["stockpiles"][1]["unit"] == "шт"


def assert_transport_indirect_equipment_cues() -> None:
    text = """
Дата - 20.05.2026
Смена - ночь
Участок - участок №8
===Перевозка===
-карьер Великий -> Мост ПК3318
/песок/
FAW б/н 718 / 36,2 м3 / 2 рейса / ЖДС
Номер 807(с527хв27)/ 71,2 м3 / 4 рейс / ООО «ЖДС»
Номер 800 (с845хк27) / 105 т / 6 рейс / ООО "ЖДС"
Номер 801 / 10 м3 / 1 рейс / «ООО ЖДС»
===Работы===
-АД12-
/Устройство насыпи из песка/
ПК ВСЖМ: ПК3117+50-ПК3117+20
Номер б/н 6041 / 228,8 м3 / экскаватор / ЖДС
"""
    parsed = parse_report_text("transport-indirect-cues.txt", text)
    assert len(parsed["transport"]) == 4
    first, second, third, fourth = parsed["transport"]
    assert first["equipment_type"] == "самосвал"
    assert first["brand_model"] == "FAW"
    assert first["unit_number"] == "718"
    assert first["trips"][0]["trips"] == 2
    assert second["unit_number"] == "807"
    assert second["plate_number"] == "с527хв27"
    assert second["owner"] == "ЖДС"
    assert third["unit_number"] == "800"
    assert third["owner"] == "ЖДС"
    assert third["plate_number"] == "с845хк27"
    assert third["trips"][0]["unit"] == "т"
    assert fourth["unit_number"] == "801"
    assert fourth["owner"] == "ЖДС"


def assert_unknown_transport_material_keeps_compact_trip() -> None:
    text = """
Дата - 23.08.2026
Смена - день
Участок - участок №3
===Перевозка===
=Карьер Васильки -> АД13 ПК2840-ПК2841=
ПРС
Номер 306 / 120 м3 / 12 рейс / ЖДС
"""
    parsed = parse_report_text("unknown-transport-material.txt", text)
    trips = [trip for unit in parsed["transport"] for trip in unit["trips"]]
    assert len(trips) == 1
    assert trips[0]["material"] == "ПРС"
    assert trips[0]["material_code"] is None
    assert trips[0]["volume"] == 120.0
    assert trips[0]["trips"] == 12


def assert_mainline_fill_status_block() -> None:
    text = """
Дата - 08.05.2026
Смена - день
Участок - участок №1
===Данные по отсыпке ОХ===
участок №1; ПК2654+50-ПК2655+20; ПРС
участок №1; ПК2655+20-ПК2656+00; Основные работы; комментарий
участок №1; ПК2656+00-ПК2657+00; ЗС №2
"""
    parsed = parse_report_text("mainline-fill.txt", text)
    assert parsed["fill_statuses"] == []
    assert parsed["mainline_fill_statuses"] == [
        {
            "section_code": "UCH_1",
            "status_type": "prep_works",
            "pk_ranges": "ПК2654+50-ПК2655+20",
            "pk_validation_override": False,
            "comment": "",
        },
        {
            "section_code": "UCH_1",
            "status_type": "main_works",
            "pk_ranges": "ПК2655+20-ПК2656+00",
            "pk_validation_override": False,
            "comment": "комментарий",
        },
        {
            "section_code": "UCH_1",
            "status_type": "protective_layer_2",
            "pk_ranges": "ПК2656+00-ПК2657+00",
            "pk_validation_override": False,
            "comment": "",
        },
    ]
    assert not parsed["warnings"]


def assert_pile_headcap_operation() -> None:
    text = """
Дата - 09.05.2026
Смена - день
Участок - участок №1
===Забивка свай===
5_1 Н - ПК2655+73.50 - забивка свай - основные - 4 шт
19_2 Т - ПК3039+20 - установка наголовников - пробные - 2 шт
31_1 Н - ПК3041+10 - монтаж наголовников - основные - 1 шт
"""
    parsed = parse_report_text("pile-headcaps.txt", text)
    assert len(parsed["piles"]) == 3
    assert parsed["piles"][0]["pile_operation"] == "driving"
    assert "work_type_code" not in parsed["piles"][0]
    assert parsed["piles"][1]["pile_operation"] == "headcap"
    assert parsed["piles"][1]["work_type_code"] == "PILE_HEADCAP_INSTALLATION"
    assert parsed["piles"][1]["pile_kind"] == "test"
    assert parsed["piles"][1]["count"] == 2
    assert parsed["piles"][2]["pile_operation"] == "headcap"
    assert parsed["piles"][2]["work_type_code"] == "PILE_HEADCAP_INSTALLATION"
    assert parsed["piles"][2]["pile_kind"] == "main"
    assert parsed["piles"][2]["count"] == 1


def assert_work_hours_and_coarse_sand() -> None:
    text = """
Дата - 09.09.2026
Смена - день
Участок - участок №1
===Перевозка===
-карьер Васильки -> АД5-
/крупный песок/
Номер б/н 913 / 111,1 м3 / 6 рейс / ЖДС / плечо 16 км / время работы 7,5 ч
===Работы===
-АД5-
/Устройство насыпи из песка/
ПК ВСЖМ: ПК2694+00
Номер б/н 6947 / 940 м3 / бульдозер / ЖДС / время работы 6 ч
"""
    parsed = parse_report_text("hours-and-coarse-sand.txt", text)
    assert parsed["transport"][0]["trips"][0]["material_code"] == "COARSE_SAND"
    assert parsed["transport"][0]["trips"][0]["work_hours"] == 7.5
    assert parsed["main_works"][0]["equipment"][0]["work_hours"] == 6
    assert infer_material_code("обычный песок") == "SAND"
    assert infer_material_code("песок средней крупности") == "SAND"


def main() -> int:
    assert_pk_variants()
    assert_unlabeled_parenthetical_pk_range_does_not_override_rail_range()
    assert_equipment_number_variants()
    assert_work_split_by_performer_volumes()
    assert_work_split_by_equipment_volume_variants()
    assert_geotextile_without_equipment_variants()
    assert_one_sheet_xlsx_geotextile_without_equipment()
    assert_real_report_volume_and_operator_variants()
    assert_section_7_8_report_variants()
    assert_compact_report_format()
    assert_fill_status_block_and_number_markers()
    assert_mainline_fill_status_block()
    assert_pile_headcap_operation()
    assert_work_hours_and_coarse_sand()
    assert_section_6_transport_variants()
    assert_tolerant_max_headings()
    assert_core_blocks_inferred_without_explicit_headings()
    assert_park_exception_statuses()
    assert_multi_report_warning()
    assert_compact_xlsx_template()
    assert_transport_marker_not_used_as_endpoint()
    assert_loose_arrow_route_without_opening_dash()
    assert_opening_sign_route_without_closing_marker()
    assert_units_stockpiles_and_own_forces_variants()
    assert_transport_indirect_equipment_cues()
    assert_unknown_transport_material_keeps_compact_trip()
    parsed = parse_report_text("sample.txt", SAMPLE)
    assert parsed["header"]["report_date"] == "2026-05-04"
    assert parsed["header"]["shift"] == "day"
    assert parsed["header"]["section_code"] == "UCH_1"
    assert "АД9" in parsed["human_summary"]["constructives"]
    assert "Карьер Васильки" in parsed["human_summary"]["delivery_info"]
    assert len(parsed["transport"]) == 1
    assert parsed["transport"][0]["unit_number"] == "42"
    assert parsed["transport"][0]["plate_number"] == "913"
    assert parsed["transport"][0]["trips"][0]["material_code"] == "SAND"
    assert len(parsed["main_works"]) == 2
    assert parsed["main_works"][1]["work_type_code"] == "EMBANKMENT_CONSTRUCTION"
    assert len(parsed["aux_works"]) == 1
    assert parsed["aux_works"][0]["work_type_code"] == "SOIL_WORK"
    assert parsed["aux_works"][0]["volume"] == 1.0
    assert parsed["aux_works"][0]["unit"] == "шт"
    assert parsed["park"] == []
    assert [row.get("plate_number") for row in parsed["park"]].count("256") == 0
    assert not any("Анипченко" in str(row) for row in parsed["park"])
    assert not any("Работа на ИССО" in warning for warning in parsed["warnings"])
    assert "personnel" not in parsed
    assert parsed["staff_counts"] == []
    assert len(parsed["stockpiles"]) == 2
    assert parsed["stockpiles"][0]["rounded_pk"] == 2715
    assert len(parsed["piles"]) == 1
    assert parsed["piles"][0]["field_code"] == "5_1 Н"
    assert parsed["piles"][0]["count"] == 4
    assert "is_composite_complete" not in parsed["piles"][0]
    assert "составная готова" not in parsed["piles"][0]["comment"].lower()
    print("ok")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
