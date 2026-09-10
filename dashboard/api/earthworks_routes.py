"""MStroy dashboard endpoints for the МСтрой tab."""
from __future__ import annotations

import base64
import hashlib
import importlib.util
import json
import os
import sys
import threading
import time
import uuid
from collections import defaultdict
from datetime import date as date_cls, datetime, timedelta, timezone
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable, Optional

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import FileResponse

from main import get_conn, query, query_one  # noqa: E402
from wip_routes import _approved_report_exists, _expand_sections, _merge_section, _parse_range  # noqa: E402
from wip_analytics_routes import WORK_TAG_EXPR, _analytics_segment_volume_expr, _effective_section_join  # noqa: E402

router = APIRouter(prefix="/api/wip/earthworks", tags=["wip-earthworks"])

METRIC_ID_PREFIX = "metric:"
NULL_UNIT = "—"
SECTION_3_SOURCE_IDS = ["UCH_3", "UCH_31", "UCH_32"]
SUPPORTED_GRANULARITIES = {"day", "week", "month"}
SUPPORTED_MODES = {"cumulative", "interval"}
SUPPORTED_GROUPINGS = {"objects", "sections"}
DETAIL_CACHE_TTL_SECONDS = int(os.getenv("VSM_MSTROY_DETAIL_CACHE_TTL_SECONDS", "90"))
_DETAIL_CACHE: dict[tuple, tuple[float, dict[str, Any]]] = {}
_DETAIL_CACHE_LOCK = threading.Lock()
MSTROY_DASHBOARD_CACHE_VERSION = "mstroy-dashboard:v1"
MSTROY_DASHBOARD_CACHE_ENDPOINT_PREFIX = "mstroy_dashboard"
MSTROY_DASHBOARD_STALE_TTL_SECONDS = int(os.getenv("VSM_MSTROY_DASHBOARD_STALE_TTL_SECONDS", "3600"))
MSTROY_DASHBOARD_JSON_CACHE_DIR = Path(os.getenv("VSM_MSTROY_DASHBOARD_JSON_CACHE_DIR", "/opt/vsm/ops/report_pipeline/report_cache/mstroy_dashboard_json"))
MSTROY_DASHBOARD_WARM_INTERVAL_SECONDS = int(os.getenv("VSM_MSTROY_DASHBOARD_WARM_INTERVAL_SECONDS", "3600"))
MSTROY_DASHBOARD_INITIAL_WARM_DELAY_SECONDS = int(os.getenv("VSM_MSTROY_DASHBOARD_INITIAL_WARM_DELAY_SECONDS", "120"))
MSTROY_DASHBOARD_BUILD_WAIT_SECONDS = float(os.getenv("VSM_MSTROY_DASHBOARD_BUILD_WAIT_SECONDS", "5"))
_MSTROY_DASHBOARD_REFRESH_LOCK = threading.Lock()
_MSTROY_DASHBOARD_REFRESHING_KEYS: set[tuple[str, str]] = set()
_MSTROY_DASHBOARD_WARMER_LOCK = threading.Lock()
_MSTROY_DASHBOARD_WARMER_STARTED = False
MSTROY_GRAPH_SCRIPT = Path(os.getenv("VSM_MSTROY_GRAPH_SCRIPT", "/opt/vsm/ops/report_pipeline/dim_report/generate_mstroi_graph_report.py"))
MSTROY_GRAPH_TEMPLATE = Path(os.getenv("VSM_MSTROY_GRAPH_TEMPLATE", "/opt/vsm/ops/report_pipeline/dim_report/templates/MStroy graph report.xlsx"))
MSTROY_SUMMARY_OBJECT_PREFIX = "type:"
MSTROY_SUMMARY_WORK_PREFIX = "summary-work:"
MSTROY_SUMMARY_PARENT_LABEL = "МСтрой / Свод"
MSTROY_PIPE_ANALYTICS_TAG = "ВрТруба"
MSTROY_VISIBLE_OBJECT_TYPE_FALLBACK = (
    "TEMP_ROAD",
    "MAIN_TRACK",
    "STATION_TRACK",
    "PILE_FIELD",
    "ISSO_ACCESS",
    "PILE_DRIVING_PAD",
    "TECH_ROAD",
    "SERVICE_ROAD",
    "TRACTION_SUBSTATION",
)
MSTROY_OBJECT_LABEL_FALLBACK = {
    "TEMP_ROAD": "Притрассовая дорога",
    "MAIN_TRACK": "Сегмент основного хода",
    "STATION_TRACK": "Станционный путь",
    "PILE_FIELD": "Свайные поля",
    "ISSO_ACCESS": "Площадка и проезд к ИССО",
    "PILE_DRIVING_PAD": "Стартовая площадка погружения свай",
    "TECH_ROAD": "Технологический проезд",
    "SERVICE_ROAD": "Содержаие",
    "TRACTION_SUBSTATION": "Тяговые подстанции",
    "PIPE": "ПЖБТ / трубы",
}
MSTROY_SUMMARY_WORK_FALLBACK = (
    {"row": 4, "label": "Снятие ПРС", "unit": "м3", "scopes": ({"segment_codes": ("TOPSOIL_STRIPPING",), "object_type_codes": MSTROY_VISIBLE_OBJECT_TYPE_FALLBACK},)},
    {"row": 18, "label": "Замена грунта (выемка) - выторфовка", "unit": "м3", "scopes": ({"segment_codes": ("PEAT_REMOVAL",), "object_type_codes": MSTROY_VISIBLE_OBJECT_TYPE_FALLBACK},)},
    {"row": 32, "label": "Замена грунта (насыпь) - песок/грунт(замена)", "unit": "м3", "scopes": ({"segment_codes": ("WEAK_SOIL_REPLACEMENT",), "object_type_codes": MSTROY_VISIBLE_OBJECT_TYPE_FALLBACK},)},
    {"row": 46, "label": "Насыпь - песок/грунт(насыпь)", "unit": "м3", "scopes": ({"segment_codes": ("EMBANKMENT_CONSTRUCTION",), "object_type_codes": MSTROY_VISIBLE_OBJECT_TYPE_FALLBACK},)},
    {"row": 60, "label": "Выемка", "unit": "м3", "scopes": ({"segment_codes": ("EARTH_EXCAVATION", "EXCAVATION_MAIN"), "object_type_codes": MSTROY_VISIBLE_OBJECT_TYPE_FALLBACK},)},
    {"row": 74, "label": "Защитный слой из песка - ПГС", "unit": "м3", "scopes": ({"segment_codes": ("ZS2",), "object_type_codes": ("MAIN_TRACK", "STATION_TRACK", "PILE_FIELD", "ISSO_ACCESS", "PILE_DRIVING_PAD", "TECH_ROAD", "SERVICE_ROAD", "TRACTION_SUBSTATION")},)},
    {"row": 86, "label": "ДСО (песок)", "unit": "м3", "scopes": ({"segment_codes": ("PAVEMENT_SANDING",), "object_type_codes": ("TEMP_ROAD", "ISSO_ACCESS", "PILE_DRIVING_PAD", "TECH_ROAD", "SERVICE_ROAD", "TRACTION_SUBSTATION")},)},
    {"row": 96, "label": "Защитный слой из ЩПС - ЩПС", "unit": "м3", "scopes": ({"segment_codes": ("ZS1", "CRUSHED_STONE_PLACEMENT"), "object_type_codes": MSTROY_VISIBLE_OBJECT_TYPE_FALLBACK},)},
    {"row": 110, "label": "Усиление сваями", "unit": "шт", "scopes": ({"pile_field_codes": ("PILE_MAIN", "PILE_TRIAL"), "object_type_codes": ("PILE_FIELD",)},)},
    {"row": 113, "label": "водопропускные трубы", "unit": "шт", "scopes": ({"pile_field_codes": ("PILE_MAIN", "PILE_TRIAL"), "pile_object_type_codes": ("PIPE",), "object_type_codes": ("PIPE",)},)},
    {"row": 114, "label": "устройство а/б", "unit": "м2", "scopes": ({"segment_codes": ("ASPHALT_PAVEMENT",), "object_type_codes": MSTROY_VISIBLE_OBJECT_TYPE_FALLBACK},)},
    {"row": 115, "label": "Гибкий ростверк - геотекстиль свай", "unit": "м2", "scopes": ({"segment_codes": ("FLEXIBLE_GRILLAGE_GEOTEXTILE",), "pile_field_codes": ("FLEXIBLE_GRILLAGE_GEOTEXTILE",), "object_type_codes": MSTROY_VISIBLE_OBJECT_TYPE_FALLBACK},)},
)

SUMMARY_COLUMNS = [
    {"key": "label", "label": "Наименование", "role": "name", "dataType": "string", "sortable": True, "conditionalFormat": "none"},
    {"key": "unit", "label": "Ед. изм.", "role": "unit", "dataType": "string", "sortable": True, "conditionalFormat": "none"},
    {"key": "plan_total", "label": "Проектный объем", "role": "plan_total", "dataType": "number", "precision": 1, "sortable": True, "conditionalFormat": "none"},
    {"key": "fact_total", "label": "Факт накопительно", "role": "fact_total", "dataType": "number", "precision": 1, "sortable": True, "conditionalFormat": "none"},
    {"key": "remaining", "label": "Остаток", "role": "remaining", "dataType": "number", "precision": 1, "sortable": True, "conditionalFormat": "none"},
    {"key": "completion_total_pct", "label": "Выполнение, %", "role": "completion_total_pct", "dataType": "percent", "precision": 1, "sortable": True, "conditionalFormat": "progress"},
]

PERIOD_COLUMNS = [
    {"key": "label", "label": "Наименование", "role": "name", "dataType": "string", "sortable": True, "conditionalFormat": "none"},
    {"key": "unit", "label": "Ед. изм.", "role": "unit", "dataType": "string", "sortable": True, "conditionalFormat": "none"},
    {"key": "plan_period", "label": "План за период", "role": "plan_period", "dataType": "number", "precision": 1, "sortable": True, "conditionalFormat": "none"},
    {"key": "fact_period", "label": "Факт за период", "role": "fact_period", "dataType": "number", "precision": 1, "sortable": True, "conditionalFormat": "none"},
    {"key": "deviation", "label": "Отклонение", "role": "deviation", "dataType": "number", "precision": 1, "sortable": True, "conditionalFormat": "deviation"},
    {"key": "completion_pct", "label": "Выполнение, %", "role": "completion_pct", "dataType": "percent", "precision": 1, "sortable": True, "conditionalFormat": "progress"},
]


def _csv_values(value: Optional[str]) -> list[str]:
    if not value or value == "all":
        return []
    return [part.strip() for part in value.split(",") if part.strip()]


def _parse_earthworks_range(date_from: Optional[str], date_to: Optional[str]) -> tuple[date_cls, date_cls]:
    d_from, d_to = _parse_range(date_from, date_to)
    if d_from > d_to:
        raise HTTPException(400, "date_from не может быть позже date_to")
    if (d_to - d_from).days > 730:
        raise HTTPException(400, "Период не должен превышать 730 дней")
    return d_from, d_to


def _metric_id(tag: object, unit: object) -> str:
    raw = f"{str(tag or '').strip()}\u241f{str(unit or NULL_UNIT).strip() or NULL_UNIT}"
    encoded = base64.urlsafe_b64encode(raw.encode("utf-8")).decode("ascii").rstrip("=")
    return f"{METRIC_ID_PREFIX}{encoded}"


def _decode_metric_id(metric_id: Optional[str]) -> tuple[Optional[str], Optional[str]]:
    if not metric_id:
        return None, None
    if not metric_id.startswith(METRIC_ID_PREFIX):
        return None, None
    payload = metric_id[len(METRIC_ID_PREFIX):]
    try:
        padded = payload + "=" * (-len(payload) % 4)
        raw = base64.urlsafe_b64decode(padded.encode("ascii")).decode("utf-8")
    except Exception:
        return None, None
    if "\u241f" not in raw:
        return None, None
    tag, unit = raw.split("\u241f", 1)
    return tag or None, unit or NULL_UNIT


def _section_label(code: Optional[str], fallback: Optional[str] = None) -> str:
    display = _merge_section(code) or code
    if display == "UCH_3":
        return "Участок №3"
    if display and display.startswith("UCH_"):
        return f"Участок №{display.split('_', 1)[1]}"
    return fallback or display or "—"


def _section_source_ids(code: Optional[str]) -> list[str]:
    return SECTION_3_SOURCE_IDS if code == "UCH_3" else ([code] if code else [])


def _round_value(value: object, precision: int = 1) -> Optional[float]:
    if value is None:
        return None
    try:
        return round(float(value), precision)
    except (TypeError, ValueError):
        return None


def _completion(fact: Optional[float], plan: Optional[float]) -> Optional[float]:
    if plan is None or plan <= 0 or fact is None:
        return None
    return round(fact / plan * 100, 1)


def _deviation(fact: Optional[float], plan: Optional[float]) -> Optional[float]:
    if fact is None or plan is None:
        return None
    return round(fact - plan, 1)


def _remaining(plan_total: Optional[float], fact_total: Optional[float]) -> Optional[float]:
    if plan_total is None or fact_total is None:
        return None
    return round(max(plan_total - fact_total, 0), 1)


def _value_key(tag: object, unit: object) -> tuple[str, str]:
    return (str(tag or "Без показателя"), str(unit or NULL_UNIT))


def _ordered_unique(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        item = str(value or "").strip()
        if not item or item in seen:
            continue
        seen.add(item)
        out.append(item)
    return out


def _spec_tuple(spec: object, attr: str) -> tuple[str, ...]:
    return tuple(str(item) for item in (getattr(spec, attr, ()) or ()) if str(item or "").strip())


def _scope_from_summary_leaf(spec: object) -> dict[str, Any]:
    return {
        "segment_codes": _spec_tuple(spec, "segment_codes"),
        "pile_field_codes": _spec_tuple(spec, "pile_field_codes"),
        "object_type_codes": _spec_tuple(spec, "object_type_codes"),
        "keyword_patterns": _spec_tuple(spec, "keyword_patterns"),
        "distinct_work_code_count": bool(getattr(spec, "distinct_work_code_count", False)),
    }


@lru_cache(maxsize=1)
def _load_mstroy_graph_module() -> Any | None:
    if not MSTROY_GRAPH_SCRIPT.exists():
        return None
    module_name = "_vsm_mstroy_graph_report_scope"
    module_dir = str(MSTROY_GRAPH_SCRIPT.parent)
    if module_dir not in sys.path:
        sys.path.insert(0, module_dir)
    spec = importlib.util.spec_from_file_location(module_name, MSTROY_GRAPH_SCRIPT)
    if spec is None or spec.loader is None:
        return None
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    try:
        spec.loader.exec_module(module)
    except Exception:
        sys.modules.pop(module_name, None)
        return None
    return module


@lru_cache(maxsize=1)
def _mstroy_summary_template_rows() -> dict[int, dict[str, Optional[str]]]:
    if not MSTROY_GRAPH_TEMPLATE.exists():
        return {}
    try:
        from openpyxl import load_workbook
    except Exception:
        return {}
    try:
        wb = load_workbook(MSTROY_GRAPH_TEMPLATE, read_only=True, data_only=False)
        sheet_names = tuple(getattr(_load_mstroy_graph_module(), "SUMMARY_SHEET_CANDIDATES", ("Свод", "Лист1")))
        sheet_name = next((name for name in sheet_names if name in wb.sheetnames), wb.sheetnames[0])
        ws = wb[sheet_name]
        rows: dict[int, dict[str, Optional[str]]] = {}
        for row in range(1, ws.max_row + 1):
            unit = str(ws[f"B{row}"].value or "").strip() or None
            label = None
            for col in ("E", "D", "C"):
                value = str(ws[f"{col}{row}"].value or "").strip()
                if value:
                    label = value
                    break
            rows[row] = {"unit": unit, "label": label}
        return rows
    except Exception:
        return {}


def _summary_row_label(row: int, fallback: str = "") -> str:
    template = _mstroy_summary_template_rows().get(row) or {}
    return str(template.get("label") or fallback or f"Строка {row}")


def _summary_row_unit(row: int, fallback: Optional[str] = None) -> Optional[str]:
    template = _mstroy_summary_template_rows().get(row) or {}
    return template.get("unit") or fallback


@lru_cache(maxsize=1)
def _mstroy_summary_scope_static() -> dict[str, Any]:
    module = _load_mstroy_graph_module()

    def scope_object_types(scope: dict[str, Any], visible: tuple[str, ...]) -> tuple[str, ...]:
        codes = tuple(scope.get("object_type_codes") or ())
        if codes:
            return codes
        if scope.get("pile_field_codes"):
            return ("PILE_FIELD",)
        return visible

    if module is None or not hasattr(module, "SUMMARY_BLOCKS"):
        visible = MSTROY_VISIBLE_OBJECT_TYPE_FALLBACK
        work_options = []
        work_nodes: list[dict[str, Any]] = []
        all_scopes: list[dict[str, Any]] = []
        work_scopes: dict[str, list[dict[str, Any]]] = {}
        object_labels = dict(MSTROY_OBJECT_LABEL_FALLBACK)
        for item in MSTROY_SUMMARY_WORK_FALLBACK:
            row = int(item["row"])
            token = f"{MSTROY_SUMMARY_WORK_PREFIX}{row}"
            scopes = [dict(scope) for scope in item.get("scopes", ())]
            work_options.append({"id": token, "row": row, "label": str(item["label"]), "unit": item.get("unit")})
            work_scopes[token] = scopes
            all_scopes.extend(scopes)
            leaves: list[dict[str, Any]] = []
            seen_leaf_types: set[str] = set()
            for scope in scopes:
                for object_type in scope_object_types(scope, visible):
                    if object_type in seen_leaf_types:
                        continue
                    seen_leaf_types.add(object_type)
                    leaf_scope = dict(scope)
                    leaf_scope["object_type_codes"] = (object_type,)
                    leaves.append({
                        "id": f"{token}:type:{object_type}",
                        "row": row,
                        "label": object_labels.get(object_type) or object_type,
                        "object_type_codes": (object_type,),
                        "scopes": [leaf_scope],
                    })
            work_nodes.append({
                "id": token,
                "row": row,
                "label": str(item["label"]),
                "unit": item.get("unit"),
                "groups": [{"id": f"{token}:objects", "row": row, "label": "Объекты", "children": leaves}] if leaves else [],
                "direct_children": [],
                "scopes": scopes,
            })
        all_object_types = _ordered_unique(
            object_type
            for scope in all_scopes
            for object_type in scope_object_types(scope, visible)
        )
        return {
            "visible_object_types": visible,
            "all_object_types": tuple(all_object_types),
            "object_labels": object_labels,
            "work_options": work_options,
            "work_scopes": work_scopes,
            "work_nodes": work_nodes,
            "all_scopes": all_scopes,
        }

    visible = tuple(str(code) for code in getattr(module, "SUMMARY_VISIBLE_OBJECT_TYPES", MSTROY_VISIBLE_OBJECT_TYPE_FALLBACK))
    object_labels = dict(MSTROY_OBJECT_LABEL_FALLBACK)
    work_options: list[dict[str, Any]] = []
    work_scopes: dict[str, list[dict[str, Any]]] = {}
    work_nodes: list[dict[str, Any]] = []
    all_scopes: list[dict[str, Any]] = []

    def register_leaf(scope: dict[str, Any], row: int, *, object_label: bool = True) -> None:
        all_scopes.append(scope)
        if not object_label:
            return
        object_codes = tuple(scope.get("object_type_codes") or ())
        if len(object_codes) == 1:
            label = _summary_row_label(row, "")
            if label:
                object_labels[object_codes[0]] = label

    def make_leaf(token: str, row: int, scope: dict[str, Any]) -> dict[str, Any]:
        object_types = scope_object_types(scope, visible)
        return {
            "id": f"{token}:leaf:{row}",
            "row": row,
            "label": _summary_row_label(row),
            "object_type_codes": object_types,
            "scopes": [scope],
        }

    for block in getattr(module, "SUMMARY_BLOCKS", ()) or ():
        row = int(getattr(block, "total_row"))
        token = f"{MSTROY_SUMMARY_WORK_PREFIX}{row}"
        leaf_specs = tuple(getattr(block, "leaf_specs", ()) or ())
        scopes = [_scope_from_summary_leaf(leaf) for leaf in leaf_specs]
        leaves_by_row = {int(getattr(leaf, "row", row)): make_leaf(token, int(getattr(leaf, "row", row)), scope) for leaf, scope in zip(leaf_specs, scopes)}
        groups: list[dict[str, Any]] = []
        grouped_rows: set[int] = set()
        for group in getattr(block, "group_specs", ()) or ():
            group_row = int(getattr(group, "row"))
            child_rows = tuple(int(child) for child in (getattr(group, "child_rows", ()) or ()))
            children = [leaves_by_row[child] for child in child_rows if child in leaves_by_row]
            grouped_rows.update(child_rows)
            groups.append({"id": f"{token}:group:{group_row}", "row": group_row, "label": _summary_row_label(group_row), "children": children})
        direct_children = [leaf for leaf_row, leaf in leaves_by_row.items() if leaf_row not in grouped_rows]
        work_options.append({"id": token, "row": row, "label": _summary_row_label(row), "unit": _summary_row_unit(row)})
        work_scopes[token] = scopes
        for leaf, scope in zip(leaf_specs, scopes):
            register_leaf(scope, int(getattr(leaf, "row", row)))
        work_nodes.append({
            "id": token,
            "row": row,
            "label": _summary_row_label(row),
            "unit": _summary_row_unit(row),
            "groups": groups,
            "direct_children": direct_children,
            "scopes": scopes,
        })
    for spec in getattr(module, "SUMMARY_DIRECT_ROWS", ()) or ():
        row = int(getattr(spec, "row"))
        token = f"{MSTROY_SUMMARY_WORK_PREFIX}{row}"
        scope = _scope_from_summary_leaf(spec)
        children: list[dict[str, Any]] = []
        for object_type in scope_object_types(scope, visible):
            child_scope = dict(scope)
            child_scope["object_type_codes"] = (object_type,)
            children.append({
                "id": f"{token}:type:{object_type}",
                "row": row,
                "label": object_labels.get(object_type) or object_type,
                "object_type_codes": (object_type,),
                "scopes": [child_scope],
            })
        work_options.append({"id": token, "row": row, "label": _summary_row_label(row), "unit": _summary_row_unit(row)})
        work_scopes[token] = [scope]
        register_leaf(scope, row, object_label=False)
        work_nodes.append({
            "id": token,
            "row": row,
            "label": _summary_row_label(row),
            "unit": _summary_row_unit(row),
            "groups": [],
            "direct_children": children,
            "scopes": [scope],
        })
    all_object_types = _ordered_unique(
        object_type
        for scope in all_scopes
        for object_type in scope_object_types(scope, visible)
    )
    return {
        "visible_object_types": visible,
        "all_object_types": tuple(all_object_types),
        "object_labels": object_labels,
        "work_options": work_options,
        "work_scopes": work_scopes,
        "work_nodes": work_nodes,
        "all_scopes": all_scopes,
    }


def _mstroy_visible_object_types() -> list[str]:
    return list(_mstroy_summary_scope_static()["visible_object_types"])


def _mstroy_all_object_types() -> list[str]:
    return list(_mstroy_summary_scope_static().get("all_object_types") or _mstroy_summary_scope_static()["visible_object_types"])


def _mstroy_pipe_work_codes() -> list[str]:
    rows = query(
        """
        SELECT code
        FROM work_types
        WHERE is_active IS NOT FALSE
          AND LOWER(BTRIM(COALESCE(analytics_tag, ''))) = LOWER(%s)
        ORDER BY code
        """,
        (MSTROY_PIPE_ANALYTICS_TAG,),
    )
    return _ordered_unique(row.get("code") for row in rows)


def _work_codes_from_scope(scope: dict[str, Any], pipe_codes: list[str]) -> list[str]:
    segment_codes = pipe_codes if scope.get("distinct_work_code_count") else list(scope.get("segment_codes") or ())
    return _ordered_unique([*segment_codes, *(scope.get("pile_field_codes") or ())])


def _mstroy_summary_work_codes() -> list[str]:
    static = _mstroy_summary_scope_static()
    pipe_codes = _mstroy_pipe_work_codes() if any(scope.get("distinct_work_code_count") for scope in static["all_scopes"]) else []
    codes: list[str] = []
    for scope in static["all_scopes"]:
        codes.extend(_work_codes_from_scope(scope, pipe_codes))
    return _ordered_unique(codes)


def _mstroy_pair_scopes(summary_work_tokens: Optional[list[str]] = None) -> list[tuple[list[str], list[str]]]:
    static = _mstroy_summary_scope_static()
    if summary_work_tokens:
        scopes = [scope for token in summary_work_tokens for scope in static["work_scopes"].get(token, [])]
    else:
        scopes = list(static["all_scopes"])
    pipe_codes = _mstroy_pipe_work_codes() if any(scope.get("distinct_work_code_count") for scope in scopes) else []
    pairs: list[tuple[list[str], list[str]]] = []
    seen: set[tuple[tuple[str, ...], tuple[str, ...]]] = set()
    for scope in scopes:
        object_types = _ordered_unique(scope.get("object_type_codes") or static["visible_object_types"])
        segment_codes = pipe_codes if scope.get("distinct_work_code_count") else _ordered_unique(scope.get("segment_codes") or ())
        pile_codes = _ordered_unique(scope.get("pile_field_codes") or ())
        pile_object_types = _ordered_unique(scope.get("pile_object_type_codes") or ["PILE_FIELD"])
        for codes, types in ((segment_codes, object_types), (pile_codes, pile_object_types)):
            if not codes or not types:
                continue
            key = (tuple(codes), tuple(types))
            if key in seen:
                continue
            seen.add(key)
            pairs.append((codes, types))
    return pairs


def _mstroy_object_options() -> list[dict[str, Any]]:
    static = _mstroy_summary_scope_static()
    labels = static["object_labels"]
    return [
        {
            "id": f"{MSTROY_SUMMARY_OBJECT_PREFIX}{code}",
            "label": labels.get(code) or code,
            "parentId": MSTROY_SUMMARY_PARENT_LABEL,
            "sourceIds": [code],
        }
        for code in static.get("all_object_types") or static["visible_object_types"]
    ]


def _mstroy_work_options() -> list[dict[str, Any]]:
    return [
        {
            "id": option["id"],
            "label": option["label"],
            "parentId": MSTROY_SUMMARY_PARENT_LABEL,
            "sourceIds": [str(option["row"])],
            "unit": option.get("unit"),
        }
        for option in _mstroy_summary_scope_static()["work_options"]
    ]


def _mstroy_metric_options() -> list[dict[str, Any]]:
    codes = _mstroy_summary_work_codes()
    if not codes:
        return []
    rows = query(
        """
        WITH code_order AS (
          SELECT code, ord
          FROM unnest(%s::text[]) WITH ORDINALITY AS src(code, ord)
        )
        SELECT COALESCE(NULLIF(BTRIM(wt.analytics_tag), ''), wt.name) AS tag,
               COALESCE(NULLIF(wt.default_unit, ''), %s) AS unit,
               MIN(code_order.ord)::int AS sort_order
        FROM code_order
        JOIN work_types wt ON wt.code = code_order.code
        WHERE wt.is_active IS NOT FALSE
          AND COALESCE(NULLIF(BTRIM(wt.analytics_tag), ''), wt.name) IS NOT NULL
        GROUP BY COALESCE(NULLIF(BTRIM(wt.analytics_tag), ''), wt.name), COALESCE(NULLIF(wt.default_unit, ''), %s)
        ORDER BY MIN(code_order.ord), tag, unit
        """,
        (codes, NULL_UNIT, NULL_UNIT),
    )
    return [
        {"id": _metric_id(row.get("tag"), row.get("unit")), "label": row.get("tag"), "unit": row.get("unit"), "precision": 1}
        for row in rows
    ]


def _object_filter_clause(filters: dict[str, Any], params: list[Any], *, object_id_sql: str = "o.id::text", object_type_sql: str = "ot.code") -> Optional[str]:
    object_ids = filters.get("object_ids") or []
    object_type_codes = filters.get("object_type_codes") or []
    if object_ids and object_type_codes:
        params.extend([object_ids, object_type_codes])
        return f"({object_id_sql} = ANY(%s) OR {object_type_sql} = ANY(%s))"
    if object_ids:
        params.append(object_ids)
        return f"{object_id_sql} = ANY(%s)"
    if object_type_codes:
        params.append(object_type_codes)
        return f"{object_type_sql} = ANY(%s)"
    return None


def _mstroy_pair_scope_clause(filters: dict[str, Any], params: list[Any], *, object_type_sql: str = "ot.code") -> Optional[str]:
    pairs = filters.get("mstroy_pair_scopes") or []
    if not pairs:
        return None
    clauses: list[str] = []
    for work_codes, object_type_codes in pairs:
        if not work_codes or not object_type_codes:
            continue
        clauses.append(f"(wt.code = ANY(%s) AND {object_type_sql} = ANY(%s))")
        params.extend([work_codes, object_type_codes])
    if not clauses:
        return None
    return "(" + " OR ".join(clauses) + ")"


def _bucket_expr(granularity: str, date_sql: str) -> str:
    if granularity == "month":
        return f"date_trunc('month', {date_sql})::date"
    if granularity == "week":
        return f"date_trunc('week', {date_sql})::date"
    return f"{date_sql}::date"


def _bucket_step(granularity: str) -> timedelta:
    if granularity == "month":
        return timedelta(days=32)
    if granularity == "week":
        return timedelta(days=7)
    return timedelta(days=1)


def _bucket_start(day: date_cls, granularity: str) -> date_cls:
    if granularity == "month":
        return day.replace(day=1)
    if granularity == "week":
        return day - timedelta(days=day.weekday())
    return day


def _next_bucket(day: date_cls, granularity: str) -> date_cls:
    if granularity == "month":
        probe = (day.replace(day=28) + timedelta(days=4)).replace(day=1)
        return probe
    return day + _bucket_step(granularity)


def _periods(d_from: date_cls, d_to: date_cls, granularity: str) -> list[dict[str, str]]:
    cur = _bucket_start(d_from, granularity)
    out: list[dict[str, str]] = []
    while cur <= d_to:
        nxt = _next_bucket(cur, granularity)
        p_from = max(cur, d_from)
        p_to = min(nxt - timedelta(days=1), d_to)
        if granularity == "month":
            label = cur.strftime("%m.%Y")
        elif granularity == "week":
            label = f"{p_from.strftime('%d.%m')}–{p_to.strftime('%d.%m')}"
        else:
            label = cur.isoformat()
        out.append({"key": cur.isoformat(), "label": label, "dateFrom": p_from.isoformat(), "dateTo": p_to.isoformat()})
        cur = nxt
    return out


def _updated_at() -> str:
    row = query_one(
        """
        SELECT GREATEST(
          COALESCE((SELECT MAX(created_at) FROM daily_reports), '-infinity'::timestamptz),
          COALESCE((SELECT MAX(updated_at) FROM planned_work_items), '-infinity'::timestamptz),
          COALESCE((SELECT MAX(created_at) FROM project_work_items), '-infinity'::timestamptz)
        ) AS updated_at
        """
    ) or {}
    value = row.get("updated_at")
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).isoformat()
    return datetime.now(timezone.utc).isoformat()


def _filter_bundle(
    *,
    section_ids: Optional[str] = None,
    work_type_ids: Optional[str] = None,
    contractor_ids: Optional[str] = None,
    object_ids: Optional[str] = None,
    metric_id: Optional[str] = None,
) -> dict[str, Any]:
    section_codes = _expand_sections(section_ids) if section_ids else None
    tag, unit = _decode_metric_id(metric_id)
    raw_work_values = _csv_values(work_type_ids)
    summary_work_tokens = [value for value in raw_work_values if value.startswith(MSTROY_SUMMARY_WORK_PREFIX)]
    concrete_work_type_ids = [value for value in raw_work_values if not value.startswith(MSTROY_SUMMARY_WORK_PREFIX)]
    raw_object_values = _csv_values(object_ids)
    object_type_codes = [value[len(MSTROY_SUMMARY_OBJECT_PREFIX):] for value in raw_object_values if value.startswith(MSTROY_SUMMARY_OBJECT_PREFIX)]
    concrete_object_ids = [value for value in raw_object_values if not value.startswith(MSTROY_SUMMARY_OBJECT_PREFIX)]
    if not raw_object_values:
        object_type_codes = _mstroy_all_object_types()
    return {
        "section_codes": section_codes or [],
        "work_type_ids": concrete_work_type_ids,
        "summary_work_tokens": summary_work_tokens,
        "mstroy_pair_scopes": [] if concrete_work_type_ids else _mstroy_pair_scopes(summary_work_tokens or None),
        "contractor_ids": _csv_values(contractor_ids),
        "object_ids": concrete_object_ids,
        "object_type_codes": object_type_codes,
        "metric_tag": tag,
        "metric_unit": unit,
    }


def _mstroy_source_hash() -> str:
    try:
        return hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    except Exception:
        return "unknown"


def _mstroy_table_exists(table_name: str) -> bool:
    if not table_name.replace("_", "").isalnum():
        return False
    row = query_one("SELECT to_regclass(%s) AS table_name", [f"public.{table_name}"])
    return bool(row and row.get("table_name"))


def _mstroy_cache_piece(table_name: str, where_sql: str = "", params: tuple[object, ...] = ()) -> dict[str, str]:
    if not _mstroy_table_exists(table_name):
        return {"table": table_name, "missing": "1", "rows": "0", "max_xmin": "0"}
    where_clause = f"WHERE {where_sql}" if where_sql else ""
    row = query_one(
        f"""
        SELECT COUNT(*)::text AS rows,
               COALESCE(MAX(xmin::text::bigint), 0)::text AS max_xmin
        FROM {table_name}
        {where_clause}
        """,
        params,
    ) or {}
    return {
        "table": table_name,
        "rows": str(row.get("rows") or "0"),
        "max_xmin": str(row.get("max_xmin") or "0"),
    }


def _mstroy_dashboard_fingerprint(d_from: date_cls, d_to: date_cls) -> str:
    pieces = [
        _mstroy_cache_piece("daily_reports", "report_date BETWEEN %s AND %s", (d_from, d_to)),
        _mstroy_cache_piece("daily_work_items", "report_date BETWEEN %s AND %s", (d_from, d_to)),
        _mstroy_cache_piece("daily_work_item_segments"),
        _mstroy_cache_piece("planned_work_items"),
        _mstroy_cache_piece("project_work_items"),
        _mstroy_cache_piece("project_work_item_segments"),
        _mstroy_cache_piece("objects"),
        _mstroy_cache_piece("object_segments"),
        _mstroy_cache_piece("object_types"),
        _mstroy_cache_piece("work_types"),
        _mstroy_cache_piece("construction_sections"),
        _mstroy_cache_piece("construction_section_versions", "is_current IS TRUE"),
        _mstroy_cache_piece("temporary_roads"),
        _mstroy_cache_piece("pile_fields"),
    ]
    stable = json.dumps(
        {
            "version": MSTROY_DASHBOARD_CACHE_VERSION,
            "date_from": d_from.isoformat(),
            "date_to": d_to.isoformat(),
            "source_sha256": _mstroy_source_hash(),
            "pieces": pieces,
        },
        ensure_ascii=False,
        sort_keys=True,
        default=str,
    )
    return f"{MSTROY_DASHBOARD_CACHE_VERSION}:{hashlib.sha256(stable.encode('utf-8')).hexdigest()}"


def _mstroy_dashboard_cache_params(endpoint: str, **params: object) -> dict[str, object]:
    return {
        "endpoint": endpoint,
        **{key: ("" if value is None else value) for key, value in sorted(params.items())},
    }


def _mstroy_dashboard_persistent_key(cache_params: dict[str, object]) -> str:
    raw = json.dumps(cache_params, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _mstroy_dashboard_json_cache_path(persistent_key: str, fingerprint: str) -> Path:
    fingerprint_hash = hashlib.sha256(fingerprint.encode("utf-8")).hexdigest()
    return MSTROY_DASHBOARD_JSON_CACHE_DIR / f"{persistent_key}_{fingerprint_hash}.json"


def _mstroy_dashboard_json_cache_glob(persistent_key: str) -> list[Path]:
    if not MSTROY_DASHBOARD_JSON_CACHE_DIR.exists():
        return []
    return sorted(
        MSTROY_DASHBOARD_JSON_CACHE_DIR.glob(f"{persistent_key}_*.json"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )


def _mstroy_dashboard_json_cache_read(path: Path, *, max_age_seconds: int) -> Optional[dict]:
    try:
        if max_age_seconds > 0 and time.time() - path.stat().st_mtime > max_age_seconds:
            return None
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None
    if not isinstance(data, dict):
        return None
    cache_info = data.get("cache") if isinstance(data.get("cache"), dict) else {}
    return {
        "payload": data,
        "fingerprint": str(cache_info.get("fingerprint") or ""),
        "created_at": str(cache_info.get("persistent_created_at") or datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat()),
    }


def _mstroy_dashboard_cache_get(persistent_key: str, fingerprint: str) -> Optional[dict]:
    entry = _mstroy_dashboard_json_cache_read(
        _mstroy_dashboard_json_cache_path(persistent_key, fingerprint),
        max_age_seconds=max(MSTROY_DASHBOARD_STALE_TTL_SECONDS, MSTROY_DASHBOARD_WARM_INTERVAL_SECONDS),
    )
    return entry["payload"] if entry else None


def _mstroy_dashboard_cache_file_response(persistent_key: str, fingerprint: str) -> Optional[FileResponse]:
    path = _mstroy_dashboard_json_cache_path(persistent_key, fingerprint)
    try:
        if not path.exists() or time.time() - path.stat().st_mtime > max(MSTROY_DASHBOARD_STALE_TTL_SECONDS, MSTROY_DASHBOARD_WARM_INTERVAL_SECONDS):
            return None
    except OSError:
        return None
    return FileResponse(path, media_type="application/json")


def _mstroy_dashboard_cache_get_latest(persistent_key: str, max_age_seconds: int) -> Optional[dict]:
    for path in _mstroy_dashboard_json_cache_glob(persistent_key):
        entry = _mstroy_dashboard_json_cache_read(path, max_age_seconds=max_age_seconds)
        if entry is not None:
            return entry
    return None


def _mstroy_dashboard_mark_cache(
    payload: dict,
    *,
    status: str,
    fingerprint: str,
    persistent_created_at: Optional[str] = None,
    refreshing: bool = False,
) -> dict:
    out = dict(payload)
    generated_at = str(out.get("updatedAt") or out.get("updated_at") or persistent_created_at or datetime.now(timezone.utc).isoformat())
    out["updatedAt"] = generated_at
    out["cache"] = {
        "status": status,
        "fingerprint": fingerprint,
        "generated_at": generated_at,
        "persistent_created_at": persistent_created_at,
        "refreshing": refreshing,
    }
    return out


def _mstroy_dashboard_cache_set(persistent_key: str, fingerprint: str, payload: dict) -> dict:
    MSTROY_DASHBOARD_JSON_CACHE_DIR.mkdir(parents=True, exist_ok=True)
    created_at = datetime.now(timezone.utc).isoformat()
    cache_path = _mstroy_dashboard_json_cache_path(persistent_key, fingerprint)
    tmp_path = cache_path.with_suffix(cache_path.suffix + f".{uuid.uuid4().hex}.tmp")
    file_payload = _mstroy_dashboard_mark_cache(
        payload,
        status="fresh",
        fingerprint=fingerprint,
        persistent_created_at=created_at,
        refreshing=False,
    )
    tmp_path.write_text(json.dumps(file_payload, ensure_ascii=False, default=str), encoding="utf-8")
    tmp_path.replace(cache_path)
    cutoff = time.time() - max(MSTROY_DASHBOARD_STALE_TTL_SECONDS, MSTROY_DASHBOARD_WARM_INTERVAL_SECONDS)
    for path in _mstroy_dashboard_json_cache_glob(persistent_key)[8:]:
        try:
            if path.stat().st_mtime < cutoff:
                path.unlink()
        except FileNotFoundError:
            pass
    return file_payload


def _mstroy_dashboard_try_build_lock(lock_key: str):
    conn = get_conn()
    conn.autocommit = True
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT pg_try_advisory_lock(hashtext(%s))", (lock_key,))
            locked = bool((cur.fetchone() or [False])[0])
        if locked:
            return conn
    except Exception:
        conn.close()
        raise
    conn.close()
    return None


def _mstroy_dashboard_release_build_lock(conn, lock_key: str) -> None:
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT pg_advisory_unlock(hashtext(%s))", (lock_key,))
    except Exception:
        pass
    finally:
        conn.close()


def _mstroy_dashboard_build_and_cache(
    *,
    endpoint: str,
    persistent_key: str,
    fingerprint: str,
    build_fn,
) -> dict:
    payload = dict(build_fn())
    payload["updatedAt"] = str(payload.get("updatedAt") or _updated_at())
    return _mstroy_dashboard_cache_set(persistent_key, fingerprint, payload)


def _mstroy_dashboard_refresh_async(
    *,
    endpoint: str,
    persistent_key: str,
    fingerprint: str,
    build_fn,
    reason: str,
) -> bool:
    refresh_key = (persistent_key, fingerprint)
    with _MSTROY_DASHBOARD_REFRESH_LOCK:
        if refresh_key in _MSTROY_DASHBOARD_REFRESHING_KEYS:
            return False
        _MSTROY_DASHBOARD_REFRESHING_KEYS.add(refresh_key)

    def run() -> None:
        lock_key = f"{MSTROY_DASHBOARD_CACHE_ENDPOINT_PREFIX}:{endpoint}:{persistent_key}:{fingerprint}"
        lock_conn = None
        try:
            lock_conn = _mstroy_dashboard_try_build_lock(lock_key)
            if lock_conn is None:
                return
            if _mstroy_dashboard_cache_get(persistent_key, fingerprint) is not None:
                return
            _mstroy_dashboard_build_and_cache(
                endpoint=endpoint,
                persistent_key=persistent_key,
                fingerprint=fingerprint,
                build_fn=build_fn,
            )
        except Exception as exc:
            print(f"[mstroy-cache] background refresh failed ({reason}): {exc}", flush=True)
        finally:
            if lock_conn is not None:
                _mstroy_dashboard_release_build_lock(lock_conn, lock_key)
            with _MSTROY_DASHBOARD_REFRESH_LOCK:
                _MSTROY_DASHBOARD_REFRESHING_KEYS.discard(refresh_key)

    thread = threading.Thread(target=run, name=f"mstroy-{endpoint}-refresh", daemon=True)
    thread.start()
    return True


def _mstroy_dashboard_cached_response(
    *,
    endpoint: str,
    d_from: date_cls,
    d_to: date_cls,
    cache_params: dict[str, object],
    build_fn,
    allow_stale: bool,
    use_file_response: bool = True,
):
    fingerprint = _mstroy_dashboard_fingerprint(d_from, d_to)
    persistent_key = _mstroy_dashboard_persistent_key(cache_params)
    if use_file_response:
        cached_file = _mstroy_dashboard_cache_file_response(persistent_key, fingerprint)
        if cached_file is not None:
            return cached_file
    cached = _mstroy_dashboard_cache_get(persistent_key, fingerprint)
    if cached is not None:
        return _mstroy_dashboard_mark_cache(cached, status="fresh", fingerprint=fingerprint)
    if allow_stale:
        stale = _mstroy_dashboard_cache_get_latest(persistent_key, MSTROY_DASHBOARD_STALE_TTL_SECONDS)
        if stale is not None:
            refreshing = _mstroy_dashboard_refresh_async(
                endpoint=endpoint,
                persistent_key=persistent_key,
                fingerprint=fingerprint,
                build_fn=build_fn,
                reason="stale_request",
            )
            return _mstroy_dashboard_mark_cache(
                stale["payload"],
                status="stale",
                fingerprint=str(stale.get("fingerprint") or ""),
                persistent_created_at=stale.get("created_at"),
                refreshing=refreshing,
            )

    lock_key = f"{MSTROY_DASHBOARD_CACHE_ENDPOINT_PREFIX}:{endpoint}:{persistent_key}:{fingerprint}"
    lock_conn = _mstroy_dashboard_try_build_lock(lock_key)
    if lock_conn is None:
        deadline = time.monotonic() + max(0.0, MSTROY_DASHBOARD_BUILD_WAIT_SECONDS)
        while time.monotonic() < deadline:
            if use_file_response:
                cached_file = _mstroy_dashboard_cache_file_response(persistent_key, fingerprint)
                if cached_file is not None:
                    return cached_file
            cached = _mstroy_dashboard_cache_get(persistent_key, fingerprint)
            if cached is not None:
                return _mstroy_dashboard_mark_cache(cached, status="fresh", fingerprint=fingerprint)
            time.sleep(0.25)
        stale = _mstroy_dashboard_cache_get_latest(persistent_key, MSTROY_DASHBOARD_STALE_TTL_SECONDS)
        if stale is not None:
            marked = _mstroy_dashboard_mark_cache(
                stale["payload"],
                status="stale",
                fingerprint=str(stale.get("fingerprint") or ""),
                persistent_created_at=stale.get("created_at"),
                refreshing=True,
            )
            marked["cache_expected_fingerprint"] = fingerprint
            return marked
        raise HTTPException(503, "МСтрой уже пересчитывается; обновите вкладку через несколько секунд.")

    try:
        if use_file_response:
            cached_file = _mstroy_dashboard_cache_file_response(persistent_key, fingerprint)
            if cached_file is not None:
                return cached_file
        cached = _mstroy_dashboard_cache_get(persistent_key, fingerprint)
        if cached is not None:
            return _mstroy_dashboard_mark_cache(cached, status="fresh", fingerprint=fingerprint)
        return _mstroy_dashboard_build_and_cache(
            endpoint=endpoint,
            persistent_key=persistent_key,
            fingerprint=fingerprint,
            build_fn=build_fn,
        )
    finally:
        _mstroy_dashboard_release_build_lock(lock_conn, lock_key)


def _mstroy_default_range() -> tuple[date_cls, date_cls]:
    date_row = query_one(
        """
        SELECT COALESCE(MAX(report_date), CURRENT_DATE)::date AS max_fact_date
        FROM daily_work_items
        WHERE is_demo IS NOT TRUE
        """
    ) or {}
    default_to = date_row.get("max_fact_date") or date_cls.today()
    if isinstance(default_to, str):
        default_to = date_cls.fromisoformat(default_to[:10])
    return default_to.replace(day=1), default_to


def _warm_mstroy_dashboard_default() -> dict[str, object]:
    d_from, d_to = _mstroy_default_range()
    filters = _filter_bundle(section_ids=None, work_type_ids=None, contractor_ids=None, object_ids=None, metric_id=None)
    overview_params = _mstroy_dashboard_cache_params(
        "overview",
        date_from=d_from.isoformat(),
        date_to=d_to.isoformat(),
        section_ids="",
        work_type_ids="",
        contractor_ids="",
        object_ids="",
        metric_id="",
        grouping="objects",
    )
    _mstroy_dashboard_cached_response(
        endpoint="overview",
        d_from=d_from,
        d_to=d_to,
        cache_params=overview_params,
        allow_stale=False,
        use_file_response=False,
        build_fn=lambda: _earthworks_overview_payload(d_from, d_to, filters, grouping="objects", metric_id=None),
    )
    timeseries_params = _mstroy_dashboard_cache_params(
        "timeseries",
        date_from=d_from.isoformat(),
        date_to=d_to.isoformat(),
        section_ids="",
        work_type_ids="",
        contractor_ids="",
        object_ids="",
        metric_id="",
        granularity="day",
        mode="cumulative",
        grouping="objects",
    )
    _mstroy_dashboard_cached_response(
        endpoint="timeseries",
        d_from=d_from,
        d_to=d_to,
        cache_params=timeseries_params,
        allow_stale=False,
        use_file_response=False,
        build_fn=lambda: _earthworks_timeseries_payload(d_from, d_to, filters, granularity="day", mode="cumulative", grouping="objects"),
    )
    return {
        "ok": True,
        "status": "warmed",
        "date_from": d_from.isoformat(),
        "date_to": d_to.isoformat(),
        "endpoints": ["overview", "timeseries"],
    }


def _mstroy_dashboard_warmer_loop() -> None:
    delay = max(0, MSTROY_DASHBOARD_INITIAL_WARM_DELAY_SECONDS)
    if delay:
        time.sleep(delay)
    while True:
        try:
            _warm_mstroy_dashboard_default()
        except Exception as exc:
            print(f"[mstroy-cache] hourly warm failed: {exc}", flush=True)
        interval = max(300, MSTROY_DASHBOARD_WARM_INTERVAL_SECONDS)
        time.sleep(interval)


def _start_mstroy_dashboard_warmer() -> bool:
    global _MSTROY_DASHBOARD_WARMER_STARTED
    if MSTROY_DASHBOARD_WARM_INTERVAL_SECONDS <= 0:
        return False
    with _MSTROY_DASHBOARD_WARMER_LOCK:
        if _MSTROY_DASHBOARD_WARMER_STARTED:
            return False
        _MSTROY_DASHBOARD_WARMER_STARTED = True
    thread = threading.Thread(target=_mstroy_dashboard_warmer_loop, name="mstroy-dashboard-warmer", daemon=True)
    thread.start()
    return True


def _metric_where(filters: dict[str, Any], unit_expr: str, tag_expr: str, params: list[Any]) -> list[str]:
    clauses: list[str] = []
    if filters.get("metric_tag"):
        clauses.append(f"{tag_expr} = %s")
        params.append(filters["metric_tag"])
    if filters.get("metric_unit"):
        clauses.append(f"{unit_expr} = %s")
        params.append(filters["metric_unit"])
    return clauses


def _fact_object_type_sql() -> str:
    return "COALESCE(ot.code, CASE WHEN pf.id IS NOT NULL THEN 'PILE_FIELD' ELSE NULL END)"


def _fact_object_id_sql() -> str:
    return "COALESCE(o.id::text, pf.id::text)"


def _fact_where(d_from: date_cls, d_to: date_cls, filters: dict[str, Any], *, before: bool = False) -> tuple[list[str], list[Any]]:
    if before:
        where = ["dwi.report_date < %s"]
        params: list[Any] = [d_from]
    else:
        where = ["dwi.report_date BETWEEN %s AND %s"]
        params = [d_from, d_to]
    unit_expr = "COALESCE(NULLIF(dwi.unit, ''), wt.default_unit, '—')"
    tag_expr = WORK_TAG_EXPR
    where.extend([
        _approved_report_exists("dwi"),
        "dwi.is_demo IS NOT TRUE",
        "wt.is_active IS NOT FALSE",
        f"{tag_expr} IS NOT NULL",
    ])
    if filters.get("section_codes"):
        where.append("cs_eff.code = ANY(%s)")
        params.append(filters["section_codes"])
    if filters.get("work_type_ids"):
        where.append("dwi.work_type_id::text = ANY(%s)")
        params.append(filters["work_type_ids"])
    object_clause = _object_filter_clause(filters, params, object_id_sql=_fact_object_id_sql(), object_type_sql=_fact_object_type_sql())
    if object_clause:
        where.append(object_clause)
    pair_scope_clause = _mstroy_pair_scope_clause(filters, params, object_type_sql=_fact_object_type_sql())
    if pair_scope_clause:
        where.append(pair_scope_clause)
    if filters.get("contractor_ids"):
        where.append("COALESCE(NULLIF(dwi.contractor_name, ''), dwi.labor_source_type, '—') = ANY(%s)")
        params.append(filters["contractor_ids"])
    where.extend(_metric_where(filters, unit_expr, tag_expr, params))
    return where, params


def _plan_section_join(source_alias: str = "pwi", *, include_source_span: bool = True) -> str:
    source_section_join = f"""
        LEFT JOIN LATERAL (
          SELECT csv.section_id
          FROM construction_section_versions csv
          WHERE csv.is_current
            AND {source_alias}.pk_start IS NOT NULL
            AND {source_alias}.pk_end IS NOT NULL
            AND csv.boundary_kind = CASE
              WHEN ot.code = 'TEMP_ROAD' THEN 'temp_roads'
              WHEN ot.code IN ('PILE_FIELD', 'PILE_DRIVING_PAD', 'ISSO_ACCESS', 'PIPE', 'BRIDGE', 'OVERPASS', 'TECH_ROAD') THEN 'piles_pipes'
              ELSE 'main'
            END
            AND (({source_alias}.pk_start + {source_alias}.pk_end) / 2.0) >= LEAST(csv.pk_start, csv.pk_end)
            AND (({source_alias}.pk_start + {source_alias}.pk_end) / 2.0) < GREATEST(csv.pk_start, csv.pk_end)
          ORDER BY csv.pk_start, csv.pk_end
          LIMIT 1
        ) effective_section ON true
    """ if include_source_span else """
        LEFT JOIN LATERAL (SELECT NULL::uuid AS section_id) effective_section ON true
    """
    return f"""
        {source_section_join}
        LEFT JOIN LATERAL (
          SELECT MIN(os.pk_start) AS pk_start, MAX(os.pk_end) AS pk_end
          FROM object_segments os
          WHERE os.object_id = o.id
        ) object_span ON true
        LEFT JOIN LATERAL (
          SELECT csv.section_id
          FROM construction_section_versions csv
          WHERE csv.is_current
            AND csv.boundary_kind = CASE
              WHEN ot.code = 'TEMP_ROAD' THEN 'temp_roads'
              WHEN ot.code IN ('PILE_FIELD', 'PILE_DRIVING_PAD', 'ISSO_ACCESS', 'PIPE', 'BRIDGE', 'OVERPASS', 'TECH_ROAD') THEN 'piles_pipes'
              ELSE 'main'
            END
            AND object_span.pk_start IS NOT NULL
            AND object_span.pk_end IS NOT NULL
            AND ((object_span.pk_start + object_span.pk_end) / 2.0) >= LEAST(csv.pk_start, csv.pk_end)
            AND ((object_span.pk_start + object_span.pk_end) / 2.0) < GREATEST(csv.pk_start, csv.pk_end)
          ORDER BY csv.pk_start, csv.pk_end
          LIMIT 1
        ) object_section ON true
        LEFT JOIN construction_sections cs_override
          ON cs_override.code = CASE WHEN o.object_code IN ('ASP_2840', 'BRIDGE_282952') THEN 'UCH_3' ELSE NULL END
         AND cs_override.is_active IS NOT FALSE
        LEFT JOIN construction_sections cs_eff
          ON cs_eff.id = COALESCE(cs_override.id, effective_section.section_id, object_section.section_id)
         AND cs_eff.is_active IS NOT FALSE
    """


def _plan_where(d_from: date_cls, d_to: date_cls, filters: dict[str, Any], *, period: bool = True) -> tuple[list[str], list[Any], str]:
    tag_expr = "COALESCE(NULLIF(BTRIM(wt.analytics_tag), ''), wt.name)"
    unit_expr = "COALESCE(NULLIF(pwi.unit, ''), wt.default_unit, '—')"
    params: list[Any] = []
    where = [
        "pwi.is_active IS NOT FALSE",
        "wt.is_active IS NOT FALSE",
        "o.is_active IS NOT FALSE",
        f"{tag_expr} IS NOT NULL",
    ]
    if period:
        where.append("((pwi.period_start IS NULL AND pwi.period_end IS NULL AND %s::date = %s::date) OR (pwi.period_start IS NOT NULL AND pwi.period_end IS NOT NULL AND pwi.period_start <= %s::date AND pwi.period_end >= %s::date))")
        params.extend([d_from, d_to, d_to, d_from])
    if filters.get("section_codes"):
        where.append("cs_eff.code = ANY(%s)")
        params.append(filters["section_codes"])
    if filters.get("work_type_ids"):
        where.append("pwi.work_type_id::text = ANY(%s)")
        params.append(filters["work_type_ids"])
    object_clause = _object_filter_clause(filters, params)
    if object_clause:
        where.append(object_clause)
    pair_scope_clause = _mstroy_pair_scope_clause(filters, params)
    if pair_scope_clause:
        where.append(pair_scope_clause)
    where.extend(_metric_where(filters, unit_expr, tag_expr, params))
    return where, params, tag_expr


def _project_where(filters: dict[str, Any]) -> tuple[list[str], list[Any], str]:
    tag_expr = "COALESCE(NULLIF(BTRIM(wt.analytics_tag), ''), wt.name)"
    unit_expr = "COALESCE(NULLIF(pwi.unit, ''), wt.default_unit, '—')"
    params: list[Any] = []
    where = ["wt.is_active IS NOT FALSE", "o.is_active IS NOT FALSE", f"{tag_expr} IS NOT NULL"]
    if filters.get("section_codes"):
        where.append("cs_eff.code = ANY(%s)")
        params.append(filters["section_codes"])
    if filters.get("work_type_ids"):
        where.append("pwi.work_type_id::text = ANY(%s)")
        params.append(filters["work_type_ids"])
    object_clause = _object_filter_clause(filters, params)
    if object_clause:
        where.append(object_clause)
    pair_scope_clause = _mstroy_pair_scope_clause(filters, params)
    if pair_scope_clause:
        where.append(pair_scope_clause)
    where.extend(_metric_where(filters, unit_expr, tag_expr, params))
    return where, params, tag_expr


def _fact_metric_rows(d_from: date_cls, d_to: date_cls, filters: dict[str, Any], *, before: bool = False, by_section: bool = False) -> list[dict]:
    where, params = _fact_where(d_from, d_to, filters, before=before)
    section_select = ", cs_eff.code AS raw_section_code" if by_section else ""
    section_group = ", cs_eff.code" if by_section else ""
    unit_expr = "COALESCE(NULLIF(dwi.unit, ''), wt.default_unit, '—')"
    seg_volume_expr = _analytics_segment_volume_expr("seg")
    return query(
        f"""
        SELECT {WORK_TAG_EXPR} AS tag,
               {unit_expr} AS unit
               {section_select},
               SUM({seg_volume_expr})::numeric AS volume,
               MAX(dwi.report_date)::text AS last_fact_date
        FROM daily_work_items dwi
        JOIN work_types wt ON wt.id = dwi.work_type_id
        LEFT JOIN objects o ON o.id = dwi.object_id
        LEFT JOIN object_types ot ON ot.id = o.object_type_id
        {_effective_section_join('seg')}
        LEFT JOIN pile_fields pf ON pf.id = seg.pile_field_id
        WHERE {' AND '.join(where)}
        GROUP BY {WORK_TAG_EXPR}, {unit_expr}{section_group}
        """,
        params,
    )


def _plan_metric_rows(d_from: date_cls, d_to: date_cls, filters: dict[str, Any], *, by_section: bool = False) -> list[dict]:
    where, params, tag_expr = _plan_where(d_from, d_to, filters, period=True)
    section_select = ", cs_eff.code AS raw_section_code" if by_section else ""
    section_group = ", cs_eff.code" if by_section else ""
    unit_expr = "COALESCE(NULLIF(pwi.unit, ''), wt.default_unit, '—')"
    return query(
        f"""
        SELECT {tag_expr} AS tag,
               {unit_expr} AS unit
               {section_select},
               SUM(
                 CASE
                   WHEN pwi.period_start IS NULL AND pwi.period_end IS NULL THEN
                     CASE WHEN %s::date = %s::date THEN COALESCE(pwi.planned_volume, 0) ELSE 0 END
                   ELSE COALESCE(pwi.planned_volume, 0)
                     * GREATEST(0, (LEAST(pwi.period_end, %s::date) - GREATEST(pwi.period_start, %s::date) + 1))::numeric
                     / GREATEST(1, (pwi.period_end - pwi.period_start + 1))::numeric
                 END
               )::numeric AS volume
        FROM planned_work_items pwi
        JOIN work_types wt ON wt.id = pwi.work_type_id
        JOIN objects o ON o.id = pwi.object_id
        LEFT JOIN object_types ot ON ot.id = o.object_type_id
        {_plan_section_join('pwi')}
        WHERE {' AND '.join(where)}
        GROUP BY {tag_expr}, {unit_expr}{section_group}
        """,
        [d_from, d_to, d_to, d_from, *params],
    )


def _project_metric_rows(filters: dict[str, Any]) -> list[dict]:
    where, params, tag_expr = _project_where(filters)
    unit_expr = "COALESCE(NULLIF(pwi.unit, ''), wt.default_unit, '—')"
    return query(
        f"""
        SELECT {tag_expr} AS tag,
               {unit_expr} AS unit,
               SUM(pwi.project_volume)::numeric AS volume
        FROM project_work_items pwi
        JOIN work_types wt ON wt.id = pwi.work_type_id
        JOIN objects o ON o.id = pwi.object_id
        LEFT JOIN object_types ot ON ot.id = o.object_type_id
        {_plan_section_join('pwi', include_source_span=False)}
        WHERE {' AND '.join(where)}
        GROUP BY {tag_expr}, {unit_expr}
        """,
        params,
    )


DetailKey = tuple[str, str, str, str]


def _detail_key(row: dict[str, Any], section_overrides: Optional[dict[tuple[str, str, str], str]] = None) -> DetailKey:
    work_code = str(row.get("work_code") or "")
    object_type_code = str(row.get("object_type_code") or "")
    object_key = str(row.get("object_id") or row.get("object_code") or "—")
    base = (work_code, object_type_code, object_key)
    raw_section = row.get("raw_section_code")
    section = str(_merge_section(raw_section) or raw_section or "—")
    if section == "—" and section_overrides:
        section = str(section_overrides.get(base) or section)
    return (*base, section)


def _fact_detail_rows(d_from: date_cls, d_to: date_cls, filters: dict[str, Any], *, before: bool = False) -> list[dict]:
    where, params = _fact_where(d_from, d_to, filters, before=before)
    unit_expr = "COALESCE(NULLIF(dwi.unit, ''), wt.default_unit, '—')"
    seg_volume_expr = _analytics_segment_volume_expr("seg")
    return query(
        f"""
        SELECT wt.code AS work_code,
               {WORK_TAG_EXPR} AS tag,
               {unit_expr} AS unit,
               COALESCE(ot.code, CASE WHEN pf.id IS NOT NULL THEN 'PILE_FIELD' ELSE NULL END, '') AS object_type_code,
               COALESCE(o.id::text, pf.id::text, '') AS object_id,
               COALESCE(NULLIF(o.object_code, ''), NULLIF(pf.field_code, ''), '') AS object_code,
               COALESCE(NULLIF(o.name, ''), NULLIF(o.object_code, ''), NULLIF(pf.field_code, ''), 'Без объекта') AS object_label,
               cs_eff.code AS raw_section_code,
               SUM({seg_volume_expr})::numeric AS volume,
               MAX(dwi.report_date)::text AS last_fact_date
        FROM daily_work_items dwi
        JOIN work_types wt ON wt.id = dwi.work_type_id
        LEFT JOIN objects o ON o.id = dwi.object_id
        LEFT JOIN object_types ot ON ot.id = o.object_type_id
        {_effective_section_join('seg')}
        LEFT JOIN pile_fields pf ON pf.id = seg.pile_field_id
        WHERE {' AND '.join(where)}
        GROUP BY wt.code, {WORK_TAG_EXPR}, {unit_expr}, COALESCE(ot.code, CASE WHEN pf.id IS NOT NULL THEN 'PILE_FIELD' ELSE NULL END, ''), COALESCE(o.id::text, pf.id::text, ''), COALESCE(NULLIF(o.object_code, ''), NULLIF(pf.field_code, ''), ''), COALESCE(NULLIF(o.name, ''), NULLIF(o.object_code, ''), NULLIF(pf.field_code, ''), 'Без объекта'), cs_eff.code
        """,
        params,
    )


def _plan_detail_rows(d_from: date_cls, d_to: date_cls, filters: dict[str, Any]) -> list[dict]:
    where, params, tag_expr = _plan_where(d_from, d_to, filters, period=True)
    unit_expr = "COALESCE(NULLIF(pwi.unit, ''), wt.default_unit, '—')"
    return query(
        f"""
        SELECT wt.code AS work_code,
               {tag_expr} AS tag,
               {unit_expr} AS unit,
               COALESCE(ot.code, '') AS object_type_code,
               o.id::text AS object_id,
               COALESCE(NULLIF(o.object_code, ''), '') AS object_code,
               COALESCE(NULLIF(o.name, ''), NULLIF(o.object_code, ''), 'Без объекта') AS object_label,
               cs_eff.code AS raw_section_code,
               SUM(
                 CASE
                   WHEN pwi.period_start IS NULL AND pwi.period_end IS NULL THEN
                     CASE WHEN %s::date = %s::date THEN COALESCE(pwi.planned_volume, 0) ELSE 0 END
                   ELSE COALESCE(pwi.planned_volume, 0)
                     * GREATEST(0, (LEAST(pwi.period_end, %s::date) - GREATEST(pwi.period_start, %s::date) + 1))::numeric
                     / GREATEST(1, (pwi.period_end - pwi.period_start + 1))::numeric
                 END
               )::numeric AS volume
        FROM planned_work_items pwi
        JOIN work_types wt ON wt.id = pwi.work_type_id
        JOIN objects o ON o.id = pwi.object_id
        LEFT JOIN object_types ot ON ot.id = o.object_type_id
        {_plan_section_join('pwi')}
        WHERE {' AND '.join(where)}
        GROUP BY wt.code, {tag_expr}, {unit_expr}, ot.code, o.id, o.object_code, o.name, cs_eff.code
        """,
        [d_from, d_to, d_to, d_from, *params],
    )


def _project_detail_rows(filters: dict[str, Any]) -> list[dict]:
    where, params, tag_expr = _project_where(filters)
    unit_expr = "COALESCE(NULLIF(pwi.unit, ''), wt.default_unit, '—')"
    return query(
        f"""
        SELECT wt.code AS work_code,
               {tag_expr} AS tag,
               {unit_expr} AS unit,
               COALESCE(ot.code, '') AS object_type_code,
               o.id::text AS object_id,
               COALESCE(NULLIF(o.object_code, ''), '') AS object_code,
               COALESCE(NULLIF(o.name, ''), NULLIF(o.object_code, ''), 'Без объекта') AS object_label,
               cs_eff.code AS raw_section_code,
               SUM(pwi.project_volume)::numeric AS volume
        FROM project_work_items pwi
        JOIN work_types wt ON wt.id = pwi.work_type_id
        JOIN objects o ON o.id = pwi.object_id
        LEFT JOIN object_types ot ON ot.id = o.object_type_id
        {_plan_section_join('pwi', include_source_span=False)}
        WHERE {' AND '.join(where)}
        GROUP BY wt.code, {tag_expr}, {unit_expr}, ot.code, o.id, o.object_code, o.name, cs_eff.code
        """,
        params,
    )


def _detail_cache_key(d_from: date_cls, d_to: date_cls, filters: dict[str, Any]) -> tuple:
    return (
        d_from.isoformat(),
        d_to.isoformat(),
        tuple(filters.get("section_codes") or ()),
        tuple(filters.get("work_type_ids") or ()),
        tuple(filters.get("summary_work_tokens") or ()),
        tuple(tuple(tuple(part) for part in pair) for pair in (filters.get("mstroy_pair_scopes") or ())),
        tuple(filters.get("contractor_ids") or ()),
        tuple(filters.get("object_ids") or ()),
        tuple(filters.get("object_type_codes") or ()),
        filters.get("metric_tag"),
        filters.get("metric_unit"),
    )


def _apply_single_section_overrides(maps: dict[str, Any]) -> None:
    concrete_sections: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    for bucket in ("fact_before", "fact_period", "plan_period"):
        for key in (maps.get(bucket) or {}).keys():
            if len(key) == 4 and key[3] != "—":
                concrete_sections[key[:3]].add(key[3])
    overrides = {base: next(iter(sections)) for base, sections in concrete_sections.items() if len(sections) == 1}
    maps["section_overrides"] = overrides
    if not overrides:
        return
    metric_keys = maps.get("metric_keys") or {}
    for bucket in ("plan_period", "plan_total"):
        values = maps.get(bucket) or {}
        for key, value in list(values.items()):
            if len(key) != 4 or key[3] != "—":
                continue
            section = overrides.get(key[:3])
            if not section:
                continue
            target = (key[0], key[1], key[2], section)
            _add_to_map(values, target, value)
            values.pop(key, None)
            maps["keys"].discard(key)
            maps["keys"].add(target)
            if key in metric_keys and target not in metric_keys:
                metric_keys[target] = metric_keys[key]
            metric_keys.pop(key, None)
            maps["section_codes"][section] = section


def _build_detail_maps(d_from: date_cls, d_to: date_cls, filters: dict[str, Any]) -> dict[str, Any]:
    maps: dict[str, Any] = {
        "keys": set(),
        "plan_period": {},
        "fact_before": {},
        "fact_period": {},
        "plan_total": {},
        "object_labels": {},
        "object_codes": {},
        "section_codes": {},
        "section_overrides": {},
        "metric_keys": {},
        "last_fact_date": {},
    }

    def add_row(bucket: str, row: dict[str, Any]) -> None:
        key = _detail_key(row)
        metric_key = _value_key(row.get("tag"), row.get("unit"))
        maps["keys"].add(key)
        maps["metric_keys"][key] = metric_key
        _add_to_map(maps[bucket], key, row.get("volume"))
        maps["object_labels"][(key[1], key[2])] = row.get("object_label") or row.get("object_code") or "Без объекта"
        maps["object_codes"][(key[1], key[2])] = row.get("object_code") or key[2]
        maps["section_codes"][key[3]] = key[3]
        if bucket == "fact_period":
            fact_date = row.get("last_fact_date")
            if fact_date:
                current = maps["last_fact_date"].get(metric_key)
                if current is None or str(fact_date) > str(current):
                    maps["last_fact_date"][metric_key] = str(fact_date)

    for row in _plan_detail_rows(d_from, d_to, filters):
        add_row("plan_period", row)
    for row in _fact_detail_rows(d_from, d_to, filters, before=True):
        add_row("fact_before", row)
    for row in _fact_detail_rows(d_from, d_to, filters):
        add_row("fact_period", row)
    for row in _project_detail_rows(filters):
        add_row("plan_total", row)
    _apply_single_section_overrides(maps)
    maps["keys"] = set(maps["keys"])
    return maps


def _build_detail_maps_cached(d_from: date_cls, d_to: date_cls, filters: dict[str, Any]) -> dict[str, Any]:
    if DETAIL_CACHE_TTL_SECONDS <= 0:
        return _build_detail_maps(d_from, d_to, filters)
    key = _detail_cache_key(d_from, d_to, filters)
    now = time.monotonic()
    with _DETAIL_CACHE_LOCK:
        cached = _DETAIL_CACHE.get(key)
        if cached and now - cached[0] <= DETAIL_CACHE_TTL_SECONDS:
            return cached[1]
    payload = _build_detail_maps(d_from, d_to, filters)
    with _DETAIL_CACHE_LOCK:
        _DETAIL_CACHE[key] = (now, payload)
        if len(_DETAIL_CACHE) > 24:
            oldest = sorted(_DETAIL_CACHE.items(), key=lambda item: item[1][0])[:8]
            for old_key, _ in oldest:
                _DETAIL_CACHE.pop(old_key, None)
    return payload


def _metric_maps_from_detail_maps(detail_maps: dict[str, Any]) -> dict[str, Any]:
    metric_keys = detail_maps.get("metric_keys") or {}
    maps: dict[str, Any] = {
        "keys": [],
        "plan_period": {},
        "fact_before": {},
        "fact_period": {},
        "plan_total": {},
        "last_fact_date": dict(detail_maps.get("last_fact_date") or {}),
    }
    for bucket in ("plan_period", "fact_before", "fact_period", "plan_total"):
        for detail_key, value in (detail_maps.get(bucket) or {}).items():
            metric_key = metric_keys.get(detail_key)
            if not metric_key:
                continue
            _add_to_map(maps[bucket], metric_key, value)
    maps["keys"] = sorted(
        set(maps["plan_period"]) | set(maps["fact_before"]) | set(maps["fact_period"]) | set(maps["plan_total"]),
        key=lambda item: (item[1], item[0]),
    )
    return maps


def _scope_pairs(scope: dict[str, Any], pipe_codes: list[str]) -> list[tuple[list[str], list[str]]]:
    object_types = _ordered_unique(scope.get("object_type_codes") or _mstroy_visible_object_types())
    segment_codes = pipe_codes if scope.get("distinct_work_code_count") else _ordered_unique(scope.get("segment_codes") or ())
    pile_codes = _ordered_unique(scope.get("pile_field_codes") or ())
    pile_object_types = _ordered_unique(scope.get("pile_object_type_codes") or ["PILE_FIELD"])
    pairs: list[tuple[list[str], list[str]]] = []
    if segment_codes and object_types:
        pairs.append((segment_codes, object_types))
    if pile_codes:
        pairs.append((pile_codes, pile_object_types))
    return pairs


def _detail_keys_for_scopes(detail_keys: Iterable[DetailKey], scopes: Iterable[dict[str, Any]], pipe_codes: list[str]) -> list[DetailKey]:
    pairs: list[tuple[set[str], set[str]]] = []
    for scope in scopes:
        pairs.extend((set(codes), set(object_types)) for codes, object_types in _scope_pairs(scope, pipe_codes))
    if not pairs:
        return []
    out = []
    for key in detail_keys:
        work_code, object_type_code, _object_key, _section = key
        if any(work_code in codes and object_type_code in object_types for codes, object_types in pairs):
            out.append(key)
    return out


def _sum_detail_values(keys: Iterable[DetailKey], detail_maps: dict[str, Any]) -> dict[str, Optional[float]]:
    key_list = list(keys)

    def total(name: str) -> float:
        src = detail_maps.get(name) or {}
        return round(sum(float(src.get(key) or 0) for key in key_list), 1)

    plan = total("plan_period")
    before = total("fact_before")
    fact = total("fact_period")
    project = total("plan_total")
    return _metric_values(plan, before, fact, project)


def _tree_row(
    *,
    row_id: str,
    parent_id: Optional[str],
    label: str,
    unit: Optional[str],
    level: int,
    keys: Iterable[DetailKey],
    detail_maps: dict[str, Any],
    has_children: bool,
    source_ids: Optional[list[str]] = None,
    include_internal_keys: bool = False,
) -> dict[str, Any]:
    key_list = list(keys)
    row = {
        "id": row_id,
        "parentId": parent_id,
        "label": label,
        "unit": unit,
        "level": level,
        "hasChildren": has_children,
        "sourceIds": source_ids or [],
        "values": _sum_detail_values(key_list, detail_maps),
    }
    if include_internal_keys:
        row["_detailKeys"] = key_list
    return row


def _section_sort_key(code: str) -> tuple[int, str]:
    if code == "—":
        return (9999, code)
    if code.startswith("UCH_"):
        try:
            return (int(code.split("_", 1)[1]), code)
        except ValueError:
            pass
    return (9998, code)


def _build_mstroy_tree_rows(detail_maps: dict[str, Any], *, grouping: str, prefix: str, filters: dict[str, Any], include_internal_keys: bool = False) -> list[dict[str, Any]]:
    static = _mstroy_summary_scope_static()
    pipe_codes = _mstroy_pipe_work_codes() if any(scope.get("distinct_work_code_count") for scope in static["all_scopes"]) else []
    detail_keys = list(detail_maps.get("keys") or [])
    allowed_work_tokens = set(filters.get("summary_work_tokens") or [])

    def object_rows(parent_id: str, parent_level: int, unit: Optional[str], keys: list[DetailKey]) -> list[dict[str, Any]]:
        grouped: dict[tuple[str, str], list[DetailKey]] = defaultdict(list)
        for key in keys:
            grouped[(key[1], key[2])].append(key)
        out: list[dict[str, Any]] = []
        for (object_type, object_key), object_keys in sorted(grouped.items(), key=lambda item: str(detail_maps["object_labels"].get(item[0], item[0][1]))):
            label = str(detail_maps["object_labels"].get((object_type, object_key)) or object_key or "Без объекта")
            code = str(detail_maps["object_codes"].get((object_type, object_key)) or object_key)
            out.append(_tree_row(
                row_id=f"{parent_id}:object:{object_type}:{object_key}",
                parent_id=parent_id,
                label=label,
                unit=unit,
                level=parent_level + 1,
                keys=object_keys,
                detail_maps=detail_maps,
                has_children=False,
                source_ids=[code],
                include_internal_keys=include_internal_keys,
            ))
        return out

    def leaf_rows(parent_id: str, level: int, unit: Optional[str], leaf: dict[str, Any], available_keys: list[DetailKey]) -> list[dict[str, Any]]:
        leaf_id = f"{parent_id}:leaf:{leaf['row']}:{leaf['id'].split(':')[-1]}"
        leaf_keys = _detail_keys_for_scopes(available_keys, leaf.get("scopes") or [], pipe_codes)
        children = object_rows(leaf_id, level, unit, leaf_keys)
        return [
            _tree_row(
                row_id=leaf_id,
                parent_id=parent_id,
                label=str(leaf.get("label") or "Объекты"),
                unit=unit,
                level=level,
                keys=leaf_keys,
                detail_maps=detail_maps,
                has_children=bool(children),
                source_ids=list(leaf.get("object_type_codes") or []),
                include_internal_keys=include_internal_keys,
            ),
            *children,
        ]

    def work_rows(parent_id: Optional[str], level: int, work: dict[str, Any], available_keys: list[DetailKey], work_id: str) -> list[dict[str, Any]]:
        token = str(work.get("id"))
        work_keys = _detail_keys_for_scopes(available_keys, work.get("scopes") or [], pipe_codes)
        children: list[dict[str, Any]] = []
        for group in work.get("groups") or []:
            group_id = f"{work_id}:group:{group['row']}"
            group_scopes = [scope for child in group.get("children") or [] for scope in (child.get("scopes") or [])]
            group_keys = _detail_keys_for_scopes(work_keys, group_scopes, pipe_codes)
            group_children: list[dict[str, Any]] = []
            for leaf in group.get("children") or []:
                group_children.extend(leaf_rows(group_id, level + 2, work.get("unit"), leaf, work_keys))
            children.append(_tree_row(
                row_id=group_id,
                parent_id=work_id,
                label=str(group.get("label") or "Объекты"),
                unit=work.get("unit"),
                level=level + 1,
                keys=group_keys,
                detail_maps=detail_maps,
                has_children=bool(group_children),
                source_ids=[],
                include_internal_keys=include_internal_keys,
            ))
            children.extend(group_children)
        for leaf in work.get("direct_children") or []:
            children.extend(leaf_rows(work_id, level + 1, work.get("unit"), leaf, work_keys))
        return [
            _tree_row(
                row_id=work_id,
                parent_id=parent_id,
                label=str(work.get("label") or token),
                unit=work.get("unit"),
                level=level,
                keys=work_keys,
                detail_maps=detail_maps,
                has_children=bool(children),
                source_ids=[str(work.get("row") or "")],
                include_internal_keys=include_internal_keys,
            ),
            *children,
        ]

    work_nodes = [work for work in (static.get("work_nodes") or []) if not allowed_work_tokens or str(work.get("id")) in allowed_work_tokens]
    if grouping == "sections":
        out: list[dict[str, Any]] = []
        section_groups: dict[str, list[DetailKey]] = defaultdict(list)
        for key in detail_keys:
            section_groups[key[3]].append(key)
        for section, section_keys in sorted(section_groups.items(), key=lambda item: _section_sort_key(item[0])):
            section_id = f"{prefix}:section:{section}"
            children: list[dict[str, Any]] = []
            for work in work_nodes:
                token = str(work.get("id"))
                children.extend(work_rows(section_id, 1, work, section_keys, f"{section_id}:{token}"))
            out.append(_tree_row(
                row_id=section_id,
                parent_id=None,
                label=_section_label(section),
                unit=None,
                level=0,
                keys=section_keys,
                detail_maps=detail_maps,
                has_children=bool(children),
                source_ids=_section_source_ids(section),
                include_internal_keys=include_internal_keys,
            ))
            out.extend(children)
        return out

    rows: list[dict[str, Any]] = []
    for work in work_nodes:
        token = str(work.get("id"))
        rows.extend(work_rows(None, 0, work, detail_keys, f"{prefix}:{token}"))
    return rows


def _metric_values(plan_period: Optional[float], fact_before: Optional[float], fact_period: Optional[float], plan_total: Optional[float]) -> dict[str, Optional[float]]:
    fact_total = None if fact_before is None and fact_period is None else round((fact_before or 0) + (fact_period or 0), 1)
    return {
        "plan_period": plan_period,
        "fact_before": fact_before,
        "fact_period": fact_period,
        "fact_total": fact_total,
        "plan_total": plan_total,
        "deviation": _deviation(fact_period, plan_period),
        "remaining": _remaining(plan_total, fact_total),
        "completion_pct": _completion(fact_period, plan_period),
        "completion_total_pct": _completion(fact_total, plan_total),
    }


def _row_values_from_maps(
    key: tuple[str, str],
    plan_period: dict[tuple[str, str], float],
    fact_before: dict[tuple[str, str], float],
    fact_period: dict[tuple[str, str], float],
    plan_total: dict[tuple[str, str], float],
) -> dict[str, Optional[float]]:
    return _metric_values(
        _round_value(plan_period.get(key)),
        _round_value(fact_before.get(key)),
        _round_value(fact_period.get(key)),
        _round_value(plan_total.get(key)),
    )


def _add_to_map(target: dict, key: tuple, value: object) -> None:
    amount = _round_value(value)
    if amount is None:
        return
    target[key] = target.get(key, 0.0) + amount


def _build_metric_maps(d_from: date_cls, d_to: date_cls, filters: dict[str, Any]) -> dict[str, Any]:
    plan_period: dict[tuple[str, str], float] = {}
    fact_before: dict[tuple[str, str], float] = {}
    fact_period: dict[tuple[str, str], float] = {}
    plan_total: dict[tuple[str, str], float] = {}
    last_fact_date: dict[tuple[str, str], Optional[str]] = {}

    for row in _plan_metric_rows(d_from, d_to, filters):
        _add_to_map(plan_period, _value_key(row.get("tag"), row.get("unit")), row.get("volume"))
    for row in _fact_metric_rows(d_from, d_to, filters, before=True):
        _add_to_map(fact_before, _value_key(row.get("tag"), row.get("unit")), row.get("volume"))
    for row in _fact_metric_rows(d_from, d_to, filters):
        key = _value_key(row.get("tag"), row.get("unit"))
        _add_to_map(fact_period, key, row.get("volume"))
        last_fact_date[key] = row.get("last_fact_date")
    for row in _project_metric_rows(filters):
        _add_to_map(plan_total, _value_key(row.get("tag"), row.get("unit")), row.get("volume"))
    keys = sorted(set(plan_period) | set(fact_before) | set(fact_period) | set(plan_total), key=lambda item: (item[1], item[0]))
    return {"keys": keys, "plan_period": plan_period, "fact_before": fact_before, "fact_period": fact_period, "plan_total": plan_total, "last_fact_date": last_fact_date}


def _metric_row(key: tuple[str, str], maps: dict[str, Any], *, level: int = 1, parent_id: Optional[str] = None) -> dict[str, Any]:
    tag, unit = key
    metric_id = _metric_id(tag, unit)
    return {
        "id": metric_id,
        "parentId": parent_id,
        "label": tag,
        "unit": unit,
        "level": level,
        "hasChildren": False,
        "sourceIds": [tag],
        "values": _row_values_from_maps(key, maps["plan_period"], maps["fact_before"], maps["fact_period"], maps["plan_total"]),
    }


def _metric_totals(metric_key: Optional[tuple[str, str]], maps: dict[str, Any], d_from: date_cls, d_to: date_cls, filters: dict[str, Any]) -> Optional[dict[str, Any]]:
    if not metric_key:
        return None
    tag, unit = metric_key
    values = _row_values_from_maps(metric_key, maps["plan_period"], maps["fact_before"], maps["fact_period"], maps["plan_total"])
    pace = _pace_7d(metric_key, d_from, d_to, filters)
    return {
        "metricId": _metric_id(tag, unit),
        "label": tag,
        "unit": unit,
        "plan": values.get("plan_period"),
        "fact": values.get("fact_period"),
        "deviation": values.get("deviation"),
        "remaining": values.get("remaining"),
        "completionPct": values.get("completion_pct"),
        "lastFactDate": maps["last_fact_date"].get(metric_key),
        "pace7d": pace,
    }


def _pace_7d(metric_key: tuple[str, str], d_from: date_cls, d_to: date_cls, filters: dict[str, Any]) -> Optional[float]:
    local_filters = dict(filters)
    local_filters["metric_tag"], local_filters["metric_unit"] = metric_key
    start = max(d_from, d_to - timedelta(days=6))
    rows = _fact_metric_rows(start, d_to, local_filters)
    total = sum(float(row.get("volume") or 0) for row in rows)
    days = max((d_to - start).days + 1, 1)
    return round(total / days, 1) if total else None


@router.get("/meta")
def earthworks_meta(
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
):
    date_row = query_one(
        """
        SELECT COALESCE(MAX(report_date), CURRENT_DATE)::date AS max_fact_date,
               COALESCE(MIN(report_date), CURRENT_DATE)::date AS min_fact_date
        FROM daily_work_items
        WHERE is_demo IS NOT TRUE
        """
    ) or {}
    default_to = date_row.get("max_fact_date") or date_cls.today()
    if isinstance(default_to, str):
        default_to = date_cls.fromisoformat(default_to[:10])
    default_from = default_to.replace(day=1)
    if date_from or date_to:
        default_from, default_to = _parse_earthworks_range(date_from, date_to)

    sections = []
    for row in query("SELECT code, name FROM construction_sections WHERE is_active IS NOT FALSE ORDER BY sort_order NULLS LAST, code"):
        code = row.get("code")
        sections.append({"id": code, "label": _section_label(code, row.get("name")), "sourceIds": _section_source_ids(code)})

    work_types = _mstroy_work_options()
    contractors = [
        {"id": row.get("contractor_id"), "label": row.get("contractor_label") or row.get("contractor_id")}
        for row in query(
            """
            SELECT COALESCE(NULLIF(contractor_name, ''), labor_source_type, '—') AS contractor_id,
                   COALESCE(NULLIF(contractor_name, ''), labor_source_type, '—') AS contractor_label,
                   COUNT(*)::int AS row_count
            FROM daily_work_items
            WHERE is_demo IS NOT TRUE
            GROUP BY COALESCE(NULLIF(contractor_name, ''), labor_source_type, '—')
            ORDER BY row_count DESC, contractor_label
            LIMIT 250
            """
        )
    ]
    objects = _mstroy_object_options()
    metrics = _mstroy_metric_options()
    return {
        "updatedAt": _updated_at(),
        "defaultDateFrom": default_from.isoformat(),
        "defaultDateTo": default_to.isoformat(),
        "dimensions": {"sections": sections, "workTypes": work_types, "contractors": contractors, "objects": objects, "metrics": metrics},
        "columns": {"summary": SUMMARY_COLUMNS, "period": PERIOD_COLUMNS, "overall": PERIOD_COLUMNS},
        "capabilities": {"exportCsv": True, "exportXlsx": False, "hasForecast": False, "source": "mstroy_graph_summary", "supportedGranularities": sorted(SUPPORTED_GRANULARITIES), "supportedModes": sorted(SUPPORTED_MODES), "supportedGroupings": sorted(SUPPORTED_GROUPINGS)},
    }


def _earthworks_overview_payload(
    d_from: date_cls,
    d_to: date_cls,
    filters: dict[str, Any],
    *,
    grouping: str,
    metric_id: Optional[str],
) -> dict[str, Any]:
    detail_maps = _build_detail_maps_cached(d_from, d_to, filters)
    maps = _metric_maps_from_detail_maps(detail_maps)
    keys: list[tuple[str, str]] = maps["keys"]

    summary_rows = _build_mstroy_tree_rows(detail_maps, grouping=grouping, prefix="summary", filters=filters)
    period_rows = _build_mstroy_tree_rows(detail_maps, grouping=grouping, prefix="period", filters=filters)
    overall_rows = [_metric_row(key, maps, level=0, parent_id=None) for key in keys]
    selected_key = None
    if metric_id:
        decoded = _decode_metric_id(metric_id)
        if decoded[0]:
            selected_key = decoded  # type: ignore[assignment]
    if selected_key not in keys:
        selected_key = keys[0] if keys else None
    return {
        "updatedAt": _updated_at(),
        "summaryRows": summary_rows,
        "periodRows": period_rows,
        "overallRows": overall_rows,
        "selectedMetricTotals": _metric_totals(selected_key, maps, d_from, d_to, filters),
    }


@router.get("/overview")
def earthworks_overview(
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    section_ids: Optional[str] = Query(None),
    work_type_ids: Optional[str] = Query(None),
    contractor_ids: Optional[str] = Query(None),
    object_ids: Optional[str] = Query(None),
    metric_id: Optional[str] = Query(None),
    grouping: str = Query("objects"),
    allow_stale: bool = Query(True),
):
    if grouping not in SUPPORTED_GROUPINGS:
        raise HTTPException(400, "grouping должен быть objects или sections")
    d_from, d_to = _parse_earthworks_range(date_from, date_to)
    filters = _filter_bundle(section_ids=section_ids, work_type_ids=work_type_ids, contractor_ids=contractor_ids, object_ids=object_ids, metric_id=metric_id)
    cache_params = _mstroy_dashboard_cache_params(
        "overview",
        date_from=d_from.isoformat(),
        date_to=d_to.isoformat(),
        section_ids=section_ids or "",
        work_type_ids=work_type_ids or "",
        contractor_ids=contractor_ids or "",
        object_ids=object_ids or "",
        metric_id=metric_id or "",
        grouping=grouping,
    )
    return _mstroy_dashboard_cached_response(
        endpoint="overview",
        d_from=d_from,
        d_to=d_to,
        cache_params=cache_params,
        allow_stale=allow_stale,
        build_fn=lambda: _earthworks_overview_payload(d_from, d_to, filters, grouping=grouping, metric_id=metric_id),
    )


def _fact_bucket_rows(d_from: date_cls, d_to: date_cls, filters: dict[str, Any], granularity: str) -> list[dict]:
    where, params = _fact_where(d_from, d_to, filters, before=False)
    unit_expr = "COALESCE(NULLIF(dwi.unit, ''), wt.default_unit, '—')"
    bucket = _bucket_expr(granularity, "dwi.report_date")
    seg_volume_expr = _analytics_segment_volume_expr("seg")
    return query(
        f"""
        SELECT {bucket} AS bucket,
               {WORK_TAG_EXPR} AS tag,
               {unit_expr} AS unit,
               SUM({seg_volume_expr})::numeric AS volume
        FROM daily_work_items dwi
        JOIN work_types wt ON wt.id = dwi.work_type_id
        LEFT JOIN objects o ON o.id = dwi.object_id
        LEFT JOIN object_types ot ON ot.id = o.object_type_id
        {_effective_section_join('seg')}
        LEFT JOIN pile_fields pf ON pf.id = seg.pile_field_id
        WHERE {' AND '.join(where)}
        GROUP BY {bucket}, {WORK_TAG_EXPR}, {unit_expr}
        ORDER BY bucket, tag, unit
        """,
        params,
    )


def _plan_bucket_rows(d_from: date_cls, d_to: date_cls, filters: dict[str, Any], granularity: str) -> list[dict]:
    where, params, tag_expr = _plan_where(d_from, d_to, filters, period=True)
    unit_expr = "COALESCE(NULLIF(pwi.unit, ''), wt.default_unit, '—')"
    bucket = _bucket_expr(granularity, "bucket_day")
    return query(
        f"""
        WITH days AS (
          SELECT generate_series(%s::date, %s::date, interval '1 day')::date AS bucket_day
        )
        SELECT {bucket} AS bucket,
               {tag_expr} AS tag,
               {unit_expr} AS unit,
               SUM(
                 CASE
                   WHEN pwi.period_start IS NULL AND pwi.period_end IS NULL THEN
                     CASE WHEN days.bucket_day = %s::date THEN COALESCE(pwi.planned_volume, 0) ELSE 0 END
                   ELSE COALESCE(pwi.planned_volume, 0) / GREATEST(1, (pwi.period_end - pwi.period_start + 1))::numeric
                 END
               )::numeric AS volume
        FROM planned_work_items pwi
        JOIN work_types wt ON wt.id = pwi.work_type_id
        JOIN objects o ON o.id = pwi.object_id
        LEFT JOIN object_types ot ON ot.id = o.object_type_id
        {_plan_section_join('pwi')}
        JOIN days ON (
          (pwi.period_start IS NULL AND pwi.period_end IS NULL AND days.bucket_day = %s::date)
          OR (pwi.period_start IS NOT NULL AND pwi.period_end IS NOT NULL AND days.bucket_day BETWEEN pwi.period_start AND pwi.period_end)
        )
        WHERE {' AND '.join(where)}
        GROUP BY {bucket}, {tag_expr}, {unit_expr}
        ORDER BY bucket, tag, unit
        """,
        [d_from, d_to, d_from, d_from, *params],
    )


def _plan_bucket_detail_rows(d_from: date_cls, d_to: date_cls, filters: dict[str, Any], granularity: str) -> list[dict]:
    where, params, tag_expr = _plan_where(d_from, d_to, filters, period=True)
    unit_expr = "COALESCE(NULLIF(pwi.unit, ''), wt.default_unit, '—')"
    bucket = _bucket_expr(granularity, "bucket_day")
    return query(
        f"""
        WITH days AS (
          SELECT generate_series(%s::date, %s::date, interval '1 day')::date AS bucket_day
        )
        SELECT {bucket} AS bucket,
               wt.code AS work_code,
               {tag_expr} AS tag,
               {unit_expr} AS unit,
               COALESCE(ot.code, '') AS object_type_code,
               o.id::text AS object_id,
               COALESCE(NULLIF(o.object_code, ''), '') AS object_code,
               cs_eff.code AS raw_section_code,
               SUM(
                 CASE
                   WHEN pwi.period_start IS NULL AND pwi.period_end IS NULL THEN
                     CASE WHEN days.bucket_day = %s::date THEN COALESCE(pwi.planned_volume, 0) ELSE 0 END
                   ELSE COALESCE(pwi.planned_volume, 0) / GREATEST(1, (pwi.period_end - pwi.period_start + 1))::numeric
                 END
               )::numeric AS volume
        FROM planned_work_items pwi
        JOIN work_types wt ON wt.id = pwi.work_type_id
        JOIN objects o ON o.id = pwi.object_id
        LEFT JOIN object_types ot ON ot.id = o.object_type_id
        {_plan_section_join('pwi')}
        JOIN days ON (
          (pwi.period_start IS NULL AND pwi.period_end IS NULL AND days.bucket_day = %s::date)
          OR (pwi.period_start IS NOT NULL AND pwi.period_end IS NOT NULL AND days.bucket_day BETWEEN pwi.period_start AND pwi.period_end)
        )
        WHERE {' AND '.join(where)}
        GROUP BY {bucket}, wt.code, {tag_expr}, {unit_expr}, ot.code, o.id, o.object_code, cs_eff.code
        ORDER BY bucket, wt.code, object_type_code, raw_section_code
        """,
        [d_from, d_to, d_from, d_from, *params],
    )


def _fact_bucket_detail_rows(d_from: date_cls, d_to: date_cls, filters: dict[str, Any], granularity: str) -> list[dict]:
    where, params = _fact_where(d_from, d_to, filters, before=False)
    bucket = _bucket_expr(granularity, "dwi.report_date")
    object_type_expr = _fact_object_type_sql()
    object_id_expr = _fact_object_id_sql()
    seg_volume_expr = _analytics_segment_volume_expr("seg")
    return query(
        f"""
        SELECT {bucket} AS bucket,
               wt.code AS work_code,
               {object_type_expr} AS object_type_code,
               COALESCE({object_id_expr}, '') AS object_id,
               COALESCE(NULLIF(o.object_code, ''), NULLIF(pf.field_code, ''), '') AS object_code,
               cs_eff.code AS raw_section_code,
               SUM({seg_volume_expr})::numeric AS volume
        FROM daily_work_items dwi
        JOIN work_types wt ON wt.id = dwi.work_type_id
        LEFT JOIN objects o ON o.id = dwi.object_id
        LEFT JOIN object_types ot ON ot.id = o.object_type_id
        {_effective_section_join('seg')}
        LEFT JOIN pile_fields pf ON pf.id = seg.pile_field_id
        WHERE {' AND '.join(where)}
        GROUP BY {bucket}, wt.code, {object_type_expr}, COALESCE({object_id_expr}, ''), COALESCE(NULLIF(o.object_code, ''), NULLIF(pf.field_code, ''), ''), cs_eff.code
        ORDER BY bucket, wt.code, object_type_code, raw_section_code
        """,
        params,
    )


def _work_scope_matchers(work_nodes: Iterable[dict[str, Any]], pipe_codes: list[str]) -> dict[str, tuple[set[tuple[str, str]], dict[str, Any]]]:
    matchers: dict[str, tuple[set[tuple[str, str]], dict[str, Any]]] = {}
    for work in work_nodes:
        pairs: set[tuple[str, str]] = set()
        for scope in work.get("scopes") or []:
            for work_codes, object_types in _scope_pairs(scope, pipe_codes):
                for work_code in work_codes:
                    for object_type in object_types:
                        pairs.add((str(work_code), str(object_type)))
        matchers[str(work.get("id"))] = (pairs, work)
    return matchers


def _normal_bucket(value: object) -> str:
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def _normal_section(value: object) -> str:
    raw = str(value or "").strip()
    return _merge_section(raw) or raw or "—"


def _mstroy_timeseries_payload(
    d_from: date_cls,
    d_to: date_cls,
    filters: dict[str, Any],
    *,
    granularity: str,
    mode: str,
    grouping: str,
) -> dict[str, Any]:
    static = _mstroy_summary_scope_static()
    allowed_work_tokens = set(filters.get("summary_work_tokens") or [])
    work_nodes = [work for work in (static.get("work_nodes") or []) if not allowed_work_tokens or str(work.get("id")) in allowed_work_tokens]
    pipe_codes = _mstroy_pipe_work_codes() if any(scope.get("distinct_work_code_count") for scope in static["all_scopes"]) else []
    matchers = _work_scope_matchers(work_nodes, pipe_codes)
    work_order = {token: idx for idx, token in enumerate(matchers)}
    fact_by_bucket_work: dict[tuple[str, str], float] = defaultdict(float)
    fact_by_bucket_section_work: dict[tuple[str, str, str], float] = defaultdict(float)
    total_by_work: dict[str, float] = defaultdict(float)
    total_by_section_work: dict[tuple[str, str], float] = defaultdict(float)

    for row in _fact_bucket_detail_rows(d_from, d_to, filters, granularity):
        bucket = _normal_bucket(row.get("bucket"))
        work_code = str(row.get("work_code") or "")
        object_type = str(row.get("object_type_code") or "")
        section = _normal_section(row.get("raw_section_code"))
        volume = _round_value(row.get("volume")) or 0.0
        if not volume:
            continue
        for token, (pairs, _work) in matchers.items():
            if (work_code, object_type) not in pairs:
                continue
            fact_by_bucket_work[(bucket, token)] += volume
            fact_by_bucket_section_work[(bucket, section, token)] += volume
            total_by_work[token] += volume
            total_by_section_work[(section, token)] += volume

    periods = _periods(d_from, d_to, granularity)
    if grouping == "sections" and filters.get("section_codes"):
        series_keys = sorted(
            [key for key, total in total_by_section_work.items() if abs(total) > 0.0005],
            key=lambda key: (_section_sort_key(key[0]), work_order.get(key[1], 9999)),
        )
        series = []
        for section, token in series_keys:
            work = matchers.get(token, (set(), {}))[1]
            series.append({
                "id": f"section:{section}:work:{token}",
                "label": f"{_section_label(section)} / {work.get('label') or token}",
                "unit": work.get("unit"),
                "groupId": token,
            })
        running: dict[tuple[str, str], float] = defaultdict(float)
        points = []
        for period in periods:
            values: dict[str, Optional[float]] = {}
            for section, token in series_keys:
                value = round(fact_by_bucket_section_work.get((period["key"], section, token), 0.0), 1)
                series_id = f"section:{section}:work:{token}"
                if mode == "cumulative":
                    running[(section, token)] += value
                    values[series_id] = round(running[(section, token)], 1) if running[(section, token)] else None
                else:
                    values[series_id] = value if value else None
            points.append({"date": period["key"], "values": values})
        return {"series": series, "points": points}

    series_tokens = [str(work.get("id")) for work in work_nodes if abs(total_by_work.get(str(work.get("id")), 0.0)) > 0.0005]
    series = []
    for token in series_tokens:
        work = matchers.get(token, (set(), {}))[1]
        series.append({"id": f"work:{token}", "label": str(work.get("label") or token), "unit": work.get("unit"), "groupId": token})
    running_by_work: dict[str, float] = defaultdict(float)
    points = []
    for period in periods:
        values: dict[str, Optional[float]] = {}
        for token in series_tokens:
            value = round(fact_by_bucket_work.get((period["key"], token), 0.0), 1)
            series_id = f"work:{token}"
            if mode == "cumulative":
                running_by_work[token] += value
                values[series_id] = round(running_by_work[token], 1) if running_by_work[token] else None
            else:
                values[series_id] = value if value else None
        points.append({"date": period["key"], "values": values})
    return {"series": series, "points": points}


def _earthworks_timeseries_payload(
    d_from: date_cls,
    d_to: date_cls,
    filters: dict[str, Any],
    *,
    granularity: str,
    mode: str,
    grouping: str,
) -> dict[str, Any]:
    dynamics = _mstroy_timeseries_payload(d_from, d_to, filters, granularity=granularity, mode=mode, grouping=grouping)
    return {"updatedAt": _updated_at(), "planFact": None, "dynamics": dynamics}


@router.get("/timeseries")
def earthworks_timeseries(
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    section_ids: Optional[str] = Query(None),
    work_type_ids: Optional[str] = Query(None),
    contractor_ids: Optional[str] = Query(None),
    object_ids: Optional[str] = Query(None),
    metric_id: Optional[str] = Query(None),
    granularity: str = Query("day"),
    mode: str = Query("cumulative"),
    grouping: str = Query("objects"),
    allow_stale: bool = Query(True),
):
    if granularity not in SUPPORTED_GRANULARITIES:
        raise HTTPException(400, "granularity должен быть day, week или month")
    if mode not in SUPPORTED_MODES:
        raise HTTPException(400, "mode должен быть cumulative или interval")
    if grouping not in SUPPORTED_GROUPINGS:
        raise HTTPException(400, "grouping должен быть objects или sections")
    d_from, d_to = _parse_earthworks_range(date_from, date_to)
    filters = _filter_bundle(section_ids=section_ids, work_type_ids=work_type_ids, contractor_ids=contractor_ids, object_ids=object_ids, metric_id=metric_id)
    cache_params = _mstroy_dashboard_cache_params(
        "timeseries",
        date_from=d_from.isoformat(),
        date_to=d_to.isoformat(),
        section_ids=section_ids or "",
        work_type_ids=work_type_ids or "",
        contractor_ids=contractor_ids or "",
        object_ids=object_ids or "",
        metric_id=metric_id or "",
        granularity=granularity,
        mode=mode,
        grouping=grouping,
    )
    return _mstroy_dashboard_cached_response(
        endpoint="timeseries",
        d_from=d_from,
        d_to=d_to,
        cache_params=cache_params,
        allow_stale=allow_stale,
        build_fn=lambda: _earthworks_timeseries_payload(d_from, d_to, filters, granularity=granularity, mode=mode, grouping=grouping),
    )


def _matrix_cell_from_amounts(plan_amount: float, fact_amount: float) -> dict[str, Optional[float]]:
    plan = round(plan_amount, 1)
    fact = round(fact_amount, 1)
    return {
        "plan": plan or None,
        "fact": fact or None,
        "deviation": _deviation(fact, plan) if plan or fact else None,
        "completionPct": _completion(fact, plan),
    }


def _matrix_cell_from_optional(plan: Optional[float], fact: Optional[float]) -> dict[str, Optional[float]]:
    plan_value = _round_value(plan)
    fact_value = _round_value(fact)
    return {
        "plan": plan_value,
        "fact": fact_value,
        "deviation": _deviation(fact_value, plan_value),
        "completionPct": _completion(fact_value, plan_value),
    }


def _bucket_detail_maps(d_from: date_cls, d_to: date_cls, filters: dict[str, Any], granularity: str, detail_maps: dict[str, Any]) -> tuple[dict[tuple[str, DetailKey], float], dict[tuple[str, DetailKey], float]]:
    section_overrides = detail_maps.get("section_overrides") or {}
    fact_by_bucket_key: dict[tuple[str, DetailKey], float] = {}
    plan_by_bucket_key: dict[tuple[str, DetailKey], float] = {}
    for row in _fact_bucket_detail_rows(d_from, d_to, filters, granularity):
        bucket = _normal_bucket(row.get("bucket"))
        _add_to_map(fact_by_bucket_key, (bucket, _detail_key(row, section_overrides)), row.get("volume"))
    for row in _plan_bucket_detail_rows(d_from, d_to, filters, granularity):
        bucket = _normal_bucket(row.get("bucket"))
        _add_to_map(plan_by_bucket_key, (bucket, _detail_key(row, section_overrides)), row.get("volume"))
    return plan_by_bucket_key, fact_by_bucket_key


def _earthworks_matrix_payload(
    d_from: date_cls,
    d_to: date_cls,
    filters: dict[str, Any],
    *,
    granularity: str,
    grouping: str,
) -> dict[str, Any]:
    detail_maps = _build_detail_maps_cached(d_from, d_to, filters)
    periods = _periods(d_from, d_to, granularity)
    plan_by_bucket_key, fact_by_bucket_key = _bucket_detail_maps(d_from, d_to, filters, granularity, detail_maps)
    tree_rows = _build_mstroy_tree_rows(detail_maps, grouping=grouping, prefix="matrix", filters=filters, include_internal_keys=True)

    rows = []
    for tree_row in tree_rows:
        key_list = [tuple(key) for key in (tree_row.get("_detailKeys") or [])]
        if not key_list:
            continue
        cells = {}
        total_plan = 0.0
        total_fact = 0.0
        for period in periods:
            plan = round(sum(float(plan_by_bucket_key.get((period["key"], key), 0.0) or 0.0) for key in key_list), 1)
            fact = round(sum(float(fact_by_bucket_key.get((period["key"], key), 0.0) or 0.0) for key in key_list), 1)
            total_plan += plan
            total_fact += fact
            cells[period["key"]] = _matrix_cell_from_amounts(plan, fact)
        values = _sum_detail_values(key_list, detail_maps)
        rows.append({
            "id": tree_row["id"],
            "parentId": tree_row.get("parentId"),
            "label": tree_row.get("label"),
            "unit": tree_row.get("unit"),
            "level": tree_row.get("level", 0),
            "hasChildren": bool(tree_row.get("hasChildren")),
            "sourceIds": tree_row.get("sourceIds") or [],
            "cumulativeTotals": _matrix_cell_from_optional(values.get("plan_total"), values.get("fact_total")),
            "totals": _matrix_cell_from_amounts(total_plan, total_fact),
            "cells": cells,
        })
    parent_ids = {row.get("parentId") for row in rows if row.get("parentId")}
    for row in rows:
        row["hasChildren"] = row["id"] in parent_ids
    return {"updatedAt": _updated_at(), "periods": periods, "rows": rows}


@router.get("/matrix")
def earthworks_matrix(
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    section_ids: Optional[str] = Query(None),
    work_type_ids: Optional[str] = Query(None),
    contractor_ids: Optional[str] = Query(None),
    object_ids: Optional[str] = Query(None),
    metric_id: Optional[str] = Query(None),
    granularity: str = Query("week"),
    grouping: str = Query("objects"),
    allow_stale: bool = Query(True),
):
    if granularity not in SUPPORTED_GRANULARITIES:
        raise HTTPException(400, "granularity должен быть day, week или month")
    if grouping not in SUPPORTED_GROUPINGS:
        raise HTTPException(400, "grouping должен быть objects или sections")
    d_from, d_to = _parse_earthworks_range(date_from, date_to)
    filters = _filter_bundle(section_ids=section_ids, work_type_ids=work_type_ids, contractor_ids=contractor_ids, object_ids=object_ids, metric_id=metric_id)
    cache_params = _mstroy_dashboard_cache_params(
        "matrix",
        date_from=d_from.isoformat(),
        date_to=d_to.isoformat(),
        section_ids=section_ids or "",
        work_type_ids=work_type_ids or "",
        contractor_ids=contractor_ids or "",
        object_ids=object_ids or "",
        metric_id=metric_id or "",
        granularity=granularity,
        grouping=grouping,
    )
    return _mstroy_dashboard_cached_response(
        endpoint="matrix",
        d_from=d_from,
        d_to=d_to,
        cache_params=cache_params,
        allow_stale=allow_stale,
        build_fn=lambda: _earthworks_matrix_payload(d_from, d_to, filters, granularity=granularity, grouping=grouping),
    )
