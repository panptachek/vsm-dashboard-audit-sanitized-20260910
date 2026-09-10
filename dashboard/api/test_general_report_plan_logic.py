from __future__ import annotations

from datetime import date
from unittest import mock

from openpyxl import Workbook

import wip_routes


START_OF_JUNE_2026 = date(2026, 6, 1)
END_OF_JUNE_2026 = date(2026, 6, 30)
REPORT_DATE = date(2026, 6, 16)


def _plan_slot(section_code: str, amount: float) -> dict[str, dict[str, dict[str, float]]]:
    slot = wip_routes._new_work_fact_slot()
    slot["total"][section_code] = amount
    return {"VYEMKA": slot}


def test_general_volume_plan_periods_use_full_month_and_derive_day_and_week_from_month() -> None:
    calls: list[tuple[date, date, list[str] | None, str]] = []

    def fake_query_general_work_plans(
        d_from: date,
        d_to: date,
        codes: list[str] | None,
        *,
        plan_scope: str = "all",
    ):
        calls.append((d_from, d_to, codes, plan_scope))
        if (d_from, d_to, plan_scope) == (START_OF_JUNE_2026, END_OF_JUNE_2026, "dated_only"):
            return _plan_slot("UCH_2", 19500.0)
        if (d_from, d_to, plan_scope) == (wip_routes.PROJECT_START_DATE, REPORT_DATE, "dated_only"):
            return _plan_slot("UCH_2", 10400.0)
        if (d_from, d_to, plan_scope) == (wip_routes.PROJECT_START_DATE, REPORT_DATE, "total_only"):
            return _plan_slot("UCH_2", 33000.0)
        raise AssertionError(f"Unexpected plan query window: {(d_from, d_to, plan_scope)}")

    with mock.patch.object(wip_routes, "_query_general_work_plans", side_effect=fake_query_general_work_plans):
        plan_periods = wip_routes._query_general_volume_plan_periods(REPORT_DATE, ["UCH_2"])

    assert calls == [
        (START_OF_JUNE_2026, END_OF_JUNE_2026, ["UCH_2"], "dated_only"),
        (wip_routes.PROJECT_START_DATE, REPORT_DATE, ["UCH_2"], "dated_only"),
        (wip_routes.PROJECT_START_DATE, REPORT_DATE, ["UCH_2"], "total_only"),
    ]
    assert plan_periods["month"]["VYEMKA"]["total"]["UCH_2"] == 19500.0
    assert plan_periods["week"]["VYEMKA"]["total"]["UCH_2"] == 4875.0
    assert plan_periods["day"]["VYEMKA"]["total"]["UCH_2"] == 650.0
    assert plan_periods["total"]["VYEMKA"]["total"]["UCH_2"] == 33000.0


def test_query_general_work_plans_sql_uses_object_type_boundary_routing_and_main_track_fallback() -> None:
    captured_sql: list[str] = []
    captured_params: list[list[object]] = []

    def fake_query(sql, params):
        captured_sql.append(sql)
        captured_params.append(list(params))
        return []

    with mock.patch.object(wip_routes, "query_one", return_value={"exists": True}), mock.patch.object(
        wip_routes, "query", side_effect=fake_query
    ):
        result = wip_routes._query_general_work_plans(START_OF_JUNE_2026, END_OF_JUNE_2026, ["UCH_2"])

    assert result == {}
    assert len(captured_sql) == 1
    sql = captured_sql[0]
    params = captured_params[0]
    assert "main_track_section_map(object_code, section_code)" in sql
    assert "('MAIN_002', 'UCH_2')" in sql
    assert "WHEN ot.code = 'TEMP_ROAD' THEN 'temp_roads'" in sql
    assert (
        "WHEN ot.code IN ('PILE_FIELD', 'PILE_DRIVING_PAD', 'ISSO_ACCESS', 'PIPE', 'BRIDGE', 'OVERPASS', 'TECH_ROAD') THEN 'piles_pipes'"
        in sql
    )
    assert "COALESCE(cs_override.code, sec_plan.section_code, sec_object.section_code, cs_mt.code) = ANY(%s)" in sql
    assert "WHERE pwi.pk_start IS NOT NULL" in sql
    assert "LEAST(GREATEST(pwi.pk_start, pwi.pk_end), GREATEST(csv.pk_start, csv.pk_end))" in sql
    assert params[4] == ["UCH_2"]


def test_general_effective_work_category_routes_pile_access_earthworks_into_pile_bucket() -> None:
    rules = wip_routes._normalize_general_report_work_rows(None)

    assert (
        wip_routes._general_effective_work_category(
            rules,
            "Выемка",
            "EARTH_EXCAVATION",
            "PILE_DRIVING_PAD",
            "PFP_001",
        )
        == "VYEMKA_PILE"
    )
    assert (
        wip_routes._general_effective_work_category(
            rules,
            "Выемка",
            "EARTH_EXCAVATION",
            "TECH_ROAD",
            "PFT_001",
        )
        == "VYEMKA_PILE"
    )
    assert (
        wip_routes._general_effective_work_category(
            rules,
            "Выемка",
            "EARTH_EXCAVATION",
            "TECH_ROAD",
            "TR_001",
        )
        == "VYEMKA"
    )
    assert (
        wip_routes._general_effective_work_category(
            rules,
            "Песок (Насыпь)",
            "EMBANKMENT_CONSTRUCTION",
            "PILE_DRIVING_PAD",
            "PFP_001",
        )
        == "SAND_PILE"
    )
    assert (
        wip_routes._general_effective_work_category(
            rules,
            "Песок (Насыпь)",
            "EMBANKMENT_CONSTRUCTION",
            "TECH_ROAD",
            "PFT_001",
        )
        == "SAND_PILE"
    )
    assert (
        wip_routes._general_effective_work_category(
            rules,
            "Песок (Насыпь)",
            "EMBANKMENT_CONSTRUCTION",
            "TECH_ROAD",
            "TR_001",
        )
        == "SAND_VAD"
    )


def test_general_plan_category_routes_zs2_to_sand_without_changing_fact_category() -> None:
    rules = wip_routes._normalize_general_report_work_rows(None)

    assert wip_routes._general_effective_work_category(rules, "ПГС", "ZS2", "MAIN_TRACK", "MAIN_006") == "SHPS"
    assert wip_routes._general_plan_work_category(rules, "ПГС", "ZS2", "MAIN_TRACK", "MAIN_006") == "SAND_OH"
    assert (
        wip_routes._general_plan_work_category(
            rules,
            "ПГС",
            "SECOND_PROTECTIVE_LAYER",
            "PILE_DRIVING_PAD",
            "PFP_001",
        )
        == "SAND_PILE"
    )
    assert wip_routes._general_plan_work_category(rules, "ПГС", "ZS2", "TEMP_ROAD", "TR_001") == "SAND_VAD"
    assert wip_routes._general_plan_work_category(rules, "ЩПГС", "ZS1", "MAIN_TRACK", "MAIN_006") == "SHPS"
    assert "ZS2" in wip_routes._ready_general_volume_tag_text("SAND_OH", "plan")
    assert "ZS2" not in wip_routes._ready_general_volume_tag_text("SHPS", "plan")
    assert "ZS2" in wip_routes._ready_general_volume_tag_text("SHPS", "fact")


def test_query_general_work_plans_routes_zs2_amount_to_sand_oh() -> None:
    rows = [
        {
            "section_code": "UCH_6",
            "wt_code": "ZS2",
            "object_type_code": "MAIN_TRACK",
            "object_code": "MAIN_006",
            "analytics_tag": "ПГС",
            "amount": 1250.0,
        }
    ]

    with mock.patch.object(wip_routes, "query_one", return_value={"exists": True}), mock.patch.object(
        wip_routes, "query", return_value=rows
    ), mock.patch.object(
        wip_routes,
        "_general_report_work_rules",
        return_value=wip_routes._normalize_general_report_work_rows(None),
    ):
        plans = wip_routes._query_general_work_plans(START_OF_JUNE_2026, END_OF_JUNE_2026, None)

    assert plans["SAND_OH"]["total"]["UCH_6"] == 1250.0
    assert "SHPS" not in plans


def test_general_pile_plan_period_uses_full_week_and_month_horizons() -> None:
    assert wip_routes._general_pile_plan_period(
        "week",
        date(2026, 6, 15),
        date(2026, 6, 16),
    ) == (date(2026, 6, 15), date(2026, 6, 21))
    assert wip_routes._general_pile_plan_period(
        "month",
        START_OF_JUNE_2026,
        REPORT_DATE,
    ) == (START_OF_JUNE_2026, END_OF_JUNE_2026)


def test_statement_general_plan_rows_keep_month_plan_and_override_cumulative_with_total_only_rows() -> None:
    days = [{"date": date(2026, 6, day).isoformat()} for day in range(1, 31)]
    dated_row = {
        "object_id": "obj-1",
        "object_code": "TRACTION_SUBSTATION_2905",
        "object_name": "Тяговая подстанция Валдай ПК2905",
        "object_type_code": "TECH_ROAD",
        "object_type_name": "Тяговая подстанция",
        "section_code": "UCH_3",
        "section_name": "Участок 3",
        "work_type_id": "wt-1",
        "work_type_code": "TOPSOIL_STRIPPING",
        "work_type_name": "Снятие ПРС",
        "analytics_tag": "Снятие ПРС",
        "unit": "м3",
        "section_pk_start": 290500.0,
        "section_pk_end": 290500.0,
        "object_pk_start": 290500.0,
        "object_pk_end": 290500.0,
        "plan_pk_start": 290500.0,
        "plan_pk_end": 290500.0,
        "pk_start": 290500.0,
        "pk_end": 290500.0,
        "pk_raw_text": "ПК2905",
        "period_start": date(2026, 6, 1),
        "period_end": date(2026, 6, 30),
        "plan_type": "manual_period",
        "planned_volume": 960.0,
    }
    total_only_row = dict(dated_row)
    total_only_row.update({
        "period_start": None,
        "period_end": None,
        "plan_type": "total",
        "planned_volume": 6365.0,
    })

    rows_by_key: dict[str, dict] = {}

    with mock.patch.object(wip_routes, "_statement_table_exists", return_value=True), mock.patch.object(
        wip_routes, "_statement_general_plan_scope_filters", return_value=([], [])
    ), mock.patch.object(wip_routes, "_statement_scope_cte", return_value=""), mock.patch.object(
        wip_routes, "query", return_value=[dated_row, total_only_row]
    ), mock.patch.object(
        wip_routes,
        "_general_plan_work_category",
        return_value="PRS",
    ), mock.patch.object(
        wip_routes,
        "_statement_clip_volume_to_section",
        side_effect=lambda item, volume, *_args: (item.get("pk_start"), item.get("pk_end"), volume, True),
    ):
        wip_routes._statement_general_add_plan_rows(
            rows_by_key,
            days=days,
            period_start=date(2026, 6, 1),
            period_end=date(2026, 6, 30),
            section_codes=None,
            object_type=None,
            object_id=None,
            q=None,
        )

    assert len(rows_by_key) == 1
    row = next(iter(rows_by_key.values()))
    assert row["month_plan"] == 960.0
    assert row["cumulative_plan"] == 6365.0
    assert row["day_values"][date(2026, 6, 1).isoformat()]["plan_volume"] == 32.0


def test_statement_general_dyntest_rows_are_object_level_and_keep_pk_search_refs() -> None:
    days = [{"date": date(2026, 8, 1).isoformat()}]
    rows_by_key: dict[str, dict] = {}

    first = wip_routes._statement_general_row_seed(
        rows_by_key,
        days=days,
        category="work:PILE_DYNTEST",
        category_label="Динамические испытания свай",
        section_code="UCH_2",
        section_name="Участок 2",
        object_id="obj-1",
        object_code="PF_25Н",
        object_name="Поле свай 25Н",
        object_type_code="PILE_FIELD",
        object_type_name="Свайное поле",
        work_type_id="wt-dyn",
        work_type_code="PILE_DYNTEST",
        work_type_name="Динамические испытания свай",
        unit="шт",
        pk_start=275500.0,
        pk_end=275500.0,
        pk_raw_text="ПК2755",
        pile_length_m=37.0,
    )
    second = wip_routes._statement_general_row_seed(
        rows_by_key,
        days=days,
        category="work:PILE_DYNTEST",
        category_label="Динамические испытания свай",
        section_code="UCH_2",
        section_name="Участок 2",
        object_id="obj-1",
        object_code="PF_25Н",
        object_name="Поле свай 25Н",
        object_type_code="PILE_FIELD",
        object_type_name="Свайное поле",
        work_type_id="wt-dyn",
        work_type_code="PILE_DYNTEST",
        work_type_name="Динамические испытания свай",
        unit="шт",
        pk_start=276000.0,
        pk_end=276000.0,
        pk_raw_text="ПК2760",
        pile_length_m=21.0,
    )

    assert first is second
    assert len(rows_by_key) == 1
    assert first["pk_start"] is None
    assert first["pk_end"] is None
    assert first["pk_raw_text"] is None
    assert first["pile_length_m"] is None

    found = wip_routes._statement_general_find_existing_row_for_fact(
        rows_by_key,
        category="work:PILE_DYNTEST",
        section_code="UCH_2",
        object_id="obj-1",
        work_type_id="wt-dyn",
        unit="шт",
        pk_start=275800.0,
        pk_end=275800.0,
        pile_length_m=16.0,
        work_type_code="PILE_DYNTEST",
    )
    assert found is first

    first["project_refs"] = [{"pk_start": 275500.0, "pk_end": 275500.0}]
    assert wip_routes._statement_general_row_overlaps_pk(first, 275500.0, 275500.0)


def test_statement_general_plan_rows_sql_skips_generated_pile_total_mirrors() -> None:
    captured_sql: list[str] = []

    def fake_query(sql, params):
        captured_sql.append(sql)
        return []

    with mock.patch.object(wip_routes, "_statement_table_exists", return_value=True), mock.patch.object(
        wip_routes, "_statement_general_plan_scope_filters", return_value=([], [])
    ), mock.patch.object(wip_routes, "_statement_scope_cte", return_value=""), mock.patch.object(
        wip_routes, "query", side_effect=fake_query
    ):
        wip_routes._statement_general_add_plan_rows(
            {},
            days=[],
            period_start=date(2026, 8, 1),
            period_end=date(2026, 8, 31),
            section_codes=None,
            object_type=None,
            object_id=None,
            q=None,
        )

    assert len(captured_sql) == 1
    sql = captured_sql[0]
    assert "PILE_MAIN" in sql
    assert "PILE_TRIAL" in sql
    assert "PILE_DYNTEST" in sql
    assert "PILE_HEADCAP_INSTALLATION" in sql
    assert "pwi.source_reference = 'project_work_items:total'" in sql
    assert "NOT (wt.code = ANY" in sql


def test_query_general_equipment_counts_dedupes_units_and_uses_master_equipment_type() -> None:
    rows = [
        {
            "section_code": "UCH_1",
            "shift": "day",
            "equipment_type": "кдм камаз",
            "equipment_name": "Самосвал",
            "brand_model": "FAW J6 6*4",
            "plate_number": "С036МЕ27",
            "unit_number": "609",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "eu-609",
        },
        {
            "section_code": "UCH_1",
            "shift": "day",
            "equipment_type": "Самосвал",
            "equipment_name": "Самосвал",
            "brand_model": "FAW J6 6*4",
            "plate_number": "С036МЕ27",
            "unit_number": "609",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "eu-609",
        },
        {
            "section_code": "UCH_1",
            "shift": "day",
            "equipment_type": "самосвал",
            "equipment_name": "Самосвал",
            "brand_model": "Shacman",
            "plate_number": "А001АА69",
            "unit_number": "610",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "eu-610",
        },
        {
            "section_code": "UCH_1",
            "shift": "day",
            "equipment_type": "самосвал",
            "equipment_name": "Самосвал",
            "brand_model": "Shacman",
            "plate_number": "А777АА69",
            "unit_number": "701",
            "ownership_type": "hired",
            "contractor_name": "ООО Наем",
            "unit_key": "eu-701",
        },
    ]

    with mock.patch.object(wip_routes, "query", return_value=rows):
        result = wip_routes._query_general_equipment_counts(REPORT_DATE, ["UCH_1"])

    dump = result["dump_truck"]
    assert dump["day"]["UCH_1"] == 2.0
    assert dump["day_non_hired"]["UCH_1"] == 2.0
    assert dump["total"]["UCH_1"] == 2.0
    assert dump["hired"]["UCH_1"] == 1.0
    assert dump["day_hired"]["UCH_1"] == 1.0


def test_general_dump_truck_sand_productivity_counts_only_non_hired_and_dedupes_units() -> None:
    eq_rows = [
        {
            "section_code": "UCH_1",
            "report_date": REPORT_DATE,
            "shift": "day",
            "equipment_type": "кдм камаз",
            "equipment_name": "Самосвал",
            "brand_model": "FAW J6 6*4",
            "plate_number": "С036МЕ27",
            "unit_number": "609",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "eu-609",
        },
        {
            "section_code": "UCH_1",
            "report_date": REPORT_DATE,
            "shift": "day",
            "equipment_type": "Самосвал",
            "equipment_name": "Самосвал",
            "brand_model": "FAW J6 6*4",
            "plate_number": "С036МЕ27",
            "unit_number": "609",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "eu-609",
        },
        {
            "section_code": "UCH_1",
            "report_date": REPORT_DATE,
            "shift": "day",
            "equipment_type": "самосвал",
            "equipment_name": "Самосвал",
            "brand_model": "Shacman",
            "plate_number": "А777АА69",
            "unit_number": "701",
            "ownership_type": "hired",
            "contractor_name": "ООО Наем",
            "unit_key": "eu-701",
        },
        {
            "section_code": "UCH_1",
            "report_date": REPORT_DATE,
            "shift": "day",
            "equipment_type": "самосвал",
            "equipment_name": "Самосвал",
            "brand_model": "Shacman",
            "plate_number": "С999МЕ27",
            "unit_number": "702",
            "ownership_type": "hired",
            "contractor_name": "АЛМАЗ",
            "unit_key": "eu-702",
        },
        {
            "section_code": "UCH_1",
            "report_date": REPORT_DATE,
            "shift": "night",
            "equipment_type": "самосвал",
            "equipment_name": "Самосвал",
            "brand_model": "Shacman",
            "plate_number": "А001АА69",
            "unit_number": "610",
            "ownership_type": "own",
            "contractor_name": "ЖДС",
            "unit_key": "eu-610",
        },
    ]
    movement_rows = [
        {
            "section_code": "UCH_1",
            "report_date": REPORT_DATE,
            "shift": "day",
            "movement_type": "pit_to_constructive",
            "contractor_name": "ЖДС",
            "volume": 100.0,
        },
        {
            "section_code": "UCH_1",
            "report_date": REPORT_DATE,
            "shift": "day",
            "movement_type": "pit_to_constructive",
            "contractor_name": "АЛМАЗ",
            "volume": 100.0,
        },
        {
            "section_code": "UCH_1",
            "report_date": REPORT_DATE,
            "shift": "night",
            "movement_type": "pit_to_constructive",
            "contractor_name": "ЖДС",
            "volume": 100.0,
        },
    ]

    with mock.patch.object(wip_routes, "query", side_effect=[eq_rows, movement_rows]), mock.patch.object(
        wip_routes,
        "DUMP_TRUCK_NORMS",
        {(1, "SAND"): {"include_pit_to_stockpile": True}},
    ):
        result = wip_routes._query_general_dump_truck_sand_productivity_plan_counts(REPORT_DATE, ["UCH_1"])

    assert result["day"]["UCH_1"] == 1.0
    assert result["night"]["UCH_1"] == 1.0
    assert result["total"]["UCH_1"] == 2.0


def test_equipment_productivity_dump_truck_plan_uses_general_non_hired_units() -> None:
    mm_rows = [
        {
            "section_code": "UCH_1",
            "report_date": REPORT_DATE,
            "shift": "day",
            "mtype": "pit_to_constructive",
            "material_code": "SAND",
            "contractor_name": "ЖДС",
            "volume": 100.0,
            "trips": 7,
        },
        {
            "section_code": "UCH_1",
            "report_date": REPORT_DATE,
            "shift": "day",
            "mtype": "pit_to_constructive",
            "material_code": "SAND",
            "contractor_name": "АЛМАЗ",
            "volume": 25.0,
            "trips": 2,
        },
    ]
    dump_units = {
        ("UCH_1", REPORT_DATE, "day"): {"eu-609", "eu-610"},
    }

    with mock.patch.object(wip_routes, "query", side_effect=[[], [], [], mm_rows]), mock.patch.object(
        wip_routes,
        "_query_general_dump_truck_non_hired_unit_keys_by_shift",
        return_value=dump_units,
    ), mock.patch.object(
        wip_routes,
        "DUMP_TRUCK_NORMS",
        {(1, "SAND"): {"trips": 9, "m3_per_trip": 16, "include_pit_to_stockpile": True, "quarry": "Карьер"}},
    ):
        result = wip_routes.equipment_productivity(REPORT_DATE.isoformat(), REPORT_DATE.isoformat(), "UCH_1")

    dump_row = next(row for row in result["rows"] if row["section_code"] == "UCH_1" and row["equipment_type"] == "самосвал")
    sand = next(item for item in dump_row["by_material"] if item["material"] == "SAND")

    assert sand["avg_units"] == 2.0
    assert sand["expected"] == 288.0
    assert sand["fact_in_norm"] == 100.0
    assert sand["fact_off_norm"] == 25.0


def test_equipment_productivity_work_facts_use_executor_report_section_sql() -> None:
    captured_sql: list[str] = []

    def fake_query(sql, params=None):
        captured_sql.append(sql)
        return []

    with mock.patch.object(wip_routes, "query", side_effect=fake_query), mock.patch.object(
        wip_routes,
        "_query_general_dump_truck_non_hired_unit_keys_by_shift",
        return_value={},
    ):
        result = wip_routes.equipment_productivity(REPORT_DATE.isoformat(), REPORT_DATE.isoformat(), "UCH_3")

    assert result["rows"] == []
    work_fact_sql = next(
        sql for sql in captured_sql
        if "SUM(dwi.volume)::numeric AS volume" in sql and "FROM daily_work_items dwi" in sql
    )
    work_unit_sql = next(sql for sql in captured_sql if "FROM work_item_equipment_usage wieu" in sql)
    for sql in (work_fact_sql, work_unit_sql):
        assert "JOIN daily_reports dr_work ON dr_work.id = dwi.daily_report_id" in sql
        assert "LEFT JOIN construction_sections cs ON cs.id = COALESCE(dwi.section_id, dr_work.section_id)" in sql
        assert "cs.code = ANY(%s)" in sql


def test_mechanization_units_work_usage_uses_executor_report_section_sql() -> None:
    captured_sql: list[str] = []

    def fake_query(sql, params=None):
        captured_sql.append(sql)
        return []

    empty_operation_metrics = {
        "by_unit": {},
        "by_unit_section": {},
        "by_section_productivity_type": {},
        "by_section_actual_type": {},
        "total": wip_routes._operation_metric_empty_totals(),
    }

    with mock.patch.object(wip_routes, "query", side_effect=fake_query), mock.patch.object(
        wip_routes,
        "_refresh_norms",
        return_value=None,
    ), mock.patch.object(
        wip_routes,
        "_query_equipment_operation_metrics",
        return_value=empty_operation_metrics,
    ):
        result = wip_routes._mechanization_units_impl(REPORT_DATE, REPORT_DATE, include_work_diagnostics=True)

    assert result["units"] == []
    work_usage_sql = next(sql for sql in captured_sql if "FROM work_item_equipment_usage wieu" in sql)
    assert "JOIN daily_reports dr_work ON dr_work.id = dwi.daily_report_id" in work_usage_sql
    assert "LEFT JOIN construction_sections cs ON cs.id = COALESCE(dwi.section_id, dr_work.section_id)" in work_usage_sql


def test_query_general_work_facts_routes_tp_local_soil_into_sand_tp_rows() -> None:
    rows = [
        {
            "section_code": "UCH_3",
            "shift": "day",
            "wt_code": "AREA_GRADING",
            "object_type_code": "TECH_ROAD",
            "object_code": "TRACTION_SUBSTATION_2905",
            "analytics_tag": "Грунт (площадка)",
            "labor_source_type": "own",
            "contractor_name": "ЖДС",
            "amount": 120.0,
        },
        {
            "section_code": "UCH_4",
            "shift": "night",
            "wt_code": "AREA_GRADING",
            "object_type_code": "TECH_ROAD",
            "object_code": "STATION_TRACK_3294",
            "analytics_tag": "Грунт (проезд)",
            "labor_source_type": "own",
            "contractor_name": "ЖДС",
            "amount": 80.0,
        },
    ]

    with mock.patch.object(wip_routes, "query", return_value=rows), mock.patch.object(
        wip_routes,
        "_general_report_work_rules",
        return_value=wip_routes._normalize_general_report_work_rows(None),
    ):
        facts = wip_routes._query_general_work_facts(REPORT_DATE, REPORT_DATE, None)

    assert facts["SOIL_LOCAL"]["total"]["UCH_3"] == 120.0
    assert facts["SOIL_LOCAL"]["total"]["UCH_4"] == 80.0
    assert facts["SAND"]["total"]["UCH_3"] == 120.0
    assert facts["SAND"]["total"]["UCH_4"] == 80.0
    assert facts["SAND_VALDAI"]["total"]["UCH_3"] == 120.0
    assert facts["SAND_VYPOLZOVO"]["total"]["UCH_4"] == 80.0


def test_fill_ready_general_equipment_writes_unique_daily_dump_truck_row() -> None:
    wb = Workbook()
    ws = wb.active
    ws.cell(row=55, column=2, value="Самосвал")
    ws.cell(row=55, column=3, value="Плановое количество самосвалов ДЕНЬ")
    ws.cell(row=56, column=3, value="Плановое количество самосвалов НОЧЬ")
    ws.cell(row=57, column=3, value="Фактическое количество самосвалов ЖДС СУТКИ")
    ws.cell(row=58, column=3, value="Фактическое количество самосвалов ЖДС ДЕНЬ")
    ws.cell(row=59, column=3, value="Фактическое количество самосвалов ЖДС НОЧЬ")
    ws.cell(row=60, column=3, value="% от норматива в сутки")
    ws.cell(row=61, column=3, value="Песок")
    ws.cell(row=61, column=4, value="Плановая выработка от кол-ва техники")
    ws.cell(row=62, column=4, value="План за день")
    ws.cell(row=63, column=4, value="План за ночь")
    ws.cell(row=64, column=4, value="Факт за сутки")
    ws.cell(row=65, column=4, value="Факт за день")
    ws.cell(row=66, column=4, value="Факт за ночь")

    count_slot = wip_routes._new_equipment_fact_slot()
    count_slot["total"]["UCH_1"] = 2.0
    count_slot["day"]["UCH_1"] = 1.0
    count_slot["night"]["UCH_1"] = 1.0

    with mock.patch.object(wip_routes, "_query_general_equipment_counts", return_value={"dump_truck": count_slot}), mock.patch.object(
        wip_routes,
        "_query_general_equipment_work_facts",
        return_value={},
    ), mock.patch.object(
        wip_routes,
        "_query_general_equipment_work_counts",
        return_value={},
    ), mock.patch.object(
        wip_routes,
        "_query_general_dump_truck_sand_productivity_plan_counts",
        return_value=wip_routes._new_equipment_fact_slot(),
    ), mock.patch.object(
        wip_routes,
        "_query_general_dump_truck_non_hired_unique_daily_total",
        return_value=2.0,
    ):
        wip_routes._fill_ready_general_equipment(ws, REPORT_DATE, ["UCH_1"], {}, {})

    assert wip_routes._find_ready_general_equipment_rows(ws)["dump_truck"]["counts"]["fact_total"] == 57
    assert wip_routes._read_section_values_at_row(ws, 57)["UCH_1"] == 2.0
    assert wip_routes._read_section_values_at_row(ws, 58)["UCH_1"] == 1.0
    assert wip_routes._read_section_values_at_row(ws, 59)["UCH_1"] == 1.0


def test_fill_ready_general_equipment_overrides_dump_truck_total_with_global_unique_daily_count() -> None:
    wb = Workbook()
    ws = wb.active
    ws.cell(row=55, column=2, value="Самосвал")
    ws.cell(row=55, column=3, value="Плановое количество самосвалов ДЕНЬ")
    ws.cell(row=56, column=3, value="Плановое количество самосвалов НОЧЬ")
    ws.cell(row=57, column=3, value="Фактическое количество самосвалов ЖДС СУТКИ")
    ws.cell(row=58, column=3, value="Фактическое количество самосвалов ЖДС ДЕНЬ")
    ws.cell(row=59, column=3, value="Фактическое количество самосвалов ЖДС НОЧЬ")
    ws.cell(row=60, column=3, value="% от норматива в сутки")
    ws.cell(row=61, column=3, value="Песок")
    ws.cell(row=61, column=4, value="Плановая выработка от кол-ва техники")
    ws.cell(row=62, column=4, value="План за день")
    ws.cell(row=63, column=4, value="План за ночь")
    ws.cell(row=64, column=4, value="Факт за сутки")
    ws.cell(row=65, column=4, value="Факт за день")
    ws.cell(row=66, column=4, value="Факт за ночь")

    count_slot = wip_routes._new_equipment_fact_slot()
    count_slot["total"]["UCH_1"] = 2.0
    count_slot["total"]["UCH_2"] = 1.0
    count_slot["day"]["UCH_1"] = 1.0
    count_slot["night"]["UCH_1"] = 1.0
    count_slot["night"]["UCH_2"] = 1.0

    with mock.patch.object(wip_routes, "_query_general_equipment_counts", return_value={"dump_truck": count_slot}), mock.patch.object(
        wip_routes,
        "_query_general_equipment_work_facts",
        return_value={},
    ), mock.patch.object(
        wip_routes,
        "_query_general_equipment_work_counts",
        return_value={},
    ), mock.patch.object(
        wip_routes,
        "_query_general_dump_truck_sand_productivity_plan_counts",
        return_value=wip_routes._new_equipment_fact_slot(),
    ), mock.patch.object(
        wip_routes,
        "_query_general_dump_truck_non_hired_unique_daily_total",
        return_value=2.0,
    ):
        wip_routes._fill_ready_general_equipment(ws, REPORT_DATE, ["UCH_1", "UCH_2"], {}, {})

    assert ws.cell(row=57, column=wip_routes.GENERAL_ANALYTICS_VALUE_START_COL).value == 2.0
    assert wip_routes._read_section_values_at_row(ws, 57)["UCH_1"] == 2.0
    assert wip_routes._read_section_values_at_row(ws, 57)["UCH_2"] == 1.0
