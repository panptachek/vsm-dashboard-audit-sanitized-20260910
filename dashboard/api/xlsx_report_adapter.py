"""Convert the VSM daily XLSX template into the deterministic text format.

The dashboard parser intentionally stays text-first. This adapter is a thin
compatibility layer for the approved workbook template:

    Отчет шаблон ВСЖМ (для текстовой части).xlsx

It performs no DB writes and does not interpret values beyond reshaping workbook
rows into the existing ``===Секция===`` report text contract.
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
import re
from typing import Any

from openpyxl import load_workbook


TEMPLATE_MARKER = "ДЕЖУРНЫЙ ОТЧЁТ ПО УЧАСТКУ"
COMPACT_TEMPLATE_MARKER = "VSM_COMPACT_REPORT_TEMPLATE_V1"
ONE_SHEET_TEMPLATE_MARKER = "VSM_DAILY_REPORT_ONE_SHEET_V2"


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\xa0", " ").strip()


def _has_value(*values: Any) -> bool:
    return any(_clean(value) and _clean(value).lower() not in {"н/д", "нд", "-"} for value in values)


def _is_example_sheet(ws) -> bool:
    return "пример" in ws.title.strip().lower()


def _sheet_score(ws) -> int:
    score = 0
    if TEMPLATE_MARKER.lower() in _clean(ws["A1"].value).lower():
        score += 100
    if _has_value(ws["A5"].value, ws["B5"].value, ws["C5"].value):
        score += 50
    data_ranges = (
        ("A", "L", 32, 76),
        ("A", "N", 80, 104),
        ("A", "M", 108, 122),
        ("A", "F", 126, 180),
        ("A", "E", 195, 204),
        ("A", "G", 208, 217),
    )
    for start_col, end_col, start_row, end_row in data_ranges:
        start_idx = ws[start_col + str(start_row)].column
        end_idx = ws[end_col + str(start_row)].column
        for row_idx in range(start_row, end_row + 1):
            if _has_value(*(ws.cell(row_idx, col_idx).value for col_idx in range(start_idx, end_idx + 1))):
                score += 1
    return score


def _content_score(ws) -> int:
    score = _sheet_score(ws)
    if TEMPLATE_MARKER.lower() in _clean(ws["A1"].value).lower():
        score -= 100
    return score


def _select_report_sheet(workbook):
    candidates = [
        ws
        for ws in workbook.worksheets
        if "справоч" not in ws.title.lower()
        and TEMPLATE_MARKER.lower() in _clean(ws["A1"].value).lower()
    ]
    if not candidates:
        candidates = [ws for ws in workbook.worksheets if "справоч" not in ws.title.lower()]
    if not candidates:
        return workbook.active
    filled = [ws for ws in candidates if _content_score(ws) > 0]
    active = workbook.active
    if active in filled and not _is_example_sheet(active):
        return active
    for ws in filled:
        if ws.title.strip().lower() == "шаблон":
            return ws
    non_example = [ws for ws in filled if not _is_example_sheet(ws)]
    if non_example:
        return max(non_example, key=_sheet_score)
    if filled:
        return max(filled, key=_sheet_score)
    return max(candidates, key=_sheet_score)


def _compact_header_sheet(workbook):
    for ws in workbook.worksheets:
        if COMPACT_TEMPLATE_MARKER.lower() in _clean(ws["A1"].value).lower():
            return ws
    return None


def _one_sheet_report_sheets(workbook):
    candidates = [
        ws
        for ws in workbook.worksheets
        if ONE_SHEET_TEMPLATE_MARKER.lower() in _clean(ws["A1"].value).lower()
    ]
    if not candidates:
        return []

    # MAX reports sometimes arrive in a workbook where the visible sheet named
    # "Отчет" is still an empty template, while the user-filled report was left on
    # a sheet named "Пример заполнения". The stock template example has the fixed
    # seed date 08.05.2026 and must be ignored; any other dated, filled example
    # sheet is a real submitted report and should be parsed. This is a generic
    # content/date rule, not a per-day workaround.
    scored: list[tuple[int, Any]] = []
    for ws in candidates:
        if _is_example_sheet(ws) and _normal_date_text(ws["B4"].value) == "08.05.2026":
            continue
        score = _one_sheet_business_score(ws)
        if score > 0:
            scored.append((score, ws))
    if not scored:
        return []
    non_examples = [ws for _score, ws in scored if not _is_example_sheet(ws)]
    if non_examples:
        return non_examples
    return [ws for _score, ws in scored]


def _normal_date_text(value: Any) -> str:
    text = _clean(value)
    m = re.match(r"^(\d{1,2})[./-](\d{1,2})(?:[./-](\d{2,4}))?$", text)
    if not m:
        return text
    day = int(m.group(1))
    month = int(m.group(2))
    year = int(m.group(3) or 0)
    if year and year < 100:
        year += 2000
    return f"{day:02d}.{month:02d}.{year:04d}" if year else f"{day:02d}.{month:02d}"


def _one_sheet_business_score(ws) -> int:
    # A real one-sheet report must at least have the three business header values.
    # Empty template sheets still contain all markers/table headers and otherwise
    # look superficially report-like.
    if not (_clean(ws["B4"].value) and _clean(ws["B5"].value) and _clean(ws["B6"].value)):
        return 0
    score = 15
    header_labels = {
        "дата", "смена", "участок", "направление", "конструктивы",
        "источник", "материал", "объем", "ед.", "откуда", "куда",
        "номер техники", "номер/число техники", "число техники", "рейсов", "организация", "объект/конструктив",
        "работа", "пк всжм", "пк ад", "тип техники", "производительность",
        "комментарий", "модель", "госномер", "статус", "категория",
        "количество", "автодорога", "дорога", "пионерка", "не в отметку",
        "дсо", "готово под щпгс", "под аб", "прс", "зс №2", "зс №1", "асфальтобетон", "накопитель", "пк",
        "свайное поле", "тип свай", "вид/тип свай", "проблема", "плечо", "плечо, км", "плечо возки, км",
        "время работы", "время работы, ч",
    }
    for row_idx in range(11, min(ws.max_row, 260) + 1):
        values = [_clean(ws.cell(row_idx, col_idx).value) for col_idx in range(1, min(ws.max_column, 12) + 1)]
        values = [value for value in values if value]
        if not values:
            continue
        first = values[0].strip().lower().replace("ё", "е")
        if first.startswith("##"):
            continue
        normalized = [value.lower().replace("ё", "е") for value in values]
        if normalized and all(value in header_labels for value in normalized):
            continue
        score += 1
    return score


def _one_sheet_report_sheet(workbook):
    non_examples = _one_sheet_report_sheets(workbook)
    if not non_examples:
        return None
    active = workbook.active
    if active in non_examples:
        return active
    for ws in non_examples:
        if ws.title.strip().lower() == "отчет":
            return ws
    return non_examples[0]


def _sheet(workbook, name: str):
    for ws in workbook.worksheets:
        if ws.title.strip().lower() == name.lower():
            return ws
    return None


def _iter_filled_rows(ws, min_row: int = 3, max_row: int | None = None):
    if ws is None:
        return
    last_row = max_row or ws.max_row
    for row in ws.iter_rows(min_row=min_row, max_row=last_row, values_only=True):
        if _has_value(*row):
            yield row


def _one_sheet_marker_rows(ws) -> dict[str, tuple[int, int]]:
    markers: list[tuple[str, int]] = []
    for row_idx in range(1, ws.max_row + 1):
        value = _clean(ws.cell(row_idx, 1).value)
        if value.startswith("##"):
            markers.append((value.lstrip("#").strip().upper(), row_idx))
    result: dict[str, tuple[int, int]] = {}
    for idx, (name, row_idx) in enumerate(markers):
        end_row = (markers[idx + 1][1] - 1) if idx + 1 < len(markers) else ws.max_row
        result[name] = (row_idx, end_row)
    return result


def _one_sheet_table(ws, markers: dict[str, tuple[int, int]], name: str) -> list[dict[str, str]]:
    bounds = markers.get(name.upper())
    if not bounds:
        return []
    marker_row, end_row = bounds
    header_row = marker_row + 1
    headers = [_clean(cell.value) for cell in ws[header_row]]
    rows: list[dict[str, str]] = []
    for row_idx in range(header_row + 1, end_row + 1):
        values = [_clean(cell.value) for cell in ws[row_idx]]
        if not _has_value(*values):
            continue
        row = {header: values[idx] for idx, header in enumerate(headers) if header}
        rows.append(row)
    return rows


def _row_value(row: dict[str, str], *names: str) -> str:
    for name in names:
        value = _clean(row.get(name))
        if value:
            return value
    return ""


def _strip_label(value: Any, *labels: str) -> str:
    text = _clean(value)
    for label in labels:
        m = re.match(rf"(?i)^\s*{re.escape(label)}\s*[:\-–—]?\s*(.+?)\s*$", text)
        if m:
            return _clean(m.group(1))
    return text


def _format_plain_number(value: Any) -> str:
    text = _clean(value).replace(",", ".")
    if not text:
        return ""
    try:
        number = float(text)
        return str(int(number)) if number.is_integer() else str(number)
    except ValueError:
        return text


def _owner_is_internal(owner: Any) -> bool:
    text = _clean(owner).lower().replace("ё", "е")
    if not text:
        return False
    text = re.sub(r"[\"'«»“”„]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    compact = re.sub(r"[^0-9a-zа-я]+", "", text)
    return (
        compact in {"ждс", "ооождс", "желдорстрой", "ооожелдорстрой"}
        or "ждс" in text
        or "желдорстрой" in text
        or "жел дор строй" in text
        or "железнодорожное строительство" in text
        or "собствен" in text
    )


def _equipment_count_value(value: Any) -> str:
    text = _clean(value).lower().replace(",", ".")
    if not text:
        return ""
    m = re.fullmatch(r"\s*(\d{1,3})(?:\.0+)?\s*(?:ед\.?|шт\.?|машин\w*|техник\w*|а\s*/\s*с)?\s*", text, re.I)
    return str(int(m.group(1))) if m else ""


def _equipment_ref_value(row: dict[str, str]) -> str:
    return _row_value(row, "Номер/число техники", "Номер техники", "Число техники")


def _work_hours_suffix(hours: Any) -> str:
    value = _clean(hours)
    return f" / время работы {value} ч" if value else ""


def _transport_equipment_line(number_or_count: Any, volume: Any, unit: Any, trips: Any, owner: Any, haul: Any, work_hours: Any = None) -> str:
    owner_text = _clean(owner) or "Н/Д"
    count = _equipment_count_value(number_or_count)
    if count and not _owner_is_internal(owner_text):
        return f"{count} а/с / {_volume(volume)} {_clean(unit) or 'м3'} / {_clean(trips) or 0} рейс / {owner_text}{_haul_suffix(haul)}{_work_hours_suffix(work_hours)}"
    return f"Номер {_compact_number_text(number_or_count) or 'Н/Д'} / {_volume(volume)} {_clean(unit) or 'м3'} / {_clean(trips) or 0} рейс / {owner_text}{_haul_suffix(haul)}{_work_hours_suffix(work_hours)}"


def _work_equipment_line(number_or_count: Any, volume: Any, unit: Any, equipment_type: Any, owner: Any, work_hours: Any = None) -> str:
    owner_text = _clean(owner) or "Н/Д"
    equipment_text = _clean(equipment_type) or "техника"
    count = _equipment_count_value(number_or_count)
    if count and not _owner_is_internal(owner_text):
        return f"{count} ед. / {_volume(volume)} {_clean(unit) or 'м3'} / {equipment_text} / {owner_text}{_work_hours_suffix(work_hours)}"
    return f"Номер {_compact_number_text(number_or_count) or 'Н/Д'} / {_volume(volume)} {_clean(unit) or 'м3'} / {equipment_text} / {owner_text}{_work_hours_suffix(work_hours)}"


def _has_work_equipment_reference(*values: Any) -> bool:
    return _has_value(*values)


def _append_work_fact_or_equipment_line(
    lines: list[str],
    *,
    number_or_count: Any = None,
    volume: Any,
    unit: Any,
    equipment_type: Any = None,
    owner: Any = None,
    model: Any = None,
    work_hours: Any = None,
) -> None:
    if _has_work_equipment_reference(number_or_count, equipment_type, owner, model):
        equipment_text = " ".join(part for part in (_clean(equipment_type), _clean(model)) if part) or "техника"
        lines.append(_work_equipment_line(number_or_count, volume, unit, equipment_text, owner, work_hours))
    elif _has_value(volume):
        lines.append(f"V = {_volume(volume)} {_clean(unit) or 'м3'}")


def _row_has_any(ws, row_idx: int, start_col: int, end_col: int) -> bool:
    return _has_value(*(ws.cell(row_idx, col_idx).value for col_idx in range(start_col, end_col + 1)))


def _max_row4_report_sheet(workbook):
    """Detect the MAX-mutated daily XLSX variant.

    Some reports sent from MAX are based on the old workbook but have inline
    labelled header values in row 4 (A4='Дата ...', B4='Смена ...', C4='Участок ...')
    and the filled tables shifted down. They do not carry the official template
    marker in A1, so the normal old-template adapter would treat header rows as
    facts. This detector is intentionally narrow and prefers the filled
    "Шаблон" sheet when present.
    """
    candidates = []
    for ws in workbook.worksheets:
        if "справоч" in ws.title.lower():
            continue
        row4 = " ".join(_clean(ws.cell(4, col_idx).value).lower() for col_idx in range(1, 6))
        if all(marker in row4 for marker in ("дата", "смен", "участ")):
            score = 10
            if ws.title.strip().lower() == "шаблон":
                score += 5
            for row_idx in (33, 90, 137, 203, 207, 220):
                if _row_has_any(ws, row_idx, 1, min(ws.max_column, 15)):
                    score += 1
            candidates.append((score, ws))
    if not candidates:
        return None
    return max(candidates, key=lambda item: item[0])[1]


def max_row4_workbook_to_report_text(workbook) -> str:
    """Convert MAX row-4 header variant into the standard report text contract."""
    ws = _max_row4_report_sheet(workbook)
    if ws is None:
        raise ValueError("MAX row-4 report variant not found")

    header_row = [_clean(ws.cell(4, col_idx).value) for col_idx in range(1, 6)]
    report_date = _strip_label(header_row[0], "Дата")
    shift_raw = _strip_label(header_row[1], "Смена")
    shift_key = shift_raw.strip().lower()
    shift = "день" if shift_key in {"1", "день", "day"} else ("ночь" if shift_key in {"2", "ночь", "night"} else shift_raw)
    section = _strip_label(header_row[2], "Участок")
    if section and not section.lower().startswith("участ"):
        section = f"участок №{section}"
    direction = _strip_label(header_row[3], "Направление")
    constructives = _strip_label(header_row[4], "Конструктивы")

    lines = [
        f"Дата - {report_date}",
        f"Смена - {shift}",
        f"Участок - {section}",
        f"Направление - {direction}",
        f"Конструктивы - {constructives}",
        f"%% XLSX source sheet: {ws.title} / MAX variant row4 header",
    ]

    source = _clean(ws["E5"].value)
    org = _clean(ws["F5"].value)
    material = _clean(ws["G5"].value)
    volume = _clean(ws["H5"].value)
    if _has_value(source, org, material, volume):
        lines.extend(["|", "Информация по завозу: " + " - ".join(part for part in (source, org, material, volume) if part)])

    lines.extend(["", "===Перевозка==="])
    current_transport_key: tuple[str, str, str] | None = None
    for row_idx in range(33, 85):
        if not _row_has_any(ws, row_idx, 1, 13):
            continue
        kind = _clean(ws.cell(row_idx, 2).value)
        model = _clean(ws.cell(row_idx, 3).value)
        number = _clean(ws.cell(row_idx, 4).value)
        owner = _clean(ws.cell(row_idx, 5).value)
        comment = _clean(ws.cell(row_idx, 6).value)
        mat = _clean(ws.cell(row_idx, 7).value)
        src = _clean(ws.cell(row_idx, 8).value)
        dst = _clean(ws.cell(row_idx, 9).value)
        trips = _format_plain_number(ws.cell(row_idx, 11).value)
        vol = _format_plain_number(ws.cell(row_idx, 12).value)
        unit = _clean(ws.cell(row_idx, 13).value) or "м3"
        haul = _clean(ws.cell(row_idx, 14).value)
        if not (mat and src and dst and (trips or vol)):
            continue
        key = (src, dst, mat)
        if key != current_transport_key:
            lines.append(f"-{src or 'Н/Д'} -> {dst or 'Н/Д'}-")
            lines.append(f"/{mat or 'материал н/д'}/")
            current_transport_key = key
        number_text = _compact_number_text(number) or "Н/Д"
        lines.append(f"Номер {number_text} / {vol or '0'} {unit} / {trips or '0'} рейс / {owner or 'Н/Д'}{_haul_suffix(haul)}")
        if comment:
            lines.append(f"%% {comment}")

    lines.extend(["", "===Работы==="])
    for row_idx in range(90, 117):
        if not _row_has_any(ws, row_idx, 1, 15):
            continue
        constructive = _clean(ws.cell(row_idx, 1).value)
        work = _clean(ws.cell(row_idx, 2).value)
        if not work or work.lower().startswith("вид работ"):
            continue
        pk_rail = _pk_range(ws.cell(row_idx, 3).value, ws.cell(row_idx, 4).value)
        pk_ad = _pk_range(ws.cell(row_idx, 5).value, ws.cell(row_idx, 6).value)
        kind = _clean(ws.cell(row_idx, 8).value)
        model = _clean(ws.cell(row_idx, 9).value)
        number = _clean(ws.cell(row_idx, 10).value)
        owner = _clean(ws.cell(row_idx, 11).value)
        vol = _format_plain_number(ws.cell(row_idx, 13).value)
        unit = _clean(ws.cell(row_idx, 14).value) or "м3"
        comment = _clean(ws.cell(row_idx, 15).value)
        if constructive:
            lines.append(f"-{constructive}-")
        lines.append(f"/{work}/")
        pk_line = f"ПК ВСЖМ: {pk_rail}"
        if pk_ad and pk_ad != "Н/Д":
            pk_line += f" (ПК АД: {pk_ad})"
        lines.append(pk_line)
        _append_work_fact_or_equipment_line(
            lines,
            number_or_count=number,
            volume=vol or "0",
            unit=unit,
            equipment_type=kind,
            model=model,
            owner=owner,
        )
        if comment:
            lines.append(f"%% {comment}")

    lines.extend(["", "===Сопутствующие работы==="])
    for row_idx in range(119, 135):
        if not _row_has_any(ws, row_idx, 1, 14):
            continue
        work = _clean(ws.cell(row_idx, 1).value)
        if not work or work.lower().startswith("вид работ"):
            continue
        pk_rail = _pk_range(ws.cell(row_idx, 2).value, ws.cell(row_idx, 3).value)
        pk_ad = _pk_range(ws.cell(row_idx, 4).value, ws.cell(row_idx, 5).value)
        kind = _clean(ws.cell(row_idx, 7).value)
        model = _clean(ws.cell(row_idx, 8).value)
        number = _clean(ws.cell(row_idx, 9).value)
        owner = _clean(ws.cell(row_idx, 10).value)
        vol = _format_plain_number(ws.cell(row_idx, 11).value)
        unit = _clean(ws.cell(row_idx, 13).value) or "м3"
        comment = _clean(ws.cell(row_idx, 14).value)
        lines.append(f"/{work}/")
        pk_line = f"ПК ВСЖМ: {pk_rail}"
        if pk_ad and pk_ad != "Н/Д":
            pk_line += f" (ПК АД: {pk_ad})"
        lines.append(pk_line)
        _append_work_fact_or_equipment_line(
            lines,
            number_or_count=number,
            volume=vol or "0",
            unit=unit,
            equipment_type=kind,
            model=model,
            owner=owner,
        )
        if comment:
            lines.append(f"%% {comment}")

    lines.extend(["", "===Парк техники==="])
    for row_idx in range(137, 194):
        if not _row_has_any(ws, row_idx, 1, 6):
            continue
        kind = _clean(ws.cell(row_idx, 1).value)
        model = _clean(ws.cell(row_idx, 2).value)
        number = _clean(ws.cell(row_idx, 3).value)
        owner = _clean(ws.cell(row_idx, 4).value)
        status = _clean(ws.cell(row_idx, 5).value)
        comment = _clean(ws.cell(row_idx, 6).value)
        if not _has_value(kind, model, number, status, comment):
            continue
        left = " ".join(part for part in (kind or "Техника", model) if part)
        tail = "; ".join(part for part in (owner or "Н/Д", status, comment) if part)
        lines.append(f"{left} (г/н {number or 'Н/Д'}); {tail}")

    lines.extend(["", "===Персонал==="])
    if _row_has_any(ws, 203, 1, 4):
        for col_idx, category in enumerate(("ДР", "ИТР", "Водители", "Механизаторы"), 1):
            count = _clean(ws.cell(203, col_idx).value)
            if count:
                lines.append(f"{category}: {count}")

    lines.extend(["", "===Накопители==="])
    for row_idx in range(207, 218):
        if not _row_has_any(ws, row_idx, 1, 5):
            continue
        name = _clean(ws.cell(row_idx, 1).value)
        pk = _pk(ws.cell(row_idx, 2).value)
        vol = _format_plain_number(ws.cell(row_idx, 3).value)
        unit = _clean(ws.cell(row_idx, 4).value) or "м3"
        comment = _clean(ws.cell(row_idx, 5).value)
        if name and pk and vol:
            line = f"{name} - {pk} - {vol} {unit}"
            if comment:
                line += f" %% {comment}"
            lines.append(line)

    lines.extend(["", "===Забивка свай==="])
    for row_idx in range(220, 243):
        if not _row_has_any(ws, row_idx, 1, 7):
            continue
        field_code = _clean(ws.cell(row_idx, 1).value)
        pk = _pk(ws.cell(row_idx, 2).value)
        pile_kind = _clean(ws.cell(row_idx, 3).value) or "основные"
        count = _format_plain_number(ws.cell(row_idx, 4).value)
        pile_type = _clean(ws.cell(row_idx, 5).value)
        comment = _clean(ws.cell(row_idx, 7).value)
        if _has_value(field_code, pk, count, pile_type, comment):
            parts = [field_code or "поле н/д", pk or "ПК н/д", pile_kind]
            if count:
                parts.append(f"{count} шт")
            if pile_type:
                parts.append(pile_type)
            if comment:
                parts.append(comment)
            lines.append(" - ".join(parts))

    return "\n".join(lines).strip() + "\n"


def _vehicle_line_with_numbers(kind: Any, model: Any, unit_number: Any, plate: Any, owner: Any) -> str:
    kind_text = _clean(kind) or "Техника"
    model_text = _clean(model)
    left = " ".join(part for part in (kind_text, model_text) if part)
    number_parts = []
    if _clean(unit_number):
        number_parts.append(f"б/н {_clean(unit_number)}")
    if _clean(plate):
        number_parts.append(f"г/н {_clean(plate)}")
    owner_text = _clean(owner)
    number_text = "; ".join(number_parts)
    if number_text:
        return f"{left} ({number_text}); {owner_text or 'Н/Д'}"
    return f"{left}; {owner_text or 'Н/Д'}"


def _compact_number_text(number: Any) -> str:
    text = _clean(number)
    if not text:
        return ""
    lowered = text.lower()
    if any(marker in lowered for marker in ("б/н", "борт", "г/н", "гос")):
        return text
    has_digit = bool(re.search(r"\d", text))
    has_letter = bool(re.search(r"[A-Za-zА-Яа-яЁё]", text))
    if has_digit and not has_letter:
        return f"б/н {text}"
    return text


def _split_pk_chunks(value: Any) -> list[str]:
    text = _clean(value)
    if not text:
        return []
    return [part.strip() for part in re.split(r"[\n;]+", text) if part.strip()]


def one_sheet_workbook_to_report_text(workbook, ws=None) -> str:
    ws = ws or _one_sheet_report_sheet(workbook)
    if ws is None:
        raise ValueError("one-sheet report template marker not found")

    lines = [
        f"Дата - {_clean(ws['B4'].value)}",
        f"Смена - {_clean(ws['B5'].value)}",
        f"Участок - {_clean(ws['B6'].value)}",
        f"Направление - {_clean(ws['B7'].value)}",
        f"Конструктивы - {_clean(ws['B8'].value)}",
        f"%% XLSX source sheet: {ws.title}",
    ]
    markers = _one_sheet_marker_rows(ws)

    delivery_rows = []
    for row in _one_sheet_table(ws, markers, "ЗАВОЗ"):
        source = _row_value(row, "Источник")
        material = _row_value(row, "Материал")
        volume = _row_value(row, "Объем")
        unit = _row_value(row, "Ед.") or "м3"
        if source and material and volume:
            delivery_rows.append(f"{source} - {material} - {_volume(volume)} {unit}")
    if delivery_rows:
        lines.extend(["|", f"Информация по завозу: {'; '.join(delivery_rows)}"])

    lines.extend(["", "===Перевозка==="])
    current_transport_key: tuple[str, str, str] | None = None
    for row in _one_sheet_table(ws, markers, "ПЕРЕВОЗКА"):
        src = _row_value(row, "Откуда")
        dst = _row_value(row, "Куда")
        material = _row_value(row, "Материал")
        number = _equipment_ref_value(row)
        volume = _row_value(row, "Объем")
        unit = _row_value(row, "Ед.") or "м3"
        trips = _row_value(row, "Рейсов")
        owner = _row_value(row, "Организация")
        haul = _row_value(row, "Плечо", "Плечо, км", "Плечо возки, км")
        work_hours = _row_value(row, "Время работы", "Время работы, ч")
        if not (src and dst and material and volume and (number or trips or owner)):
            continue
        key = (src or "Н/Д", dst or "Н/Д", material or "материал н/д")
        if key != current_transport_key:
            lines.append(f"-{key[0]} -> {key[1]}-")
            lines.append(f"/{key[2]}/")
            current_transport_key = key
        lines.append(_transport_equipment_line(number, volume, unit, trips, owner, haul, work_hours))

    work_rows: dict[bool, list[dict[str, str]]] = {True: [], False: []}
    for row in _one_sheet_table(ws, markers, "РАБОТЫ"):
        work_name = _row_value(row, "Работа")
        number = _equipment_ref_value(row)
        volume = _row_value(row, "Объем")
        if not (
            _has_value(work_name)
            and _has_value(
                number,
                volume,
                _row_value(row, "ПК ВСЖМ"),
                _row_value(row, "ПК АД"),
                _row_value(row, "Тип техники"),
                _row_value(row, "Организация"),
                _row_value(row, "Комментарий"),
            )
        ):
            continue
        work_rows[_bool_flag(_row_value(row, "Производительность"), True)].append(row)

    for enabled, title in ((True, "===Работы==="), (False, "===Сопутствующие работы===")):
        rows = work_rows[enabled]
        if not rows:
            continue
        lines.extend(["", title])
        current_work_key: tuple[str, str, str, str] | None = None
        for row in rows:
            constructive = _row_value(row, "Объект/конструктив")
            work_name = _row_value(row, "Работа")
            pk_rail = _row_value(row, "ПК ВСЖМ")
            pk_ad = _row_value(row, "ПК АД")
            number = _equipment_ref_value(row)
            volume = _row_value(row, "Объем")
            unit = _row_value(row, "Ед.") or "м3"
            equipment_type = _row_value(row, "Тип техники")
            owner = _row_value(row, "Организация")
            work_hours = _row_value(row, "Время работы", "Время работы, ч")
            comment = _row_value(row, "Комментарий")
            key = (constructive, work_name, pk_rail, pk_ad)
            if key != current_work_key:
                if constructive:
                    lines.append(f"-{constructive}-")
                lines.append(f"/{work_name or 'Работа н/д'}/")
                pk_line = f"ПК ВСЖМ: {pk_rail or 'Н/Д'}"
                if pk_ad:
                    pk_line += f" (ПК АД: {pk_ad})"
                lines.append(pk_line)
                current_work_key = key
            _append_work_fact_or_equipment_line(
                lines,
                number_or_count=number,
                volume=volume,
                unit=unit,
                equipment_type=equipment_type,
                owner=owner,
                work_hours=work_hours,
            )
            if comment:
                lines.append(f"%% {comment}")

    # Old one-sheet templates may still contain Park/Personnel blocks.
    # They are deprecated and intentionally ignored; transport/works are the source of equipment facts.

    fill_status_columns = [
        ("pioneer_fill", "Пионерка", ("Пионерка",)),
        ("subgrade_not_to_grade", "не в отметку", ("не в отметку", "ЗП в работе")),
        ("dso", "ДСО", ("ДСО", "устройство СДО", "устройство ДСО")),
        ("ready_for_shpgs", "готово под ЩПГС", ("готово под ЩПГС", "Под ЩПГС")),
        ("shpgs_done", "под аб", ("под аб", "ЩПГС уложен", "ЗП готово")),
    ]
    fill_lines: list[str] = []
    for row in _one_sheet_table(ws, markers, "ДАННЫЕ ПО ОТСЫПКЕ"):
        road = _row_value(row, "Автодорога", "Дорога")
        comment = _row_value(row, "Примечание", "Комментарий")
        if not road:
            continue
        for _field, status_label, headers in fill_status_columns:
            for pk_text in _split_pk_chunks(_row_value(row, *headers)):
                line = f"{road}; {pk_text}; {status_label}"
                if comment:
                    line = f"{line}; {comment}"
                fill_lines.append(line)
    if fill_lines:
        lines.extend(["", "===Данные по отсыпке===", *fill_lines])

    mainline_fill_status_columns = [
        ("prep_works", "ПРС", ("ПРС", "Подготовительные работы")),
        ("main_works", "Основные работы", ("Основные работы", "ЗП")),
        ("protective_layer_2", "ЗС №2", ("ЗС №2", "Защитный слой №2")),
        ("protective_layer_1", "ЗС №1", ("ЗС №1", "Защитный слой №1", "готово под ЩПГС")),
        ("asphalt_layer", "Асфальтобетон", ("Асфальтобетон", "под аб", "под АБ")),
    ]
    mainline_fill_lines: list[str] = []
    for row in _one_sheet_table(ws, markers, "ДАННЫЕ ПО ОТСЫПКЕ ОХ"):
        section = _row_value(row, "Участок", "Секция")
        comment = _row_value(row, "Примечание", "Комментарий")
        if not section:
            continue
        for _field, status_label, headers in mainline_fill_status_columns:
            for pk_text in _split_pk_chunks(_row_value(row, *headers)):
                line = f"{section}; {pk_text}; {status_label}"
                if comment:
                    line = f"{line}; {comment}"
                mainline_fill_lines.append(line)
    if mainline_fill_lines:
        lines.extend(["", "===Данные по отсыпке ОХ===", *mainline_fill_lines])

    lines.extend(["", "===Накопители==="])
    for row in _one_sheet_table(ws, markers, "НАКОПИТЕЛИ"):
        name = _row_value(row, "Накопитель")
        pk = _row_value(row, "ПК")
        volume = _row_value(row, "Объем")
        unit = _row_value(row, "Ед.") or "м3"
        comment = _row_value(row, "Комментарий")
        if name and pk and volume:
            line = f"{name or 'Накопитель'} - {_pk(pk) or 'Н/Д'} - {_volume(volume)} {unit}"
            if comment:
                line = f"{line} %% {comment}"
            lines.append(line)

    lines.extend(["", "===Забивка свай==="])
    for row in _one_sheet_table(ws, markers, "ЗАБИВКА СВАЙ"):
        field_code = _row_value(row, "Свайное поле")
        pk = _row_value(row, "ПК")
        operation = _row_value(row, "Операция") or _row_value(row, "Работа")
        pile_kind = _row_value(row, "Тип свай") or "основные"
        count = _row_value(row, "Количество")
        pile_type = _row_value(row, "Вид/тип свай")
        comment = _row_value(row, "Комментарий")
        if not _has_value(field_code, pk, count, pile_type, comment):
            continue
        parts = [field_code or "поле н/д", _pk(pk)]
        if operation:
            parts.append(operation)
        parts.append(pile_kind)
        if count:
            parts.append(f"{count} шт")
        if pile_type:
            parts.append(pile_type)
        if comment:
            parts.append(comment)
        lines.append(" - ".join(part for part in parts if part))

    lines.extend(["", "===Проблемные вопросы==="])
    for row in _one_sheet_table(ws, markers, "ПРОБЛЕМНЫЕ ВОПРОСЫ"):
        text = _row_value(row, "Проблема")
        if text:
            lines.append(text)

    return "\n".join(lines).strip() + "\n"


def compact_workbook_to_report_text(workbook) -> str:
    header = _compact_header_sheet(workbook)
    if header is None:
        raise ValueError("compact report template marker not found")

    lines = [
        f"Дата - {_clean(header['B3'].value)}",
        f"Смена - {_clean(header['B4'].value)}",
        f"Участок - {_clean(header['B5'].value)}",
        f"Направление - {_clean(header['B6'].value)}",
        f"Конструктивы - {_clean(header['B7'].value)}",
    ]
    delivery_rows: list[str] = []
    for source, material, volume, unit, *_ in _iter_filled_rows(header, 12, 18):
        if not (_has_value(source) and _has_value(material) and _has_value(volume)):
            continue
        delivery_rows.append(f"{_clean(source)} - {_clean(material)} - {_volume(volume)} {_clean(unit) or 'м3'}")
    if delivery_rows:
        lines.extend(["|", f"Информация по завозу: {'; '.join(delivery_rows)}"])

    transport_ws = _sheet(workbook, "Перевозка")
    lines.extend(["", "===Перевозка==="])
    current_transport_key: tuple[str, str, str] | None = None
    for src, dst, material, number, volume, unit, trips, owner, haul, *_ in _iter_filled_rows(transport_ws, 3):
        if not (_has_value(src) and _has_value(dst) and _has_value(material) and _has_value(volume) and _has_value(number, trips, owner)):
            continue
        key = (_clean(src), _clean(dst), _clean(material))
        if key != current_transport_key:
            lines.append(f"-{key[0] or 'Н/Д'} -> {key[1] or 'Н/Д'}-")
            lines.append(f"/{key[2] or 'материал н/д'}/")
            current_transport_key = key
        lines.append(
            f"Номер {_compact_number_text(number) or 'Н/Д'} / {_volume(volume)} {_clean(unit) or 'м3'} / {_clean(trips) or 0} рейс / {_clean(owner) or 'Н/Д'}{_haul_suffix(haul)}"
        )

    works_ws = _sheet(workbook, "Работы")
    work_rows: dict[bool, list[tuple[str, str, str, str, str, str, str, str, str, str]]] = {True: [], False: []}
    for constructive, work_name, pk_rail, pk_ad, number, volume, unit, equipment_type, owner, productivity, comment, *_ in _iter_filled_rows(works_ws, 3):
        if not (_has_value(work_name) and _has_value(number, volume, pk_rail, pk_ad, equipment_type, owner, comment)):
            continue
        enabled = _bool_flag(productivity, True)
        work_rows[enabled].append((
            _clean(constructive),
            _clean(work_name),
            _clean(pk_rail),
            _clean(pk_ad),
            _clean(number),
            _volume(volume),
            _clean(unit) or "м3",
            _clean(equipment_type),
            _clean(owner),
            _clean(comment),
        ))

    for enabled, title in ((True, "===Работы==="), (False, "===Сопутствующие работы===")):
        rows = work_rows[enabled]
        if not rows:
            continue
        lines.extend(["", title])
        current_work_key: tuple[str, str, str, str] | None = None
        for constructive, work_name, pk_rail, pk_ad, number, volume, unit, equipment_type, owner, comment in rows:
            key = (constructive, work_name, pk_rail, pk_ad)
            if key != current_work_key:
                if key[0]:
                    lines.append(f"-{key[0]}-")
                lines.append(f"/{key[1] or 'Работа н/д'}/")
                pk_line = f"ПК ВСЖМ: {pk_rail or 'Н/Д'}"
                if key[3]:
                    pk_line += f" (ПК АД: {key[3]})"
                lines.append(pk_line)
                current_work_key = key
            _append_work_fact_or_equipment_line(
                lines,
                number_or_count=number,
                volume=volume,
                unit=unit,
                equipment_type=equipment_type,
                owner=owner,
            )
            if comment:
                lines.append(f"%% {comment}")

    staff_ws = _sheet(workbook, "Персонал")
    lines.extend(["", "===Персонал==="])
    for category, count, *_ in _iter_filled_rows(staff_ws, 3):
        if _has_value(category, count):
            lines.append(f"{_clean(category)}: {_clean(count)}")

    stockpile_ws = _sheet(workbook, "Накопители")
    lines.extend(["", "===Накопители==="])
    for name, pk, volume, unit, *_ in _iter_filled_rows(stockpile_ws, 3):
        if _has_value(name, pk, volume):
            lines.append(f"{_clean(name)} - {_pk(pk)} - {_volume(volume)} {_clean(unit) or 'м3'}")

    piles_ws = _sheet(workbook, "Забивка свай")
    if piles_ws is not None:
        lines.extend(["", "===Забивка свай==="])
        for field_code, pk, pile_kind, count, pile_type, composite, comment, *_ in _iter_filled_rows(piles_ws, 3):
            if not _has_value(field_code, pk, pile_kind, count, pile_type, comment):
                continue
            parts = [_clean(field_code) or "поле н/д", _pk(pk), _clean(pile_kind) or "основные"]
            if _clean(count):
                parts.append(f"{_clean(count)} шт")
            if _clean(pile_type):
                parts.append(_clean(pile_type))
            if _clean(comment):
                parts.append(_clean(comment))
            lines.append(" - ".join(part for part in parts if part))

    problem_ws = _sheet(workbook, "Проблемы")
    lines.extend(["", "===Проблемные вопросы==="])
    for text, *_ in _iter_filled_rows(problem_ws, 3):
        if _clean(text):
            lines.append(_clean(text))

    return "\n".join(lines).strip() + "\n"


def _vehicle_line(kind: Any, model: Any, plate: Any, owner: Any) -> str:
    kind_text = _clean(kind) or "Техника"
    model_text = _clean(model)
    left = " ".join(part for part in (kind_text, model_text) if part)
    owner_text = _clean(owner) or "Н/Д"
    return f"{left} (б/н ; г/н {_clean(plate)}); {owner_text}"


def _pk(value: Any) -> str:
    text = _clean(value)
    if not text or text.lower() in {"н/д", "нд", "-"}:
        return ""
    return text if text.lower().startswith("пк") else f"ПК{text}"


def _pk_range(start: Any, end: Any) -> str:
    left = _pk(start)
    right = _pk(end)
    if left and right:
        return f"{left}-{right}"
    return left or right or "Н/Д"


def _volume(value: Any) -> str:
    text = _clean(value)
    return text if text else "Н/Д"


def _haul_suffix(value: Any) -> str:
    text = _clean(value)
    return f" / плечо {text} км" if text else ""


def _bool_flag(value: Any, default: bool = True) -> bool:
    text = _clean(value).lower().replace("ё", "е")
    if not text:
        return default
    if text in {"да", "д", "yes", "y", "true", "1", "+", "учитывать"}:
        return True
    if text in {"нет", "н", "no", "false", "0", "-", "не учитывать"}:
        return False
    return default


def _append_comment(lines: list[str], comment: Any) -> None:
    text = _clean(comment)
    if text:
        lines.append(f"%% {text}")


def _append_transport(lines: list[str], ws) -> None:
    lines.extend(["", "===Перевозка==="])
    current_key: tuple[str, str, str, str, str] | None = None
    for row_idx in range(32, 77):
        unit_marker = _clean(ws[f"A{row_idx}"].value)
        kind = _clean(ws[f"B{row_idx}"].value)
        model = _clean(ws[f"C{row_idx}"].value)
        plate = _clean(ws[f"D{row_idx}"].value)
        owner = _clean(ws[f"E{row_idx}"].value)
        comment = _clean(ws[f"F{row_idx}"].value)
        material = _clean(ws[f"G{row_idx}"].value)
        source = _clean(ws[f"H{row_idx}"].value)
        destination = _clean(ws[f"I{row_idx}"].value)
        volume = _clean(ws[f"J{row_idx}"].value)
        trips = _clean(ws[f"K{row_idx}"].value)
        unit = _clean(ws[f"L{row_idx}"].value) or "м3"
        haul = _clean(ws[f"M{row_idx}"].value)
        if not _has_value(unit_marker, kind, model, plate, owner, material, source, destination, volume, trips):
            continue

        key = (kind, model, plate, owner, comment)
        if key != current_key:
            lines.extend(["", "Единица техники:"])
            lines.append(_vehicle_line(kind, model, plate, owner))
            _append_comment(lines, comment)
            current_key = key
        if material:
            lines.append(f"/{material}/")
        if source or destination:
            lines.append(f"{source or 'Н/Д'} → {destination or 'Н/Д'}")
        if volume or trips:
            lines.append(f"{_volume(volume)} {unit} / {trips or 0} рейс{_haul_suffix(haul)}")


def _append_work_row(
    lines: list[str],
    *,
    constructive: Any = None,
    work_name: Any,
    rail_start: Any,
    rail_end: Any,
    ad_start: Any,
    ad_end: Any,
    operator: Any,
    equipment_kind: Any,
    model: Any,
    plate: Any,
    owner: Any,
    volume: Any,
    unit: Any,
    comment: Any,
) -> None:
    if not _has_value(work_name, rail_start, rail_end, ad_start, ad_end, operator, equipment_kind, model, plate, owner, volume, unit, comment):
        return
    constructive_text = _clean(constructive)
    if constructive_text:
        lines.append(f"-{constructive_text}-")
    lines.append(f"/{_clean(work_name) or 'Н/Д'}/")
    lines.append(f"ПК ВСЖМ: {_pk_range(rail_start, rail_end)}")
    lines.append(f"ПК АД: {_pk_range(ad_start, ad_end)}")
    if _clean(operator):
        lines.append("Исполнитель:")
    if _has_value(equipment_kind, model, plate, owner):
        lines.append(_vehicle_line(equipment_kind, model, plate, owner))
    if _has_value(volume, unit):
        lines.append(f"V = {_volume(volume)} {_clean(unit) or 'м3'}")
    _append_comment(lines, comment)
    lines.append("")


def _append_main_works(lines: list[str], ws) -> None:
    lines.extend(["", "===Основные работы==="])
    for row_idx in range(80, 105):
        _append_work_row(
            lines,
            constructive=ws[f"A{row_idx}"].value,
            work_name=ws[f"B{row_idx}"].value,
            rail_start=ws[f"C{row_idx}"].value,
            rail_end=ws[f"D{row_idx}"].value,
            ad_start=ws[f"E{row_idx}"].value,
            ad_end=ws[f"F{row_idx}"].value,
            operator=ws[f"G{row_idx}"].value,
            equipment_kind=ws[f"H{row_idx}"].value,
            model=ws[f"I{row_idx}"].value,
            plate=ws[f"J{row_idx}"].value,
            owner=ws[f"K{row_idx}"].value,
            volume=ws[f"L{row_idx}"].value,
            unit=ws[f"M{row_idx}"].value,
            comment=ws[f"N{row_idx}"].value,
        )


def _append_aux_works(lines: list[str], ws) -> None:
    lines.extend(["", "===Сопутствующие работы==="])
    for row_idx in range(108, 123):
        _append_work_row(
            lines,
            work_name=ws[f"A{row_idx}"].value,
            rail_start=ws[f"B{row_idx}"].value,
            rail_end=ws[f"C{row_idx}"].value,
            ad_start=ws[f"D{row_idx}"].value,
            ad_end=ws[f"E{row_idx}"].value,
            operator=ws[f"F{row_idx}"].value,
            equipment_kind=ws[f"G{row_idx}"].value,
            model=ws[f"H{row_idx}"].value,
            plate=ws[f"I{row_idx}"].value,
            owner=ws[f"J{row_idx}"].value,
            volume=ws[f"K{row_idx}"].value,
            unit=ws[f"L{row_idx}"].value,
            comment=ws[f"M{row_idx}"].value,
        )


def _append_park(lines: list[str], ws) -> None:
    lines.extend(["", "===Парк техники==="])
    for row_idx in range(126, 181):
        kind = ws[f"A{row_idx}"].value
        model = ws[f"B{row_idx}"].value
        plate = ws[f"C{row_idx}"].value
        owner = ws[f"D{row_idx}"].value
        status = _clean(ws[f"E{row_idx}"].value)
        comment = _clean(ws[f"F{row_idx}"].value)
        if not _has_value(kind, model, plate, owner, status, comment):
            continue
        tail = "; ".join(part for part in (_clean(owner) or "Н/Д", status, comment) if part)
        left = " ".join(part for part in (_clean(kind) or "Техника", _clean(model)) if part)
        lines.append(f"{left} (б/н ; г/н {_clean(plate)}); {tail}")


def _append_problems(lines: list[str], ws) -> None:
    lines.extend(["", "===Проблемные вопросы==="])
    for row_idx in range(183, 188):
        text = _clean(ws[f"A{row_idx}"].value)
        if text:
            lines.append(text)


def _append_staff_counts(lines: list[str], ws) -> None:
    lines.extend(["", "===Персонал==="])
    categories = (
        ("ИТР", "A191"),
        ("ДР", "B191"),
        ("Механизаторы", "C191"),
        ("Водители", "D191"),
    )
    for category, cell in categories:
        value = _clean(ws[cell].value)
        if value:
            lines.append(f"{category}: {value}")


def _append_stockpiles(lines: list[str], ws) -> None:
    lines.extend(["", "===Накопители==="])
    for row_idx in range(195, 205):
        name = _clean(ws[f"A{row_idx}"].value)
        pk = _pk(ws[f"B{row_idx}"].value)
        volume = _clean(ws[f"C{row_idx}"].value)
        unit = _clean(ws[f"D{row_idx}"].value) or "м3"
        comment = _clean(ws[f"E{row_idx}"].value)
        if not _has_value(name, pk, volume, comment):
            continue
        line = f"{name or 'Накопитель'} - {pk or 'Н/Д'} - {_volume(volume)} {unit}"
        if comment:
            line = f"{line} %% {comment}"
        lines.append(line)


def _append_pile_driving(lines: list[str], ws) -> None:
    lines.extend(["", "===Забивка свай==="])
    for row_idx in range(208, 218):
        field_code = _clean(ws[f"A{row_idx}"].value)
        pk = _pk(ws[f"B{row_idx}"].value)
        pile_kind_raw = _clean(ws[f"C{row_idx}"].value)
        count = _clean(ws[f"D{row_idx}"].value)
        pile_type = _clean(ws[f"E{row_idx}"].value)
        comment_header = _clean(ws["F207"].value).lower()
        comment = _clean(ws[f"G{row_idx}"].value if "состав" in comment_header else ws[f"F{row_idx}"].value)
        if not _has_value(field_code, pk, count, pile_type, comment):
            continue
        pile_kind = pile_kind_raw or "основные"
        parts = [field_code or "поле н/д", pk or "ПК н/д", pile_kind]
        if count:
            parts.append(f"{count} шт")
        if pile_type:
            parts.append(pile_type)
        if comment:
            parts.append(comment)
        lines.append(" - ".join(parts))


def workbook_to_report_text(workbook) -> str:
    ws = _select_report_sheet(workbook)
    lines = [
        f"Дата - {_clean(ws['A5'].value)}",
        f"Смена - {_clean(ws['B5'].value)}",
        f"Участок - {_clean(ws['C5'].value)}",
        f"Направление - {_clean(ws['D5'].value)}",
        f"Конструктивы - {_clean(ws['E5'].value)}",
        f"%% XLSX source sheet: {ws.title}",
    ]
    delivery_info = _clean(ws["B6"].value)
    if delivery_info:
        lines.extend(["|", f"Информация по завозу: {delivery_info}"])
    _append_staff_counts(lines, ws)
    _append_transport(lines, ws)
    _append_main_works(lines, ws)
    _append_aux_works(lines, ws)
    _append_park(lines, ws)
    _append_problems(lines, ws)
    _append_stockpiles(lines, ws)
    _append_pile_driving(lines, ws)
    return "\n".join(lines).strip() + "\n"


def xlsx_bytes_to_report_text(blob: bytes) -> str:
    workbook = load_workbook(BytesIO(blob), data_only=True, read_only=False)
    one_sheet_reports = _one_sheet_report_sheets(workbook)
    if one_sheet_reports:
        if len(one_sheet_reports) > 1:
            return "\n\n".join(one_sheet_workbook_to_report_text(workbook, ws) for ws in one_sheet_reports).strip() + "\n"
        return one_sheet_workbook_to_report_text(workbook)
    if _compact_header_sheet(workbook) is not None:
        return compact_workbook_to_report_text(workbook)
    if _max_row4_report_sheet(workbook) is not None:
        return max_row4_workbook_to_report_text(workbook)
    return workbook_to_report_text(workbook)
